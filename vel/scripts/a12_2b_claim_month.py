"""A12 (Pawan 23-09): fill 'GSTR 3B Month' (column E) on the FY 24-25 2B sheet 'GSTR-2B ITC Data'.

The claim month is the GSTR-3B month in which the credit was actually taken, which lives in the ITC registers, not in the
2B. Each 2B document is looked up by supplier GSTIN + zero-insensitive invoice number:
  1. this year's register (FY 24-25 invoices claimed in FY 25-26 carry KEY2 'PY:<gstin><invoice>' and a 3B Claim Month),
  2. last year's ITC Register 2024-25 in the LY 9C file (3B Claim Month text such as '12 Mar 2025'),
  3. otherwise left blank (never claimed in either year, or matched by amount only).
Written as first-of-month dates (mmm-yy). Values only; the existing 'Claim month (derived)' column AA is left alone."""
import datetime as dt, re, shutil, sys, time, warnings, openpyxl, win32com.client as win32, pythoncom
from openpyxl.utils import column_index_from_string as CI
warnings.filterwarnings("ignore"); sys.path.insert(0, r"C:\PROJECTS\gst-audit-engine"); from vel.scripts.reco_lib import zkey, norm, S
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"; B = P[:-5] + ".xlsb"
try: open(P, "r+b").close()
except PermissionError: raise SystemExit("LOCKED - close in Excel without saving: " + P)
def xlsb_free():
    try: open(B, "r+b").close(); return True
    except PermissionError: return False
shutil.copy(P, "master2_snapshot_before_a12.xlsx")
MON = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7, "aug": 8, "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dec": 12}
def month_of(v):
    """'12 Mar 2025' / '06 Sept 2024' / datetime -> first-of-month datetime, else None."""
    if isinstance(v, dt.datetime): return dt.datetime(v.year, v.month, 1)
    m = re.search(r"([A-Za-z]{3,5})\s*'?(\d{2,4})", S(v))
    if not m or m.group(1).lower() not in MON: return None
    y = int(m.group(2)); y = y + 2000 if y < 100 else y
    return dt.datetime(y, MON[m.group(1).lower()], 1)
E0 = dt.datetime(1899, 12, 30); ser = lambda d: (d - E0).days
t0 = time.time()
# 1. this year's register: PY keys -> claim month
wb = openpyxl.load_workbook(P, read_only=True, data_only=True)
r = wb["ITC Register 2025-26"]; H = {c.value: i for i, c in enumerate(next(r.iter_rows(min_row=5, max_row=5))) if c.value}
cy = {}
for x in r.iter_rows(min_row=6, values_only=True):
    k = S(x[H["KEY2 (matched 2B key)"]])
    if k.startswith("PY:") and isinstance(x[H["3B Claim  Month"]], dt.datetime): cy.setdefault(k[3:], x[H["3B Claim  Month"]])
# 2. last year's register
ly_wb = openpyxl.load_workbook("ly_9c_fy2425.xlsx", read_only=True, data_only=True); lr = ly_wb["ITC Register 2024-25"]
LH = {c.value: i for i, c in enumerate(next(lr.iter_rows(min_row=5, max_row=5))) if c.value}
ly = {}
for x in lr.iter_rows(min_row=6, values_only=True):
    if not x[3]: continue
    m = month_of(x[LH["3B Claim  Month"]])
    if not m: continue
    for g, inv in ((x[LH["Vendor GSTIN"]], x[LH["Invoice No."]]), (x[LH.get("Correct GSTIN", LH["Vendor GSTIN"])], x[LH.get("Correct Invoice no.", LH["Invoice No."])])):
        g = S(g).upper()
        if len(g) == 15 and S(inv): ly.setdefault(g + zkey(inv), m)
ly_wb.close()
# 3. the 2B base
o = wb["GSTR-2B ITC Data"]; OH = {c.value: i for i, c in enumerate(next(o.iter_rows(min_row=5, max_row=5))) if c.value}
rows = list(o.iter_rows(min_row=6, values_only=True)); N = 5 + max(i for i, x in enumerate(rows, 1) if x[0]); rows = rows[:N - 5]
out = []; src = {"this year's register": 0, "last year's register": 0, "blank": 0}
for x in rows:
    g = S(x[OH["GSTIN of supplier"]]).upper(); inv = x[OH["Invoice number"]]
    k_norm, k_z = g + norm(inv), g + zkey(inv)
    m = cy.get(k_norm) or cy.get(k_z)
    if m: src["this year's register"] += 1
    else:
        m = ly.get(k_z)
        if m: src["last year's register"] += 1
        else: src["blank"] += 1
    out.append([ser(dt.datetime(m.year, m.month, 1)) if m else None])
wb.close(); print("read %.0fs | CY PY keys %d | LY register keys %d | 2B base rows %d | filled from %s" % (time.time() - t0, len(cy), len(ly), len(rows), src))
# ---- write
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); xl.Calculation = -4135; ws = w.Worksheets("GSTR-2B ITC Data")
    assert S(ws.Cells(5, CI("E")).Value) == "GSTR 3B Month", ws.Cells(5, 5).Value
    rng = ws.Range("E6:E%d" % N); rng.Value = out; rng.NumberFormat = "mmm-yy"
    ws.Cells(5, CI("E")).AddComment("A12 (23-09): 3B month the credit was claimed in - from this year's register (PY-matched lines) or last year's ITC Register 2024-25, by supplier GSTIN + invoice. Blank = not found in either register.")
    xl.Calculation = -4105; xl.CalculateFullRebuild()
    e = 0
    for sh in w.Worksheets:
        try: e += sh.UsedRange.SpecialCells(-4123, 16).Count
        except Exception: pass
    rg = w.Worksheets("ITC Register 2025-26")
    print("written rows 6..%d | error cells %d | golden %.2f | E6 shows %r" % (N, e, rg.Range("V4").Value, ws.Range("E6").Text))
    assert e == 0 and abs(rg.Range("V4").Value - 1069969542.15) < 0.01
    w.Save(); w.Close(False); print("saved (%.0fs)" % (time.time() - t0))
finally: xl.Quit()
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); print("re-open in a fresh instance: OK (%.0fs)" % (time.time() - t0))
    if xlsb_free(): w.SaveAs(B, FileFormat=50); print("xlsb exported")
    else: print("xlsb OPEN in Excel - export skipped")
    w.Close(False)
finally: xl.Quit()
