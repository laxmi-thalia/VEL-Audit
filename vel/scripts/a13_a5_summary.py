"""A13 + A5 (Pawan 23-09) on ITC Summary.

A13: 'GSTR 3B Gross ITC 6A' (D:F) was blank. GSTR-9 Table 6A is the sum total of GSTR-3B Table 4A, the figure the portal
auto-populates, so it is taken from the FILED 3B returns on '3B Data' (19 GSTINs x 12 months): 4A(3) + 4A(4) + 4A(5) per head
(the client has no 4A(1)/4A(2) import columns in its 3B). Live SUMIFS by GSTIN.

A5: 'Table 8C - ITC availed in 26-27' (BF:BH) summed the 2B sheet as last year; M1 ruled it comes from the FY 26-27 ITC
register filtered to invoices dated FY 25-26 - which is exactly what that sheet now holds (A14) - Category ITC only, since
8C excludes reverse-charge supplies. Table 8D (BI:BK) and everything after it recompute from it."""
import shutil, time, warnings, openpyxl, win32com.client as win32, pythoncom
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
warnings.filterwarnings("ignore")
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"; B = P[:-5] + ".xlsb"
try: open(P, "r+b").close()
except PermissionError: raise SystemExit("LOCKED - close in Excel without saving: " + P)
def xlsb_free():
    try: open(B, "r+b").close(); return True
    except PermissionError: return False
shutil.copy(P, "master2_snapshot_before_a13.xlsx")
S = lambda v: "" if v is None else str(v).strip(); n = lambda v: float(v) if isinstance(v, (int, float)) else 0.0
# expected totals from the source sheets
wb = openpyxl.load_workbook(P, read_only=True, data_only=True)
d3 = [r for r in wb["3B Data"].iter_rows(min_row=3, values_only=True) if r[1]]
want6a = [round(sum(n(r[CI(c) - 1]) for r in d3 for c in cols), 2) for cols in (("P", "S", "V"), ("Q", "T", "W"), ("R", "U", "X"))]
r27 = wb["ITC Register 2026-27"]; H = {c.value: i for i, c in enumerate(next(r27.iter_rows(min_row=5, max_row=5))) if c.value}
d27 = [r for r in r27.iter_rows(min_row=6, values_only=True) if r[3] and S(r[H["Category"]]).upper() == "ITC"]
want8c = [round(sum(n(r[H[k]]) for r in d27), 2) for k in ("IGST", "CGST", "SGST")]
wb.close(); print("expected 6A (3B Data 4A3+4A4+4A5) %s | expected 8C (26-27 register, ITC) %s" % (want6a, want8c))
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); xl.Calculation = -4135
    s = w.Worksheets("ITC Summary"); rg = w.Worksheets("ITC Register 2025-26")
    snap = lambda: [round(s.Cells(25, c).Value or 0, 2) for c in list(range(7, 22)) + [49, 50, 51, 83, 84, 85]]; before = snap()
    gst_rows = [r for r in range(6, 40) if S(s.Cells(r, 3).Value)[:2].isdigit()]; a, z = gst_rows[0], gst_rows[-1]
    bi_before = [round(s.Cells(r, CI("BI")).Value or 0, 2) for r in gst_rows]
    for col, srcs in (("D", ("P", "S", "V")), ("E", ("Q", "T", "W")), ("F", ("R", "U", "X"))):
        s.Range("%s%d:%s%d" % (col, a, col, z)).Formula = "=" + "+".join(
            "SUMIFS('3B Data'!$%s$3:$%s$5000,'3B Data'!$B$3:$B$5000,$C%d)" % (c, c, a) for c in srcs)
    for col, src in (("BF", "S"), ("BG", "T"), ("BH", "U")):
        s.Range("%s%d:%s%d" % (col, a, col, z)).Formula = (
            "=SUMIFS('ITC Register 2026-27'!$%s$6:$%s$5000,'ITC Register 2026-27'!$D$6:$D$5000,$C%d,"
            "'ITC Register 2026-27'!$F$6:$F$5000,\"ITC\")" % (src, src, a))
    s.Cells(3, CI("BF")).Value = "From the FY 26-27 ITC register: invoices dated FY 25-26 availed in 26-27, Category ITC (A5, M1 23-09; was the 2B sheet)"
    xl.Calculation = -4105; xl.CalculateFullRebuild()
    e = 0
    for ws in w.Worksheets:
        try: e += ws.UsedRange.SpecialCells(-4123, 16).Count
        except Exception: pass
    got6a = [round(sum(s.Cells(r, CI(c)).Value or 0 for r in gst_rows), 2) for c in ("D", "E", "F")]
    got8c = [round(sum(s.Cells(r, CI(c)).Value or 0 for r in gst_rows), 2) for c in ("BF", "BG", "BH")]
    bi_after = [round(s.Cells(r, CI("BI")).Value or 0, 2) for r in gst_rows]
    after = snap()
    print("6A block D:F now %s (expected %s) | 8C block BF:BH now %s (expected %s) | 8D IGST moved on %d GSTINs by %.2f in total"
          % (got6a, want6a, got8c, want8c, sum(1 for x, y in zip(bi_before, bi_after) if abs(x - y) > 0.01), round(sum(bi_after) - sum(bi_before), 2)))
    print("error cells %d | golden %.2f | 6A1 blocks unchanged %s" % (e, rg.Range("V4").Value, before == after))
    assert e == 0 and all(abs(x - y) < 1 for x, y in zip(got6a, want6a)) and all(abs(x - y) < 1 for x, y in zip(got8c, want8c)) and before == after and abs(rg.Range("V4").Value - 1069969542.15) < 0.01, "VERIFICATION FAILED"
    w.Save(); w.Close(False); print("saved (%.0fs)" % (time.time() - t0))
finally: xl.Quit()
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); print("re-open in a fresh instance: OK (%.0fs)" % (time.time() - t0))
    if xlsb_free(): w.SaveAs(B, FileFormat=50); print("xlsb exported")
    else: print("xlsb OPEN in Excel - export skipped")
    w.Close(False)
finally: xl.Quit()
