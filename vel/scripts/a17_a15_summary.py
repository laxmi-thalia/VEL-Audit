"""A17 + A15 (Pawan 23-09) on ITC Summary.

A17: the 'GST Portal -> As per Tax Comparison Report' block (BL:BN) summed a pivot at 'Tax comp report'!AR:AU that no
longer exists, so it read 0 everywhere and the 'Difference' block beside it (BO:BR) was overstated. Last year's pivot was
the per-GSTIN sum of the tax comp's own Shortfall/Excess columns L:N (verified against the LY 9C file to the paisa), so the
block now SUMIFS those columns directly - live, no pivot. Every ITC Summary formula that read 'Tax comp report' rows 8..211
is widened to 8..5000 so rows the CAs add later are picked up (Priyesh: "the formulas shouldn't shift").

A15: a per-GSTIN text column appended at the end of the sheet joining the tax comp's Reasons (column Q, except 'Matched')
and its GSTR-9C notes (AI, the 'Yes - ...' ones) month by month, as VALUES (remarks are values - CA ruling)."""
import collections, datetime as dt, re, shutil, time, warnings, openpyxl, win32com.client as win32, pythoncom
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
warnings.filterwarnings("ignore")
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"; B = P[:-5] + ".xlsb"
try: open(P, "r+b").close()
except PermissionError: raise SystemExit("LOCKED - close in Excel without saving: " + P)
def xlsb_free():
    try: open(B, "r+b").close(); return True
    except PermissionError: return False
shutil.copy(P, "master2_snapshot_before_a17.xlsx")
S = lambda v: "" if v is None else str(v).strip()
# ---- read the tax comp reasons per GSTIN (values) for A15
wb = openpyxl.load_workbook(P, read_only=True, data_only=True); tc = wb["Tax comp report"]
rows = [r for r in tc.iter_rows(min_row=8, values_only=True) if r[0] is not None and r[1] is not None]
Q, R_, AI, C_, B_, Lc, Mc, Nc = CI("Q") - 1, CI("R") - 1, CI("AI") - 1, CI("C") - 1, CI("B") - 1, CI("L") - 1, CI("M") - 1, CI("N") - 1
per: dict[str, list[str]] = collections.defaultdict(list); pend: dict[str, int] = collections.Counter(); lsum = collections.defaultdict(lambda: [0.0, 0.0, 0.0])
def mon(v): return v.strftime("%b-%y") if isinstance(v, dt.datetime) else S(v)[:6]
for r in rows:
    g = S(r[B_]).upper(); q = S(r[Q]); note = S(r[AI])
    for k, col in enumerate((Lc, Mc, Nc)):
        v = r[col]; lsum[g][k] += float(v) if isinstance(v, (int, float)) else 0.0
    if q and q.lower() != "matched": per[g].append("%s: %s" % (mon(r[C_]), q))
    if note.startswith("Yes"): per[g].append("%s: 9C note - %s" % (mon(r[C_]), note))
    if S(r[R_]).lower() == "pending": pend[g] += 1
wb.close(); print("tax comp rows %d | GSTINs with reasons %d | pending rows %d" % (len(rows), len(per), sum(pend.values())))
# ---- COM
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); xl.Calculation = -4135
    s = w.Worksheets("ITC Summary"); rg = w.Worksheets("ITC Register 2025-26")
    snap = lambda: [round(s.Cells(25, c).Value or 0, 2) for c in list(range(7, 22)) + [49, 50, 51, 83, 84, 85]]; before = snap()
    gst_rows = [r for r in range(6, 40) if S(s.Cells(r, 3).Value)[:2].isdigit()]; last = gst_rows[-1]
    # A17: widen every reference into the tax comp sheet and rebase the portal block on L:N
    ur = s.UsedRange; f2d = ur.Formula; changed = 0; base_r, base_c = ur.Row, ur.Column
    for i, row in enumerate(f2d):
        for j, f in enumerate(row):
            if isinstance(f, str) and "'Tax comp report'!" in f and "$211" in f:
                s.Cells(base_r + i, base_c + j).Formula = re.sub(r"\$(\d+)\b", lambda m: "$5000" if m.group(1) == "211" else m.group(0), f); changed += 1
    for col, src in (("BL", "L"), ("BM", "M"), ("BN", "N")):
        s.Range("%s%d:%s%d" % (col, gst_rows[0], col, last)).Formula = (
            "=SUMIFS('Tax comp report'!$%s$8:$%s$5000,'Tax comp report'!$B$8:$B$5000,$C%d)" % (src, src, gst_rows[0]))
    # A15: reasons text column at the end
    lastcol = s.Cells(5, s.Columns.Count).End(-4159).Column; nc = lastcol + 1
    h = s.Cells(5, nc); h.Value = "Reasons – Tax comp report (per GSTIN)"; h.Font.Bold = True; h.Font.Color = 0xFFFFFF; h.Interior.Color = 0x4F3F33
    b3 = s.Cells(3, nc); b3.Value = "A15 (23-09): the tax comp sheet's Reasons (col Q, other than Matched) and its 9C notes (col AI), month-wise, as values"; b3.Font.Italic = True
    for r in gst_rows:
        g = S(s.Cells(r, 3).Value).upper(); txt = "; ".join(per.get(g, []))
        if pend.get(g): txt = (txt + "; " if txt else "") + "%d month(s) pending classification" % pend[g]
        s.Cells(r, nc).Value = txt or "Matched"
    s.Columns(nc).ColumnWidth = 90; s.Range(s.Cells(gst_rows[0], nc), s.Cells(last, nc)).WrapText = False
    xl.Calculation = -4105; xl.CalculateFullRebuild()
    # ---- verify
    e = 0
    for ws in w.Worksheets:
        try: e += ws.UsedRange.SpecialCells(-4123, 16).Count
        except Exception: pass
    got = [round(sum(s.Cells(r, CI(c)).Value or 0 for r in gst_rows), 2) for c in ("BL", "BM", "BN")]
    want = [round(sum(v[k] for v in lsum.values()), 2) for k in range(3)]
    after = snap()
    print("formulas widened %d | portal block BL:BN totals %s vs tax comp L:N %s | new reasons column %s | error cells %d | golden %.2f | 6A1 blocks unchanged %s"
          % (changed, got, want, L(nc), e, rg.Range("V4").Value, before == after))
    print("sample:", [(S(s.Cells(r, 3).Value), s.Cells(r, CI("BL")).Value, s.Cells(r, CI("BO")).Value, S(s.Cells(r, nc).Value)[:60]) for r in gst_rows[:3]])
    assert e == 0 and all(abs(a - b) < 1 for a, b in zip(got, want)) and before == after and abs(rg.Range("V4").Value - 1069969542.15) < 0.01, "VERIFICATION FAILED"
    w.Save(); w.Close(False); print("saved (%.0fs)" % (time.time() - t0))
finally: xl.Quit()
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); print("re-open in a fresh instance: OK (%.0fs)" % (time.time() - t0))
    if xlsb_free(): w.SaveAs(B, FileFormat=50); print("xlsb exported")
    else: print("xlsb OPEN in Excel - export skipped")
    w.Close(False)
finally: xl.Quit()
