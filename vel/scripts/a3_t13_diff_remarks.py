"""A3 (M1 after the break, Pawan 23-09) on 'Table 13 & 6A1 differences'.

Last year every GSTIN row carried a Note (P) - the standard sentence 'The ITC pertaining to invoices of FY 20xx-yy ...' - and
a Description (Q) from a three-value set: Matched / Reasons identified / Reasons not identifiable. This year P is blank on
all 19 rows and Q never uses 'Reasons identified' (that is the reason found missing when the two sets were compared).
This script (a) carries last year's Note sentence forward, re-dated to FY 2025-26, onto every row whose Total difference is
not zero, as a VALUE the CA can edit, and (b) puts the three-value Description vocabulary as a dropdown on Q. It does NOT
reclassify anything: a row keeps its Description; the per-GSTIN reasons now available on ITC Summary (A15) are what the
CA will use to move rows from 'not identifiable' to 'identified'."""
import re, shutil, time, warnings, openpyxl, win32com.client as win32, pythoncom
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
warnings.filterwarnings("ignore")
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"; B = P[:-5] + ".xlsb"
try: open(P, "r+b").close()
except PermissionError: raise SystemExit("LOCKED - close in Excel without saving: " + P)
def xlsb_free():
    try: open(B, "r+b").close(); return True
    except PermissionError: return False
shutil.copy(P, "master2_snapshot_before_a3.xlsx")
S = lambda v: "" if v is None else str(v).strip(); SHEET = "Table 13 & 6A1 differences"
ly = openpyxl.load_workbook("ly_9c_fy2425.xlsx", read_only=True, data_only=True)[SHEET]
notes = [S(r[CI("P") - 1]) for r in ly.iter_rows(min_row=7, values_only=True) if r[0] and S(r[0])[:2].isdigit() and S(r[CI("P") - 1])]
template = max(set(notes), key=notes.count); print("LY note template (%d rows): %s" % (len(notes), template[:200]))
new_note = re.sub(r"FY\s*20\d\d-\d\d", "FY 2025-26", template); new_note = re.sub(r"\b2024-25\b", "2025-26", new_note)
DESC = ["Matched", "Reasons identified", "Reasons not identifiable"]
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); xl.Calculation = -4135; ws = w.Worksheets(SHEET)
    hdr = {S(ws.Cells(7, c).Value): c for c in range(1, 40) if ws.Cells(7, c).Value}; assert hdr.get("Note") == CI("P") and hdr.get("Description") == CI("Q") and hdr.get("Total") == CI("O"), hdr
    rows = [r for r in range(8, 40) if S(ws.Cells(r, 1).Value)[:2].isdigit()]
    filled = 0
    for r in rows:
        tot = ws.Cells(r, CI("O")).Value
        if isinstance(tot, (int, float)) and abs(tot) >= 1 and not S(ws.Cells(r, CI("P")).Value):
            ws.Cells(r, CI("P")).Value = new_note; filled += 1
    v = ws.Range("Q8:Q%d" % rows[-1]).Validation; v.Delete(); v.Add(3, 1, 1, ",".join(DESC)); v.ShowError = False; v.InCellDropdown = True
    ws.Cells(6, CI("P")).Value = "Note carried from FY 24-25 (A3, 23-09) - edit per GSTIN"; ws.Cells(6, CI("P")).Font.Italic = True
    ws.Cells(6, CI("Q")).Value = "Description set: " + " / ".join(DESC); ws.Cells(6, CI("Q")).Font.Italic = True
    xl.Calculation = -4105; xl.CalculateFullRebuild()
    e = 0
    for sh in w.Worksheets:
        if sh.Name == "T6A1 Extract - 24-25": continue   # its 2B row pointers are #REF! until cascade_fix rebuilds the extract
        try: e += sh.UsedRange.SpecialCells(-4123, 16).Count
        except Exception: pass
    print("rows %d | notes filled %d | error cells %d | golden %.2f" % (len(rows), filled, e, w.Worksheets("ITC Register 2025-26").Range("V4").Value))
    assert e == 0
    w.Save(); w.Close(False); print("saved (%.0fs)" % (time.time() - t0))
finally: xl.Quit()
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); print("re-open in a fresh instance: OK (%.0fs)" % (time.time() - t0))
    if xlsb_free(): w.SaveAs(B, FileFormat=50); print("xlsb exported")
    else: print("xlsb OPEN in Excel - export skipped")
    w.Close(False)
finally: xl.Quit()
