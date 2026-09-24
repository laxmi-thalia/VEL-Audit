"""A4 (M1, Pawan 23-09): the RCM lines of FY 25-26 whose credit was claimed in FY 26-27 go into ITC Register 2026-27 with
Category = RCM, in the same layout as the ITC lines, so Table 13 / 12C carry RCM as well as ITC.

Source: RCM Register rows whose column BK 'GSTR 3B Claim month' reads April 2026 (392; all paid Mar-26, all dated FY 25-26 -
measured 24-09, the earlier 33-row gap was the payment month stored as a number). Columns are mapped by name onto the
25-26 layout the sheet now has; the reco block and the sheet's own formula columns are extended by copying row 6's
R1C1 formulas. Rows are INSERTED at row 7, inside the data range, so every dependent range grows with them (never insert
at the first data row). fy2627_reco.py + remarks_mirror.py run afterwards to stamp Countif / KEY2 / remarks."""
import datetime as dt, shutil, time, warnings, openpyxl, win32com.client as win32, pythoncom
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
warnings.filterwarnings("ignore")
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
try: open(P, "r+b").close()
except PermissionError: raise SystemExit("LOCKED - close in Excel without saving: " + P)
shutil.copy(P, "master2_snapshot_before_a4.xlsx")
S = lambda v: "" if v is None else str(v).strip(); E0 = dt.datetime(1899, 12, 30)
def xv(v):
    if isinstance(v, dt.datetime): return (v - E0).days
    return v
def fy(d):
    if not isinstance(d, dt.datetime): return ""
    y = d.year if d.month >= 4 else d.year - 1
    return "%d-%02d" % (y, (y + 1) % 100)
wb = openpyxl.load_workbook(P, read_only=True, data_only=True)
r = wb["RCM Register"]; RH = {c.value: i for i, c in enumerate(next(r.iter_rows(min_row=5, max_row=5))) if c.value}
g = lambda x, h: x[RH[h]] if h in RH else None
src = [x for x in r.iter_rows(min_row=6, values_only=True) if (x[1] or x[3]) and S(g(x, "GSTR 3B Claim month")).lower().startswith(("april 2026", "apr-26", "2026-04"))]
n27 = wb["ITC Register 2026-27"]; hd = [c.value for c in next(n27.iter_rows(min_row=5, max_row=5))]; H = {h: i + 1 for i, h in enumerate(hd) if h}
n_before = sum(1 for x in n27.iter_rows(min_row=6, values_only=True) if x[3]); r6 = next(n27.iter_rows(min_row=6, max_row=6, values_only=True))
wb.close()
def inv_date(x):
    d = g(x, "Inv. Date"); return d if isinstance(d, dt.datetime) else g(x, "Document Date")
CLAIM = dt.datetime(2026, 4, 1)
def build(x):
    row = {h: None for h in H}
    d = inv_date(x); pd_ = g(x, "Posting Date")
    row.update({
        "Company": "VEL", "Business place": g(x, "Business place"), "State Name": g(x, "State"), "VEL GSTIN": S(g(x, "MY GSTN")).upper(),
        "3B Claim  Month": CLAIM, "Category": "RCM", "Document Type": g(x, "Document type"), "Document Number": g(x, "Document Number"),
        "Posting Date": pd_, "Posting Year": g(x, "Posting Year") or fy(pd_), "Invoice No.": g(x, "Inv. No.") or g(x, "Reference"),
        "Invoice Date": d, "Invoice Year": fy(d) or g(x, "Doc year"), "Vendor Code": g(x, "Vendor Code"), "Vendor GSTIN": S(g(x, "GSTN")).upper() or "Missing",
        "Vendor Name/RCM Category": g(x, "Vendor Name"), "Taxable Value": g(x, "Taxable Value as per SAP"), "IGST": g(x, "IGST AS PER SAP"),
        "CGST": g(x, "CGST AS PER SAP"), "SGST": g(x, "SGST AS PER SAP"), "GSTR 9C_Reporting": "12C",
        "Reasons": "RCM of FY 25-26 (paid Mar-26) - ITC claimed in Apr-26, FY 26-27", "Nature of Services": g(x, "Nature of Services"),
        "G/L Account": g(x, "G/L Account"), "G/L Account Text": g(x, "GL Name"), "SAP Period": dt.datetime(pd_.year, pd_.month, 1) if isinstance(pd_, dt.datetime) else None,
        "Profit Center": g(x, "Profit Center"), "TYPE": "RCM", "GST CREDIT YES/NO": "Yes", "F.Y/Booking Year": g(x, "Posting Year") or fy(pd_), "Company Code": 1000,
        "GSTR 9_Reporting": "13", "Reasons (GSTR 9)": "RCM of FY 25-26 claimed in FY 26-27 - Table 13", "Source": "RCM Register (BK = April 2026) - A4, 23-09-2026",
    })
    return [xv(row[h]) for h in hd if h]
new = [build(x) for x in src]
tax = round(sum((v or 0) for row in new for v in (row[H["IGST"] - 1], row[H["CGST"] - 1], row[H["SGST"] - 1]) if isinstance(v, (int, float))), 2)
print("RCM Register rows claimed Apr-26: %d | tax %.2f | 26-27 register rows before: %d" % (len(new), tax, n_before))
assert new, "no RCM rows selected"
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); xl.Calculation = -4135; ws = w.Worksheets("ITC Register 2026-27")
    isum = w.Worksheets("ITC Summary"); gst_rows = [rr for rr in range(6, 40) if S(isum.Cells(rr, 3).Value)[:2].isdigit()]
    t13_before = [round(sum(isum.Cells(rr, c).Value or 0 for rr in gst_rows), 2) for c in (CI("CH"), CI("CI"), CI("CJ"))]
    assert ws.Cells(ws.Rows.Count, 4).End(-4162).Row == 5 + n_before
    formula_cols = [c for c in range(1, len(hd) + 1) if isinstance(ws.Cells(6, c).Formula, str) and ws.Cells(6, c).Formula.startswith("=")]
    r1c1 = {c: ws.Cells(6, c).FormulaR1C1 for c in formula_cols}
    k = len(new); ws.Rows("7:%d" % (6 + k)).Insert()
    ws.Range(ws.Cells(7, 1), ws.Cells(6 + k, len(hd))).Value = new
    for c, f in r1c1.items(): ws.Range(ws.Cells(7, c), ws.Cells(6 + k, c)).FormulaR1C1 = f
    for hname, fmt in (("3B Claim  Month", "dd-mm-yy"), ("Posting Date", "dd-mm-yy"), ("Invoice Date", "dd-mm-yy"), ("SAP Period", "dd-mm-yy")):
        ws.Range(ws.Cells(7, H[hname]), ws.Cells(6 + k, H[hname])).NumberFormat = fmt
    last = ws.Cells(ws.Rows.Count, 4).End(-4162).Row
    xl.Calculation = -4105; xl.CalculateFullRebuild()
    e = 0
    for sh in w.Worksheets:
        if sh.Name == "T6A1 Extract - 24-25": continue   # its 2B row pointers are #REF! until cascade_fix rebuilds the extract
        try: e += sh.UsedRange.SpecialCells(-4123, 16).Count
        except Exception: pass
    t13_after = [round(sum(isum.Cells(rr, c).Value or 0 for rr in gst_rows), 2) for c in (CI("CH"), CI("CI"), CI("CJ"))]
    cats = {}
    for v in ws.Range(ws.Cells(6, H["Category"]), ws.Cells(last, H["Category"])).Value: cats[S(v[0])] = cats.get(S(v[0]), 0) + 1
    print("inserted %d rows at 7 | rows now %d | categories %s | error cells %d | Table 13 block (ITC Summary CH:CJ) %s -> %s | golden %.2f"
          % (k, last - 5, cats, e, t13_before, t13_after, w.Worksheets("ITC Register 2025-26").Range("V4").Value))
    assert e == 0 and last - 5 == n_before + k, "VERIFICATION FAILED"
    w.Save(); w.Close(False); print("saved (%.0fs)" % (time.time() - t0))
finally: xl.Quit()
