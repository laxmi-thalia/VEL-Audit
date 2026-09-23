"""A6 (M2, Pawan 23-09): rebuild 'RCM Paid vs ITC Claimed' so the two months are side by side and every side is split by
Taxable / IGST / CGST / SGST instead of total tax only.

One row per GSTIN x PAID month M (Apr-25..Mar-26). Columns:
  A GSTIN | B State | C Paid month (M) | D Claimed month (M+1)
  E:H  RCM paid per RCM Register (Final 3B Month = M): Taxable / IGST / CGST / SGST  (SAP figures)
  I:L  RCM liability per filed GSTR-3B Table 3.1(d), month M: Taxable / IGST / CGST / SGST
  M:O  ITC claimed per filed GSTR-3B Table 4A(3), month M+1: IGST / CGST / SGST
  P:S  ITC claimed per ITC Register RCM lines (3B Claim Month = M+1): Taxable / IGST / CGST / SGST
  T:W  Difference, RCM paid (register) minus 4A(3) claimed in M+1: IGST / CGST / SGST / Total
  X    DPS Remarks (values, carried over from the old sheet by GSTIN + month)
Live SUMIFS throughout; SUBTOTAL totals on row 5; the old sheet is kept renamed '(old)' for one round of review."""
import datetime as dt, shutil, time, warnings, openpyxl, win32com.client as win32, pythoncom
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
warnings.filterwarnings("ignore")
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"; B = P[:-5] + ".xlsb"
try: open(P, "r+b").close()
except PermissionError: raise SystemExit("LOCKED - close in Excel without saving: " + P)
def xlsb_free():
    try: open(B, "r+b").close(); return True
    except PermissionError: return False
shutil.copy(P, "master2_snapshot_before_a6.xlsx")
S = lambda v: "" if v is None else str(v).strip(); E0 = dt.datetime(1899, 12, 30); ser = lambda d: (d - E0).days
NAME = "RCM Paid vs ITC Claimed"
wb = openpyxl.load_workbook(P, read_only=True, data_only=True)
old = wb[NAME]; oh = {c.value: i for i, c in enumerate(next(old.iter_rows(min_row=5, max_row=5))) if c.value}
remarks = {}
for x in old.iter_rows(min_row=6, values_only=True):
    if x[0] and S(x[oh.get("DPS Remarks", 8)]): remarks[(S(x[0]).upper(), S(x[oh.get("Month (3B)", 2)]))] = S(x[oh.get("DPS Remarks", 8)])
isum = wb["ITC Summary"]; states = [(S(r[0]), S(r[2])) for r in isum.iter_rows(min_row=6, max_row=40, values_only=True) if r[2] and S(r[2])[:2].isdigit()]
rr = wb["RCM Register"]; RN = 5 + sum(1 for x in rr.iter_rows(min_row=6, values_only=True) if x[1] or x[3])
reg = wb["ITC Register 2025-26"]; RH = {c.value: i + 1 for i, c in enumerate(next(reg.iter_rows(min_row=5, max_row=5))) if c.value}; RGN = 5 + sum(1 for x in reg.iter_rows(min_row=6, values_only=True) if x[3])
d3 = wb["3B Data"]; D3N = 2 + sum(1 for x in d3.iter_rows(min_row=3, values_only=True) if x[1])
# '3B Data' months are text like 'Apr-25' — read the exact spellings so the criteria match
d3_months = sorted({S(x[2]) for x in d3.iter_rows(min_row=3, values_only=True) if x[1]})
wb.close()
months = [dt.datetime(2025, m, 1) for m in range(4, 13)] + [dt.datetime(2026, m, 1) for m in range(1, 4)]
mtxt = lambda d: d.strftime("%b-%y"); nxt = lambda d: dt.datetime(d.year + (d.month == 12), d.month % 12 + 1, 1)
assert all(mtxt(m) in d3_months for m in months), ("3B Data month spellings", d3_months[:5])
rc = lambda h: L(RH[h])
print("old remarks kept: %d | RCM Register rows 6..%d | register rows 6..%d | 3B Data rows 3..%d | GSTINs %d" % (len(remarks), RN, RGN, D3N, len(states)))
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); xl.Calculation = -4135
    for sh in w.Worksheets:
        if sh.Name == NAME + " (old)": sh.Delete()
    w.Worksheets(NAME).Name = NAME + " (old)"
    ws = w.Worksheets.Add(After=w.Worksheets(NAME + " (old)")); ws.Name = NAME
    ws.Cells(2, 1).Value = "VIKRAN ENGINEERING LIMITED"; ws.Cells(2, 1).Font.Bold = True
    ws.Cells(3, 1).Value = ("RCM paid in month M (RCM Register / GSTR-3B 3.1(d)) vs RCM ITC claimed in month M+1 (GSTR-3B 4A(3) / ITC Register RCM lines). "
                            "VEL claims one month behind payment (A6, Pawan 23-09: both months shown, split by tax head).")
    bands = [(5, 8, "RCM paid – RCM Register (SAP), month M"), (9, 12, "RCM liability – GSTR-3B 3.1(d), month M"), (13, 15, "ITC claimed – GSTR-3B 4A(3), month M+1"),
             (16, 19, "ITC claimed – ITC Register RCM lines, month M+1"), (20, 23, "Difference: RCM paid (register) – 4A(3) claimed M+1")]
    for a, b, txt in bands:
        c = ws.Cells(4, a); c.Value = txt; c.Font.Bold = True; ws.Range(ws.Cells(4, a), ws.Cells(4, b)).Interior.Color = 0xB09784; c.Font.Color = 0xFFFFFF
    hdr = ["GSTIN", "State", "Paid month (M)", "Claimed month (M+1)", "Taxable", "IGST", "CGST", "SGST", "Taxable", "IGST", "CGST", "SGST",
           "IGST", "CGST", "SGST", "Taxable", "IGST", "CGST", "SGST", "IGST", "CGST", "SGST", "Total", "DPS Remarks"]
    for j, h in enumerate(hdr, 1):
        c = ws.Cells(6, j); c.Value = h; c.Font.Bold = True; c.Font.Color = 0xFFFFFF; c.Interior.Color = 0x4F3F33
    rows = []; r = 7
    for st, g in states:
        for m in months:
            nm = nxt(m); ws.Cells(r, 1).Value = g; ws.Cells(r, 2).Value = st; ws.Cells(r, 3).Value = ser(m); ws.Cells(r, 4).Value = ser(nm)
            R = lambda col: "'RCM Register'!$%s$6:$%s$%d" % (col, col, RN)
            for j, col in enumerate(("AR", "AS", "AT", "AU")):
                ws.Cells(r, 5 + j).Formula = "=SUMIFS(%s,%s,$A%d,%s,$C%d)" % (R(col), R("D"), r, R("B"), r)
            D = lambda col: "'3B Data'!$%s$3:$%s$%d" % (col, col, D3N)
            for j, col in enumerate(("L", "M", "N", "O")):
                ws.Cells(r, 9 + j).Formula = "=SUMIFS(%s,%s,$A%d,%s,TEXT($C%d,\"mmm-yy\"))" % (D(col), D("B"), r, D("C"), r)
            for j, col in enumerate(("P", "Q", "R")):
                ws.Cells(r, 13 + j).Formula = "=IF($D%d>DATE(2026,3,1),\"(next FY)\",SUMIFS(%s,%s,$A%d,%s,TEXT($D%d,\"mmm-yy\")))" % (r, D(col), D("B"), r, D("C"), r)
            G = lambda h: "'ITC Register 2025-26'!$%s$6:$%s$%d" % (rc(h), rc(h), RGN)
            for j, h in enumerate(("Taxable Value", "IGST", "CGST", "SGST")):
                ws.Cells(r, 16 + j).Formula = "=SUMIFS(%s,%s,$A%d,%s,$D%d,%s,\"RCM\")" % (G(h), G("VEL GSTIN"), r, G("3B Claim  Month"), r, G("Category"))
            for j in range(3):
                ws.Cells(r, 20 + j).Formula = "=IF(ISNUMBER(%s%d),%s%d-%s%d,\"\")" % (L(13 + j), r, L(6 + j), r, L(13 + j), r)
            ws.Cells(r, 23).Formula = "=IF(ISNUMBER(T%d),SUM(T%d:V%d),\"\")" % (r, r, r)
            rm = remarks.get((g, m.strftime("%d %b %Y").replace(" 0", " ")), "") or remarks.get((g, mtxt(m)), "")
            if rm: ws.Cells(r, 24).Value = rm
            r += 1
    last = r - 1
    for col in range(5, 24):
        t = ws.Cells(5, col); t.Formula = "=SUBTOTAL(9,%s7:%s%d)" % (L(col), L(col), last); t.Font.Bold = True; t.NumberFormat = "#,##0.00"
    ws.Range("C7:D%d" % last).NumberFormat = "mmm-yy"; ws.Range("E7:W%d" % last).NumberFormat = "#,##0.00"
    ws.Range("A6:X%d" % last).AutoFilter(); ws.Columns(1).ColumnWidth = 18; ws.Columns(2).ColumnWidth = 18; ws.Columns(24).ColumnWidth = 50
    for col in range(3, 24): ws.Columns(col).ColumnWidth = 13
    ws.Activate(); xl.ActiveWindow.FreezePanes = False; ws.Range("E7").Select(); xl.ActiveWindow.FreezePanes = True
    xl.Calculation = -4105; xl.CalculateFullRebuild()
    e = 0
    for sh in w.Worksheets:
        try: e += sh.UsedRange.SpecialCells(-4123, 16).Count
        except Exception: pass
    tot = lambda cols: [round(ws.Cells(5, c).Value or 0, 2) for c in cols]
    print("rows 7..%d | paid per register (tax) %s | 3.1(d) tax %s | 4A(3) M+1 %s | register claims M+1 %s | diff %s | error cells %d | golden %.2f"
          % (last, tot((6, 7, 8)), tot((10, 11, 12)), tot((13, 14, 15)), tot((17, 18, 19)), tot((20, 21, 22, 23)), e, w.Worksheets("ITC Register 2025-26").Range("V4").Value))
    assert e == 0 and abs(w.Worksheets("ITC Register 2025-26").Range("V4").Value - 1069969542.15) < 0.01
    w.Save(); w.Close(False); print("saved (%.0fs)" % (time.time() - t0))
finally: xl.Quit()
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); print("re-open in a fresh instance: OK (%.0fs)" % (time.time() - t0))
    if xlsb_free(): w.SaveAs(B, FileFormat=50); print("xlsb exported")
    else: print("xlsb OPEN in Excel - export skipped")
    w.Close(False)
finally: xl.Quit()
