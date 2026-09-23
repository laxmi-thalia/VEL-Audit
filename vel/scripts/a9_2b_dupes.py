"""A9 + A11 (M1, Pawan 23-09): duplicate 2B rows.

Practice (M1): where an amendment (B2BA) exists for a document, drop the original B2B row and keep the revised one. And a
document that arrives twice with identical amounts (the same statement exported twice) is counted once.
  'GSTR-2B Apr25-Aug26' (Octa, header row 2, 'Is Amendment' column): per supplier GSTIN + Doc No key with >1 row -
     revised present  -> delete the '' / 'Yes (Original)' rows of that key;
     no revision, all rows identical in tax -> keep the first, delete the rest;
     different amounts, no revision -> left alone (a supplier may reuse a number: not for a script to decide).
  'GSTR-2B ITC Data' (FY 24-25 PAN export, header row 5, no amendment flag): identical-amount repeats only.
Rows are deleted bottom-up inside the data range so dependent ranges shrink with them. The register's 2B_ SUMIFS then pull
the surviving row only; the '1 –' tie is re-read after the change."""
import collections, shutil, time, warnings, openpyxl, win32com.client as win32, pythoncom
from openpyxl.utils import column_index_from_string as CI
warnings.filterwarnings("ignore")
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"; B = P[:-5] + ".xlsb"
try: open(P, "r+b").close()
except PermissionError: raise SystemExit("LOCKED - close in Excel without saving: " + P)
def xlsb_free():
    try: open(B, "r+b").close(); return True
    except PermissionError: return False
shutil.copy(P, "master2_snapshot_before_a9.xlsx")
S = lambda v: "" if v is None else str(v).strip(); n = lambda v: float(v) if isinstance(v, (int, float)) else 0.0
wb = openpyxl.load_workbook(P, read_only=True, data_only=True)
plan = {}   # sheet -> [(row, reason, tax)]
def scan(sheet, hr, gcol, dcol, taxcols, amcol=None):
    ws = wb[sheet]; hdr = [c.value for c in next(ws.iter_rows(min_row=hr, max_row=hr))]; H = {h: i for i, h in enumerate(hdr) if h}
    rows = [(i, x) for i, x in enumerate(ws.iter_rows(min_row=hr + 1, values_only=True), hr + 1) if x[0]]
    by = collections.defaultdict(list)
    for i, x in rows: by[(S(x[H[gcol]]).upper(), S(x[H[dcol]]).upper())].append((i, x))
    kill = []
    for k, grp in by.items():
        if len(grp) < 2: continue
        tax = lambda x: tuple(round(n(x[H[c]]), 2) for c in taxcols)
        if amcol:
            flags = [S(x[H[amcol]]) for _, x in grp]
            if any(f.startswith("Yes (Revised") for f in flags):
                for (i, x), f in zip(grp, flags):
                    if not f.startswith("Yes (Revised"): kill.append((i, "original superseded by amendment", sum(tax(x))))
                continue
        if len({tax(x) for _, x in grp}) == 1:
            for i, x in grp[1:]: kill.append((i, "identical repeat", sum(tax(x))))
    plan[sheet] = (sorted(kill), len(rows), hr)
    print("%-22s rows %5d | to delete %3d (%s) | tax removed %.2f" % (sheet, len(rows), len(kill), dict(collections.Counter(r for _, r, _ in kill)), sum(t for _, _, t in kill)))
scan("GSTR-2B Apr25-Aug26", 2, "Supplier GSTIN", "Doc No", ("IGST (Net)", "CGST (Net)", "SGST (Net)"), "Is Amendment")
scan("GSTR-2B ITC Data", 5, "GSTIN of supplier", "Invoice number", ("Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)"))
wb.close()
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); xl.Calculation = -4135
    rg = w.Worksheets("ITC Register 2025-26"); H = {rg.Cells(5, c).Value: c for c in range(1, 90) if rg.Cells(5, c).Value}
    RN = rg.Cells(rg.Rows.Count, 4).End(-4162).Row
    col = lambda h: [v[0] for v in rg.Range(rg.Cells(6, H[h]), rg.Cells(RN, H[h])).Value]
    def tie():
        rem, cf = col("Reco Remarks"), col("Countif")
        return [round(sum(n(v) for v, m, c in zip(col(h), rem, cf) if c == "Consider" and S(m).startswith("1 –")), 2) for h in ("2B_IGST", "2B_CGST", "2B_SGST")]
    t_before = tie(); tot_before = round(sum(n(v) for v in col("2B_Total GST")), 2)
    for sheet, (kill, nrows, hr) in plan.items():
        ws = w.Worksheets(sheet)
        for i, _, _ in sorted(kill, reverse=True): ws.Rows(i).Delete()
        print("%s: deleted %d rows, last row now %d" % (sheet, len(kill), ws.Cells(ws.Rows.Count, 1).End(-4162).Row))
    xl.Calculation = -4105; xl.CalculateFullRebuild()
    e = 0
    for sh in w.Worksheets:
        try: e += sh.UsedRange.SpecialCells(-4123, 16).Count
        except Exception: pass
    t_after = tie(); tot_after = round(sum(n(v) for v in col("2B_Total GST")), 2)
    b2 = w.Worksheets("GSTR-2B Apr25-Aug26"); BH = {b2.Cells(2, c).Value: c for c in range(1, 70) if b2.Cells(2, c).Value}; NB = b2.Cells(b2.Rows.Count, 1).End(-4162).Row
    rem = [v[0] for v in b2.Range(b2.Cells(3, BH["Reco Remarks"]), b2.Cells(NB, BH["Reco Remarks"])).Value]
    sheet_1 = [round(sum(n(v[0]) for v, m in zip(b2.Range(b2.Cells(3, BH[h]), b2.Cells(NB, BH[h])).Value, rem) if S(m).startswith("1 –") and "(ITCR 26-27)" not in S(m)), 2) for h in ("IGST (Net)", "CGST (Net)", "SGST (Net)")]
    print("error cells %d | golden %.2f | register 2B_ total %.2f -> %.2f | '1 –' tie register %s vs 2B sheet %s -> %s" % (e, rg.Range("V4").Value, tot_before, tot_after, t_after, sheet_1, "TIES" if all(abs(a - b) < 1 for a, b in zip(t_after, sheet_1)) else "DIFFERS"))
    assert e == 0 and abs(rg.Range("V4").Value - 1069969542.15) < 0.01 and all(abs(a - b) < 1 for a, b in zip(t_after, sheet_1)), "VERIFICATION FAILED"
    w.Save(); w.Close(False); print("saved (%.0fs)" % (time.time() - t0))
finally: xl.Quit()
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); print("re-open in a fresh instance: OK (%.0fs)" % (time.time() - t0))
    if xlsb_free(): w.SaveAs(B, FileFormat=50); print("xlsb exported")
    else: print("xlsb OPEN in Excel - export skipped")
    w.Close(False)
finally: xl.Quit()
