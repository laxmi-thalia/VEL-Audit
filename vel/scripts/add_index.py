"""INDEX sheet: hyperlink navigation for the master. First tab; grouped by step; live HYPERLINK formulas."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
try:
    f=open(P,'r+b'); f.close()
except Exception as e:
    print("LOCKED - close the master first:", e); raise SystemExit(1)
wb = openpyxl.load_workbook(P)
GROUPS = [
 ("SOURCE & REGISTER", [
   ("SR_2025-26","The sales register - single source every reco reads from (27,002 rows)"),
   ("GSTR-1 Data","GSTR-1 as filed + Match Status / Matched SR Doc No (vice-versa)"),
   ("3B Data","GSTR-3B 3.1(a) by GSTIN-month + live books diff & vice-versa status"),
   ("GL Data","SAP GL output-tax rows + Matched with SR (vice-versa)"),
   ("FS Revenue Data","Client's tagged FS revenue extract (state level)"),
 ]),
 ("STEP 1 - REGISTER CHECKS", [
   ("Step 1 Flags","Live gate: blocking / review flags with detail list"),
   ("Sample Selection","144 samples + PO/SO refs + source of each PO"),
   ("SAP Enrichment","How SAP fields were joined onto the register (ODN)"),
   ("Queries (Draft)","Draft client queries - last year's bank re-tested + new"),
   ("Load Log","Which monthly file each row came from"),
   ("Format Legend (preserved)","CA's original format notes"),
 ]),
 ("STEP 2 - SR vs GSTR-1", [
   ("S2 Reco (CA format)","Totality reco in the CA's grid + auto remarks"),
   ("S2 Pivot Month-on-Month","PivotTable view - double-click a figure to drill down; right-click > Refresh after edits"),
   ("S2 Month-on-Month","State-month-type detail; type remarks here"),
   ("S2 Exceptions","The named documents behind every difference"),
 ]),
 ("STEP 3 - GSTR-1 vs 3B & SR vs 3B", [
   ("S3 1 vs 3B","GSTR-1 vs GSTR-3B totality"),
   ("S3 Month-on-Month","GSTR-1 vs 3B by state-month"),
   ("S3 SR vs 3B","Books straight to 3B - state level"),
   ("S3 SR vs 3B MoM","Books straight to 3B - month level"),
   ("S3 Amendment Check","FY 26-27 amendment window - pending exports"),
 ]),
 ("STEP 4 - GL vs SR", [
   ("S4 GL vs SR","GL output tax vs register by state"),
   ("S4 Exceptions","The 3+2+2 documents behind the Rs 4.7L difference"),
 ]),
 ("STEP 5 - HSN / RATE", [
   ("S5 HSN Summary","HSN x rate x UQC (GSTR-9 Table 17) + master check"),
   ("S5 Rate-wise","Rate-wise taxable & tax (9C Table 9 feed)"),
 ]),
 ("STEP 6 - ADVANCES", [
   ("S6 Advances Control","Opening + received - adjusted = closing, per state"),
   ("S6 Month-on-Month","The control account by state-month; opening rolls forward"),
   ("S6 GL Advance Check","Advance ledger vs books, all 8 states tie"),
 ]),
 ("STEP 7 / 9 - CN & FS", [
   ("S7 CN Time-bar","Sec 34(2) window per credit note - live from register BF/BG"),
   ("S9 Sales Reco","FS vs GST turnover in last year's 9C layout (Table 5 refs)"),
 ]),
 ("OPEN ITEMS", [
   ("Open Points (all)","All 38 open points, merged"),
   ("Open Points","Step 1 open points (original)"),
 ]),
]
if "INDEX" in wb.sheetnames: del wb["INDEX"]
ix = wb.create_sheet("INDEX", 0)
ix.sheet_view.showGridLines = False
ix["B2"] = "VEL - GST Audit FY 2025-26 - MASTER WORKBOOK"
ix["B2"].font = Font(bold=True, size=16, color="1F4E79")
ix["B3"] = "Click a sheet name to jump. All recos read live from SR_2025-26."
ix["B3"].font = Font(italic=True, color="808080")
r = 5
missing = []
LINK = Font(color="0563C1", underline="single")
HDRF = Font(bold=True, color="FFFFFF"); HDRB = PatternFill("solid", fgColor="1F4E79")
for grp, items in GROUPS:
    c = ix.cell(r, 2, grp); c.font = HDRF; c.fill = HDRB
    ix.cell(r, 3, "").fill = HDRB
    r += 1
    for nm, desc in items:
        if nm not in wb.sheetnames:
            missing.append(nm); continue
        cell = ix.cell(r, 2)
        cell.value = '=HYPERLINK("#\'%s\'!A1","%s")' % (nm.replace('"','""'), nm)
        cell.font = LINK
        d = ix.cell(r, 3, desc); d.font = Font(color="404040")
        r += 1
    r += 1
ix.column_dimensions["A"].width = 2
ix.column_dimensions["B"].width = 30
ix.column_dimensions["C"].width = 72
# every sheet listed?
listed = {nm for _, items in GROUPS for nm, _ in items}
unlisted = [s for s in wb.sheetnames if s not in listed and s != "INDEX"]
wb.save(P)
print("INDEX added as first tab | groups:", len(GROUPS), "| missing sheet names:", missing, "| sheets not on index:", unlisted)
