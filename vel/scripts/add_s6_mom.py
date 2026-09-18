"""S6 Month-on-Month: advances control at state x month. Opening chains from the
annual S6 opening (first month) then each month's closing rolls forward. All live."""
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f=open(P,'r+b'); f.close()
HF=Font(bold=True,color="FFFFFF"); HB=PatternFill("solid",fgColor="1F4E79"); TOT=Font(bold=True)
wb = openpyxl.load_workbook(P)
ws0 = wb["SR_2025-26"]
H = {S(ws0.cell(4,c).value): L(c) for c in range(1, ws0.max_column+1)}
SR = lambda n: "'SR_2025-26'!$%s$5:$%s$27006" % (H[n], H[n])
ng = wb["GSTR-1 Data"].max_row
G1 = lambda c: "'GSTR-1 Data'!$%s$2:$%s$%d" % (c, c, ng)
ac = wb["S6 Advances Control"]
states = [(S(ac.cell(r,1).value), S(ac.cell(r,2).value)) for r in range(4,23)]
MONTHS = ["Apr-25","May-25","Jun-25","Jul-25","Aug-25","Sep-25","Oct-25","Nov-25","Dec-25","Jan-26","Feb-26","Mar-26"]
NM = "S6 Month-on-Month"
if NM in wb.sheetnames: del wb[NM]
pos = wb.sheetnames.index("S6 Advances Control") + 1
ws = wb.create_sheet(NM, pos)
ws["A1"] = "S6 - Advances control, month on month. Opening chains from the annual opening; every figure live. Type remarks in K."
ws["A1"].font = TOT
hdr = ["State","GSTIN","Month","Opening","Received (books)","Adjusted (books, |sum|)","Closing",
       "GSTR-1 net advances","Books net","Books vs GSTR-1","Remark (type here)"]
for c,h in enumerate(hdr,1):
    x=ws.cell(3,c,h); x.font=HF; x.fill=HB
r = 3
for st, g_ in states:
    for i,m in enumerate(MONTHS):
        r += 1
        ws.cell(r,1,st); ws.cell(r,2,g_); ws.cell(r,3,m)
        if i == 0:
            ws.cell(r,4).value = "=IFERROR(INDEX('S6 Advances Control'!$C$4:$C$22,MATCH($B%d,'S6 Advances Control'!$B$4:$B$22,0)),0)" % r
        else:
            ws.cell(r,4).value = "=G%d" % (r-1)
        ws.cell(r,5).value = '=SUMIFS(%s,%s,$B%d,%s,$C%d,%s,"Received")' % (SR("Taxable Value"),SR("My GSTIN"),r,SR("Month"),r,SR("Adv Bucket (helper)"))
        ws.cell(r,6).value = '=-SUMIFS(%s,%s,$B%d,%s,$C%d,%s,"Adjusted")' % (SR("Taxable Value"),SR("My GSTIN"),r,SR("Month"),r,SR("Adv Bucket (helper)"))
        ws.cell(r,7).value = "=D%d+E%d-F%d" % (r,r,r)
        ws.cell(r,8).value = ('=SUMIFS(%s,%s,$B%d,%s,$C%d,%s,"Advance received")+SUMIFS(%s,%s,$B%d,%s,$C%d,%s,"Advance adjusted")'
            % (G1("F"),G1("A"),r,G1("K"),r,G1("J"), G1("F"),G1("A"),r,G1("K"),r,G1("J")))
        ws.cell(r,9).value = "=E%d-F%d" % (r,r)
        ws.cell(r,10).value = "=I%d-H%d" % (r,r)
        for cc in range(4,11): ws.cell(r,cc).number_format="#,##0.00"
last = r
r += 1
ws.cell(r,1,"Total").font=TOT
ws.cell(r,4).value = '=SUMIFS(D4:D%d,$C$4:$C$%d,"Apr-25")' % (last,last)
for cc in (5,6,8,9,10): ws.cell(r,cc).value = "=SUM(%s4:%s%d)" % (L(cc),L(cc),last)
ws.cell(r,7).value = '=SUMIFS(G4:G%d,$C$4:$C$%d,"Mar-26")' % (last,last)
for cc in range(4,11): ws.cell(r,cc).number_format="#,##0.00"; ws.cell(r,cc).font=TOT
for c_,w in zip("ABCDEFGHIJK",[20,18,9,16,17,19,16,17,15,15,32]): ws.column_dimensions[c_].width=w
ws.freeze_panes = "A4"
ws.auto_filter.ref = "A3:K%d" % last
wb.save(P)
print("added '%s': %d rows, totals row %d" % (NM, last-3, r))
