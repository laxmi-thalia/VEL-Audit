"""Add 'S3 SR vs 3B MoM' to the master: state x month, SR (register incl advances, live SUMIFS)
vs GSTR-3B 3.1(a) (live SUMIFS over 3B Data), taxable + CGST, diff, typed-remark column."""
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
try:
    f=open(P,'r+b'); f.close()
except Exception as e:
    print("LOCKED - close the master first:", e); raise SystemExit(1)
HF=Font(bold=True,color="FFFFFF"); HB=PatternFill("solid",fgColor="1F4E79"); TOT=Font(bold=True)
wb = openpyxl.load_workbook(P)
ws0 = wb["SR_2025-26"]
H = {S(ws0.cell(4,c).value): L(c) for c in range(1, ws0.max_column+1)}
SR = lambda n: "'SR_2025-26'!$%s$5:$%s$27006" % (H[n], H[n])
n3 = wb["3B Data"].max_row
B3 = lambda c: "'3B Data'!$%s$2:$%s$%d" % (c, c, n3)
GSTINS = [("37AAECR0503Q1Z7","Andhra Pradesh"),("12AAECR0503Q1ZJ","Arunachal Pradesh"),("18AAECR0503Q1Z7","Assam"),
 ("10AAECR0503Q1ZN","Bihar"),("22AAECR0503Q1ZI","Chhattisgarh"),("24AAECR0503Q1ZE","Gujarat"),("06AAECR0503Q1ZC","Haryana"),
 ("01AAECR0503Q1ZM","Jammu & Kashmir"),("20AAECR0503Q1ZM","Jharkhand"),("29AAECR0503Q1Z4","Karnataka"),
 ("32AAECR0503Q1ZH","Kerala"),("23AAECR0503Q1ZG","Madhya Pradesh"),("27AAECR0503Q1Z8","Maharashtra"),
 ("03AAECR0503Q1ZI","Punjab"),("08AAECR0503Q1Z8","Rajasthan"),("33AAECR0503Q1ZF","Tamil Nadu"),
 ("36AAECR0503Q1Z9","Telangana"),("09AAECR0503Q1Z6","Uttar Pradesh"),("19AAECR0503Q1Z5","West Bengal")]
MONTHS = ["Apr-25","May-25","Jun-25","Jul-25","Aug-25","Sep-25","Oct-25","Nov-25","Dec-25","Jan-26","Feb-26","Mar-26"]
NM = "S3 SR vs 3B MoM"
if NM in wb.sheetnames: del wb[NM]
pos = wb.sheetnames.index("S3 SR vs 3B") + 1
ws = wb.create_sheet(NM, pos)
ws["A1"] = "S3 - Sales Register (books incl advances) vs GSTR-3B 3.1(a), month on month. Both sides live; type remarks in J."
ws["A1"].font = TOT
hdr = ["State","GSTIN","Month","SR Taxable (incl adv)","3B Taxable","Diff Taxable","SR CGST","3B CGST","Diff CGST","Remark (type here)"]
for c,h in enumerate(hdr,1):
    x = ws.cell(3,c,h); x.font=HF; x.fill=HB
r = 3
for g_, st in GSTINS:
    for m in MONTHS:
        r += 1
        ws.cell(r,1,st); ws.cell(r,2,g_); ws.cell(r,3,m)
        ws.cell(r,4).value = '=SUMIFS(%s,%s,$B%d,%s,$C%d)' % (SR("Taxable Value"),SR("My GSTIN"),r,SR("Month"),r)
        ws.cell(r,5).value = '=SUMIFS(%s,%s,$B%d,%s,$C%d)' % (B3("D"),B3("B"),r,B3("C"),r)
        ws.cell(r,6).value = "=D%d-E%d" % (r,r)
        ws.cell(r,7).value = '=SUMIFS(%s,%s,$B%d,%s,$C%d)' % (SR("CGST Amount"),SR("My GSTIN"),r,SR("Month"),r)
        ws.cell(r,8).value = '=SUMIFS(%s,%s,$B%d,%s,$C%d)' % (B3("F"),B3("B"),r,B3("C"),r)
        ws.cell(r,9).value = "=F%d-G%d" % (r,r)  # placeholder fixed below
        ws.cell(r,9).value = "=G%d-H%d" % (r,r)
        for cc in range(4,10): ws.cell(r,cc).number_format = "#,##0.00"
last = r
r += 1
ws.cell(r,1,"Total").font = TOT
for cc in range(4,10):
    ws.cell(r,cc).value = "=SUM(%s4:%s%d)" % (L(cc),L(cc),last)
    ws.cell(r,cc).number_format = "#,##0.00"; ws.cell(r,cc).font = TOT
for c_,w in zip("ABCDEFGHIJ",[20,18,9,18,18,16,15,15,14,40]): ws.column_dimensions[c_].width = w
ws.freeze_panes = "A4"
ws.auto_filter.ref = "A3:J%d" % last
wb.save(P)
print("added '%s': %d state-month rows + total row at %d" % (NM, last-3, r))
