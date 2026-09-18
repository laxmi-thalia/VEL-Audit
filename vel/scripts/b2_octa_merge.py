"""Batch-1 step A: merge the Octa GSTR-2B exports (14 state files Apr-25..Mar-26 + all-states FY 26-27 file,
capped at Aug-26) into the master as 'GSTR-2B Apr25-Aug26' (Purchase-Net, 45 raw cols verbatim + our cols LAST)
and '2B ISD Apr25-Aug26' (ISD, 19 raw cols + tags). Row 1 reserved; header row 2; data from row 3.
Our columns (last): Source file | FY (2B period) | Doc FY (invoice date) | KEY (supplier GSTIN & normalised doc no)
| 3B Claim Month | Reco Remarks | 6A1 mark | Table 8A | 8A Reco | Table 8C | Table 13 | Final Remarks | GSTR-9/9C
| Permanent Reversals | Reclaim - Table 6H | Query   (classification cols filled in later steps)."""
import os, re, warnings, datetime as dt, collections, pickle
from copy import copy
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
B = "//192.168.1.69/gst folder/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/DPS Workings/Portal Reports/GSTR-2B/"
SP = os.path.dirname(os.path.abspath(__file__))
open(P, "r+b").close()
CAP = dt.datetime(2026, 8, 31)
def fy(d):
    if not isinstance(d, dt.datetime): return ""
    y = d.year if d.month >= 4 else d.year - 1
    return "%d-%02d" % (y, (y + 1) % 100)
def ninv(v): return re.sub(r"[^A-Z0-9]", "", S(v).upper().lstrip("'"))
PN_H = ISD_H = None; pn_rows = []; isd_rows = []; hdr_style = None
files = sorted(f for f in os.listdir(B) if f.endswith(".xlsx") and not f.startswith("~$"))
for fn in files:
    wb = openpyxl.load_workbook(B + fn, data_only=True)
    tag = fn.replace("GSTR2B-Net-VIKRAN ENGINEERING LIMITED-", "").replace("GSTR2B-VIKRAN ENGINEERING LIMITED-", "").replace(".xlsx", "")
    for sn, store, key in (("Purchase-Net", pn_rows, "PN"), ("ISD", isd_rows, "ISD")):
        if sn not in wb.sheetnames: print("  no", sn, "in", fn); continue
        ws = wb[sn]
        H = [S(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
        while H and H[-1] == "": H.pop()
        if key == "PN":
            if PN_H is None: PN_H = H; hdr_style = copy(ws.cell(1, 1).font), copy(ws.cell(1, 1).fill)
            elif H != PN_H: raise SystemExit("header drift in %s: %s" % (fn, H))
        else:
            if ISD_H is None: ISD_H = H
            elif H != ISD_H: raise SystemExit("ISD header drift in %s" % fn)
        n = 0
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or not r[0]: continue
            per = r[1]
            if isinstance(per, dt.datetime) and per > CAP: continue
            store.append((tag, list(r[:len(H)]))); n += 1
        print("  %-52s %-12s %5d rows" % (tag, sn, n))
    wb.close()
print("Purchase-Net rows:", len(pn_rows), "| ISD rows:", len(isd_rows))
iPer, iDate, iSup, iDoc = PN_H.index("Tax Period"), PN_H.index("Doc Date"), PN_H.index("Supplier GSTIN"), PN_H.index("Doc No")
OUR = ["Source file", "FY (2B period)", "Doc FY (doc date)", "KEY", "3B Claim Month", "Reco Remarks", "6A1 mark", "Table 8A", "8A Reco", "Table 8C",
       "Table 13", "Final Remarks", "GSTR-9/9C", "Permanent Reversals", "Reclaim - Table 6H", "Query"]
wb = openpyxl.load_workbook(P)
for nm in ("GSTR-2B Apr25-Aug26", "2B ISD Apr25-Aug26"):
    if nm in wb.sheetnames: del wb[nm]
pos = wb.sheetnames.index("GSTR-2B ITC Data")
ws = wb.create_sheet("GSTR-2B Apr25-Aug26", pos)
ws.row_dimensions[1].height = 21
HF, HB = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
OF, OB = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="7F6000")
for c, h in enumerate(PN_H + OUR, 1):
    x = ws.cell(2, c, h)
    if c <= len(PN_H): x.font = copy(hdr_style[0]); x.fill = copy(hdr_style[1])
    else: x.font = OF; x.fill = OB
    x.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[2].height = 32
pn_rows.sort(key=lambda t: (t[1][iPer] if isinstance(t[1][iPer], dt.datetime) else dt.datetime(2099, 1, 1), S(t[1][0])))
r = 2
NF = "#,##0.00"
for tag, row in pn_rows:
    r += 1
    for c, v in enumerate(row, 1):
        cell = ws.cell(r, c, v)
        if isinstance(v, dt.datetime): cell.number_format = "DD-MM-YYYY"
        elif isinstance(v, float): cell.number_format = NF
    base = len(PN_H)
    ws.cell(r, base + 1, tag); ws.cell(r, base + 2, fy(row[iPer])); ws.cell(r, base + 3, fy(row[iDate]))
    ws.cell(r, base + 4, S(row[iSup]).upper() + ninv(row[iDoc]))
N = r
NC = len(PN_H) + len(OUR)
for c in range(1, NC + 1): ws.column_dimensions[L(c)].width = 14
for c, w in ((1, 18), (5, 18), (7, 18), (8, 28), (len(PN_H) + 1, 30), (len(PN_H) + 4, 30), (len(PN_H) + 6, 34)): ws.column_dimensions[L(c)].width = w
ws.freeze_panes = "A3"; ws.auto_filter.ref = "A2:%s%d" % (L(NC), N)
ws.cell(1, 3, "Octa GSTR-2B exports merged: 14 state files Apr-25..Mar-26 + all-states FY 26-27 file (Apr..Jul-26 available; capped at Aug-26). Raw Octa columns A..%s verbatim; DPS columns from %s. Missing FY 25-26 state files: Tamil Nadu, Telangana, West Bengal, Haryana; Kerala file empty." % (L(len(PN_H)), L(len(PN_H) + 1))).font = Font(italic=True, size=9, color="808080")
# ISD sheet
ws2 = wb.create_sheet("2B ISD Apr25-Aug26", pos + 1)
ws2.row_dimensions[1].height = 21
OUR2 = ["Source file", "FY (2B period)", "3B Claim Month", "Reco Remarks"]
for c, h in enumerate(ISD_H + OUR2, 1):
    x = ws2.cell(2, c, h); x.font = HF if c <= len(ISD_H) else OF; x.fill = HB if c <= len(ISD_H) else OB; x.alignment = Alignment(wrap_text=True, vertical="center")
jPer = ISD_H.index("Tax Period")
isd_rows.sort(key=lambda t: (t[1][jPer] if isinstance(t[1][jPer], dt.datetime) else dt.datetime(2099, 1, 1), S(t[1][0])))
r = 2
for tag, row in isd_rows:
    r += 1
    for c, v in enumerate(row, 1):
        cell = ws2.cell(r, c, v)
        if isinstance(v, dt.datetime): cell.number_format = "DD-MM-YYYY"
        elif isinstance(v, float): cell.number_format = NF
    ws2.cell(r, len(ISD_H) + 1, tag); ws2.cell(r, len(ISD_H) + 2, fy(row[jPer]))
N2 = r
for c in range(1, len(ISD_H) + len(OUR2) + 1): ws2.column_dimensions[L(c)].width = 14
ws2.freeze_panes = "A3"; ws2.auto_filter.ref = "A2:%s%d" % (L(len(ISD_H) + len(OUR2)), N2)
# INDEX rows
ix = wb["INDEX"]
from openpyxl.styles import Border, Side
thin = Side(style="thin"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
have = {S(ix.cell(r2, 2).value) for r2 in range(5, ix.max_row + 1)}
for desc, sheet, lbl in (("GSTR-2B Apr-25 to Aug-26 (Octa exports merged; base for 6A1 / 8A / 8C)", "GSTR-2B Apr25-Aug26", "2B Apr25-Aug26"),
                         ("GSTR-2B ISD Apr-25 to Aug-26 (Octa)", "2B ISD Apr25-Aug26", "2B ISD Apr25-Aug26")):
    if desc in have: continue
    r2 = ix.max_row + 1
    ix.cell(r2, 1, "ITC"); ix.cell(r2, 2, desc)
    x = ix.cell(r2, 7); x.value = '=HYPERLINK("#\'%s\'!A1","%s")' % (sheet, lbl); x.font = Font(color="0563C1", underline="single")
    for c in range(1, 11): ix.cell(r2, c).border = BD
wb.save(P)
pickle.dump({"PN_H": PN_H, "OUR": OUR, "N": N, "ISD_H": ISD_H, "N2": N2}, open(os.path.join(SP, "b2_octa_meta.pkl"), "wb"))
print("saved: GSTR-2B Apr25-Aug26 rows 3..%d (%d cols) | 2B ISD rows 3..%d" % (N, NC, N2))
