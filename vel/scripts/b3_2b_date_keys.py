"""B3 on the 2B side (Pawan 24-09): 'GSTR-2B ITC Data' carries 8 invoice numbers the Octa export turned into dates.
(a) The sheet's own KEY formula concatenated the date's serial number; it now keys a date-typed invoice number as m/yy,
    the same rule cascade_fix applies on the books side, so both sides can meet.
(b) Those rows get a 'verify' flag in the Reason column (appended after any existing text), mirroring the register flag."""
import datetime as dt, shutil, time, warnings, openpyxl, win32com.client as win32, pythoncom
from openpyxl.utils import get_column_letter as L
warnings.filterwarnings("ignore")
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
try: open(P, "r+b").close()
except PermissionError: raise SystemExit("LOCKED - close in Excel without saving: " + P)
shutil.copy(P, "master2_snapshot_before_b3_2b.xlsx")
S = lambda v: "" if v is None else str(v).strip(); SHEET = "GSTR-2B ITC Data"; FLAG = "invoice no. is a DATE in the 2B export, verify"
wb = openpyxl.load_workbook(P, read_only=True, data_only=True); ws = wb[SHEET]
H = {S(c.value): i + 1 for i, c in enumerate(next(ws.iter_rows(min_row=5, max_row=5))) if c.value}
KEY = next(h for h in H if h.startswith("KEY (")); INV, GST, REASON = H["Invoice number"], H["GSTIN of supplier"], H["Reason"]
rows = []; int_keys = []; last = 5
for i, x in enumerate(ws.iter_rows(min_row=6, values_only=True), 6):
    if not x[0]: continue
    last = i
    if isinstance(x[INV - 1], dt.datetime): rows.append((i, x[INV - 1], S(x[REASON - 1])))
    elif isinstance(x[INV - 1], int) and len(int_keys) < 50: int_keys.append((i, S(x[H[KEY] - 1])))
wb.close(); print("rows 6..%d | KEY col %s | Invoice col %s | Reason col %s | date-typed invoice numbers: %d" % (last, L(H[KEY]), L(INV), L(REASON), len(rows)))
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); w = xl.Workbooks.Open(P); xl.Calculation = -4135; sh = w.Worksheets(SHEET)
    f = sh.Cells(6, H[KEY]).FormulaR1C1; inv_ref = "RC%d" % INV
    assert f.count(inv_ref) == 1 and "CELL(" not in f, f
    # only a DATE-formatted cell is converted: ISNUMBER would also catch the 2,276 plain integer invoice numbers
    f2 = f.replace(inv_ref, 'IF(LEFT(CELL("format",%s),1)="D",TEXT(%s,"m/yy"),%s)' % (inv_ref, inv_ref, inv_ref))
    sh.Range(sh.Cells(6, H[KEY]), sh.Cells(last, H[KEY])).FormulaR1C1 = f2
    print("KEY formula now:", sh.Cells(6, H[KEY]).Formula[:170])
    for i, d, reason in rows:
        sh.Cells(i, REASON).Value = (reason + " | " + FLAG) if reason else FLAG
    sh.Cells(5, REASON).AddComment("Rows flagged 'invoice no. is a DATE in the 2B export' (B3, 24-09): the Octa export holds a date, the original text is lost; keyed as m/yy.") if sh.Cells(5, REASON).Comment is None else None
    xl.Calculation = -4105; xl.CalculateFullRebuild()
    e = 0
    for s_ in w.Worksheets:
        if s_.Name == "T6A1 Extract - 24-25": continue
        try: e += s_.UsedRange.SpecialCells(-4123, 16).Count
        except Exception: pass
    keys = [sh.Cells(i, H[KEY]).Value for i, _, _ in rows]
    same = sum(1 for i, k in int_keys if S(sh.Cells(i, H[KEY]).Value) == k); print('integer-invoice keys unchanged: %d of %d' % (same, len(int_keys)))
    assert same == len(int_keys)
    print("flagged %d rows | sample keys now: %s | error cells %d | golden %.2f" % (len(rows), [S(k)[:24] for k in keys[:3]], e, w.Worksheets("ITC Register 2025-26").Range("V4").Value))
    assert e == 0 and all("/" not in S(k).split("|")[0] and not S(k).split("|")[0].endswith("0000") for k in keys)
    w.Save(); w.Close(False); print("saved (%.0fs)" % (time.time() - t0))
finally: xl.Quit()
