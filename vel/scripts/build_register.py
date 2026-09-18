import os, shutil, warnings, datetime as dt
warnings.filterwarnings("ignore")
import pandas as pd, openpyxl
from openpyxl.utils import get_column_letter

SP=os.path.dirname(os.path.abspath(__file__))
BASE="//server/GST FOLDER/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/Clients Data/05.08.2026 Main Data"
FMT ="//server/GST FOLDER/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/DPS Workings/1. Sales/Sales Register_FY 2025-26 Format.xlsx"
OUT =r"C:\Users\pawar\Downloads\VEL_Sales_Register_FY2025-26_DRAFT.xlsx"
JULY=os.path.join(SP,"July2025_converted.xlsx")
MONTHS=[("Apr-25","01 April 2025","Cleartax Sales April 2025.xlsx","Working",1),
 ("May-25","02 May 2025","Cleartax sales may 2025.xlsx","working",2),
 ("Jun-25","03 June 2025","Cleartax Sales Register June 2025.xlsx","Working",1),
 ("Jul-25",None,JULY,"Working",2),
 ("Aug-25","05 Aug 2025","Cleartax Sales Report August 2025.xlsx","WORKING",2),
 ("Sep-25","06 Sep 2025","Cleartax sales register sep 2025.xlsx","Working",2),
 ("Oct-25","07 Oct 2025","Cleartax sales Register Oct 2025.xlsx","Working",2),
 ("Nov-25","08 Nov 2025","Cleartax Sales Reg Nov 2025.xlsx","Working",2),
 ("Dec-25","09 Dec 2025","Cleartax Sales Register Dec 2025.xlsx","Working",2),
 ("Jan-26","10 Jan 2026","Cleartax sales register Jan 2026.xlsx","Working",2),
 ("Feb-26","11 Feb 2026","Cleartax Sales Register Feb 2026.xlsx","Working",2),
 ("Mar-26","12 Mar 2026","Cleartax sales register Mar 2026.xlsx","Working",1)]

STATE={"Madya Pradesh":"Madhya Pradesh","Gujrat":"Gujarat","Chattisgarh":"Chhattisgarh",
 "Tamilnadu":"Tamil Nadu","UttarPradesh":"Uttar Pradesh","West bengal":"West Bengal"}
DOCT={"INV":"INV","CRN":"CRN","CRN( Manual)":"CRN","DBN":"DBN",
 "MOB ADV REC":"MOB ADV REC","MOB ADV RECEIVED":"MOB ADV REC","Mob Adv Rec":"MOB ADV REC","MOB ADV REC INT":"MOB ADV REC",
 "MOB ADV ADJ":"MOB ADV ADJ","MOB ADV ADJUSTMENT":"MOB ADV ADJ","MOB ADV ADJUSMENT":"MOB ADV ADJ",
 "MOB ADV ADJUSTME":"MOB ADV ADJ","Mob advance adj":"MOB ADV ADJ",
 "MOB ADV REV":"MOB ADV REV","MOB ADV REVERSE":"MOB ADV REV"}
# target col letter -> source header (lowercased) ; None = derived/formula
SRC={"B":"state","C":"gstin","D":"seller legal name","E":"document date","F":"document number",
 "H":"document type","I":"supply type","J":"buyer gstin","K":"buyer legal name","L":"buyer place of supply",
 "M":"item unit price","N":"item total amount","O":"item gst rate","Q":"item assessable amount",
 "R":"item igst amount","S":"item cgst amount","T":"item sgst amount","U":"item cess amount",
 "X":"invoice reference no","Z":"status","AA":"cancel date","AB":"document status",
 "AC":"shipping gstin","AD":"shipping legalname","AE":"item product description",
 "AG":"item hsn code","AH":"item product description","AI":"item quantity","AJ":"item unit"}

def n(c): return str(c).strip().lower()
def clean(v):
    if v is None: return None
    if isinstance(v,float) and pd.isna(v): return None
    if v is pd.NaT: return None
    if isinstance(v,pd.Timestamp): return v.to_pydatetime()
    if isinstance(v,str) and v.strip() in ("","nan","NaT"): return None
    return v

frames=[]; log=[]
for lbl,m,f,s,h in MONTHS:
    p = f if m is None else os.path.join(BASE,m,f)
    df=pd.read_excel(p, sheet_name=s, header=h)
    cmap={n(c):c for c in df.columns}
    out=pd.DataFrame(index=range(len(df)))
    for tgt,src in SRC.items():
        out[tgt]=df[cmap[src]].values if src in cmap else None
    out["A"]=lbl                       # month label from the source file
    out["srcmonth"]=lbl
    frames.append(out); log.append((lbl, os.path.basename(p), s, h+1, len(df)))
    print(f"  loaded {lbl}: {len(df)} rows")

D=pd.concat(frames, ignore_index=True)
D["B"]=D["B"].astype(str).str.strip().map(lambda x: STATE.get(x,x))
raw_dt=D["H"].astype(str).str.strip()
D["H"]=raw_dt.map(lambda x: DOCT.get(x,"??"))
D["G"]=D["H"].map(lambda x: "Credit Notes" if x=="CRN" else "Sales")
D["I"]=D["I"].astype(str).str.strip().replace({"nan":None,"":None,"None":None})
# advances carry no Supply Type -> fill from the (normalised) document type
D.loc[D["I"].isna() & D["H"].str.startswith("MOB"), "I"]=D["H"]
# month from the document date where available, else the source-file month
dd=pd.to_datetime(D["E"], errors="coerce")
D["A"]=dd.dt.strftime("%b-%y").fillna(D["srcmonth"])
print(f"TOTAL rows: {len(D)}   unmapped doc types: {(D['H']=='??').sum()}")

shutil.copyfile(FMT, OUT)
wb=openpyxl.load_workbook(OUT); ws=wb["SR_2025-26"]
# preserve the CA's legend rows 5-9 on their own sheet, then clear them
lg=wb.create_sheet("Format Legend (preserved)")
lg.append(["Preserved from the format file rows 5-9 before data was written"])
lg.append([get_column_letter(c) for c in range(1, ws.max_column+1)])
lg.append([ws.cell(4,c).value for c in range(1, ws.max_column+1)])
for r in range(5,10):
    lg.append([ws.cell(r,c).value for c in range(1, ws.max_column+1)])
for r in range(5, ws.max_row+1):
    for c in range(1, ws.max_column+1): ws.cell(r,c).value=None

L={get_column_letter(i):i for i in range(1,50)}
FORMULA={"P":'=IF(N({q}5)=0,"",IF(ABS(ROUND(({r}5+{s}5+{t}5)/{q}5*100,2)-{o}5)<0.01,"OK","CHECK"))'}
r0=5
for i,row in enumerate(D.itertuples(index=False), start=r0):
    d=row._asdict()
    for tgt in SRC: ws.cell(i, L[tgt]).value=clean(d[tgt])
    ws.cell(i,L["A"]).value=d["A"]; ws.cell(i,L["G"]).value=d["G"]
    ws.cell(i,L["V"]).value=f"=SUM(R{i}:T{i})"
    ws.cell(i,L["W"]).value=f"=SUM(Q{i}:U{i})"
    ws.cell(i,L["P"]).value=f'=IF(Q{i}=0,"",IF(ABS(ROUND((R{i}+S{i}+T{i})/Q{i}*100,2)-O{i})<0.01,"OK","CHECK"))'
    ws.cell(i,L["Y"]).value=f'=IF(X{i}="","",LEN(X{i}))'
    ws.cell(i,L["AF"]).value=f'=IF(LEFT(TEXT(AG{i},"0"),2)="99","S","G")'
    ws.cell(i,L["AN"]).value=f'=LEFT(C{i},2)'
    ws.cell(i,L["AO"]).value=f'=IF(J{i}="","",LEFT(J{i},2))'
    ws.cell(i,L["AP"]).value=(f'=IF(OR(H{i}="MOB ADV REC",H{i}="MOB ADV ADJ",H{i}="MOB ADV REV",J{i}=""),"NA",'
                              f'IF(AN{i}=AO{i},IF(AND(R{i}=0,S{i}+T{i}>0),"OK","CHECK"),'
                              f'IF(AND(R{i}>0,S{i}+T{i}=0),"OK","CHECK")))')
    ws.cell(i,L["AQ"]).value=f'=IF(H{i}="INV",IF(I{i}="B2C","B2C","B2B"),IF(H{i}="CRN","Credit note",IF(H{i}="DBN","Debit note","Advance")))'
    ws.cell(i,L["AV"]).value=f'=IF(H{i}="INV",IF(I{i}="B2C","4A","4B"),IF(H{i}="CRN","4I",IF(H{i}="DBN","4J","4F")))'
    ws.cell(i,50).value=d["srcmonth"]
last=r0+len(D)-1
ws.cell(4,50).value="Source File Month"
for col in ["Q","R","S","T","U","V","AI"]:
    ws.cell(2,L[col]).value=f"=SUBTOTAL(9,{col}{r0}:{col}{last})"
lo=wb.create_sheet("Load Log")
lo.append(["Month","File","Sheet","Header row","Rows loaded"])
for x in log: lo.append(list(x))
lo.append(["TOTAL","","","",sum(x[4] for x in log)])
lo.append(["Written to register rows", r0, "to", last, len(D)])
wb.save(OUT)
print("WROTE", OUT, "| data rows", r0, "-", last)
