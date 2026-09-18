import os, warnings; warnings.filterwarnings("ignore")
import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
SP=os.path.dirname(os.path.abspath(__file__))
R=pd.read_pickle(os.path.join(SP,"register.pkl"))
G=pd.read_pickle(os.path.join(SP,"gstr1.pkl"))
S=pd.read_pickle(os.path.join(SP,"g1_summary.pkl"))
def s(x): return "" if pd.isna(x) else str(x).strip().upper()
GSTINS=[("01AAECR0503Q1ZM","Jammu & Kashmir"),("03AAECR0503Q1ZI","Punjab"),("06AAECR0503Q1ZC","Haryana"),
 ("08AAECR0503Q1Z8","Rajasthan"),("10AAECR0503Q1ZN","Bihar"),("12AAECR0503Q1ZJ","Arunachal Pradesh"),
 ("18AAECR0503Q1Z7","Assam"),("19AAECR0503Q1Z5","West Bengal"),("20AAECR0503Q1ZM","Jharkhand"),
 ("23AAECR0503Q1ZG","Madhya Pradesh"),("24AAECR0503Q1ZE","Gujarat"),("27AAECR0503Q1Z8","Maharashtra"),
 ("32AAECR0503Q1ZH","Kerala"),("33AAECR0503Q1ZF","TamilNadu"),("36AAECR0503Q1Z9","Telangana"),
 ("37AAECR0503Q1Z7","Andhra Pradesh"),("09AAECR0503Q1Z6","Uttar Pradesh"),("22AAECR0503Q1ZI","Chhattisgarh"),
 ("29AAECR0503Q1Z4","Karnataka")]
GROUPS=["B2B","B2C","Credit note","Debit note","Advance received","Advance adjusted","Total liability"]
MEAS=["Taxable value","IGST","CGST","SGST"]

# ---------- BOOKS side ----------
R["g"]=R["gstin"].map(s); R["dt"]=R["dtc"].astype(str).str.strip()
R["sp"]=R["sup"].astype(str).str.strip().str.upper()
def bucket(r):
    if r["dt"]=="INV":  return "B2C" if r["sp"]=="B2C" else "B2B"
    if r["dt"]=="CRN":  return "Credit note"
    if r["dt"]=="DBN":  return "Debit note"
    # Advances are classified by their ACTUAL EFFECT on the advance balance,
    # with the sign following the GSTR-1 convention (an adjustment/reversal
    # REDUCES the balance; it is never shown as another positive receipt).
    if r["dt"]=="MOB ADV REC": return "Advance received"
    if r["dt"]=="MOB ADV ADJ": return "Advance adjusted"
    if r["dt"]=="MOB ADV REV":
        # a positive REV restores/creates advance balance -> receipt side;
        # a negative REV reverses a received advance -> adjustment side.
        return "Advance received" if r["tax"]>0 else "Advance adjusted"
    return None
R["bk"]=R.apply(bucket,axis=1)
books={}
for (g,b),x in R.dropna(subset=["bk"]).groupby(["g","bk"]):
    books[(g,b)]=(x["tax"].sum(),x["igst"].sum(),x["cgst"].sum(),x["sgst"].sum())
# ---------- GSTR-1 side ----------
TY={"Invoice":"B2B","Credit Note":"Credit note","Debit Note":"Debit note"}
G["g"]=G["Company GSTIN"].map(s); G["bk"]=G["Doc Type"].map(TY)
for c in ["Taxable Value (Net)","IGST (Net)","CGST (Net)","SGST (Net)"]: G[c]=pd.to_numeric(G[c],errors="coerce").fillna(0)
g1={}
for (g,b),x in G.dropna(subset=["bk"]).groupby(["g","bk"]):
    g1[(g,b)]=(x["Taxable Value (Net)"].sum(),x["IGST (Net)"].sum(),x["CGST (Net)"].sum(),x["SGST (Net)"].sum())
SM={"B2CS Sales":"B2C","Advance Received":"Advance received","Advance Adjusted":"Advance adjusted"}
S["g"]=S["Company GSTIN"].map(s); S["bk"]=S["Summary Type"].map(SM)
for c in ["Taxable Value (Net)","IGST (Net)","CGST (Net)","SGST (Net)"]: S[c]=pd.to_numeric(S[c],errors="coerce").fillna(0)
for (g,b),x in S.dropna(subset=["bk"]).groupby(["g","bk"]):
    p=g1.get((g,b),(0,0,0,0))
    g1[(g,b)]=(p[0]+x["Taxable Value (Net)"].sum(),p[1]+x["IGST (Net)"].sum(),
               p[2]+x["CGST (Net)"].sum(),p[3]+x["SGST (Net)"].sum())

P=r"C:\Users\pawar\Downloads\VEL_Step2_Reco_GSTR1_DRAFT.xlsx"
wb=openpyxl.load_workbook(P)
if "SR vs GSTR-1 (CA format)" in wb.sheetnames: del wb["SR vs GSTR-1 (CA format)"]
ws=wb.create_sheet("SR vs GSTR-1 (CA format)", 0)
HF=Font(bold=True,color="FFFFFF"); HB=PatternFill("solid",fgColor="1F4E79")
GF=PatternFill("solid",fgColor="DDEBF7"); TOT=Font(bold=True)
thin=Side(style="thin",color="B0B0B0"); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
ws["A1"]="Vikran Engineering Limited"; ws["A1"].font=Font(bold=True,size=13)
ws["A2"]="Sales Register vs GSTR-1"; ws["A2"].font=Font(bold=True,size=12)
ws["A3"]="FY 2025-26   |   Difference block = Sales Register minus GSTR-1 (live formulas)"

def header(r0, title):
    ws.cell(r0,1,title).font=Font(bold=True,size=12)
    for gi,gname in enumerate(GROUPS):
        c=3+gi*4
        ws.cell(r0+1,c,gname).font=Font(bold=True); ws.cell(r0+1,c).alignment=Alignment(horizontal="center")
        ws.merge_cells(start_row=r0+1,start_column=c,end_row=r0+1,end_column=c+3)
        ws.cell(r0+1,c).fill=GF
    ws.cell(r0+2,1,"GSTIN"); ws.cell(r0+2,2,"State")
    for gi in range(len(GROUPS)):
        for mi,mname in enumerate(MEAS): ws.cell(r0+2,3+gi*4+mi,mname)
    for c in range(1,31):
        cell=ws.cell(r0+2,c); cell.font=HF; cell.fill=HB; cell.alignment=Alignment(horizontal="center",wrap_text=True)

def block(r0, title, data=None, diff=False, src=(7,31)):
    header(r0,title); first=r0+3
    for i,(g,st) in enumerate(GSTINS):
        r=first+i
        ws.cell(r,1,g); ws.cell(r,2,st)
        for gi,gname in enumerate(GROUPS):
            for mi in range(4):
                c=3+gi*4+mi; L=get_column_letter(c)
                if diff:
                    ws.cell(r,c,f"={L}{src[1]+i}-{L}{src[0]+i}")
                elif gname=="Total liability":
                    parts="+".join(f"{get_column_letter(3+k*4+mi)}{r}" for k in range(6))
                    ws.cell(r,c,f"={parts}")
                else:
                    ws.cell(r,c, round(data.get((g,gname),(0,0,0,0))[mi],2))
                ws.cell(r,c).number_format="#,##0.00"; ws.cell(r,c).border=BD
        ws.cell(r,1).border=BD; ws.cell(r,2).border=BD
    tr=first+len(GSTINS)
    ws.cell(tr,1,"Total").font=TOT
    for c in range(3,31):
        L=get_column_letter(c); ws.cell(tr,c,f"=SUM({L}{first}:{L}{tr-1})")
        ws.cell(tr,c).font=TOT; ws.cell(tr,c).number_format="#,##0.00"; ws.cell(tr,c).border=BD
    return tr

t1=block(4,"As per GSTR-1", g1)
t2=block(t1+2,"Sales Register", books)
t3=block(t2+2,"Difference", diff=True, src=(7, t1+5))
ws.column_dimensions["A"].width=20; ws.column_dimensions["B"].width=20
for c in range(3,31): ws.column_dimensions[get_column_letter(c)].width=15
ws.freeze_panes="C7"
wb.save(P)
print(f"WROTE CA-format sheet | GSTR-1 total row {t1} | Sales Register total row {t2} | Difference total row {t3}")
