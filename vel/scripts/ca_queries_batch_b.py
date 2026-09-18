"""CA query batch 1, part B:
Q11 register gains 'Original Invoice Number' + 'Original Invoice Date' (from ClearTax source, doc-level)
Q5  Step 3 file gains 'SR vs GSTR-3B' (books straight to 3B, state x month, live) + 'SR Advances' data
"""
import os, warnings; warnings.filterwarnings("ignore")
import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L
SP = os.path.dirname(os.path.abspath(__file__))
def S(v): return "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip()
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79")
def locked(p):
    try: f = open(p, "r+b"); f.close(); return False
    except Exception: return True

# ---------- Q11: originals from source (all doc types), doc-level ----------
BASE = "//192.168.1.69/GST FOLDER/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/Clients Data/05.08.2026 Main Data"
JULY = os.path.join(SP, "July2025_converted.xlsx")
MONTHS = [("01 April 2025", "Cleartax Sales April 2025.xlsx", "Working", 1), ("02 May 2025", "Cleartax sales may 2025.xlsx", "working", 2),
 ("03 June 2025", "Cleartax Sales Register June 2025.xlsx", "Working", 1), (None, JULY, "Working", 2),
 ("05 Aug 2025", "Cleartax Sales Report August 2025.xlsx", "WORKING", 2), ("06 Sep 2025", "Cleartax sales register sep 2025.xlsx", "Working", 2),
 ("07 Oct 2025", "Cleartax sales Register Oct 2025.xlsx", "Working", 2), ("08 Nov 2025", "Cleartax Sales Reg Nov 2025.xlsx", "Working", 2),
 ("09 Dec 2025", "Cleartax Sales Register Dec 2025.xlsx", "Working", 2), ("10 Jan 2026", "Cleartax sales register Jan 2026.xlsx", "Working", 2),
 ("11 Feb 2026", "Cleartax Sales Register Feb 2026.xlsx", "Working", 2), ("12 Mar 2026", "Cleartax sales register Mar 2026.xlsx", "Working", 1)]
def n(c): return str(c).strip().lower()
orig = {}
for m, f, s, h in MONTHS:
    p = f if m is None else os.path.join(BASE, m, f)
    d = pd.read_excel(p, sheet_name=s, header=h); cm = {n(c): c for c in d.columns}
    for _, r in d.iterrows():
        o = S(r[cm["original invoice number"]]); pdt = r[cm["preceding invoice date"]]
        if o and o.lower() != "nan":
            key = (S(r[cm["gstin"]]).upper(), S(r[cm["document number"]]).upper())
            orig.setdefault(key, (o, pdt))
print("documents with an original-invoice reference in source:", len(orig))
REG = r"C:\Users\pawar\Downloads\VEL_Step1_Sales_Register_FY2025-26_DRAFT.xlsx"
if locked(REG): print("LOCKED:", REG)
else:
    wb = openpyxl.load_workbook(REG); ws = wb["SR_2025-26"]
    H = {S(ws.cell(4, c).value): c for c in range(1, ws.max_column + 1)}
    if "Original Invoice Number" in H:
        c1, c2 = H["Original Invoice Number"], H["Original Invoice Date"]
    else:
        c1, c2 = ws.max_column + 1, ws.max_column + 2
        for cc, t in ((c1, "Original Invoice Number"), (c2, "Original Invoice Date")):
            x = ws.cell(4, cc, t); x.font = HF; x.fill = HB
        ws.column_dimensions[L(c1)].width = 20; ws.column_dimensions[L(c2)].width = 14
    cG, cF = H["My GSTIN"], H["Document Number"]
    filled = 0
    for i in range(5, 27007):
        key = (S(ws.cell(i, cG).value).upper(), S(ws.cell(i, cF).value).upper())
        v = orig.get(key)
        if v:
            ws.cell(i, c1).value = v[0]
            ws.cell(i, c2).value = v[1] if pd.notna(v[1]) else None
            ws.cell(i, c2).number_format = "dd-mmm-yy"
            filled += 1
    wb.save(REG)
    print("Q11 done: originals on %d register rows (cols %s/%s)" % (filled, L(c1), L(c2)))

# ---------- Q5: SR vs GSTR-3B in the Step 3 file ----------
P3 = r"C:\Users\pawar\Downloads\VEL_Step3_Reco_GSTR1_vs_3B_DRAFT.xlsx"
if locked(P3): print("LOCKED:", P3)
else:
    R = pd.read_pickle(os.path.join(SP, "register.pkl"))
    R["advf"] = R["dtc"].astype(str).str.startswith("MOB")
    R["mon"] = R.apply(lambda r: pd.to_datetime(r["ddate"]).strftime("%b-%y") if pd.notna(r["ddate"]) else S(r["src"]), axis=1)
    wb = openpyxl.load_workbook(P3)
    for nm in ("SR Data (all rows)", "SR vs GSTR-3B"):
        if nm in wb.sheetnames: del wb[nm]
    sr = wb.create_sheet("SR Data (all rows)")
    sr.append(["My GSTIN", "State", "Month", "Advance?", "Taxable", "IGST", "CGST", "SGST"])
    for c in sr[1]: c.font = HF; c.fill = HB
    g = R.groupby([R["gstin"].map(S), R["state"].map(S), "mon", "advf"])[["tax", "igst", "cgst", "sgst"]].sum().round(2).reset_index()
    for _, r in g.iterrows():
        sr.append([r["gstin"], r["state"], r["mon"], "Y" if r["advf"] else "N", r["tax"], r["igst"], r["cgst"], r["sgst"]])
    NS = sr.max_row
    for c_, w in zip("ABCDEFGH", [18, 18, 9, 9, 16, 13, 14, 14]): sr.column_dimensions[c_].width = w
    SR_ = lambda c: "'SR Data (all rows)'!$%s$2:$%s$%d" % (c, c, NS)
    b3 = wb["3B Data"]; N3 = b3.max_row
    B3 = lambda c: "'3B Data'!$%s$2:$%s$%d" % (c, c, N3)
    ws = wb.create_sheet("SR vs GSTR-3B", 2)
    ws["A1"] = "Sales Register (books, incl. advances) vs GSTR-3B 3.1(a) — direct reco, state x month (live). Difference = SR minus 3B."
    ws["A1"].font = Font(bold=True)
    ws.append([]); ws.append(["State", "GSTIN", "Month", "SR Taxable (incl adv)", "3B Taxable", "Diff Taxable",
                              "SR CGST", "3B CGST", "Diff CGST", "SR IGST", "3B IGST", "Diff IGST", "Remark"])
    for c in ws[3]: c.font = HF; c.fill = HB; c.alignment = Alignment(horizontal="center", wrap_text=True)
    GSTINS = sorted({(S(a), S(b)) for a, b in zip(R["gstin"], R["state"])})
    MON = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26"]
    r = 3
    for gstin, st in GSTINS:
        for mo in MON:
            r += 1
            ws.cell(r, 1, st); ws.cell(r, 2, gstin); ws.cell(r, 3, mo)
            ws.cell(r, 4, "=SUMIFS(%s,%s,$B%d,%s,$C%d)" % (SR_("E"), SR_("A"), r, SR_("C"), r))
            ws.cell(r, 5, "=SUMIFS(%s,%s,$B%d,%s,$C%d)" % (B3("D"), B3("B"), r, B3("C"), r))
            ws.cell(r, 6, "=D%d-E%d" % (r, r))
            ws.cell(r, 7, "=SUMIFS(%s,%s,$B%d,%s,$C%d)" % (SR_("G"), SR_("A"), r, SR_("C"), r))
            ws.cell(r, 8, "=SUMIFS(%s,%s,$B%d,%s,$C%d)" % (B3("F"), B3("B"), r, B3("C"), r))
            ws.cell(r, 9, "=G%d-H%d" % (r, r))
            ws.cell(r, 10, "=SUMIFS(%s,%s,$B%d,%s,$C%d)" % (SR_("F"), SR_("A"), r, SR_("C"), r))
            ws.cell(r, 11, "=SUMIFS(%s,%s,$B%d,%s,$C%d)" % (B3("E"), B3("B"), r, B3("C"), r))
            ws.cell(r, 12, "=J%d-K%d" % (r, r))
            for c in range(4, 13): ws.cell(r, c).number_format = "#,##0.00"
    last = r
    r += 1
    ws.cell(r, 1, "Total").font = Font(bold=True)
    for c in range(4, 13):
        ws.cell(r, c, "=SUM(%s4:%s%d)" % (L(c), L(c), last)); ws.cell(r, c).font = Font(bold=True); ws.cell(r, c).number_format = "#,##0.00"
    ws.conditional_formatting.add("F4:F%d" % last, CellIsRule(operator="notBetween", formula=["-1", "1"], fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.auto_filter.ref = "A3:M%d" % last
    for c_, w in zip("ABCDEFGHIJKLM", [18, 18, 9] + [15] * 9 + [46]): ws.column_dimensions[c_].width = w
    ws.freeze_panes = "D4"
    wb.save(P3)
    print("Q5 done: SR vs GSTR-3B sheet, %d state-month rows over %d SR data rows" % (last - 3, NS - 1))
