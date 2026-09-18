"""Step 1 (openpyxl): embed Cost Centres_New.xlsx Sheet4 A(PC)->F(CC) as hidden 'CC Master'."""
import openpyxl
from openpyxl.styles import Font, PatternFill
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f=open(P,'r+b'); f.close()
src = openpyxl.load_workbook("cost_centres_new.xlsx", read_only=True, data_only=True)
s4 = src["Sheet4"]
pairs=[]; dups=0; seen={}
for i,r in enumerate(s4.iter_rows(values_only=True)):
    if i==0: continue
    pc = r[0]; cc = S(r[5]) if len(r)>5 else ""
    if pc is None or not S(str(pc)): continue
    try: key=float(pc)
    except Exception: continue
    if key in seen:
        if seen[key]!=cc: dups+=1
        continue
    seen[key]=cc; pairs.append((key,cc))
src.close()
wb = openpyxl.load_workbook(P)
if "CC Master" in wb.sheetnames: del wb["CC Master"]
ws = wb.create_sheet("CC Master")
ws.append(["PC (code)","CC (narration)"])
for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E79")
for k,v in pairs: ws.append([k,v])
ws.column_dimensions["A"].width=12; ws.column_dimensions["B"].width=40
ws.sheet_state="hidden"
wb.save(P)
print("CC Master embedded (hidden): %d PC->CC pairs | conflicting duplicate PCs skipped-first-kept: %d" % (len(pairs),dups))
print("sample narrations:", [v for _,v in pairs[5:12]])
