"""Master (2): rebuild the three external-data sheets in their SOURCE formats, our columns at the end,
then remap every dependent formula to the new column letters.
  GSTR-1 Data : Octa Sales-Net 34 cols (+3 summary-only cols) + [Source, Bucket, Month, Match Status, Matched SR Doc No]
  3B Data     : State|GSTIN|Month + every Octa GSTR-3B table in form order + [SR Taxable, Diff, Vice-versa status]
  GL Data     : FBL3N 35 cols A..AI + derived live cols [State, Month, Tax head, Sales-origin?, Amount(+liab)] + [Matched with SR]
Row 1 stays reserved (button); headers row 2; row counts unchanged so only column letters move."""
import os, re, sys, warnings, datetime as dt
from copy import copy
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.styles import Font, PatternFill, Alignment, Border

P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
V = "//192.168.1.69/gst folder/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/"
G1D = V + "DPS Workings/Portal Reports/GSTR-1/"
B3D = V + "DPS Workings/Portal Reports/GSTR-3B/"
GLP = V + "Clients Data/GLs/Outward Tax/GST Output.xlsx"
f = open(P, "r+b"); f.close()

def S(v): return "" if v is None else str(v).strip()
def copystyle(src, dst):
    if src.has_style:
        dst.font = copy(src.font); dst.fill = copy(src.fill); dst.border = copy(src.border)
        dst.alignment = copy(src.alignment); dst.number_format = src.number_format

wb = openpyxl.load_workbook(P)
log = []

# =====================================================================================
# 1. GSTR-1 Data
# =====================================================================================
old = wb["GSTR-1 Data"]
old_rows = [r for r in old.iter_rows(min_row=3, values_only=True) if r[0]]
old_match = {}          # key -> (Match Status, Matched SR Doc No)
for r in old_rows:
    if S(r[2]).startswith("Sales-Net"):
        old_match[("D", S(r[0]), S(r[4]), S(r[11]), round(float(r[5] or 0), 2))] = (r[13], r[14])
print("old GSTR-1 Data rows:", len(old_rows), "| doc keys:", len(old_match))

files = sorted(x for x in os.listdir(G1D) if x.endswith(".xlsx") and not x.startswith("~$"))
src0 = openpyxl.load_workbook(G1D + files[[i for i, x in enumerate(files) if "Maharashtra" in x][0]])
sn0, ss0 = src0["Sales-Net"], src0["SalesSummary-Net"]
SN_H = [S(sn0.cell(1, c).value) for c in range(1, sn0.max_column + 1) if sn0.cell(1, c).value]
SS_H = [S(ss0.cell(1, c).value) for c in range(1, ss0.max_column + 1) if ss0.cell(1, c).value]
assert len(SN_H) == 34 and len(SS_H) == 20, (len(SN_H), len(SS_H))
SN_NF = {}
for c, h in enumerate(SN_H, 1):
    SN_NF[h] = sn0.cell(2, c).number_format if sn0.max_row >= 2 else "General"
SN_W = {h: (sn0.column_dimensions[L(c)].width or 14) for c, h in enumerate(SN_H, 1)}
EXTRA_SUM = ["GST Rate", "Ecommerce GSTIN", "Original Ecommerce GSTIN"]   # summary-only fields with no Sales-Net home
OURS = ["Source", "Bucket (formula)", "Month (formula)", "Match Status", "Matched SR Doc No"]
NEW_H = SN_H + ["Summary: " + h for h in EXTRA_SUM] + OURS
docs, sums = [], []
for x in files:
    w = openpyxl.load_workbook(G1D + x, data_only=True)
    sn, ss = w["Sales-Net"], w["SalesSummary-Net"]
    hs = [S(sn.cell(1, c).value) for c in range(1, sn.max_column + 1)]
    assert hs[:34] == SN_H, x
    for r in sn.iter_rows(min_row=2, values_only=True):
        if r[0]: docs.append(list(r[:34]))
    hss = [S(ss.cell(1, c).value) for c in range(1, ss.max_column + 1)]
    assert hss[:20] == SS_H, x
    for r in ss.iter_rows(min_row=2, values_only=True):
        if r[0]: sums.append(dict(zip(SS_H, r[:20])))
    w.close()
print("source docs:", len(docs), "| summaries:", len(sums))
assert len(docs) + len(sums) == len(old_rows), "row count changed vs old sheet"

del wb["GSTR-1 Data"]
ws = wb.create_sheet("GSTR-1 Data", 5)
ncol = len(NEW_H)
for c, h in enumerate(NEW_H, 1):
    cell = ws.cell(2, c, h); copystyle(sn0.cell(1, 1), cell)
iC, iSrc, iB, iM, iMS, iMD = CI("C"), ncol - 4, ncol - 3, ncol - 2, ncol - 1, ncol
r = 2
miss = 0
for d in docs:
    r += 1
    for c, v in enumerate(d, 1):
        cell = ws.cell(r, c, v); cell.number_format = SN_NF[SN_H[c - 1]]
    key = ("D", S(d[0]), S(d[4]), S(d[32]), round(float(d[21] or 0), 2))
    ms = old_match.get(key)
    if ms is None: miss += 1; ms = ("", "")
    ws.cell(r, iSrc, "Sales-Net (document)")
    ws.cell(r, iB, '=IF(C{n}="Invoice","B2B",IF(C{n}="Credit Note","Credit note",IF(C{n}="Debit Note","Debit note",IF(C{n}="B2CS Sales","B2C",IF(C{n}="Advance Received","Advance received","Advance adjusted")))))'.format(n=r))
    ws.cell(r, iM, '=TEXT(B{n},"mmm-yy")'.format(n=r))
    ws.cell(r, iMS, ms[0]); ws.cell(r, iMD, ms[1])
SUMMAP = {"Company GSTIN": "Company GSTIN", "Tax Period": "Tax Period", "Summary Type": "Doc Type", "Place of Supply": "Place of Supply",
          "Taxable Value": "Taxable Value", "IGST": "IGST", "CGST": "CGST", "SGST": "SGST", "Cess": "Cess", "Is Amendment": "Is Amendment",
          "Taxable Value (Net)": "Taxable Value (Net)", "IGST (Net)": "IGST (Net)", "CGST (Net)": "CGST (Net)", "SGST (Net)": "SGST (Net)",
          "Cess (Net)": "Cess (Net)", "Original/Revised Period": "Original/Revised Period", "GSTR-1A": "GSTR-1A"}
for s_ in sums:
    r += 1
    for sh, dh in SUMMAP.items():
        c = SN_H.index(dh) + 1
        cell = ws.cell(r, c, s_[sh]); cell.number_format = SN_NF[dh]
    for k, h in enumerate(EXTRA_SUM):
        ws.cell(r, 35 + k, s_[h])
    ws.cell(r, iSrc, "SalesSummary-Net (summary)")
    ws.cell(r, iB, '=IF(C{n}="Invoice","B2B",IF(C{n}="Credit Note","Credit note",IF(C{n}="Debit Note","Debit note",IF(C{n}="B2CS Sales","B2C",IF(C{n}="Advance Received","Advance received","Advance adjusted")))))'.format(n=r))
    ws.cell(r, iM, '=TEXT(B{n},"mmm-yy")'.format(n=r))
    ws.cell(r, iMS, "Summary level"); ws.cell(r, iMD, "")
NG1 = r
print("GSTR-1 Data written rows 3..%d | match carry-over misses: %d" % (NG1, miss))
assert miss == 0
for c, h in enumerate(NEW_H, 1):
    ws.column_dimensions[L(c)].width = SN_W.get(h, 16 if h.startswith("Summary") else 20)
ws.freeze_panes = "A3"
ws.auto_filter.ref = "A2:%s%d" % (L(ncol), NG1)
G1MAP = {"A": "A", "B": "B", "C": L(iSrc), "D": "C", "E": "E", "F": L(SN_H.index("Taxable Value (Net)") + 1),
         "G": L(SN_H.index("IGST (Net)") + 1), "H": L(SN_H.index("CGST (Net)") + 1), "I": L(SN_H.index("SGST (Net)") + 1),
         "J": L(iB), "K": L(iM), "L": L(SN_H.index("IRN") + 1), "M": L(SN_H.index("Is Amendment") + 1), "N": L(iMS), "O": L(iMD)}
log.append(("GSTR-1 Data", G1MAP))

# =====================================================================================
# 2. 3B Data
# =====================================================================================
old = wb["3B Data"]
old_trip = [(r[0], r[1], r[2]) for r in old.iter_rows(min_row=3, values_only=True) if r[1]]
old_vals = {}   # (gstin, month) -> old row tuple (for tie check)
for r in old.iter_rows(min_row=3, values_only=True):
    if r[1]: old_vals[(S(r[1]), S(r[2]))] = r
print("old 3B rows:", len(old_trip))
TABLES = [  # (header prefix, Octa Section startswith, [types])
    ("Total Liability (Other than reverse charge)", "Total Liability (Other than reverse charge)", ["Value"]),
    ("Total Liability (Reverse Charge)", "Total Liability (Reverse Charge)", ["Value"]),
    ("Paid using ITC", "Paid using ITC", ["Value"]),
    ("Paid using Cash", "Paid using Cash", ["Value"]),
    ("3.1(a) Outward taxable supplies (excl. zero rated)", "3.1.A", ["Supply Value", "Integrated Tax", "Central Tax", "State/UT Tax"]),
    ("3.1(d) Inward supplies liable to reverse charge", "3.1.D", ["Supply Value", "Integrated Tax", "Central Tax", "State/UT Tax"]),
    ("4A(3) Inward supplies liable to reverse charge", "4.A.3", ["Integrated Tax", "Central Tax", "State/UT Tax"]),
    ("4A(4) Inward supplies from ISD", "4.A.4", ["Integrated Tax", "Central Tax", "State/UT Tax"]),
    ("4A(5) All other ITC", "4.A.5", ["Integrated Tax", "Central Tax", "State/UT Tax"]),
    ("4B(1) Reversed - rules 38,42,43 & sec 17(5)", "4.B.1", ["Integrated Tax", "Central Tax", "State/UT Tax"]),
    ("4B(2) Reversed - Others", "4.B.2", ["Integrated Tax", "Central Tax", "State/UT Tax"]),
    ("4(C) Net ITC Available", "Net ITC", ["Integrated Tax", "Central Tax", "State/UT Tax"]),
    ("4D(1) ITC reclaimed (reversed under 4B(2) earlier)", "4.D.1", ["Integrated Tax", "Central Tax", "State/UT Tax"]),
    ("4D(2) Ineligible ITC u/s 16(4) & PoS rules", "4.D.2", ["Integrated Tax", "Central Tax", "State/UT Tax"]),
    ("Payment - Integrated Tax", "Integrated Tax", ["Integrated Tax ITC", "Cash"]),
    ("Payment - Central Tax", "Central Tax", ["Integrated Tax ITC", "Central Tax ITC", "Cash"]),
    ("Payment - State/UT Tax", "State/UT Tax", ["Integrated Tax ITC", "State/UT Tax ITC", "Cash"]),
]
TYPE_SHORT = {"Supply Value": "Taxable", "Integrated Tax": "IGST", "Central Tax": "CGST", "State/UT Tax": "SGST", "Value": "Value",
              "Integrated Tax ITC": "paid via IGST ITC", "Central Tax ITC": "paid via CGST ITC", "State/UT Tax ITC": "paid via SGST ITC", "Cash": "paid in Cash"}
MONS = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26"]
octa = {}   # gstin -> {(section-key, type): [12 values]}
hdr_style_src = None
for x in sorted(os.listdir(B3D)):
    if not x.lower().endswith(".xlsx") or x.startswith("~$"): continue
    w = openpyxl.load_workbook(B3D + x, data_only=True)
    ov = w["Overview"]; gstin = None
    for row in ov.iter_rows(values_only=True):
        if len(row) > 2 and S(row[1]).lower() == "gstin": gstin = S(row[2]).split("(")[0].strip()
    m = w["GSTR-3B"]
    if hdr_style_src is None: hdr_style_src = copy(m.cell(1, 2).font), copy(m.cell(1, 2).fill), copy(m.cell(1, 2).alignment)
    mh = [S(m.cell(1, c).value) for c in range(4, 16)]
    assert mh[0].startswith("Apr") and mh[11].startswith("Mar"), mh
    d = {}
    for row in m.iter_rows(min_row=2, values_only=True):
        sec, typ = S(row[1]), S(row[2])
        if not sec: continue
        vals = [float(v) if isinstance(v, (int, float)) else 0.0 for v in row[3:15]]
        for hp, pre, types in TABLES:
            if sec == pre or (pre[0].isdigit() and sec.startswith(pre + " ")):
                d[(hp, typ)] = vals
    octa[gstin] = d
    w.close()
print("Octa 3B files read:", len(octa))
del wb["3B Data"]
ws = wb.create_sheet("3B Data", 6)
H = ["State", "GSTIN", "Month"]
COLS = []
for hp, pre, types in TABLES:
    for t in types:
        H.append("%s - %s" % (hp, TYPE_SHORT[t])); COLS.append((hp, t))
OUR3 = ["SR Taxable (books, live)", "Diff (3B - books)", "Vice-versa status"]
H += OUR3
n3 = len(H)
HFo, HBo, HAo = hdr_style_src
for c, h in enumerate(H, 1):
    cell = ws.cell(2, c, h); cell.font = copy(HFo); cell.fill = copy(HBo); cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
ws.row_dimensions[2].height = 45
r = 2
iSR, iDF, iVV = n3 - 2, n3 - 1, n3
iTax = 3 + COLS.index(("3.1(a) Outward taxable supplies (excl. zero rated)", "Supply Value")) + 1
tie_bad = 0
OLDCHK = {"D": ("3.1(a) Outward taxable supplies (excl. zero rated)", "Supply Value"), "E": ("3.1(a) Outward taxable supplies (excl. zero rated)", "Integrated Tax"),
          "F": ("3.1(a) Outward taxable supplies (excl. zero rated)", "Central Tax"), "G": ("3.1(a) Outward taxable supplies (excl. zero rated)", "State/UT Tax"),
          "K": ("3.1(d) Inward supplies liable to reverse charge", "Supply Value"), "L": ("3.1(d) Inward supplies liable to reverse charge", "Integrated Tax"),
          "M": ("3.1(d) Inward supplies liable to reverse charge", "Central Tax"), "N": ("3.1(d) Inward supplies liable to reverse charge", "State/UT Tax"),
          "O": ("4(C) Net ITC Available", "Integrated Tax"), "P": ("4(C) Net ITC Available", "Central Tax"), "Q": ("4(C) Net ITC Available", "State/UT Tax"),
          "R": ("4A(3) Inward supplies liable to reverse charge", "Integrated Tax"), "S": ("4A(3) Inward supplies liable to reverse charge", "Central Tax"), "T": ("4A(3) Inward supplies liable to reverse charge", "State/UT Tax"),
          "U": ("4A(4) Inward supplies from ISD", "Integrated Tax"), "V": ("4A(4) Inward supplies from ISD", "Central Tax"), "W": ("4A(4) Inward supplies from ISD", "State/UT Tax")}
for st, g, mon in old_trip:
    r += 1
    ws.cell(r, 1, st); ws.cell(r, 2, g); ws.cell(r, 3, mon)
    mi = MONS.index(S(mon))
    d = octa.get(S(g), {})
    for k, (hp, t) in enumerate(COLS):
        v = d.get((hp, t), [0.0] * 12)[mi]
        cell = ws.cell(r, 4 + k, v); cell.number_format = "#,##0.00"
    orow = old_vals[(S(g), S(mon))]
    for oc, key in OLDCHK.items():
        ov = orow[CI(oc) - 1]; ov = float(ov) if isinstance(ov, (int, float)) else 0.0
        nv = d.get(key, [0.0] * 12)[mi]
        if abs(ov - nv) > 0.01: tie_bad += 1
    ws.cell(r, iSR, "=SUMIFS('SR_2025-26'!$T$6:$T$27007,'SR_2025-26'!$C$6:$C$27007,$B%d,'SR_2025-26'!$A$6:$A$27007,$C%d)" % (r, r)).number_format = "#,##0.00"
    ws.cell(r, iDF, "=%s%d-%s%d" % (L(iTax), r, L(iSR), r)).number_format = "#,##0.00"
    ws.cell(r, iVV, '=IF(ABS(%s%d)<1,"Matched with books",IF(%s%d=0,"In books, NOT reported in 3B",IF(%s%d=0,"Reported in 3B, NOT in books","3B differs from books - see S3 SR vs 3B MoM")))' % (L(iDF), r, L(iSR), r, L(iTax), r))
N3 = r
print("3B Data written rows 3..%d | old-vs-Octa value mismatches: %d" % (N3, tie_bad))
assert tie_bad == 0
ws.column_dimensions["A"].width = 20; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 9
for c in range(4, n3 + 1): ws.column_dimensions[L(c)].width = 16
ws.column_dimensions[L(iVV)].width = 40
ws.freeze_panes = "D3"; ws.auto_filter.ref = "A2:%s%d" % (L(n3), N3)
def c3(hp, t): return L(4 + COLS.index((hp, t)))
B3MAP = {"A": "A", "B": "B", "C": "C", "H": L(iSR), "I": L(iDF), "J": L(iVV)}
for oc, key in OLDCHK.items(): B3MAP[oc] = c3(*key)
log.append(("3B Data", B3MAP))

# =====================================================================================
# 3. GL Data
# =====================================================================================
old = wb["GL Data"]
old_gl = [r for r in old.iter_rows(min_row=3, values_only=True) if r[3] is not None or r[2]]
print("old GL rows:", len(old_gl))
src = openpyxl.load_workbook(GLP, data_only=True); sd = src["Data"]
GL_H = [S(sd.cell(6, c).value) for c in range(1, 36)]
assert GL_H[1] == "Document Number" and GL_H[21] == "Year/Month" and GL_H[31] == "G/L Account" and GL_H[32] == "Business place", GL_H
GL_NF = [sd.cell(7, c).number_format for c in range(1, 36)]
GL_W = [(sd.column_dimensions[L(c)].width or 12) for c in range(1, 36)]
FY = {"2025/%02d" % m for m in range(1, 13)}
rows = [list(r[:35]) for r in sd.iter_rows(min_row=7, values_only=True) if r[1] is not None and S(r[21]) in FY]
print("source GL FY rows:", len(rows))
assert len(rows) == len(old_gl)
# carry 'Matched with SR' by position after key check (FI doc, reference, G/L, amount)
bad = 0; carried = []
for o, n in zip(old_gl, rows):
    ok = (S(o[3]) == S(n[1]).lstrip("0") or S(o[3]) == S(n[1])) and S(o[2]).upper() == S(n[30]).upper() and abs(float(o[9] or 0) + float(n[6] or 0)) < 0.01
    if not ok: bad += 1
    carried.append(o[13])
print("GL positional key mismatches:", bad)
assert bad == 0
del wb["GL Data"]
ws = wb.create_sheet("GL Data", 7)
DER = ["State", "Month", "Tax head", "Sales-origin?", "Amount (+ = liability)"]
H = GL_H + DER + ["Matched with SR"]
for c, h in enumerate(H, 1):
    cell = ws.cell(2, c, h); copystyle(sd.cell(6, min(c, 35)), cell)
    if c > 35: cell.font = copy(sd.cell(6, 1).font); cell.fill = copy(sd.cell(6, 1).fill)
# BP -> State lookup block in hidden CC Master (cols D:E)
cm = wb["CC Master"]
BP2STATE = {"AP01": "Andhra Pradesh", "AR01": "Arunachal Pradesh", "AS01": "Assam", "BR01": "Bihar", "CG01": "Chhattisgarh",
            "GU01": "Gujarat", "HR01": "Haryana", "JH01": "Jharkhand", "JK01": "Jammu & Kashmir", "KA01": "Karnataka",
            "KL01": "Kerala", "MH01": "Maharashtra", "MP01": "Madhya Pradesh", "PB01": "Punjab", "RJ01": "Rajasthan",
            "TG01": "Telangana", "TN01": "Tamil Nadu", "UP01": "Uttar Pradesh", "WB01": "West Bengal"}
cm.cell(1, 4, "Business place"); cm.cell(1, 5, "State")
for i, (k, v) in enumerate(sorted(BP2STATE.items()), 2): cm.cell(i, 4, k); cm.cell(i, 5, v)
NBP = 1 + len(BP2STATE)
iSt, iMo, iTH, iSO, iAm, iMt = 36, 37, 38, 39, 40, 41
r = 2
for n, mt in zip(rows, carried):
    r += 1
    for c, v in enumerate(n, 1):
        cell = ws.cell(r, c, v); cell.number_format = GL_NF[c - 1]
    ws.cell(r, iSt, "=IFERROR(INDEX('CC Master'!$E$2:$E$%d,MATCH(AG%d,'CC Master'!$D$2:$D$%d,0)),AG%d)" % (NBP, r, NBP, r))
    ws.cell(r, iMo, '=TEXT(DATE(VALUE(LEFT(V%d,4)),VALUE(RIGHT(V%d,2))+3,1),"mmm-yy")' % (r, r))
    ws.cell(r, iTH, '=IF(AF%d="2610080100","CGST",IF(AF%d="2610080101","SGST",IF(AF%d="2610080102","IGST",IF(AF%d="2610080103","UTGST",""))))' % (r, r, r, r))
    ws.cell(r, iSO, '=IF(OR(D%d="RV",D%d="DA",D%d="DG",D%d="DR",D%d="XD"),"Y","N")' % (r, r, r, r, r))
    ws.cell(r, iAm, "=-G%d" % r).number_format = "#,##0.00"
    ws.cell(r, iMt, mt)
NGL = r
print("GL Data written rows 3..%d" % NGL)
for c in range(1, 36): ws.column_dimensions[L(c)].width = GL_W[c - 1]
for c, w_ in zip(range(36, 42), [18, 9, 9, 12, 16, 18]): ws.column_dimensions[L(c)].width = w_
ws.freeze_panes = "A3"; ws.auto_filter.ref = "A2:%s%d" % (L(iMt), NGL)
GLMAP = {"A": L(iSt), "F": L(iSO), "I": L(iTH), "J": L(iAm)}
log.append(("GL Data", GLMAP))

# =====================================================================================
# 4. Remap every dependent formula
# =====================================================================================
MAPS = {n: m for n, m in log}
pat = re.compile(r"'(GSTR-1 Data|3B Data|GL Data)'!\$([A-Z]{1,3})\$(\d+)(?::\$([A-Z]{1,3})\$(\d+))?")
def sub(m):
    sh, c1, r1, c2, r2 = m.groups()
    mp = MAPS[sh]
    if c1 not in mp or (c2 and c2 not in mp): raise KeyError((sh, c1, c2))
    out = "'%s'!$%s$%s" % (sh, mp[c1], r1)
    if c2: out += ":$%s$%s" % (mp[c2], r2)
    return out
n_cells = 0
for nm in wb.sheetnames:
    if nm in ("GSTR-1 Data", "3B Data", "GL Data"): continue
    w = wb[nm]
    for row in w.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("=") and "Data'!" in v:
                nv = pat.sub(sub, v)
                if nv != v: cell.value = nv; n_cells += 1
print("formulas remapped:", n_cells)
wb.save(P)
print("SAVED", P)
for n, m in log: print(n, m)
