"""Append Books/GSTR-1/Diff IGST + SGST columns (P..U) to S2 Month-on-Month so the pivot
can carry the same 12 value fields as last time. Existing columns untouched."""
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f=open(P,'r+b'); f.close()
wb = openpyxl.load_workbook(P)
ws0 = wb["SR_2025-26"]
H = {S(ws0.cell(4,c).value): L(c) for c in range(1, ws0.max_column+1)}
SR = lambda n: "'SR_2025-26'!$%s$5:$%s$27006" % (H[n], H[n])
ng = wb["GSTR-1 Data"].max_row
G1 = lambda c: "'GSTR-1 Data'!$%s$2:$%s$%d" % (c, c, ng)
mm = wb["S2 Month-on-Month"]
hdrs = {S(mm.cell(3,c).value): c for c in range(1, mm.max_column+1)}
assert "Books Taxable" in hdrs and mm.max_column == 15, (hdrs, mm.max_column)
HF=Font(bold=True,color="FFFFFF"); HB=PatternFill("solid",fgColor="1F4E79")
# how Books Taxable is computed, to mirror criteria: read its formula pattern from row 4
pat = str(mm.cell(4, hdrs["Books Taxable"]).value)
print("books-taxable pattern:", pat[:150])
patG = str(mm.cell(4, hdrs["GSTR-1 Taxable"]).value)
print("g1-taxable pattern:", patG[:150])
newcols = [("Books IGST","IGST Amount","E"),("GSTR-1 IGST",None,"G"),("Diff IGST",None,None),
           ("Books SGST","SGST Amount","F"),("GSTR-1 SGST",None,"I"),("Diff SGST",None,None)]
