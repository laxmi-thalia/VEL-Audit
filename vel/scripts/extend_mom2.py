"""Append 6 IGST/SGST columns to S2 Month-on-Month by cloning each row's own
Books/GSTR-1 taxable formulas with the value-column swapped."""
import openpyxl
from openpyxl.styles import Font, PatternFill
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f=open(P,'r+b'); f.close()
wb = openpyxl.load_workbook(P)
mm = wb["S2 Month-on-Month"]
assert mm.max_column == 15
HF=Font(bold=True,color="FFFFFF"); HB=PatternFill("solid",fgColor="1F4E79")
heads = ["Books IGST","GSTR-1 IGST","Diff IGST","Books SGST","GSTR-1 SGST","Diff SGST"]
for i,h in enumerate(heads):
    x = mm.cell(3, 16+i, h); x.font=HF; x.fill=HB
    mm.column_dimensions[openpyxl.utils.get_column_letter(16+i)].width = 14
n=0
for r in range(4, mm.max_row+1):
    if not S(mm.cell(r,1).value): continue
    fb = str(mm.cell(r,5).value); fg = str(mm.cell(r,6).value)   # Books Taxable / GSTR-1 Taxable
    if not fb.startswith("="): continue
    mm.cell(r,16).value = fb.replace("'SR_2025-26'!$R$5:$R$27006", "'SR_2025-26'!$S$5:$S$27006", 1)
    mm.cell(r,17).value = fg.replace("'GSTR-1 Data'!$F$2:$F$2080", "'GSTR-1 Data'!$G$2:$G$2080", 1)
    mm.cell(r,18).value = "=P%d-Q%d" % (r,r)
    mm.cell(r,19).value = fb.replace("'SR_2025-26'!$R$5:$R$27006", "'SR_2025-26'!$U$5:$U$27006", 1)
    mm.cell(r,20).value = fg.replace("'GSTR-1 Data'!$F$2:$F$2080", "'GSTR-1 Data'!$I$2:$I$2080", 1)
    mm.cell(r,21).value = "=S%d-T%d" % (r,r)
    for c in range(16,22): mm.cell(r,c).number_format="#,##0.00"
    n+=1
print("extended %d rows to 21 cols" % n)
wb.save(P)
