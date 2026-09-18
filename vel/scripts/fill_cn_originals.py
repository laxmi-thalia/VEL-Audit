"""Fill SR_2025-26 BF/BG (Original Invoice Number/Date) from the client's Credit Note Statement.
Statement: filter P contains 'Credit Note', key = Bill No (H), original ref parsed from Y."""
import openpyxl, re, datetime, json
from collections import defaultdict
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f=open(P,'r+b'); f.close()
sw = openpyxl.load_workbook("cn_statement.xlsx", read_only=True, data_only=True)
ss = sw["FY 2025-26"]
rows = list(ss.iter_rows(values_only=True)); sw.close()
PAT = re.compile(r"^(.*?)\s+[Dd]ated\s+(\d{1,2})\.(\d{1,2})\.(\d{4})\s*$")
refs = {}          # cn bill no (upper) -> (orig_no, date)
unparsed = []      # (bill no, Y text)
conflicts = []
for r in rows[3:]:
    if len(r) < 25: continue
    p = S(r[15]).lower()
    if "credit note" not in p: continue
    bill = S(r[7]).upper()
    y = S(r[24])
    m = PAT.match(y)
    if not (bill and m):
        unparsed.append((bill, y[:60])); continue
    orig = m.group(1).strip().rstrip(",/")
    try: dt = datetime.date(int(m.group(4)), int(m.group(3)), int(m.group(2)))
    except ValueError:
        unparsed.append((bill, y[:60])); continue
    if bill in refs and refs[bill] != (orig, dt):
        conflicts.append((bill, refs[bill], (orig, dt))); continue
    refs[bill] = (orig, dt)
print("statement: %d CN refs parsed | unparseable: %d | conflicting: %d" % (len(refs), len(unparsed), len(conflicts)))
wb = openpyxl.load_workbook(P)
ws = wb["SR_2025-26"]
from openpyxl.utils import get_column_letter as L
H = {S(ws.cell(4,c).value): c for c in range(1, ws.max_column+1)}
iD, iT = H["Document Number"], H["Document Type Code"]
iN, iDt = H["Original Invoice Number"], H["Original Invoice Date"]
filled_docs=set(); filled_rows=0; kept=0; overwrote_conf=[]
reg_cn_docs=set()
for r in range(5, 27007):
    if S(ws.cell(r,iT).value) != "CRN": continue
    doc = S(ws.cell(r,iD).value).upper()
    reg_cn_docs.add(doc)
    ref = refs.get(doc)
    if not ref: continue
    cur_n = S(ws.cell(r,iN).value)
    if cur_n and cur_n.upper() != ref[0].upper():
        overwrote_conf.append((doc, cur_n, ref[0])); continue
    if cur_n: kept += 1
    ws.cell(r,iN).value = ref[0]
    c = ws.cell(r,iDt); c.value = ref[1]; c.number_format = "DD.MM.YYYY"
    filled_docs.add(doc); filled_rows += 1
unmatched = sorted(set(refs) - filled_docs - {d for d,_,_ in overwrote_conf})
json.dump({"unparsed":unparsed,"stmt_conflicts":[[a,list(map(str,b)),list(map(str,c))] for a,b,c in conflicts],
           "reg_conflicts":overwrote_conf,"stmt_not_in_register":unmatched}, open("cn_fill_log.json","w"), indent=1, default=str)
print("register: %d CN docs total | filled %d docs / %d rows | already-had matching ref: %d" % (len(reg_cn_docs), len(filled_docs), filled_rows, kept))
print("register-side conflicts (existing != statement):", len(overwrote_conf), overwrote_conf[:3])
print("statement CNs not found in register:", len(unmatched), unmatched[:8])
wb.save(P)
