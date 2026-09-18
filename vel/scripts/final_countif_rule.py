"""FINAL Countif rule (Pawan 18-09): de-dupe on the register's OWN invoice (vendor GSTIN or name + normalised invoice no).
First line = Consider, others = Not consider, RCM/ISD blank. B_ = SUMIFS of the register by own KEY + vendor name (Consider lines);
2B_ = SUMIFS of the 2B sheet by KEY2, 0 when KEY2 blank; D_ = B_ - 2B_. Run AFTER cascade_fix.py (which stamps the older labels).
2B sheet bound is read from the sheet (never hardcoded)."""
import pickle, collections, re, time, openpyxl, warnings, win32com.client as win32, pythoncom
from openpyxl.utils import get_column_letter as L
warnings.filterwarnings("ignore")
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
try: open(P, "r+b").close()
except PermissionError: raise SystemExit("MASTER IS OPEN IN EXCEL - close it (without saving) and rerun")
def S(v): return "" if v is None else str(v).strip()
norm = lambda s: re.sub(r"[ \-/.'_]", "", S(s).upper())
wb = openpyxl.load_workbook(P, read_only=True)
ws = wb["ITC Register 2025-26"]; H = {S(ws.cell(5, c).value): c for c in range(1, ws.max_column + 1)}
rows = [r for r in ws.iter_rows(min_row=6, values_only=True) if r[3]]
b2 = wb["GSTR-2B Apr25-Aug26"]; B2H = {S(b2.cell(2, c).value): c for c in range(1, b2.max_column + 1)}
NB = 2 + sum(1 for r in b2.iter_rows(min_row=3, values_only=True) if r[0]); wb.close()
assert B2H["KEY"] == 49 and B2H["IGST (Net)"] == 22, B2H   # AW / V-X as in the original patch
g = lambda r, h: r[H[h] - 1]
seen = set(); labels = []
for r in rows:
    if g(r, "Category") != "ITC": labels.append(None); continue
    vg = S(g(r, "Vendor GSTIN")).upper(); k = (vg if re.match(r"^\d{2}[A-Z0-9]{13}$", vg) else "NM:" + S(g(r, "Vendor Name/RCM Category")).upper()) + "|" + norm(g(r, "Invoice No."))
    labels.append("Not consider" if k in seen else "Consider"); seen.add(k)
print("2B last row:", NB, "| Countif:", dict(collections.Counter(labels)))
m = pickle.load(open("itc_b1_meta.pkl", "rb")); m["consider_labels"] = labels; m["consider"] = ["Consider" if l == "Consider" else "NA" for l in labels]; pickle.dump(m, open("itc_b1_meta.pkl", "wb"))
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); wbx = xl.Workbooks.Open(P); xl.Calculation = -4135
    rg = wbx.Worksheets("ITC Register 2025-26"); c = lambda h: L(H[h]); n = 5 + len(rows)
    rg.Range("%s6:%s%d" % (c("Countif"), c("Countif"), n)).Value = [[v] for v in labels]
    KEY, VN, K2, CF = c("KEY"), c("Vendor Name/RCM Category"), c("KEY2 (matched 2B key)"), c("Countif")
    for b, src in (("B_IGST", "IGST"), ("B_CGST", "CGST"), ("B_SGST", "SGST")):
        rg.Range("%s6:%s%d" % (c(b), c(b), n)).Formula = '=IF($%s6="Consider",SUMIFS($%s$6:$%s$%d,$%s$6:$%s$%d,$%s6,$%s$6:$%s$%d,$%s6),"NA")' % (CF, c(src), c(src), n, KEY, KEY, n, KEY, VN, VN, n, VN)
    rg.Range("%s6:%s%d" % (c("B_Total GST"), c("B_Total GST"), n)).Formula = '=IF($%s6="Consider",%s6+%s6+%s6,"NA")' % (CF, c("B_IGST"), c("B_CGST"), c("B_SGST"))
    for t, col in (("2B_IGST", "V"), ("2B_CGST", "W"), ("2B_SGST", "X")):
        rg.Range("%s6:%s%d" % (c(t), c(t), n)).Formula = '=IF($%s6="Consider",IF($%s6="",0,SUMIFS(\'GSTR-2B Apr25-Aug26\'!$%s$3:$%s$%d,\'GSTR-2B Apr25-Aug26\'!$AW$3:$AW$%d,$%s6)),"NA")' % (CF, K2, col, col, NB, NB, K2)
    rg.Range("%s6:%s%d" % (c("2B_Total GST"), c("2B_Total GST"), n)).Formula = '=IF($%s6="Consider",%s6+%s6+%s6,"NA")' % (CF, c("2B_IGST"), c("2B_CGST"), c("2B_SGST"))
    for d, b, t in (("D_IGST", "B_IGST", "2B_IGST"), ("D_CGST", "B_CGST", "2B_CGST"), ("D_SGST", "B_SGST", "2B_SGST"), ("D_Total GST", "B_Total GST", "2B_Total GST")):
        rg.Range("%s6:%s%d" % (c(d), c(d), n)).Formula = '=IF($%s6="Consider",%s6-%s6,"NA")' % (CF, c(b), c(t))
    xl.Calculation = -4105; xl.CalculateFullRebuild()
    e = 0
    for w in wbx.Worksheets:
        try: e += w.UsedRange.SpecialCells(-4123, 16).Count
        except Exception: pass
    cs = lambda h: sum(v[0] for v in rg.Range("%s6:%s%d" % (c(h), c(h), n)).Value if isinstance(v[0], (int, float)))
    itc_tax = sum(float(g(r, "IGST") or 0) + float(g(r, "CGST") or 0) + float(g(r, "SGST") or 0) for r in rows if g(r, "Category") == "ITC")
    print("error cells %d | golden %.2f | B_Total (all ITC docs) %.2f vs ITC-category tax %.2f | 2B_Total %.2f | D_Total %.2f (%.0fs)" % (e, rg.Cells(4, H["Total GST"]).Value, cs("B_Total GST"), itc_tax, cs("2B_Total GST"), cs("D_Total GST"), time.time() - t0))
    rr = collections.Counter(str(v[0])[:34] for v in rg.Range("%s6:%s%d" % (c("Reco Remarks"), c("Reco Remarks"), n)).Value if v[0]); print("Reco Remarks:", dict(rr.most_common(8)))
    print("ITC Summary Net-ITC diff:", [round(wbx.Worksheets("ITC Summary").Cells(25, 49 + j).Value or 0, 2) for j in range(3)], "| ITCR vs 3B %.2f" % wbx.Worksheets("ITCR vs 3B Net ITC").Range("AG234").Value)
    wbx.Save(); wbx.Close(False)
finally: xl.Quit()
print("saved")
