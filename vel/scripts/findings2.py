import openpyxl, warnings, datetime as dt, collections
warnings.filterwarnings("ignore")
from openpyxl.styles import Font, PatternFill
P=r"C:\Users\pawar\Downloads\VEL_Sales_Register_FY2025-26_DRAFT.xlsx"
wb=openpyxl.load_workbook(P)
if "Step 1 Findings" in wb.sheetnames: del wb["Step 1 Findings"]
ws=wb["SR_2025-26"]; R0,R1=5,27006
F=[]; FY0,FY1=dt.datetime(2025,4,1),dt.datetime(2026,3,31)
docmap=collections.defaultdict(set); nrows=0
data=[]
for i,r in enumerate(ws.iter_rows(min_row=R0,max_row=R1,min_col=1,max_col=50,values_only=True),start=R0):
    data.append((i,r)); nrows+=1
    if r[5] not in (None,""): docmap[str(r[5]).strip()].add(str(r[2]))
for i,r in data:
    gstin,ddate,dnum = r[2],r[4],r[5]
    dtc,sup,bg = r[7],r[8],r[9]
    tax,igst,cgst,sgst = r[16],r[17],r[18],r[19]
    rate,irn,hsn,qty,uom = r[14],r[23],r[32],r[34],r[35]
    is_adv = str(dtc).startswith("MOB")
    if sup in (None,""): F.append(("Blank Supply Type",i,f"doc type {dtc}"))
    if ddate in (None,""): F.append(("Blank Document Date",i,f"doc type {dtc} - cannot assign a GSTR-1 month"))
    elif isinstance(ddate,dt.datetime) and not (FY0<=ddate<=FY1):
        F.append(("Document Date outside FY 25-26",i,str(ddate.date())))
    if irn not in (None,"") and len(str(irn).strip())!=64:
        F.append(("IRN length not 64",i,f"len={len(str(irn).strip())}"))
    oc = str(gstin)[:2] if gstin else None
    cc = str(bg)[:2] if bg not in (None,"") else None
    if oc and cc and not is_adv:
        I,C,S=(igst or 0),(cgst or 0),(sgst or 0)
        if oc==cc:
            if abs(I)>0.01: F.append(("Intra-state row carrying IGST",i,f"{oc}->{cc}"))
            if abs(C-S)>0.01: F.append(("Intra-state CGST <> SGST",i,f"diff {round(C-S,2)}"))
        elif abs(C)+abs(S)>0.01: F.append(("Inter-state row carrying CGST/SGST",i,f"{oc}->{cc}"))
    if not is_adv:
        if hsn in (None,""): F.append(("Blank HSN/SAC (non-advance row)",i,f"doc type {dtc}"))
        if qty in (None,""): F.append(("Blank Quantity (non-advance row)",i,f"doc type {dtc}"))
        if uom in (None,""): F.append(("Blank Unit of Measurement (non-advance row)",i,f"doc type {dtc}"))
    if tax not in (None,"") and rate not in (None,"") and abs(tax)>0.01:
        eff=round(((igst or 0)+(cgst or 0)+(sgst or 0))/tax*100,2)
        if abs(eff-float(rate))>0.05: F.append(("Effective rate <> stated GST Rate",i,f"computed {eff}% vs stated {rate}%"))
    if dnum not in (None,"") and len(docmap[str(dnum).strip()])>1:
        F.append(("Same Document Number under >1 GSTIN",i,f"{len(docmap[str(dnum).strip()])} GSTINs"))
cnt=collections.Counter(f[0] for f in F)
print(f"rows checked: {nrows}")
print("FINDINGS")
for k,v in cnt.most_common(): print(f"  {v:6d}  {k}")
print(f"  {len(F):6d}  total flags")
for chk in ["Intra-state row carrying IGST","Intra-state CGST <> SGST","Inter-state row carrying CGST/SGST",
            "IRN length not 64","Document Date outside FY 25-26","Same Document Number under >1 GSTIN"]:
    if chk not in cnt: print(f"       0  {chk}  <- CLEAN")
sh=wb.create_sheet("Step 1 Findings")
sh.append(["Check","Register row","Detail"])
for c in sh[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="C00000")
sh.append([f"SUMMARY - {nrows} rows checked","",""])
for k,v in cnt.most_common(): sh.append([k,v,""])
sh.append(["","",""]); sh.append(["DETAIL","",""])
for chk,i,d in F: sh.append([chk,i,d])
for col,w in zip("ABC",[44,14,60]): sh.column_dimensions[col].width=w
sh.freeze_panes="A2"
wb.save(P); print("saved")
