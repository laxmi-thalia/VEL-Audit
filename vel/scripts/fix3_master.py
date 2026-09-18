"""Three approved fixes on MASTER (2):
 1. S3 sheets: every '... Total Tax' formula = IGST+CGST+SGST of its own block (3B block was H+I+J incl. taxable).
 2. RCM Register: interest floored at 0.
 3. RCM Register: replace Conso rows with monthly-working rows for Gujarat Apr/Jun/Aug-25, TN Oct-25, Telangana Jan-26."""
import openpyxl, warnings, pickle, re, os, datetime as dt, collections
warnings.filterwarnings("ignore")
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
def num(v):
    try: return float(v)
    except Exception: return 0.0
SP = os.path.dirname(os.path.abspath(__file__))
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
open(P, "r+b").close()
wb = openpyxl.load_workbook(P)
# ---- fix 1
fixed = collections.Counter()
for nm in ("S3 SR vs 3B", "S3 1 vs 3B", "S3 SR vs 3B MoM", "S3 Month-on-Month"):
    ws = wb[nm]
    hdr_row = next(r for r in range(2, 8) if any("Total Tax" in S(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)))
    H = {S(ws.cell(hdr_row, c).value): c for c in range(1, ws.max_column + 1)}
    for tot_name in [k for k in H if k.endswith("Total Tax")]:
        heads = [tot_name.replace("Total Tax", h) for h in ("IGST", "CGST", "SGST")]
        if not all(h in H for h in heads): continue
        ct = H[tot_name]; hc = [L(H[h]) for h in heads]
        for r in range(hdr_row + 1, ws.max_row + 1):
            v = ws.cell(r, ct).value
            if isinstance(v, str) and v.startswith("=") and "+" in v and "SUM(" not in v.upper():
                want = "=%s%d+%s%d+%s%d" % (hc[0], r, hc[1], r, hc[2], r)
                if v.replace(" ", "") != want:
                    ws.cell(r, ct).value = want; fixed[(nm, tot_name)] += 1
print("fix1 S3 total-tax cells corrected:", dict(fixed))
# ---- fix 2
rr = wb["RCM Register"]
RH = {}
for c in range(1, rr.max_column + 1):
    h = S(rr.cell(5, c).value)
    if h and h not in RH: RH[h] = c
ci = RH["Interest @18% p.a."]; n2 = 0
for r in range(6, rr.max_row + 1):
    v = rr.cell(r, ci).value
    if isinstance(v, str) and v.startswith("=IF(") and "MAX(0," not in v:
        rr.cell(r, ci).value = v.replace("ROUND(", "MAX(0,ROUND(", 1).replace("/365,2))", "/365,2)))", 1); n2 += 1
print("fix2 interest formulas floored:", n2, "| sample:", rr.cell(6, ci).value)
# ---- fix 3
pk = pickle.load(open(os.path.join(SP, "rcm_monthly.pkl"), "rb")); D = pk["data"]
G2ST = {"24AAECR0503Q1ZE": "Gujarat", "33AAECR0503Q1ZF": "Tamil Nadu", "36AAECR0503Q1Z9": "Telangana"}
TARGET = [("24AAECR0503Q1ZE", "Apr-25", 4), ("24AAECR0503Q1ZE", "Jun-25", 6), ("24AAECR0503Q1ZE", "Aug-25", 8), ("33AAECR0503Q1ZF", "Oct-25", 10), ("36AAECR0503Q1Z9", "Jan-26", 1)]
GLN = {"2610080300": "CGST Output RCM", "2610080301": "SGST Output RCM", "2610080302": "IGST Output RCM"}
M3B = {4: "01 Apr 2025", 6: "03 June 2025", 8: "05 Aug 2025", 10: "07 Oct 2025", 1: "10 Jan 2026"}
ADATE = {4: dt.datetime(2025, 4, 1), 6: dt.datetime(2025, 6, 3), 8: dt.datetime(2025, 8, 5), 10: dt.datetime(2025, 10, 7), 1: dt.datetime(2026, 1, 10)}
NEXT = {"01 Apr 2025": "02 May 2025", "03 June 2025": "04 July 2025", "05 Aug 2025": "06 Sep 2025", "07 Oct 2025": "08 Nov 2025", "10 Jan 2026": "11 Feb 2026"}
def fy(d):
    if not isinstance(d, dt.datetime): return ""
    y = d.year if d.month >= 4 else d.year - 1
    return "%d-%02d" % (y, (y + 1) % 100)
R0 = 6
NC = max(RH.values())
OLDN = max(r for r in range(R0, rr.max_row + 1) if rr.cell(r, RH["MY GSTN"]).value or rr.cell(r, RH["Business place"]).value)
col_formula = {c: rr.cell(R0, c).value for c in range(1, NC + 1) if isinstance(rr.cell(R0, c).value, str) and str(rr.cell(R0, c).value).startswith("=")}
def rel(fx, r): return re.sub(r"(?<![A-Z!])(\$?[A-Z]{1,3}\$?)%d(?!\d)" % R0, lambda m: m.group(1) + str(r), fx)
rows = []
for r in range(R0, OLDN + 1):
    rows.append({c: rr.cell(r, c).value for c in range(1, NC + 1) if c not in col_formula})
def month_of(row):
    a = row.get(1)
    return a.month if isinstance(a, dt.datetime) else None
swap_log = []; new_rows = []
for g, mlbl, mon in TARGET:
    old = [x for x in rows if x.get(RH["MY GSTN"]) == g and month_of(x) == mon]
    old_tax = sum(num(x.get(RH["Taxable Value as per SAP"])) for x in old)
    rec = D[(g, mlbl)]; hi = {}
    for i, h in enumerate(rec["header"]):
        if h and h not in hi: hi[h] = i
    built = []
    for r_ in rec["rows"]:
        d = lambda h: r_[hi[h]] if h in hi and hi[h] < len(r_) else None
        gl = S(d("G/L Account")); tv = num(d("Taxable Value")); ig, cg, sg = num(d("IGST")), num(d("CGST")), num(d("SGST"))
        rate = num(d("Rate")) if S(d("Rate")) else (round((ig + cg + sg) / tv * 100, 2) if tv else None)
        base = {1: ADATE[mon], RH["Business place"]: S(d("Business place")), RH["MY GSTN"]: g, RH["State"]: G2ST[g], RH["Fiscal Year"]: d("Fiscal Year"),
                RH["Year/Month"]: S(d("Year/Month")), RH["Document type"]: S(d("Document type")), RH["Document Number"]: d("Document Number"),
                RH["Posting Date"]: d("Posting Date"), RH["Posting Year"]: fy(d("Posting Date")), RH["Assignment"]: S(d("Assignment")), RH["Reference"]: S(d("Reference")),
                RH["Document Date"]: d("Document Date"), RH["Doc year"]: fy(d("Document Date")), RH["Vendor Code"]: d("Vendor Code"), RH["GSTN"]: S(d("Vendor GSTN")),
                RH["Vendor Name"]: S(d("Vendor Name")), RH["Posting Key"]: d("Posting Key"), RH["Local Currency"]: S(d("Local Currency")) or "INR", RH["Tax Code"]: S(d("Tax Code")),
                RH["Clearing Document"]: d("Clearing Document"), RH["Profit Center"]: d("Profit Center"), RH["Text"]: S(d("Text")), RH["Offsetting Account"]: d("Offsetting Account"),
                RH["GST Rate"]: rate, RH["Nature of Services"]: S(d("Category")), RH["GSTR 3B Claim month"]: NEXT[M3B[mon]],
                RH["Source"]: "Monthly working %s (RCM sheet) - replaces Conso RCM per Pawan 17-09" % mlbl}
        if gl.startswith("2610080302") or (ig and not cg):
            x = dict(base); x.update({RH["G/L Account"]: "2610080302", RH["GL Name"]: GLN["2610080302"], RH["Amount in Local Currency"]: -ig, RH["Taxable Value as per SAP"]: tv, RH["IGST AS PER SAP"]: ig, RH["CGST AS PER SAP"]: 0.0, RH["SGST AS PER SAP"]: 0.0}); built.append(x)
        else:
            for acct, amt, tag in (("2610080300", cg, "C"), ("2610080301", sg, "S")):
                x = dict(base); x.update({RH["G/L Account"]: acct, RH["GL Name"]: GLN[acct], RH["Amount in Local Currency"]: -amt, RH["Taxable Value as per SAP"]: tv / 2, RH["IGST AS PER SAP"]: 0.0, RH["CGST AS PER SAP"]: cg if tag == "C" else 0.0, RH["SGST AS PER SAP"]: sg if tag == "S" else 0.0}); built.append(x)
    new_tax = sum(num(x[RH["Taxable Value as per SAP"]]) for x in built)
    swap_log.append((G2ST[g], mlbl, len(old), round(old_tax, 2), len(built), round(new_tax, 2)))
    rows = [x for x in rows if not (x.get(RH["MY GSTN"]) == g and month_of(x) == mon)]
    new_rows.extend(built)
rows.extend(new_rows)
rows.sort(key=lambda x: (x.get(1) if isinstance(x.get(1), dt.datetime) else dt.datetime(2099, 1, 1), S(x.get(RH["MY GSTN"]))))
NEWN = R0 + len(rows) - 1
for r in range(R0, max(OLDN, NEWN) + 1):
    for c in range(1, NC + 1): rr.cell(r, c).value = None
for i, x in enumerate(rows):
    r = R0 + i
    for c in range(1, NC + 1):
        rr.cell(r, c).value = rel(col_formula[c], r) if c in col_formula else x.get(c)
for c in range(1, NC + 1):
    v = rr.cell(4, c).value
    if isinstance(v, str) and "SUBTOTAL" in v:
        rr.cell(4, c).value = re.sub(r"(\$?[A-Z]{1,3}\$?)%d\)" % OLDN, lambda m: "%s%d)" % (m.group(1), NEWN), v)
rr.auto_filter.ref = "A5:%s%d" % (L(NC), NEWN)
pat = re.compile(r"('RCM Register'!\$[A-Z]{1,3}\$)(\d+)(:\$[A-Z]{1,3}\$)(\d+)")
n = 0
for nm in wb.sheetnames:
    if nm == "RCM Register": continue
    for rowc in wb[nm].iter_rows():
        for cell in rowc:
            v = cell.value
            if isinstance(v, str) and v.startswith("=") and "RCM Register" in v:
                nv = pat.sub(lambda m: "%s%s%s%d" % (m.group(1), m.group(2), m.group(3), NEWN) if int(m.group(4)) == OLDN else m.group(0), v)
                if nv != v: cell.value = nv; n += 1
print("fix3 swaps (state, month, old rows, old taxable, new rows, new taxable):")
for x in swap_log: print("   ", x)
print("register rows %d -> %d | dependents repointed %d" % (OLDN - R0 + 1, NEWN - R0 + 1, n))
wb.save(P); print("saved")
