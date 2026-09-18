import openpyxl
from openpyxl.styles import Font, PatternFill
HF=Font(bold=True,color="FFFFFF"); HB=PatternFill("solid",fgColor="1F4E79"); AMB=PatternFill("solid",fgColor="FFF2CC")
S=lambda v: "" if v is None else str(v).strip()
MONTHS=["Apr-25","May-25","Jun-25","Jul-25","Aug-25","Sep-25","Oct-25","Nov-25","Dec-25","Jan-26","Feb-26","Mar-26"]
p=r"C:\Users\pawar\Downloads\VEL_Step2_Reco_GSTR1_DRAFT.xlsx"
wb=openpyxl.load_workbook(p)
mm=wb["Month-on-Month"]; last=mm.max_row
NSR=wb["SR Data"].max_row; NG=wb["GSTR-1 Data"].max_row
SRr=lambda c:"'SR Data'!$%s$2:$%s$%d"%(c,c,NSR); G1r=lambda c:"'GSTR-1 Data'!$%s$2:$%s$%d"%(c,c,NG)
AUTO={("Bihar","Jan-26"):"34 no-IRN credit notes in GSTR-1 only - not in books, not in 3B",
 ("Telangana","May-25"):"e-invoiced docs not uploaded to GSTR-1; tax paid via 3B",
 ("Telangana","Jun-25"):"e-invoiced docs not uploaded to GSTR-1; tax paid via 3B",
 ("Telangana","Aug-25"):"e-invoiced docs not uploaded to GSTR-1; tax paid via 3B",
 ("Madhya Pradesh","Dec-25"):"Rs 9 advance rounding"}
# restore headers 17-20, add 21 typed remark + 22 key
HD={17:"Auto: books docs not in GSTR-1",18:"Auto: GSTR-1 docs not in books",19:"UNEXPLAINED residual",
    20:"(unused)",21:"Remark (type here, beside the difference)",22:"key"}
for c,t in HD.items():
    x=mm.cell(3,c,t); x.font=HF; x.fill=HB
for r in range(4,last+1):
    st,mo=S(mm.cell(r,1).value),S(mm.cell(r,3).value)
    if not st: continue
    mm.cell(r,17).value='=SUMIFS(%s,%s,$B%d,%s,$C%d,%s,$D%d,%s,"Not in GSTR-1")'%(SRr("F"),SRr("A"),r,SRr("B"),r,SRr("J"),r,SRr("K"))
    mm.cell(r,17).fill=PatternFill(fill_type=None)
    mm.cell(r,18).value='=-SUMIFS(%s,%s,$B%d,%s,$C%d,%s,$D%d,%s,"Not in books")'%(G1r("F"),G1r("A"),r,G1r("K"),r,G1r("J"),r,G1r("L"))
    mm.cell(r,19).value="=M%d-Q%d-R%d"%(r,r,r)
    mm.cell(r,20).value=None
    cur=S(mm.cell(r,21).value)
    mm.cell(r,21).value=(cur if cur and not cur.startswith("=") else None) or AUTO.get((st,mo))
    mm.cell(r,21).fill=AMB
    mm.cell(r,22).value='=A%d&"|"&C%d'%(r,r)
    for c in (17,18,19): mm.cell(r,c).number_format="#,##0.00"
mm.column_dimensions["U"].width=46; mm.column_dimensions["V"].hidden=True
mm.auto_filter.ref="A3:V%d"%last
ws=wb["SR vs GSTR-1"]
def concat(state_cell):
    terms=[]
    for mo in MONTHS:
        idx="INDEX('Month-on-Month'!$U$4:$U$%d,MATCH(%s&\"|%s\",'Month-on-Month'!$V$4:$V$%d,0))"%(last,state_cell,mo,last)
        terms.append('IFERROR(IF(LEN(%s)>0,"%s: "&%s&"; ",""),"")'%(idx,mo,idx))
    return "=TRIM("+"&".join(terms)+")"
ws.cell(54,35).value="Remarks (auto from Month-on-Month typed remarks)"
for i in range(19):
    r=55+i; ws.cell(r,35).value=concat("$B%d"%r)
wb.save(p); print("step 2 remark columns repaired (auto-explain back in Q-S, typed remark U, key V)")
