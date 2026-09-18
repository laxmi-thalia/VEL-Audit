"""Pull FY 2026-27 3B (Apr-Jul 2026, Octa annual reports) into '3B Data FY26-27' (same layout as 3B Data),
and close the Mar-26 '(next FY)' gap on
'RCM Paid vs ITC Claimed' (col F = Apr-26 4A(3) from the new sheet). Row 1 reserved."""
import os, re, warnings, datetime as dt
from copy import copy
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
A = "//192.168.1.69/gst folder/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/Audit data of FY 2026-27/"
f = open(P, "r+b"); f.close()
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79"); TOT = Font(bold=True); NF = "#,##0.00"
wb = openpyxl.load_workbook(P)
b3 = wb["3B Data"]
H = [S(b3.cell(2, c).value) for c in range(1, b3.max_column + 1)]
iVV = H.index("Vice-versa status") + 1; iSR = H.index("SR Taxable (books, live)") + 1
DATA_H = H[3:iSR - 1]           # the Octa table columns
# header -> (section prefix, type)
TABLES = [("Total Liability (Other than reverse charge)", "Total Liability (Other than reverse charge)"), ("Total Liability (Reverse Charge)", "Total Liability (Reverse Charge)"),
          ("Paid using ITC", "Paid using ITC"), ("Paid using Cash", "Paid using Cash"), ("3.1(a) Outward taxable supplies (excl. zero rated)", "3.1.A"),
          ("3.1(d) Inward supplies liable to reverse charge", "3.1.D"), ("4A(3) Inward supplies liable to reverse charge", "4.A.3"), ("4A(4) Inward supplies from ISD", "4.A.4"),
          ("4A(5) All other ITC", "4.A.5"), ("4B(1) Reversed - rules 38,42,43 & sec 17(5)", "4.B.1"), ("4B(2) Reversed - Others", "4.B.2"), ("4(C) Net ITC Available", "Net ITC"),
          ("4D(1) ITC reclaimed (reversed under 4B(2) earlier)", "4.D.1"), ("4D(2) Ineligible ITC u/s 16(4) & PoS rules", "4.D.2"),
          ("Payment - Integrated Tax", "Integrated Tax"), ("Payment - Central Tax", "Central Tax"), ("Payment - State/UT Tax", "State/UT Tax")]
TYPE = {"Taxable": "Supply Value", "IGST": "Integrated Tax", "CGST": "Central Tax", "SGST": "State/UT Tax", "Value": "Value",
        "paid via IGST ITC": "Integrated Tax ITC", "paid via CGST ITC": "Central Tax ITC", "paid via SGST ITC": "State/UT Tax ITC", "paid in Cash": "Cash"}
def key_of(h):
    hp, t = h.rsplit(" - ", 1); pre = dict(TABLES)[hp]; return pre, TYPE[t]
KEYS = [key_of(h) for h in DATA_H]
G2ST = {b3.cell(r, 2).value: b3.cell(r, 1).value for r in range(3, b3.max_row + 1) if b3.cell(r, 2).value}
MON = ["Apr-26", "May-26", "Jun-26", "Jul-26", "Aug-26", "Sep-26", "Oct-26", "Nov-26", "Dec-26", "Jan-27", "Feb-27", "Mar-27"]
octa = {}; months_present = set()
for x in sorted(os.listdir(A + "GSTR-3B")):
    if not x.lower().endswith(".xlsx") or x.startswith("~$"): continue
    w = openpyxl.load_workbook(A + "GSTR-3B/" + x, read_only=True, data_only=True)
    ov = w["Overview"]; gstin = None
    for row in ov.iter_rows(values_only=True):
        if len(row) > 2 and S(row[1]).lower() == "gstin": gstin = S(row[2]).split("(")[0].strip()
    m = w["GSTR-3B"]; rows = list(m.iter_rows(values_only=True))
    mh = [S(v) for v in rows[0][3:15]]
    assert mh[0].startswith("Apr 2026"), (x, mh[:2])
    d = {}
    for row in rows[1:]:
        sec, typ = S(row[1]), S(row[2])
        if not sec: continue
        vals = []
        for j, v in enumerate(row[3:15]):
            vals.append(float(v) if isinstance(v, (int, float)) else 0.0)
            if isinstance(v, (int, float)) and v != 0: months_present.add(j)
        for hp, pre in TABLES:
            if sec == pre or (pre[0].isdigit() and sec.startswith(pre + " ")): d[(pre, typ)] = vals
    octa[gstin] = d; w.close()
nmon = max(months_present) + 1
print("3B files:", len(octa), "| months with data:", nmon, MON[:nmon])
NM = "3B Data FY26-27"
if NM in wb.sheetnames: del wb[NM]
ws = wb.create_sheet(NM, wb.sheetnames.index("3B Data") + 1)
ws.row_dimensions[1].height = 21
for c, h in enumerate(H[:iSR - 1], 1):
    x = ws.cell(2, c, h); x.font = copy(b3.cell(2, c).font); x.fill = copy(b3.cell(2, c).fill); x.alignment = copy(b3.cell(2, c).alignment)
ws.row_dimensions[2].height = 45
r = 2
for g in [g for g in G2ST if g in octa]:
    for mi in range(nmon):
        r += 1
        ws.cell(r, 1, G2ST[g]); ws.cell(r, 2, g); ws.cell(r, 3, MON[mi])
        for k, key in enumerate(KEYS):
            ws.cell(r, 4 + k, octa[g].get(key, [0.0] * 12)[mi]).number_format = NF
N = r
ws.freeze_panes = "D3"; ws.auto_filter.ref = "A2:%s%d" % (L(iSR - 1), N)
ws.column_dimensions["A"].width = 20; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 9
for c in range(4, iSR): ws.column_dimensions[L(c)].width = 16
missing3b = [G2ST[g] for g in G2ST if g not in octa]
print("3B Data FY26-27 rows 3..%d | states without FY26-27 3B file: %s" % (N, missing3b))
# ---- close the Mar-26 gap on RCM Paid vs ITC Claimed
pc = wb["RCM Paid vs ITC Claimed"]
i43 = [c for c, h in enumerate(H, 1) if h.startswith("4A(3)")]
assert len(i43) == 3
NB = lambda col: "'%s'!$%s$3:$%s$%d" % (NM, col, col, N)
n = 0
for r in range(6, pc.max_row + 1):
    if S(pc.cell(r, 3).value) == "12 Mar 2026":
        pc.cell(r, 11).value = "Apr-26"
        pc.cell(r, 6).value = "=" + "+".join("SUMIFS(%s,%s,$A%d,%s,$K%d)" % (NB(L(c)), NB("B"), r, NB("C"), r) for c in i43)
        pc.cell(r, 6).number_format = NF
        pc.cell(r, 8).value = "=IF($F%d=\"\",\"(next FY)\",D%d-F%d)" % (r, r, r)
        n += 1
pc.cell(3, 1).value = S(pc.cell(3, 1).value).replace("Mar-26 payment claims in Apr-26 (FY 26-27).", "Mar-26 payment claims in Apr-26 (FY 26-27) - now read live from '3B Data FY26-27' (Octa annual report FY 26-27). ITC-register column stays blank for Mar-26 (no FY 26-27 ITC register yet).")
print("Mar-26 rows re-pointed to FY26-27 4A(3):", n)
# INDEX rows
ix = wb["INDEX"]
from openpyxl.styles import Border, Side
thin = Side(style="thin"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
have = {S(ix.cell(r2, 2).value) for r2 in range(5, ix.max_row + 1)}
for arena, desc, sheet, lbl in (("RCM", "GSTR-3B FY 2026-27 (Apr-Jul 2026) - for the Mar-26 RCM claim spill", NM, "3B Data FY26-27"),):
    if desc in have: continue
    r2 = ix.max_row + 1
    ix.cell(r2, 1, arena); ix.cell(r2, 2, desc)
    x = ix.cell(r2, 7); x.value = '=HYPERLINK("#\'%s\'!A1","%s")' % (sheet, lbl); x.font = Font(color="0563C1", underline="single")
    for c in range(1, 11): ix.cell(r2, c).border = BD
wb.save(P); print("saved")
