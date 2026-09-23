"""ITC Register 2026-27: 2B reconciliation in the FY 25-26 format (Pawan 21-09) - KEY / Countif / B_ / KEY2 / 2B_ / D_ / Reco Remarks
appended after the last header (ITC Summary references this sheet by column letter - never insert in the middle). Matching via
reco_lib (numbered vocabulary) against 'GSTR-2B Apr25-Aug26' (all periods). Single COM session; B_ must equal the sheet's ITC tax."""
import re, sys, os, time, shutil, collections, warnings, openpyxl, win32com.client as win32, pythoncom
from openpyxl.utils import get_column_letter as L
sys.path.insert(0, r"C:\PROJECTS\gst-audit-engine"); from vel.scripts.reco_lib import match_register, norm, S, num
warnings.filterwarnings("ignore")
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"; B = P[:-5] + ".xlsb"
try: open(P, "r+b").close()
except PermissionError: raise SystemExit("LOCKED - close in Excel without saving: " + P)
def xlsb_free():
    try: open(B, "r+b").close(); return True
    except PermissionError: return False
shutil.copy(P, "master2_snapshot_before_fy2627reco.xlsx")
wb = openpyxl.load_workbook(P, read_only=True)
n = wb["ITC Register 2026-27"]; hd = [c.value for c in next(n.iter_rows(min_row=5, max_row=5))]; H = {h: i for i, h in enumerate(hd) if h}
rows = [r for r in n.iter_rows(min_row=6, values_only=True) if r[3]]; g = lambda r, h: r[H[h]]
b2 = wb["GSTR-2B Apr25-Aug26"]; BH = {c.value: i for i, c in enumerate(next(b2.iter_rows(min_row=2, max_row=2))) if c.value}
B2 = [r for r in b2.iter_rows(min_row=3, values_only=True) if r[0]]; NB = 2 + len(B2); wb.close()
reg = [{"vendor_gstin": g(r, "Vendor GSTIN"), "invoice": g(r, "Invoice No."), "invoice_date": g(r, "Invoice Date"), "invoice_year": g(r, "Invoice Year"), "category": g(r, "Category"),
        "vel_gstin": g(r, "VEL GSTIN"), "igst": g(r, "IGST"), "cgst": g(r, "CGST"), "sgst": g(r, "SGST")} for r in rows]
b2r = [{"supplier_gstin": r[BH["Supplier GSTIN"]], "doc_no": r[BH["Doc No"]], "doc_date": r[BH["Doc Date"]], "company_gstin": r[BH["Company GSTIN"]], "key": S(r[BH["Supplier GSTIN"]]).upper() + norm(r[BH["Doc No"]]),
        "igst": r[BH["IGST (Net)"]], "cgst": r[BH["CGST (Net)"]], "sgst": r[BH["SGST (Net)"]]} for r in B2]
out = match_register(reg, b2r, set(), set())
GST = re.compile(r"^\d{2}[A-Z0-9]{13}$"); seen = set(); labels = []
for r in rows:
    if g(r, "Category") != "ITC": labels.append(None); continue
    vg = S(g(r, "Vendor GSTIN")).upper(); k = (vg if GST.match(vg) else "NM:" + S(g(r, "Vendor Name/RCM Category")).upper()) + "|" + norm(g(r, "Invoice No."))
    labels.append("Not consider" if k in seen else "Consider"); seen.add(k)

def consolidate(rows_, verdict_, key2_, cat_of, gstin_of, name_of, inv_of):
    """Sheet remarks per document (Countif grouping = vendor GSTIN or NM:name + normalised invoice, ITC lines only):
    first line of the group = Consider line -> best match of the group (lowest remark number, prior-year keys count);
    every other line of the group -> blank remark and blank KEY2. RCM/ISD lines untouched (each is its own document)."""
    import re as _re
    _G = _re.compile(r"^\d{2}[A-Z0-9]{13}$"); first = {}; members = {}
    for i, r in enumerate(rows_):
        if cat_of(r) != "ITC": continue
        vg = S(gstin_of(r)).upper(); k = (vg if _G.match(vg) else "NM:" + S(name_of(r)).upper()) + "|" + norm(inv_of(r))
        first.setdefault(k, i); members.setdefault(k, []).append(i)
    sv, sk = list(verdict_), list(key2_); moved = 0
    for k, idx in members.items():
        c = first[k]; matched = [i for i in idx if key2_[i]]
        if matched:
            best = min(matched, key=lambda i: (int(verdict_[i].split(" – ")[0]) if verdict_[i].split(" – ")[0].isdigit() else 99, i))
            if best != c and key2_[best] != key2_[c]: moved += 1
            sv[c], sk[c] = verdict_[best], key2_[best]
        for i in idx:
            if i != c: sv[i], sk[i] = "", ""
    print("remarks consolidated: %d ITC documents, %d Consider lines took a match found on another line, %d Not-consider lines blanked" % (len(members), moved, sum(len(v) - 1 for v in members.values())))
    return sv, sk
sheet_verdict, sheet_key2 = consolidate(rows, [o["verdict"] for o in out], [o["key2"] for o in out], lambda r: g(r, "Category"), lambda r: g(r, "Vendor GSTIN"), lambda r: g(r, "Vendor Name/RCM Category"), lambda r: g(r, "Invoice No."))
print("rows %d | verdicts %s | Countif %s" % (len(rows), dict(collections.Counter(o["verdict"].split(" – ")[0] + " – " + o["verdict"].split(" – ")[1] for o in out)), dict(collections.Counter(labels))))
NEW = ["KEY", "Countif", "B_IGST", "B_CGST", "B_SGST", "B_Total GST", "KEY2 (matched 2B key)", "2B_IGST", "2B_CGST", "2B_SGST", "2B_Total GST", "D_IGST", "D_CGST", "D_SGST", "D_Total GST", "Reco Remarks"]
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); wbx = xl.Workbooks.Open(P); xl.Calculation = -4135; sh = wbx.Worksheets("ITC Register 2026-27")
    HX = {sh.Cells(5, c).Value: c for c in range(1, 120) if sh.Cells(5, c).Value}; RN = sh.Cells(sh.Rows.Count, 4).End(-4162).Row; assert RN - 5 == len(rows), (RN, len(rows))
    if all(h in HX for h in NEW): start = HX["KEY"]
    else:
        start = max(HX.values()) + 1
        for i, h in enumerate(NEW):
            x = sh.Cells(5, start + i); x.Value = h; x.Font.Bold = True; x.Font.Color = 0xFFFFFF; x.Interior.Color = 0x4F3F33 if not h.startswith(("B_", "2B_", "D_")) else 0xB09784
    C = {h: L(start + i) for i, h in enumerate(NEW)}; c = lambda h: L(HX[h])
    rg = lambda h: sh.Range("%s6:%s%d" % (C[h], C[h], RN))
    rg("KEY").Formula = '=UPPER(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE($%s6&$%s6," ",""),"-",""),"/",""),".",""),"\'",""),"_",""))' % (c("Vendor GSTIN"), c("Invoice No."))
    rg("Countif").Value = [[v] for v in labels]; rg("KEY2 (matched 2B key)").Value = [[v or None] for v in sheet_key2]; rg("Reco Remarks").Value = [[v or None] for v in sheet_verdict]
    KEY, VN, K2, CF = C["KEY"], c("Vendor Name/RCM Category"), C["KEY2 (matched 2B key)"], C["Countif"]
    for b, src in (("B_IGST", "IGST"), ("B_CGST", "CGST"), ("B_SGST", "SGST")):
        rg(b).Formula = '=IF($%s6="Consider",SUMIFS($%s$6:$%s$%d,$%s$6:$%s$%d,$%s6,$%s$6:$%s$%d,$%s6),"NA")' % (CF, c(src), c(src), RN, KEY, KEY, RN, KEY, VN, VN, RN, VN)
    rg("B_Total GST").Formula = '=IF($%s6="Consider",%s6+%s6+%s6,"NA")' % (CF, C["B_IGST"], C["B_CGST"], C["B_SGST"])
    for t, col in (("2B_IGST", "V"), ("2B_CGST", "W"), ("2B_SGST", "X")):
        rg(t).Formula = '=IF($%s6="Consider",IF($%s6="",0,SUMIFS(\'GSTR-2B Apr25-Aug26\'!$%s$3:$%s$%d,\'GSTR-2B Apr25-Aug26\'!$AW$3:$AW$%d,$%s6)),"NA")' % (CF, K2, col, col, NB, NB, K2)
    rg("2B_Total GST").Formula = '=IF($%s6="Consider",%s6+%s6+%s6,"NA")' % (CF, C["2B_IGST"], C["2B_CGST"], C["2B_SGST"])
    for d, b, t in (("D_IGST", "B_IGST", "2B_IGST"), ("D_CGST", "B_CGST", "2B_CGST"), ("D_SGST", "B_SGST", "2B_SGST"), ("D_Total GST", "B_Total GST", "2B_Total GST")):
        rg(d).Formula = '=IF($%s6="Consider",%s6-%s6,"NA")' % (CF, C[b], C[t])
    for h in ("B_IGST", "B_CGST", "B_SGST", "B_Total GST", "2B_IGST", "2B_CGST", "2B_SGST", "2B_Total GST", "D_IGST", "D_CGST", "D_SGST", "D_Total GST"): rg(h).NumberFormat = "#,##0.00"
    sh.Columns(C["Reco Remarks"]).ColumnWidth = 60; pass  # layout/autofilter owned by fy2627_relayout.py (25-26 layout, header row 5)
    xl.Calculation = -4105; xl.CalculateFullRebuild()
    cs = lambda h: sum(v[0] for v in rg(h).Value if isinstance(v[0], (int, float)))
    itc_tax = sum(num(g(r, "IGST")) + num(g(r, "CGST")) + num(g(r, "SGST")) for r in rows if g(r, "Category") == "ITC")
    e = 0
    for w in wbx.Worksheets:
        try: e += w.UsedRange.SpecialCells(-4123, 16).Count
        except Exception: pass
    reg25 = wbx.Worksheets("ITC Register 2025-26"); H25 = {reg25.Cells(5, k).Value: k for k in range(1, 90) if reg25.Cells(5, k).Value}
    print("B_Total %.2f vs ITC tax %.2f | 2B_Total %.2f | D_Total %.2f | error cells %d | FY 25-26 golden %.2f" % (cs("B_Total GST"), itc_tax, cs("2B_Total GST"), cs("D_Total GST"), e, reg25.Cells(4, H25["Total GST"]).Value))
    assert e == 0 and abs(cs("B_Total GST") - itc_tax) < 1
    wbx.Save()
    if xlsb_free(): wbx.SaveAs(B, FileFormat=50); print("xlsb exported")
    else: print("xlsb OPEN in Excel - export skipped")
    wbx.Close(False); print("saved (%.0fs)" % (time.time() - t0))
finally: xl.Quit()
