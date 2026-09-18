import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
SP=os.path.dirname(os.path.abspath(__file__))
C=json.load(open(os.path.join(SP,"counts.json")))
HDR=Font(bold=True,color="FFFFFF"); FILL=PatternFill("solid",fgColor="1F4E79")
ASK=PatternFill("solid",fgColor="FFF2CC")
wb=Workbook(); ws=wb.active; ws.title="Column Map"
cols=["Target Col","Target Header (format file)","Kind","Source column in ClearTax 'Working' sheet","Resolves in all 12 months?","Notes / open question","CA confirm"]
ws.append(cols)
M=[("A","Month","Derived","Document Date","Yes","Month name taken from Document Date",""),
("B","My State","Mapped","State","Yes","Needs state-name standardisation - see 'State Names' sheet",""),
("C","My GSTIN","Mapped","GSTIN","Yes","Clean - 17 distinct GSTINs. Safer key than state name.",""),
("D","Supplier Legal Name","Mapped","Seller Legal Name","Yes","",""),
("E","Document Date","Mapped","Document Date","Yes","",""),
("F","Document Number","Mapped","Document Number","Yes","",""),
("G","Invoice Type","Derived","Document Type","Yes","'Sales' or 'Credit Notes', from document type",""),
("H","Document Type Code","Mapped","Document Type","Yes","Needs standardisation - see 'Document Types' sheet",""),
("I","Supply Type Code","Mapped","Supply Type","Yes","BLANK on all advance rows - needs filling",""),
("J","Recipient GSTIN","Mapped","Buyer GSTIN","Yes","",""),
("K","Recipient Legal Name","Mapped","Buyer Legal Name","Yes","",""),
("L","Place of Supply","Mapped","Buyer Place of Supply","Yes","",""),
("M","Item Price","Mapped","Item Unit Price","Yes","",""),
("N","Gross Amount","Mapped","Item Total Amount","Yes","",""),
("O","GST Rate","Mapped","Item GST Rate","Yes","",""),
("P","GST rate check","Derived","-","-","Tax / taxable must land on a valid GST rate",""),
("Q","Taxable Value","Mapped","Item Assessable Amount","Yes","",""),
("R","IGST Amount","Mapped","Item IGST Amount","Yes","",""),
("S","CGST Amount","Mapped","Item CGST Amount","Yes","",""),
("T","SGST Amount","Mapped","Item SGST Amount","Yes","",""),
("U","Cess Amount","Mapped","Item CESS Amount","Yes","",""),
("V","Total GST","Formula","-","-","=SUM(R:T). Source also has 'Total Tax' - will cross-check",""),
("W","Total Invoice Value","Formula","-","-","=SUM(Q:U)",""),
("X","Invoice Reference No","Mapped","Invoice Reference No","Yes","The IRN",""),
("Y","IRN (len)","Derived","-","-","=LEN(X). Must be 64, or blank for non-e-invoiced rows",""),
("Z","Status","Mapped","Status","Yes","",""),
("AA","Cancel Date","Mapped","Cancel Date","Yes","",""),
("AB","Document Status","Mapped","Document Status","Yes","",""),
("AC","Ship To GSTIN","Mapped","Shipping GSTIN","Yes","",""),
("AD","Ship To LegalName","Mapped","Shipping LegalName","Yes","",""),
("AE","Item Description","Mapped","Item Product Description","Yes","",""),
("AF","GOOD (G) or SERVICE(S)","Derived","Item HSN Code","Yes","SAC starting 99 = Service. ClearTax also has 'Item Is Service' - will cross-check",""),
("AG","HSN or SAC Code","Mapped","Item HSN Code","Yes","Plain HSN code - confirmed by CA (not the 'for Uploading' variant)",""),
("AH","Item Description for uploading","Mapped","Item Product Description","Yes","Plain description - confirmed by CA",""),
("AI","Quantity for uploading","Mapped","Item Quantity","Yes","",""),
("AJ","Unit of Measurement for uploading","Mapped","Item Unit","Yes","Plain unit - kept consistent with AG (raw ClearTax values, not client-edited)",""),
("AK","GL Name","No source","-","-","Comes from the SAP GL - not received yet",""),
("AL","Profit Centre","No source","-","-","Comes from the SAP GL - not received yet",""),
("AM","Cost Centre","No source","-","-","Comes from the SAP GL - not received yet","")]
for r in M: ws.append(list(r))
for i,c in enumerate(ws[1],1): c.font=HDR; c.fill=FILL
for row in ws.iter_rows(min_row=2):
    if "QUESTION" in str(row[5].value): 
        for c in row: c.fill=ASK
for col,w in zip("ABCDEFG",[10,36,11,42,14,62,14]): ws.column_dimensions[col].width=w
ws.freeze_panes="A2"

# State names
s2=wb.create_sheet("State Names")
s2.append(["Spelling in ClearTax data","Rows in year","Standardise to","CA confirm"])
FIX={"Madya Pradesh":"Madhya Pradesh","Gujrat":"Gujarat","Chattisgarh":"Chhattisgarh",
     "Tamilnadu":"Tamil Nadu","UttarPradesh":"Uttar Pradesh","West bengal":"West Bengal"}
for k,v in sorted(C["states"].items(), key=lambda x:-x[1]):
    s2.append([k,v,FIX.get(k,k),""])
for c in s2[1]: c.font=HDR; c.fill=FILL
for row in s2.iter_rows(min_row=2):
    if row[0].value!=row[2].value:
        for c in row: c.fill=ASK
for col,w in zip("ABCD",[30,14,24,14]): s2.column_dimensions[col].width=w
s2.freeze_panes="A2"

# Doc types
s3=wb.create_sheet("Document Types")
s3.append(["Value in ClearTax data","Rows in year","Standardise to","Proposed GSTR-9 table","CA confirm"])
DT={"INV":("INV","4B if B2B / 4A if B2C"),"CRN":("CRN","4I"),"CRN( Manual)":("CRN","4I"),"DBN":("DBN","4J"),
 "MOB ADV REC":("MOB ADV REC","4F"),"MOB ADV RECEIVED":("MOB ADV REC","4F"),"Mob Adv Rec":("MOB ADV REC","4F"),
 "MOB ADV REC INT":("MOB ADV REC","4F"),
 "MOB ADV ADJ":("MOB ADV ADJ","4F (negative)"),"MOB ADV ADJUSTMENT":("MOB ADV ADJ","4F (negative)"),
 "MOB ADV ADJUSMENT":("MOB ADV ADJ","4F (negative)"),"MOB ADV ADJUSTME":("MOB ADV ADJ","4F (negative)"),
 "Mob advance adj":("MOB ADV ADJ","4F (negative)"),
 "MOB ADV REV":("MOB ADV REV","4F (negative)"),"MOB ADV REVERSE":("MOB ADV REV","4F (negative)")}
for k,v in sorted(C["doctypes"].items(), key=lambda x:-x[1]):
    std,tbl=DT.get(k,("??","??")); s3.append([k,v,std,tbl,""])
for c in s3[1]: c.font=HDR; c.fill=FILL
for row in s3.iter_rows(min_row=2):
    if row[0].value!=row[2].value:
        for c in row: c.fill=ASK
for col,w in zip("ABCDE",[26,14,20,24,14]): s3.column_dimensions[col].width=w
s3.freeze_panes="A2"
out=os.path.join(SP,"VEL_Sales_Column_Map_DRAFT.xlsx"); wb.save(out)
print("WROTE", out)
print(f"rows in year: {C['rows']}, state spellings: {len(C['states'])}, doc types: {len(C['doctypes'])}")
