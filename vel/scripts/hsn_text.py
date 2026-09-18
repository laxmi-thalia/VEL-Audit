import openpyxl
P=r"C:\Users\pawar\Downloads\VEL_Sales_Register_FY2025-26_DRAFT.xlsx"
wb=openpyxl.load_workbook(P); ws=wb["SR_2025-26"]
AG,AF=33,32   # column indexes
conv=blank=0
for i in range(5,27007):
    c=ws.cell(i,AG); v=c.value
    if v is None or (isinstance(v,str) and not v.strip()):
        c.value=None; blank+=1
    else:
        s=str(int(v)) if isinstance(v,(int,float)) and float(v).is_integer() else str(v).strip()
        c.value=s; c.number_format="@"; conv+=1
    # AF must cope with a TEXT hsn now
    ws.cell(i,AF).value=f'=IF(AG{i}="","",IF(LEFT(AG{i}&"",2)="99","S","G"))'
wb.save(P)
print(f"HSN converted to text: {conv} cells | left blank: {blank}")
