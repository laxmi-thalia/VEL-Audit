import openpyxl
from openpyxl.utils import get_column_letter
p="//server/GST FOLDER/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/DPS Workings/1. Sales/Sales Register_FY 2025-26 Format.xlsx"
wb=openpyxl.load_workbook(p, data_only=False)
ws=wb["SR_2025-26"]
print("sheets:",wb.sheetnames," dims:",ws.dimensions," merged:",list(ws.merged_cells.ranges)[:10])
for r in range(1,10):
    cells=[f"{get_column_letter(c)}={ws.cell(r,c).value!r}" for c in range(1,ws.max_column+1) if ws.cell(r,c).value is not None]
    print(f"row{r}: "+(" | ".join(cells) if cells else "(empty)"))
