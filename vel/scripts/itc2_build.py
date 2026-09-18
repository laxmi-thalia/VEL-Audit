"""ITC step 2: Net-ITC columns on 3B Data (+4A3 RCM, 4A4 ISD) + 'ITCR vs 3B Net ITC' MoM
sheet (GSTIN x month, live SUMIFS both sides, standard tax-head sequence, DPS Remarks).
Octa 4B rows are NEGATIVE as reported -> Net = 4A + 4B."""
import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f = open(P, 'r+b'); f.close()
T = pd.read_pickle("b3itc.pkl"); T["g"] = T["gstin"].map(S)
HF=Font(bold=True,color="FFFFFF"); HB=PatternFill("solid",fgColor="1F4E79"); TOT=Font(bold=True); NF="#,##0.00"
wb = openpyxl.load_workbook(P)
b3 = wb["3B Data"]
hdr = {S(b3.cell(1,c).value): c for c in range(1, b3.max_column+1)}
base = b3.max_column
names = ["Net ITC 4(C) IGST","Net ITC 4(C) CGST","Net ITC 4(C) SGST",
         "4A(3) RCM ITC IGST","4A(3) RCM ITC CGST","4A(3) RCM ITC SGST",
         "4A(4) ISD ITC IGST","4A(4) ISD ITC CGST","4A(4) ISD ITC SGST"]
cols = []
for i, nm in enumerate(names):
    c = hdr.get(nm, base+1+i)
    x = b3.cell(1, c, nm); x.font=HF; x.fill=HB
    b3.column_dimensions[L(c)].width = 17
    cols.append(c)
MON = {m: pd.to_datetime(m, format="%b-%y").strftime("%b %Y") for m in
       ["Apr-25","May-25","Jun-25","Jul-25","Aug-25","Sep-25","Oct-25","Nov-25","Dec-25","Jan-26","Feb-26","Mar-26"]}
for r in range(2, b3.max_row+1):
    g = S(b3.cell(r,2).value); m = MON.get(S(b3.cell(r,3).value))
    row = T[T["g"]==g]
    def gv(grp, meas):
        if not len(row) or not m: return 0.0
        return float(row.iloc[0].get("%s|%s|%s" % (m, grp, meas), 0) or 0)
    vals = [gv("a","igst")+gv("b","igst"), gv("a","cgst")+gv("b","cgst"), gv("a","sgst")+gv("b","sgst"),
            gv("a3","igst"), gv("a3","cgst"), gv("a3","sgst"),
            gv("a4","igst"), gv("a4","cgst"), gv("a4","sgst")]
    for i, v in enumerate(vals):
        cell = b3.cell(r, cols[i]); cell.value = round(v,2); cell.number_format = NF
n3 = b3.max_row
B3C = {k: "'3B Data'!$%s$2:$%s$%d" % (L(c), L(c), n3) for k, c in zip(("ni","nc","ns"), cols[:3])}
B3G = "'3B Data'!$B$2:$B$%d" % n3
B3M = "'3B Data'!$C$2:$C$%d" % n3
# register ranges
rr = wb["ITC Register 2025-26"]
RH = {}
for c in range(1, rr.max_column+1):
    h = S(rr.cell(5,c).value)
    if h and h not in RH: RH[h] = c
R0, R1 = 6, 6+41508-1
REG = lambda n: "'ITC Register 2025-26'!$%s$%d:$%s$%d" % (L(RH[n]), R0, L(RH[n]), R1)
GSTINS = [("01AAECR0503Q1ZM","Jammu & Kashmir"),("03AAECR0503Q1ZI","Punjab"),("06AAECR0503Q1ZC","Haryana"),
("08AAECR0503Q1Z8","Rajasthan"),("09AAECR0503Q1Z6","Uttar Pradesh"),("10AAECR0503Q1ZN","Bihar"),
("12AAECR0503Q1ZJ","Arunachal Pradesh"),("18AAECR0503Q1Z7","Assam"),("19AAECR0503Q1Z5","West Bengal"),
("20AAECR0503Q1ZM","Jharkhand"),("22AAECR0503Q1ZI","Chhattisgarh"),("23AAECR0503Q1ZG","Madhya Pradesh"),
("24AAECR0503Q1ZE","Gujarat"),("27AAECR0503Q1Z8","Maharashtra"),("29AAECR0503Q1Z4","Karnataka"),
("32AAECR0503Q1ZH","Kerala"),("33AAECR0503Q1ZF","Tamil Nadu"),("36AAECR0503Q1Z9","Telangana"),
("37AAECR0503Q1Z7","Andhra Pradesh")]
MONTHS = ["01 Apr 2025","02 May 2025","03 June 2025","04 July 2025","05 Aug 2025","06 Sep 2025",
"07 Oct 2025","08 Nov 2025","09 Dec 2025","10 Jan 2026","11 Feb 2026","12 Mar 2026"]
M2TOK = dict(zip(MONTHS, ["Apr-25","May-25","Jun-25","Jul-25","Aug-25","Sep-25","Oct-25","Nov-25","Dec-25","Jan-26","Feb-26","Mar-26"]))
NM = "ITCR vs 3B Net ITC"
if NM in wb.sheetnames: del wb[NM]
pos = wb.sheetnames.index("GSTR-2B ITC Data") + 1
ws = wb.create_sheet(NM, pos)
ws.row_dimensions[1].height = 21
ws.cell(2,1,"VIKRAN ENGINEERING LIMITED").font = TOT
ws.cell(3,1,"ITC Register (claims by 3B month) vs GSTR-3B Net ITC (Table 4A + 4B as reported) - month on month, live. DPS Remarks: type here.").font = TOT
hdrs = ["State","GSTIN","Month",
 "Register IGST","Register CGST","Register SGST","Register Total",
 "3B Net ITC IGST","3B Net ITC CGST","3B Net ITC SGST","3B Net ITC Total",
 "Diff IGST","Diff CGST","Diff SGST","Diff Total","DPS Remarks"]
for c,h in enumerate(hdrs,1):
    x=ws.cell(5,c,h); x.font=HF; x.fill=HB
r = 5
for g_, st in GSTINS:
    for m in MONTHS:
        r += 1
        ws.cell(r,1,st); ws.cell(r,2,g_); ws.cell(r,3,m); ws.cell(r,18,M2TOK[m])
        for i, meas in enumerate(("IGST","CGST","SGST")):
            ws.cell(r,4+i).value = '=SUMIFS(%s,%s,$B%d,%s,$C%d)' % (REG(meas), REG("VEL GSTIN"), r, REG("3B Claim  Month"), r)
        ws.cell(r,7).value = "=SUM(D%d:F%d)" % (r,r)
        for i, k in enumerate(("ni","nc","ns")):
            ws.cell(r,8+i).value = '=SUMIFS(%s,%s,$B%d,%s,$R%d)' % (B3C[k], B3G, r, B3M, r)
        ws.cell(r,11).value = "=SUM(H%d:J%d)" % (r,r)
        for i in range(3):
            ws.cell(r,12+i).value = "=%s%d-%s%d" % (L(4+i), r, L(8+i), r)
        ws.cell(r,15).value = "=G%d-K%d" % (r,r)
        for c in range(4,16): ws.cell(r,c).number_format = NF
last = r
r += 1
ws.cell(r,1,"Grand Total").font = TOT
for c in range(4,16):
    ws.cell(r,c).value = "=SUM(%s6:%s%d)" % (L(c),L(c),last)
    ws.cell(r,c).number_format = NF; ws.cell(r,c).font = TOT
for c_,w in zip(range(1,17),[18,17,12]+[14]*12+[36]): ws.column_dimensions[L(c_)].width = w
ws.column_dimensions["R"].hidden = True
ws.freeze_panes = "D6"
ws.auto_filter.ref = "A5:P%d" % last
ix = wb["INDEX"]
have = {S(ix.cell(r2,2).value) for r2 in range(5, ix.max_row+1)}
if "ITC Register vs GSTR-3B Net ITC" not in have:
    thin=Side(style="thin"); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
    r2 = ix.max_row+1
    ix.cell(r2,1,"ITC"); ix.cell(r2,2,"ITC Register vs GSTR-3B Net ITC")
    x=ix.cell(r2,7); x.value='=HYPERLINK("#\'%s\'!A1","ITCR vs 3B")' % NM
    x.font=Font(color="0563C1",underline="single")
    for c in range(1,11): ix.cell(r2,c).border=BD
wb.save(P)
print("built %s + 9 cols on 3B Data" % NM)
