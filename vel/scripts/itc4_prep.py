"""Prep for GSTIN-level + receivables recos: (a) parse Receivables Open Items (FBL3N style),
(b) pull last year's ITC register claims (vendor GSTIN + invoice) from the xlsb via COM."""
import openpyxl, re, pandas as pd, os
import win32com.client as win32, pythoncom
def S(v):
    if v is None: return ""
    if isinstance(v, float) and pd.isna(v): return ""
    return str(v).strip()
def ninv(v): return re.sub(r"[^A-Z0-9]", "", S(v).upper().lstrip("'"))
# (a) receivables open items
wb = openpyxl.load_workbook("recv_open.xlsx", read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
raw = list(ws.iter_rows(values_only=True)); wb.close()
hdr_i = next((i for i, r in enumerate(raw) if any(S(v) == "Document Number" for v in r)), None)
H = {S(v): j for j, v in enumerate(raw[hdr_i]) if S(v)}
print("open-items headers:", list(H.keys()))
rows = []
for i in range(hdr_i + 1, len(raw)):
    r = raw[i]
    doc = S(r[H["Document Number"]]) if len(r) > H["Document Number"] else ""
    if not doc.isdigit(): continue
    rows.append({k: (r[j] if len(r) > j else None) for k, j in H.items()})
ro = pd.DataFrame(rows)
print("open-item rows:", len(ro))
ro.to_pickle("recv_open.pkl")
# (b) last year's claims from xlsb
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    wb2 = xl.Workbooks.Open(os.path.abspath("VEL_2425_fresh.xlsb"), ReadOnly=True)
    sh = wb2.Worksheets("ITC Register 2024-25")
    ur = sh.UsedRange
    nr = ur.Rows.Count
    # cols: L=12 Invoice No., R=18 Vendor GSTIN, E=5 3B Claim Month (header row 5, data 6+)
    inv = sh.Range(sh.Cells(6, 12), sh.Cells(nr, 12)).Value
    vgs = sh.Range(sh.Cells(6, 18), sh.Cells(nr, 18)).Value
    wb2.Close(SaveChanges=False)
finally:
    xl.Quit()
ly = set()
for a, b in zip(inv, vgs):
    i_, g_ = ninv(a[0]), S(b[0]).upper()
    if g_ and i_: ly.add((g_, i_))
print("last-year claim keys:", len(ly))
pd.to_pickle(ly, "ly_claims.pkl")
