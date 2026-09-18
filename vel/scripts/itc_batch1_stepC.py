"""Batch-1 step C:
 1. 'GSTR-2B Apr25-Aug26' vice-versa + classification columns (all live): 3B Claim Month, Reco Remarks (from the register
    by KEY2), 6A1 mark, Table 8A (Octa's own 'GSTR-9 (8A) ITC Available' + FY), 8A Reco, Table 8C, Table 13, Final Remarks,
    GSTR-9/9C. Permanent Reversals / Reclaim 6H / Query left for the CA (amber).
 2. 'ITCR vs 3B Net ITC' rebuilt with ISD / RCM / Other split on both sides (Rashid), DPS Remarks last.
 3. 'Unclaimed ITC candidates' sheet removed (+ INDEX row); INDEX rows refreshed."""
import os, re, warnings
from copy import copy
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
open(P, "r+b").close()
NF = "#,##0.00"; HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79"); AMB = PatternFill("solid", fgColor="FFF2CC"); TOT = Font(bold=True)
wb = openpyxl.load_workbook(P)
# ---------- 1. 2B sheet columns ----------
b2 = wb["GSTR-2B Apr25-Aug26"]
BH = {S(b2.cell(2, c).value): c for c in range(1, b2.max_column + 1)}
NB = max(r for r in range(3, b2.max_row + 1) if b2.cell(r, 1).value)
reg = wb["ITC Register 2025-26"]
RH = {S(reg.cell(5, c).value): c for c in range(1, reg.max_column + 1)}
RN = max(r for r in range(6, reg.max_row + 1) if reg.cell(r, 4).value)
RG = lambda h: "'ITC Register 2025-26'!$%s$6:$%s$%d" % (L(RH[h]), L(RH[h]), RN)
c = lambda h: L(BH[h])
for r in range(3, NB + 1):
    k = "$%s%d" % (c("KEY"), r)
    b2.cell(r, BH["3B Claim Month"]).value = '=IFERROR(INDEX(%s,MATCH(%s,%s,0)),"")' % (RG("3B Claim  Month"), k, RG("KEY2 (matched 2B key)"))
    b2.cell(r, BH["3B Claim Month"]).number_format = "mmm-yy"
    b2.cell(r, BH["Reco Remarks"]).value = '=IF($%s%d="","NOT IN ITC REGISTER (FY 25-26 claims)","Matched with ITC Register - claimed "&TEXT($%s%d,"mmm-yy"))' % (c("3B Claim Month"), r, c("3B Claim Month"), r)
    b2.cell(r, BH["6A1 mark"]).value = ('=IF($%s%d="2024-25",IF($%s%d="","6A1 - Unclaimed (24-25 dated, no 3B claim month)",IF($%s%d="2024-25","6A1 - 24-25 inv in 24-25 2B, claimed 25-26","6A1 - 24-25 inv in 25-26 2B, claimed 25-26")),"")'
                                        % (c("Doc FY (doc date)"), r, c("3B Claim Month"), r, c("FY (2B period)"), r))
    b2.cell(r, BH["Table 8A"]).value = ('=IF($%s%d<>"2025-26","No - 2B of "&$%s%d,IF($%s%d="Yes","No - RCM",IF($%s%d="No","No - ITC not available",IF($%s%d="Yes","No - Amendment",IF($%s%d<>"2025-26","No - "&$%s%d&" dated","Yes")))))'
                                        % (c("FY (2B period)"), r, c("FY (2B period)"), r, c("Reverse Charge"), r, c("GSTR-9 (8A) ITC Available"), r, c("Is Amendment"), r, c("Doc FY (doc date)"), r, c("Doc FY (doc date)"), r))
    b2.cell(r, BH["8A Reco"]).value = '=IF($%s%d="Yes","To be considered","Ignore")' % (c("Table 8A"), r)
    b2.cell(r, BH["Table 8C"]).value = '=IF(AND($%s%d="2025-26",$%s%d="2026-27"),"Table 8C of GSTR-9","")' % (c("Doc FY (doc date)"), r, c("FY (2B period)"), r)
    b2.cell(r, BH["Table 13"]).value = '=IF($%s%d<>"","Table 13 - 25-26 ITC availed in 26-27 (confirm with FY 26-27 register)","")' % (c("Table 8C"), r)
    b2.cell(r, BH["Final Remarks"]).value = '=IF($%s%d<>"",$%s%d,IF($%s%d<>"",$%s%d,$%s%d))' % (c("6A1 mark"), r, c("6A1 mark"), r, c("Table 8C"), r, c("Table 8C"), r, c("Reco Remarks"), r)
    b2.cell(r, BH["GSTR-9/9C"]).value = ('=IF(LEFT($%s%d,3)="6A1","Table 6A1 of GSTR-9 - "&IF(ISNUMBER(SEARCH("Unclaimed",$%s%d)),"Unclaimed","Claimed"),IF($%s%d<>"","Table 8C of GSTR-9",IF($%s%d="Yes","Table 8A / 6B of GSTR-9","")))'
                                         % (c("6A1 mark"), r, c("6A1 mark"), r, c("Table 8C"), r, c("Table 8A"), r))
for h in ("Permanent Reversals", "Reclaim - Table 6H", "Query"):
    b2.cell(2, BH[h]).fill = AMB; b2.cell(2, BH[h]).font = Font(bold=True)
b2.cell(1, 3).value = S(b2.cell(1, 3).value) + "  |  Vice-versa: 3B Claim Month + Reco Remarks read from the register by KEY2; 6A1 / 8A / 8C / 13 marks are live formulas; amber columns are for the CA."
print("2B sheet: formulas written rows 3..%d" % NB)
# ---------- 2. ITCR vs 3B split ----------
old = wb["ITCR vs 3B Net ITC"]
trip = [(old.cell(r, 1).value, old.cell(r, 2).value, old.cell(r, 3).value, old.cell(r, 18).value) for r in range(6, old.max_row + 1) if old.cell(r, 2).value]
rem = {(t[1], t[2]): old.cell(r, 16).value for r, t in zip(range(6, 6 + len(trip)), trip)}
pos = wb.sheetnames.index("ITCR vs 3B Net ITC"); del wb["ITCR vs 3B Net ITC"]
ws = wb.create_sheet("ITCR vs 3B Net ITC", pos)
ws.row_dimensions[1].height = 21
ws.cell(2, 1, "VIKRAN ENGINEERING LIMITED").font = TOT
ws.cell(3, 1, "ITC Register (claims by 3B month) vs GSTR-3B - month on month, live, SPLIT into RCM (4A3) / ISD (4A4) / Other (4A5 + 4B reversals as reported) per CA ruling 17-09. Net = sum of the three. DPS Remarks: type in the last column.").font = Font(italic=True, color="808080")
b3 = wb["3B Data"]; B3H = {S(b3.cell(2, cc).value): cc for cc in range(1, b3.max_column + 1)}
N3 = max(r for r in range(3, b3.max_row + 1) if b3.cell(r, 2).value)
def b3c(pre, tail): return L(next(cc for h, cc in B3H.items() if h.startswith(pre) and h.endswith(tail)))
B3 = lambda col: "'3B Data'!$%s$3:$%s$%d" % (col, col, N3)
blocks = [("RCM (4A3)", "RCM", [b3c("4A(3)", t) for t in ("IGST", "CGST", "SGST")], None),
          ("ISD (4A4)", "ISD", [b3c("4A(4)", t) for t in ("IGST", "CGST", "SGST")], None),
          ("Other (4A5 + 4B)", "ITC", [b3c("4A(5)", t) for t in ("IGST", "CGST", "SGST")], [[b3c("4B(1)", t), b3c("4B(2)", t)] for t in ("IGST", "CGST", "SGST")])]
hdr = ["State", "GSTIN", "Month"]
for lbl, cat, cols, adds in blocks:
    hdr += ["Reg %s %s" % (lbl, t) for t in ("IGST", "CGST", "SGST")] + ["3B %s %s" % (lbl, t) for t in ("IGST", "CGST", "SGST")] + ["Diff %s %s" % (lbl, t) for t in ("IGST", "CGST", "SGST")]
hdr += ["Register Net Total", "3B Net ITC Total (4A+4B)", "Diff Net Total", "DPS Remarks (type here)", "Month (helper)"]
for cc, h in enumerate(hdr, 1):
    x = ws.cell(5, cc, h); x.font = HF; x.fill = HB; x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
ws.row_dimensions[5].height = 32
MON = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26"]
import datetime as dt
r = 5
mh = len(hdr)
for st, g, m, mlbl in trip:
    r += 1
    ws.cell(r, 1, st); ws.cell(r, 2, g); ws.cell(r, 3, m); ws.cell(r, mh, mlbl)
    mi = MON.index(S(mlbl)); mdate = dt.datetime(2025 + (1 if mi >= 9 else 0), (mi + 3) % 12 + 1, 1)
    ws.cell(r, mh + 1, mdate).number_format = "DD.MM.YYYY"   # register claim month is a date
    col = 4
    regtot = []; b3tot = []
    for lbl, cat, cols, adds in blocks:
        for j, t in enumerate(("IGST", "CGST", "SGST")):
            ws.cell(r, col + j, "=SUMIFS(%s,%s,$B%d,%s,$%s%d,%s,\"%s\")" % (RG(t), RG("VEL GSTIN"), r, RG("3B Claim  Month"), L(mh + 1), r, RG("Category"), cat))
            f = "SUMIFS(%s,%s,$B%d,%s,$%s%d)" % (B3(cols[j]), B3("B"), r, B3("C"), L(mh), r)
            if adds: f += "+" + "+".join("SUMIFS(%s,%s,$B%d,%s,$%s%d)" % (B3(a), B3("B"), r, B3("C"), L(mh), r) for a in adds[j])
            ws.cell(r, col + 3 + j, "=" + f)
            ws.cell(r, col + 6 + j, "=%s%d-%s%d" % (L(col + j), r, L(col + 3 + j), r))
            regtot.append("%s%d" % (L(col + j), r)); b3tot.append("%s%d" % (L(col + 3 + j), r))
        col += 9
    ws.cell(r, col, "=" + "+".join(regtot)); ws.cell(r, col + 1, "=" + "+".join(b3tot)); ws.cell(r, col + 2, "=%s%d-%s%d" % (L(col), r, L(col + 1), r))
    ws.cell(r, col + 3).fill = AMB
    if rem.get((g, m)): ws.cell(r, col + 3, rem[(g, m)])
    for cc in range(4, col + 3): ws.cell(r, cc).number_format = NF
RN2 = r; r += 1
ws.cell(r, 1, "Grand Total").font = TOT
for cc in range(4, mh - 1):
    ws.cell(r, cc, "=SUM(%s6:%s%d)" % (L(cc), L(cc), RN2)); ws.cell(r, cc).font = TOT; ws.cell(r, cc).number_format = NF
ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 12
for cc in range(4, mh): ws.column_dimensions[L(cc)].width = 14
ws.column_dimensions[L(mh - 1)].width = 34; ws.column_dimensions[L(mh)].hidden = True; ws.column_dimensions[L(mh + 1)].hidden = True
ws.freeze_panes = "D6"; ws.auto_filter.ref = "A5:%s%d" % (L(mh - 1), RN2)
print("ITCR vs 3B rebuilt: %d rows, %d cols" % (len(trip), mh + 1))
# ---------- 3. remove Unclaimed sheet + INDEX ----------
if "Unclaimed ITC candidates" in wb.sheetnames: del wb["Unclaimed ITC candidates"]
ix = wb["INDEX"]
for rr in range(ix.max_row, 4, -1):
    if "Unclaimed ITC" in S(ix.cell(rr, 2).value): ix.delete_rows(rr, 1); print("INDEX row removed", rr)
thin = Side(style="thin"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
for rr in range(5, ix.max_row + 1):
    v = S(ix.cell(rr, 2).value)
    if v.startswith("ITC Register vs GSTR-3B") or "vs 3B Net ITC" in v: ix.cell(rr, 2).value = "ITC Register vs GSTR-3B - month on month, split RCM / ISD / Other (live)"
wb.save(P); print("saved")
