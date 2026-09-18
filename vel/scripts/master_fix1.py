import openpyxl
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
wb = openpyxl.load_workbook(P)
ws0 = wb["SR_2025-26"]
H = {S(ws0.cell(4, c).value): L(c) for c in range(1, ws0.max_column + 1)}
R0, R1 = 5, 27006
SR = lambda name: "'SR_2025-26'!$%s$%d:$%s$%d" % (H[name], R0, H[name], R1)
NG = wb["GSTR-1 Data"].max_row
G1 = lambda c: "'GSTR-1 Data'!$%s$2:$%s$%d" % (c, c, NG)
# ---- fix 1: S2 MoM K/L per-type criteria
mm = wb["S2 Month-on-Month"]
def bucket_crit(row):
    # returns extra SUMIFS criteria (register) for the row's type in D
    return (',%s,IF($D{r}="B2B","INV",IF($D{r}="B2C","INV",IF($D{r}="Credit note","CRN",IF($D{r}="Debit note","DBN","~none"))))'
            % SR("Document Type Code")).replace("{r}", str(row))
for r in range(4, mm.max_row + 1):
    if not S(mm.cell(r, 1).value): continue
    # books-not-in-G1, restricted to the row's type; advances rows get 0 (they are summary-level, never 'Not in GSTR-1')
    mm.cell(r, 11).value = ('=IF(OR($D{r}="Advance received",$D{r}="Advance adjusted"),0,'
        'SUMIFS(%s,%s,$B{r},%s,$C{r},%s,"Not in GSTR-1",'
        '%s,IF($D{r}="B2B","INV",IF($D{r}="B2C","INV",IF($D{r}="Credit note","CRN","DBN"))),'
        '%s,IF($D{r}="B2B","<>B2C",IF($D{r}="B2C","B2C","*"))))'
        % (SR("Taxable Value"), SR("My GSTIN"), SR("Month"), SR("Matched with GSTR-1"),
           SR("Document Type Code"), SR("Supply Type Code"))).replace("{r}", str(r))
    mm.cell(r, 12).value = ('=-SUMIFS(%s,%s,$B{r},%s,$C{r},%s,"Not in books",%s,$D{r})'
        % (G1("F"), G1("A"), G1("K"), G1("N"), G1("J"))).replace("{r}", str(r))
# ---- fix 2: S4 Tamil Nadu label
s4 = wb["S4 GL vs SR"]
for r in range(4, 24):
    if S(s4.cell(r, 2).value) == "TamilNadu": s4.cell(r, 2).value = "Tamil Nadu"
# ---- fix 3: S7 blank-as-zero guards
s7 = wb["S7 CN Time-bar"]
Hn, Hd, Hf = H["Original Invoice Number"], H["Original Invoice Date"], H["Document Number"]
for r in range(4, s7.max_row + 1):
    if not S(s7.cell(r, 2).value): continue
    s7.cell(r, 5).value = ('=IF(N(INDEX(\'SR_2025-26\'!$%s:$%s,MATCH($B{r},\'SR_2025-26\'!$%s:$%s,0)))=0,'
        'IFERROR(INDEX(\'SR_2025-26\'!$%s:$%s,MATCH($B{r},\'SR_2025-26\'!$%s:$%s,0))&"",""),'
        'INDEX(\'SR_2025-26\'!$%s:$%s,MATCH($B{r},\'SR_2025-26\'!$%s:$%s,0)))'
        % (Hn, Hn, Hf, Hf, Hn, Hn, Hf, Hf, Hn, Hn, Hf, Hf)).replace("{r}", str(r))
    s7.cell(r, 5).value = ('=IFERROR(IF(INDEX(\'SR_2025-26\'!$%s:$%s,MATCH($B{r},\'SR_2025-26\'!$%s:$%s,0))=0,"",'
        'INDEX(\'SR_2025-26\'!$%s:$%s,MATCH($B{r},\'SR_2025-26\'!$%s:$%s,0))),"")'
        % (Hn, Hn, Hf, Hf, Hn, Hn, Hf, Hf)).replace("{r}", str(r))
    s7.cell(r, 6).value = ('=IFERROR(IF(N(INDEX(\'SR_2025-26\'!$%s:$%s,MATCH($B{r},\'SR_2025-26\'!$%s:$%s,0)))=0,"",'
        'INDEX(\'SR_2025-26\'!$%s:$%s,MATCH($B{r},\'SR_2025-26\'!$%s:$%s,0))),"")'
        % (Hd, Hd, Hf, Hf, Hd, Hd, Hf, Hf)).replace("{r}", str(r))
    s7.cell(r, 7).value = '=IF($F%d="","",DATE(IF(MONTH($F%d)>=4,YEAR($F%d)+1,YEAR($F%d)),11,30))' % (r, r, r, r)
    s7.cell(r, 8).value = ('=IF($F{0}="","Original ref not captured - cannot test",IF($F{0}>=DATE(2025,4,1),"OK - original in FY 25-26",'
                           'IF($C{0}>$G{0},"BEYOND Sec 34(2) WINDOW","OK - within window")))').format(r)
wb.save(P); print("fixes applied")
