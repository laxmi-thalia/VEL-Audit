"""Add an 'Open Points' sheet to every step workbook: missing inputs, doubts, and decisions
owed, each with owner and where to look. Owners: CLIENT (data to fix/supply), CA (judgment /
confirmation), FILES (a download that has not arrived), INFO (context, nothing owed)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
D = r"C:\Users\pawar\Downloads"
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79")
FILL = {"CLIENT": PatternFill("solid", fgColor="FFC7CE"), "CA": PatternFill("solid", fgColor="FFF2CC"),
        "FILES": PatternFill("solid", fgColor="DDEBF7"), "INFO": PatternFill(fill_type=None)}
POINTS = {
 "VEL_Step1_Sales_Register_FY2025-26_DRAFT.xlsx": [
  ("CLIENT", "3 blocking flag rows keep the gate BLOCKED", "Manual credit note 1200017782 has no Document Date and no IRN; one Aug-25 advance row is tagged B2B. Fix in SR_2025-26 (rows highlighted red) or type YES in the override on 'Step 1 Flags'."),
  ("CA", "152 rows where the rate actually charged differs from the stated rate", "See column Q 'Actual Rate Charged' - values like 17, 19, 27, 33%. Client errors or genuine mixed-rate invoices? Decides whether they become queries."),
  ("CA", "Format has no 'Original Invoice Number' column", "Asked twice; E/F/G fix did not add it. The credit-note time-bar check (Step 7 file) works from ClearTax source meanwhile. Add to the frozen format or leave as-is?"),
  ("CA", "Confirm the new 'Actual Rate Charged (%)' column at position Q", "Inserting it shifted every later column by one letter. Format is declared frozen - please bless the position once."),
  ("CA", "Sample density: 5-23 documents per state", "You said 5-8 per state; the rule takes EVERY invoice >= Rs 5 crore, so Bihar gets 23. Cap at 8 or keep all >= 5 crore?"),
  ("CA", "MOB ADV REV split by sign (positive = received, negative = adjusted)", "This reconciles advances to Rs 10 against GSTR-1 and exactly against the advance ledger - but it is my classification; please confirm."),
  ("CA", "Sign-off on the two normalisation tables", "Document types (15 spellings -> 6 codes) and state names (23 spellings -> 17 states, incl. Madya->Madhya Pradesh) - in VEL_Sales_Column_Map_DRAFT.xlsx sheets 2-3."),
  ("CLIENT", "12 exact duplicate item lines exist in the client's own ClearTax files", "Same doc+HSN+value repeated (e.g. JK2500048298). Looks like legitimate repeated line items - client to confirm."),
  ("FILES", "New-party sampling needs the FY 24-25 sales register", "The comparison last-year-vs-this-year parties (new parties get sampled) cannot run until last year's register is provided."),
  ("INFO", "July source file is legacy .xls", "Register rebuilds depend on a converted copy in a temp folder. Saving '04 July 2025\\Cleartax Sales Reg July 2025.xls' once as .xlsx on the server removes the fragility.")],
 "VEL_Step2_Reco_GSTR1_DRAFT.xlsx": [
  ("CLIENT", "34 credit notes in GSTR-1 only - Rs 1,13,98,101.92 (Bihar, Jan-26, all without IRN)", "Not in the books, not taken in 3B (see Step 3). What are they? Listed on 'Exceptions' with status IN GSTR-1 ONLY."),
  ("CLIENT", "47 nil duplicate credit-note entries in GSTR-1", "Each doc number filed twice at zero value (Apr+May, Nov+Mar) while the real CN is matched. Why re-reported at nil - portal amendments? 'Exceptions' status NIL DUPLICATE."),
  ("CA", "9 Telangana documents never uploaded to GSTR-1 - Rs 54,78,582.96", "Valid IRNs exist; tax was paid via 3B (Step 3 evidence). Treat as GSTR-9 disclosure/correction? Contiguous numbering suggests a failed upload batch."),
  ("INFO", "Advances and B2C reconcile at summary level only", "GSTR-1 has no invoice-level entry for them; see the CA-format sheet columns S-AB and Month-on-Month.")],
 "VEL_Step3_Reco_GSTR1_vs_3B_DRAFT.xlsx": [
  ("CA", "Bihar Aug-25: 3B taxable mis-keyed (5,73,80,517 entered as 57,38,017), tax correct to Rs 0.39", "Same failure as FY 24-25 ('Reporting Error in GSTR-3B'). Confirm treatment and GSTR-9 correction. Evidence: 'Exceptions' section A."),
  ("FILES", "FY 26-27 GSTR-1 exports missing - amendment check half-pending", "Portal Reports\\GSTR-1 has only Apr-25..Mar-26 files. Drop the 26-27 exports in and the 'Amendment Check' sheet method runs. Within-FY amendments: zero."),
  ("INFO", "3.1.1 (e-commerce) taken as nil", "No 9(5) supplies appear in any month's 3B; confirm nothing operates through an e-commerce operator.")],
 "VEL_Step4_Reco_GL_vs_SR_DRAFT.xlsx": [
  ("CA", "Document-type filter used RV / DA / DG / DR / XD", "You quoted 'XT' in session 1; the ledger contains XD, no XT. Please confirm XD was meant."),
  ("CLIENT", "Two invoices in the GL but not in the sales register (Gujarat)", "GU2500027054 (+20,076 CGST) and GU2500027053 (-5,837). Listed on 'Exceptions'."),
  ("CLIENT", "Manual credit note 1200017782 (Rajasthan) is missing from the output GL", "Its CGST 46,406 is the whole Rajasthan difference. Same document that lacks date/IRN in Step 1."),
  ("CA", "Advance components Gujarat (+3,65,858) and MP (+43,888) vs books", "GL advance postings carry text references (MOB.../COLLECTION), so compared state-level. Confirm the residual advance differences are acceptable / explainable."),
  ("INFO", "Register carries doc+R-suffix reversal pairs (e.g. TS2500085472/R, KA2500053609/R) that the GL books net", "Zero net effect; both sides listed on 'Exceptions'."),
  ("INFO", "GL extract runs past the FY (to Aug-26) and has ~10 trailer rows", "Filtered to Year/Month 2025/01-12; trailer rows without a document number dropped.")],
 "VEL_Step5_HSN_Rate_Summary_DRAFT.xlsx": [
  ("CA", "Blank UQC written as '-' on service lines", "Technical necessity (an empty SUMIFS criterion means 'equals 0'). Confirm the presentation is acceptable for Table 17, or specify e.g. 'NA'/'OTH'."),
  ("CA", "Advances excluded from the HSN summary", "They carry no HSN and are not HSN-line supplies; confirm Table 17 treatment."),
  ("INFO", "The 152 off-rate rows from Step 1 flow into these summaries at their ACTUAL charged amounts", "If the CA rules them client errors, the summaries recalcuate once the register is corrected.")],
 "VEL_Step6_Advances_Control_DRAFT.xlsx": [
  ("CA", "Bihar closing is -4.81", "Adjustment consumed the FY 24-25 opening to the last rupee; the 4.81 is rounding between the two years' files. Confirm ignore."),
  ("CA", "MOB ADV REV classification by sign", "Same confirmation as Step 1 - it makes books tie GSTR-1 to Rs 9.49 and the advance ledger exactly."),
  ("CLIENT", "'MOB ADV REC INT' (one Dec-25 row) and 'STO' (in the client's own Dec pivot)", "Confirmed 'advance received' for REC INT; STO (stock transfer orders) - are any inter-GSTIN branch transfers happening, and are they being billed?"),
  ("INFO", "Advance rows have no Place of Supply / recipient GSTIN", "Expected - advances are summary-level; POS left blank in the register.")],
 "VEL_Step7_CN_TimeBar_DRAFT.xlsx": [
  ("CA", "3 credit notes beyond the Section 34(2) window", "MP2500068031/32/33, dated 20.03.2026 against a 28.11.2023 invoice (deadline was 30.11.2024). Taxable -11,182. GST reduction prima facie impermissible - please examine."),
  ("CLIENT", "280 of 288 credit notes carry NO original-invoice reference", "The ClearTax fields exist but are unfilled. Ask the client to supply them; fill columns G/H and the Status column retests each row live."),
  ("INFO", "The check reads ClearTax source directly", "Because the register format has no original-invoice column (see Step 1 open point).")],
 "VEL_Step9_FS_vs_GST_DRAFT.xlsx": [
  ("CA", "Four reconciler columns (G-J) to fill - the Rs 377.16 crore gap", "Essentially all Maharashtra/HO: unbilled/uncertified revenue (booked on work done, GST on billing). Opening/closing unbilled balances are in the Annual Report notes (same folder). Residual column K falls live as you fill."),
  ("CA", "Kerala FS revenue is NEGATIVE (-19,27,636) with zero GST turnover", "Some reversal booked in the financials - what is it?"),
  ("CA", "West Bengal FS revenue NEGATIVE (-24,33,213) vs GST +5,87,281", "Same question."),
  ("CA", "Other income Rs 16.96 crore (not state-allocated)", "Confirm the non-GST nature of its components (interest, forex, etc.) for Table 5."),
  ("INFO", "FS side reproduced from the client's own tagged SAP extract to Rs 0.15", "'FS Revenue Data' sheet; drillable to the 178k-row extract on the server.")],
}
import sys
done = []
for fn, pts in POINTS.items():
    p = D + "\\" + fn
    try:
        f = open(p, "r+b"); f.close()
    except Exception:
        print("LOCKED, skipped:", fn); continue
    wb = openpyxl.load_workbook(p)
    if "Open Points" in wb.sheetnames: del wb["Open Points"]
    ws = wb.create_sheet("Open Points", 0)
    ws["A1"] = "OPEN POINTS — what is missing / doubtful in this file, and who it waits on (as at 21-Aug-2026)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Owner: CLIENT = data to fix or supply | CA = judgment or confirmation | FILES = a download not yet received | INFO = context only."
    ws.append([]); ws.append(["#", "Owner", "Point", "Detail / where to look", "Resolved? (write here)"])
    for c in ws[4]: c.font = HF; c.fill = HB; c.alignment = Alignment(horizontal="center")
    r = 4
    for k, (own, pt, det) in enumerate(pts, 1):
        r += 1
        ws.cell(r, 1, k); ws.cell(r, 2, own); ws.cell(r, 3, pt); ws.cell(r, 4, det)
        for c in range(1, 6):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
            if FILL[own].fill_type: ws.cell(r, c).fill = FILL[own]
    for c_, w in zip("ABCDE", [4, 9, 52, 95, 22]): ws.column_dimensions[c_].width = w
    ws.freeze_panes = "A5"
    wb.save(p); done.append((fn, len(pts)))
for fn, n in done: print("added %2d points -> %s" % (n, fn))
