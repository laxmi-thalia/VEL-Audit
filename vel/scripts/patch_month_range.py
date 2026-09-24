"""Claim-month criteria: the register's '3B Claim  Month' holds the claim DATE (any day of the month), so an exact
'= 1st of month' SUMIFS criterion misses those lines. Use a month range instead: >= first day, < EDATE(first day, 1)."""
import openpyxl, warnings, collections, datetime as dt, io, ast
warnings.filterwarnings("ignore")
wb = openpyxl.load_workbook("master2_snapshot_before_a3.xlsx", read_only=True, data_only=True)
reg = wb["ITC Register 2025-26"]; H = {c.value: i for i, c in enumerate(next(reg.iter_rows(min_row=5, max_row=5))) if c.value}
k = collections.Counter()
for x in reg.iter_rows(min_row=6, values_only=True):
    if x[3] and str(x[H["Category"]]).strip().upper() == "ITC":
        m = x[H["3B Claim  Month"]]; k["day1" if isinstance(m, dt.datetime) and m.day == 1 else ("other day" if isinstance(m, dt.datetime) else type(m).__name__)] += 1
print("ITC lines 3B Claim Month:", dict(k))
b2 = wb["GSTR-2B Apr25-Aug26"]; BH = {c.value: i for i, c in enumerate(next(b2.iter_rows(min_row=2, max_row=2))) if c.value}; k = collections.Counter()
for x in b2.iter_rows(min_row=3, values_only=True):
    if x[0]:
        m = x[BH["3B Claim Month"]]; k["day1" if isinstance(m, dt.datetime) and m.day == 1 else ("other day" if isinstance(m, dt.datetime) else type(m).__name__)] += 1
print("2B Apr25-Aug26 3B Claim Month:", dict(k)); wb.close()

def sub(path, pairs):
    s = io.open(path, encoding="utf-8").read()
    for old, new in pairs:
        assert s.count(old) == 1, (path, old[:60], s.count(old)); s = s.replace(old, new)
    ast.parse(s); io.open(path, "w", encoding="utf-8", newline="\n").write(s); print("patched", path)
    return s

sub("a6_rcm_paid_vs_claimed.py", [
    (r'''ws.Cells(r, 16 + j).Formula = "=SUMIFS(%s,%s,$A%d,%s,$D%d,%s,\"RCM\")" % (G(h), G("VEL GSTIN"), r, G("3B Claim  Month"), r, G("Category"))''',
     r'''ws.Cells(r, 16 + j).Formula = "=SUMIFS(%s,%s,$A%d,%s,\">=\"&$D%d,%s,\"<\"&EDATE($D%d,1),%s,\"RCM\")" % (G(h), G("VEL GSTIN"), r, G("3B Claim  Month"), r, G("3B Claim  Month"), r, G("Category"))'''),
    ("  P:S  ITC claimed per ITC Register RCM lines (3B Claim Month = M+1)",
     "  P:S  ITC claimed per ITC Register RCM lines (3B Claim Month inside M+1 - the register holds the claim DATE, any day)"),
])
s = sub("a7_reco_sheet.py", [
    # register side
    (r'''$%s$%d,$B%d,'ITC Register 2025-26'!$%s$6:$%s$%d,\"ITC\"''',
     r'''$%s$%d,\">=\"&$B%d,'ITC Register 2025-26'!$%s$6:$%s$%d,\"<\"&EDATE($B%d,1),'ITC Register 2025-26'!$%s$6:$%s$%d,\"ITC\"'''),
    ('c("State Name"), c("State Name"), RN, r, c("3B Claim  Month"), c("3B Claim  Month"), RN, r,\n',
     'c("State Name"), c("State Name"), RN, r, c("3B Claim  Month"), c("3B Claim  Month"), RN, r, c("3B Claim  Month"), c("3B Claim  Month"), RN, r,\n'),
    # 2B FY 25-26 side
    (r'''$%s$%d,$B%d,'GSTR-2B Apr25-Aug26'!$%s$3:$%s$%d,\"No\"''',
     r'''$%s$%d,\">=\"&$B%d,'GSTR-2B Apr25-Aug26'!$%s$3:$%s$%d,\"<\"&EDATE($B%d,1),'GSTR-2B Apr25-Aug26'!$%s$3:$%s$%d,\"No\"'''),
    ('bc("Company GSTIN"), bc("Company GSTIN"), NB, r, bc("3B Claim Month"), bc("3B Claim Month"), NB, r,\n',
     'bc("Company GSTIN"), bc("Company GSTIN"), NB, r, bc("3B Claim Month"), bc("3B Claim Month"), NB, r, bc("3B Claim Month"), bc("3B Claim Month"), NB, r,\n'),
    # 2B FY 24-25 side
    (r'''$%s$%d,$B%d,"''' + "\n" + r'''            "'GSTR-2B ITC Data'!$%s$6:$%s$%d,\"N\"''',
     r'''$%s$%d,\">=\"&$B%d,'GSTR-2B ITC Data'!$%s$6:$%s$%d,\"<\"&EDATE($B%d,1),"''' + "\n" + r'''            "'GSTR-2B ITC Data'!$%s$6:$%s$%d,\"N\"'''),
    ('oc("State folder"), oc("State folder"), NO, r, oc("GSTR 3B Month"), oc("GSTR 3B Month"), NO, r,\n',
     'oc("State folder"), oc("State folder"), NO, r, oc("GSTR 3B Month"), oc("GSTR 3B Month"), NO, r, oc("GSTR 3B Month"), oc("GSTR 3B Month"), NO, r,\n'),
])
# smoke-render one formula of each kind against the snapshot (no Excel)
head = s.split("pythoncom.CoInitialize()")[0].replace('shutil.copy(P, "master2_snapshot_before_a7.xlsx")', "")
head = head.replace(r'P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"', 'P = "master2_snapshot_before_a3.xlsx"')
ns = {}; exec(compile(head, "a7head", "exec"), ns)
for f in (ns["reg_f"]("IGST", "2025-26", 8), ns["b2_f"]("IGST (Net)", 8), ns["od_f"]("Integrated Tax(₹)", 8)): print(f)
