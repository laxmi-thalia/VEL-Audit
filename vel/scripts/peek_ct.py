import openpyxl, sys
from openpyxl.utils import get_column_letter
p = "//server/GST FOLDER/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/Clients Data/05.08.2026 Main Data/01 April 2025/Cleartax Sales April 2025.xlsx"
wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"\n=== SHEET {ws.title!r} max_row={ws.max_row} max_col={ws.max_column} ===")
    rows = ws.iter_rows(min_row=1, max_row=3, values_only=True)
    for i, row in enumerate(rows, 1):
        vals = [(get_column_letter(c+1), v) for c, v in enumerate(row) if v is not None]
        if i == 1:
            print(f" HEADERS ({len(vals)} non-empty of {ws.max_column}):")
            for col, v in vals:
                print(f"   {col}: {v!r}")
        else:
            print(f" row{i}: {len(vals)} non-empty cells (data suppressed)")
    break
wb.close()
