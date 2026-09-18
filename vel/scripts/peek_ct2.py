import openpyxl
from openpyxl.utils import get_column_letter
p = "//server/GST FOLDER/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/Clients Data/05.08.2026 Main Data/01 April 2025/Cleartax Sales April 2025.xlsx"
wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
for ws in wb.worksheets:
    print(f"\n=== {ws.title!r} max_row={ws.max_row} max_col={ws.max_column} ===")
    if ws.title == "Summary":
        continue
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        vals = [(get_column_letter(c+1), v) for c, v in enumerate(row) if v is not None]
        print(f"  header row: {len(vals)} non-empty")
        for col, v in vals:
            print(f"   {col}: {v!r}")
        break
wb.close()
