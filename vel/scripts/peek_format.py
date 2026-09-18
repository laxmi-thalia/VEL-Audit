import openpyxl
from openpyxl.utils import get_column_letter
p = "//server/GST FOLDER/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/DPS Workings/1. Sales/Sales Register_FY 2025-26 Format.xlsx"
wb = openpyxl.load_workbook(p, data_only=False)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"\n=== SHEET {ws.title!r}  max_row={ws.max_row} max_col={ws.max_column} ===")
    for r in range(1, min(ws.max_row, 8) + 1):
        cells = [f"{get_column_letter(c)}={ws.cell(r,c).value!r}"
                 for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None]
        if cells:
            print(f" row{r}: " + " | ".join(cells))
