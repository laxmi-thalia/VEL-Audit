"""RCM step 2 (liability leg): 3.1(d) columns onto 3B Data + 'Statewise RCM vs 3B' and
'Month wise RCM vs 3B' sheets in last year's three-block layout, all live SUMIFS."""
import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f = open(P, 'r+b'); f.close()
T = pd.read_pickle("b3d.pkl")
T["g"] = T["gstin"].map(S)
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79"); TOT = Font(bold=True)
NF = "#,##0.00"
wb = openpyxl.load_workbook(P)

# ---- 1. 3B Data: RCM 3.1(d) columns K..N (values = source data)
b3 = wb["3B Data"]
hdr = {S(b3.cell(1, c).value): c for c in range(1, b3.max_column + 1)}
base = b3.max_column
names = ["RCM 3.1(d) Taxable", "RCM 3.1(d) IGST", "RCM 3.1(d) CGST", "RCM 3.1(d) SGST"]
cols = []
for i, nm in enumerate(names):
    c = hdr.get(nm, base + 1 + i)
    x = b3.cell(1, c, nm); x.font = HF; x.fill = HB
    b3.column_dimensions[L(c)].width = 16
    cols.append(c)
MON = {m: pd.to_datetime(m, format="%b-%y").strftime("%b %Y") for m in
       ["Apr-25","May-25","Jun-25","Jul-25","Aug-25","Sep-25","Oct-25","Nov-25","Dec-25","Jan-26","Feb-26","Mar-26"]}
filled = 0
for r in range(2, b3.max_row + 1):
    g = S(b3.cell(r, 2).value); m = S(b3.cell(r, 3).value)
    row = T[T["g"] == g]
    key = MON.get(m)
    for i, meas in enumerate(("taxable", "igst", "cgst", "sgst")):
        v = float(row.iloc[0].get("%s|%s" % (key, meas), 0) or 0) if len(row) and key else 0.0
        cell = b3.cell(r, cols[i]); cell.value = round(v, 2); cell.number_format = NF
    filled += 1
print("3B Data: 3.1(d) cols on %d rows" % filled)

# ---- register ranges (resolve by header, row 5)
rr = wb["RCM Register"]
RH = {S(rr.cell(5, c).value): c for c in range(1, rr.max_column + 1)}
R0, R1 = 6, 6 + 3181 - 1
def RC(name, occurrence=1):
    seen = 0
    for c in range(1, rr.max_column + 1):
        if S(rr.cell(5, c).value) == name:
            seen += 1
            if seen == occurrence: return "'RCM Register'!$%s$%d:$%s$%d" % (L(c), R0, L(c), R1)
    raise KeyError(name)
REG = {"g": RC("MY GSTN"), "m": RC("3B Month"),
       "tax": RC("Taxable Value as per SAP"), "i": RC("IGST AS PER SAP"),
       "c": RC("CGST AS PER SAP"), "s": RC("SGST AS PER SAP")}
n3 = b3.max_row
B3D = {k: "'3B Data'!$%s$2:$%s$%d" % (L(c), L(c), n3) for k, c in zip(("tax","i","c","s"), cols)}
B3G = "'3B Data'!$B$2:$B$%d" % n3
B3M = "'3B Data'!$C$2:$C$%d" % n3
GSTINS = [("01AAECR0503Q1ZM","Jammu & Kashmir"),("03AAECR0503Q1ZI","Punjab"),("06AAECR0503Q1ZC","Haryana"),
("08AAECR0503Q1Z8","Rajasthan"),("09AAECR0503Q1Z6","Uttar Pradesh"),("10AAECR0503Q1ZN","Bihar"),
("12AAECR0503Q1ZJ","Arunachal Pradesh"),("18AAECR0503Q1Z7","Assam"),("19AAECR0503Q1Z5","West Bengal"),
("20AAECR0503Q1ZM","Jharkhand"),("22AAECR0503Q1ZI","Chhattisgarh"),("23AAECR0503Q1ZG","Madhya Pradesh"),
("24AAECR0503Q1ZE","Gujarat"),("27AAECR0503Q1Z8","Maharashtra"),("29AAECR0503Q1Z4","Karnataka"),
("32AAECR0503Q1ZH","Kerala"),("33AAECR0503Q1ZF","Tamil Nadu"),("36AAECR0503Q1Z9","Telangana"),
("37AAECR0503Q1Z7","Andhra Pradesh")]
MONTHS = ["01 Apr 2025","02 May 2025","03 June 2025","04 July 2025","05 Aug 2025","06 Sep 2025",
"07 Oct 2025","08 Nov 2025","09 Dec 2025","10 Jan 2026","11 Feb 2026","12 Mar 2026"]
M2TOK = {"01 Apr 2025":"Apr-25","02 May 2025":"May-25","03 June 2025":"Jun-25","04 July 2025":"Jul-25",
"05 Aug 2025":"Aug-25","06 Sep 2025":"Sep-25","07 Oct 2025":"Oct-25","08 Nov 2025":"Nov-25",
"09 Dec 2025":"Dec-25","10 Jan 2026":"Jan-26","11 Feb 2026":"Feb-26","12 Mar 2026":"Mar-26"}

def blockhead(ws, r):
    for c, t in ((1, "RCM reg"), (7, "GSTR-3B"), (13, "Difference")):
        x = ws.cell(r, c, t); x.font = TOT

# ---- 2. Statewise RCM vs 3B
NM = "Statewise RCM vs 3B"
if NM in wb.sheetnames: del wb[NM]
pos = wb.sheetnames.index("RCM Register") + 1
ws = wb.create_sheet(NM, pos)
ws.row_dimensions[1].height = 21
ws.cell(2, 1, "VIKRAN ENGINEERING LIMITED").font = TOT
ws.cell(3, 1, "Statewise RCM vs GSTR-3B 3.1(d) — live over 'RCM Register' / '3B Data'. DPS Remarks: type here.").font = TOT
blockhead(ws, 4)
H1 = ["State","Sum of Taxable Value","Sum of IGST AS PER SAP","Sum of CGST AS PER SAP","Sum of SGST AS PER SAP","",
      "Row Labels","Taxable Value","Integrated Tax","Central Tax","State Tax","",
      "Row Labels","Taxable Value","Integrated Tax","Central Tax","State Tax","DPS Remarks"]
for c, h in enumerate(H1, 1):
    if h: x = ws.cell(5, c, h); x.font = HF; x.fill = HB
r = 5
for g_, st in GSTINS:
    r += 1
    ws.cell(r, 1, st); ws.cell(r, 19, g_)
    for i, k in enumerate(("tax","i","c","s")):
        ws.cell(r, 2 + i).value = '=SUMIFS(%s,%s,$S%d)' % (REG[k], REG["g"], r)
    ws.cell(r, 7, st)
    for i, k in enumerate(("tax","i","c","s")):
        ws.cell(r, 8 + i).value = '=SUMIFS(%s,%s,$S%d)' % (B3D[k], B3G, r)
    ws.cell(r, 13, st)
    for i in range(4):
        ws.cell(r, 14 + i).value = "=%s%d-%s%d" % (L(2 + i), r, L(8 + i), r)
    for c in list(range(2, 6)) + list(range(8, 12)) + list(range(14, 18)):
        ws.cell(r, c).number_format = NF
r += 1
ws.cell(r, 1, "HOIS / unmapped (ISD)").font = TOT
for i, k in enumerate(("tax","i","c","s")):
    ws.cell(r, 2 + i).value = '=SUMIFS(%s,%s,"")' % (REG[k], REG["g"])
    ws.cell(r, 2 + i).number_format = NF
hois = r
r += 1
for cset, lbl in ((1, "Grand Total"), (7, "Grand Total"), (13, "Grand Total")):
    ws.cell(r, cset, lbl).font = TOT
for c in list(range(2, 6)) + list(range(8, 12)) + list(range(14, 18)):
    ws.cell(r, c).value = "=SUM(%s6:%s%d)" % (L(c), L(c), r - 1)
    ws.cell(r, c).number_format = NF; ws.cell(r, c).font = TOT
ws.column_dimensions["A"].width = 22
for c_ in "BCDE HIJK NOPQ".replace(" ",""): ws.column_dimensions[c_].width = 15
ws.column_dimensions["G"].width = 20; ws.column_dimensions["M"].width = 20; ws.column_dimensions["R"].width = 34
ws.column_dimensions["S"].hidden = True
ws.freeze_panes = "B6"
st_total_row = r

# ---- 3. Month wise RCM vs 3B
NM2 = "Month wise RCM vs 3B"
if NM2 in wb.sheetnames: del wb[NM2]
ws2 = wb.create_sheet(NM2, pos + 1)
ws2.row_dimensions[1].height = 21
ws2.cell(2, 1, "VIKRAN ENGINEERING LIMITED").font = TOT
ws2.cell(3, 1, "RCM vs GSTR 3B — month on month (3.1(d)), live. DPS Remarks: type here.").font = TOT
blockhead(ws2, 4)
H2 = ["Row Labels","Sum of Taxable Value","Sum of IGST AS PER SAP","Sum of CGST AS PER SAP","Sum of SGST AS PER SAP","",
      "Taxable Value","Integrated Tax","Central Tax","State Tax","Total Tax","",
      "Taxable Value","Integrated Tax","Central Tax","State Tax","Total Tax","DPS Remarks"]
for c, h in enumerate(H2, 1):
    if h: x = ws2.cell(5, c, h); x.font = HF; x.fill = HB
r = 5
for m in MONTHS:
    r += 1
    ws2.cell(r, 1, m); ws2.cell(r, 19, M2TOK[m])
    for i, k in enumerate(("tax","i","c","s")):
        ws2.cell(r, 2 + i).value = '=SUMIFS(%s,%s,$A%d)' % (REG[k], REG["m"], r)
    for i, k in enumerate(("tax","i","c","s")):
        ws2.cell(r, 7 + i).value = '=SUMIFS(%s,%s,$S%d)' % (B3D[k], B3M, r)
    ws2.cell(r, 11).value = "=SUM(H%d:J%d)" % (r, r)
    for i in range(4):
        ws2.cell(r, 13 + i).value = "=%s%d-%s%d" % (L(2 + i), r, L(7 + i), r)
    ws2.cell(r, 17).value = "=SUM(N%d:P%d)" % (r, r)
    for c in list(range(2, 6)) + list(range(7, 12)) + list(range(13, 18)):
        ws2.cell(r, c).number_format = NF
r += 1
ws2.cell(r, 1, "Grand Total").font = TOT
for c in list(range(2, 6)) + list(range(7, 12)) + list(range(13, 18)):
    ws2.cell(r, c).value = "=SUM(%s6:%s%d)" % (L(c), L(c), r - 1)
    ws2.cell(r, c).number_format = NF; ws2.cell(r, c).font = TOT
ws2.column_dimensions["A"].width = 16
for c_ in "BCDE GHIJK MNOPQ".replace(" ",""): ws2.column_dimensions[c_].width = 15
ws2.column_dimensions["R"].width = 34
ws2.column_dimensions["S"].hidden = True
ws2.freeze_panes = "B6"

# ---- 4. INDEX rows (RCM arena)
ix = wb["INDEX"]
have = set()
for r in range(5, ix.max_row + 1):
    have.add(S(ix.cell(r, 2).value))
add = [("Statewise RCM vs GSTR-3B", "Statewise RCM vs 3B", NM),
       ("Month wise RCM vs GSTR-3B", "Month wise RCM vs 3B", NM2)]
from openpyxl.styles import Border, Side
thin = Side(style="thin"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
r = ix.max_row
for desc, part, sheet in add:
    if desc in have: continue
    r += 1
    ix.cell(r, 1, "RCM")
    ix.cell(r, 2, desc)
    x = ix.cell(r, 7)
    x.value = '=HYPERLINK("#\'%s\'!A1","%s")' % (sheet, part)
    x.font = Font(color="0563C1", underline="single")
    for c in range(1, 11): ix.cell(r, c).border = BD
wb.save(P)
print("built: %s + %s | statewise total row %d" % (NM, NM2, st_total_row))
