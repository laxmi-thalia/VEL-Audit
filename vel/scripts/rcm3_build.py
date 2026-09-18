"""RCM step 3: Time of Supply (Sec 13(3): 61st day from invoice; both doc & posting bases)
+ interest @18% (Sec 50). Columns appended at the end of RCM Register + summary sheet."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f = open(P, 'r+b'); f.close()
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79"); TOT = Font(bold=True)
NF = "#,##0.00"
wb = openpyxl.load_workbook(P)
rr = wb["RCM Register"]
RH = {}
for c in range(1, rr.max_column + 1):
    h = S(rr.cell(5, c).value)
    if h and h not in RH: RH[h] = c
R0, R1 = 6, 6 + 3181 - 1
cB = L(RH["Final 3B Month"]); cN = L(RH["Posting Date"]); cR = L(RH["Document Date"]); cAV = L(RH["Total GST"])
base = rr.max_column
NEW = ["ToS (Doc date+61)", "ToS (Posting+61)", "ToS (earlier)", "RCM due by (20th after ToS month)",
       "Paid via 3B due date", "Delay days", "Interest @18% p.a.", "ToS Status"]
start = base + 1
for i, h in enumerate(NEW):
    x = rr.cell(5, start + i, h); x.font = HF; x.fill = HB
    rr.column_dimensions[L(start + i)].width = 16
a, b, c_, d, e, g, h_, s_ = [L(start + i) for i in range(8)]
for r in range(R0, R1 + 1):
    rr.cell(r, start + 0).value = '=IF(%s%d="","",%s%d+61)' % (cR, r, cR, r)
    rr.cell(r, start + 1).value = '=IF(%s%d="","",%s%d+61)' % (cN, r, cN, r)
    rr.cell(r, start + 2).value = '=IF(AND(%s%d="",%s%d=""),"",MIN(IF(%s%d="",%s%d,%s%d),IF(%s%d="",%s%d,%s%d)))' % (a, r, b, r, a, r, b, r, a, r, b, r, a, r, b, r)
    rr.cell(r, start + 3).value = '=IF(%s%d="","",DATE(YEAR(%s%d),MONTH(%s%d)+1,20))' % (c_, r, c_, r, c_, r)
    rr.cell(r, start + 4).value = '=IF(%s%d="","",DATE(YEAR(%s%d),MONTH(%s%d)+1,20))' % (cB, r, cB, r, cB, r)
    rr.cell(r, start + 5).value = '=IF(OR(%s%d="",%s%d=""),"",MAX(0,%s%d-%s%d))' % (d, r, e, r, e, r, d, r)
    rr.cell(r, start + 6).value = '=IF(OR(%s%d="",%s%d=0),0,ROUND(%s%d*0.18*%s%d/365,2))' % (g, r, g, r, cAV, r, g, r)
    rr.cell(r, start + 7).value = '=IF(%s%d="","No dates - cannot test",IF(%s%d>0,"LATE - interest applies","OK"))' % (g, r, g, r)
    for i in (0, 1, 2, 3, 4):
        rr.cell(r, start + i).number_format = "DD.MM.YYYY"
    rr.cell(r, start + 5).number_format = "0"
    rr.cell(r, start + 6).number_format = NF
rr.cell(4, start + 6).value = "=SUBTOTAL(9,%s%d:%s%d)" % (h_, R0, h_, R1)
rr.cell(4, start + 6).font = TOT; rr.cell(4, start + 6).number_format = NF
REG = lambda name: "'RCM Register'!$%s$%d:$%s$%d" % (L(RH[name]), R0, L(RH[name]), R1)
INT = "'RCM Register'!$%s$%d:$%s$%d" % (h_, R0, h_, R1)
STA = "'RCM Register'!$%s$%d:$%s$%d" % (s_, R0, s_, R1)
TAXC = "'RCM Register'!$%s$%d:$%s$%d" % (cAV, R0, cAV, R1)

# ---- summary sheet
NM = "RCM ToS & Interest"
if NM in wb.sheetnames: del wb[NM]
pos = wb.sheetnames.index("Month wise RCM vs 3B") + 1
ws = wb.create_sheet(NM, pos)
ws.row_dimensions[1].height = 21
ws.cell(2, 1, "VIKRAN ENGINEERING LIMITED").font = TOT
ws.cell(3, 1, "RCM Time of Supply (Sec 13(3): 61st day from invoice; doc & posting bases, earlier taken) and interest u/s 50 @18% p.a.").font = TOT
ws.cell(4, 1, "Payment dates not available in books - due dates approximated as the 20th following the month (3B cycle). DPS to confirm approach.").font = Font(italic=True, color="808080")
hdrs = ["State", "GSTIN", "Rows", "LATE rows", "Tax on LATE rows", "Interest @18%", "No-date rows", "DPS Remarks"]
for c, h in enumerate(hdrs, 1):
    x = ws.cell(6, c, h); x.font = HF; x.fill = HB
GSTINS = [("01AAECR0503Q1ZM","Jammu & Kashmir"),("03AAECR0503Q1ZI","Punjab"),("06AAECR0503Q1ZC","Haryana"),
("08AAECR0503Q1Z8","Rajasthan"),("09AAECR0503Q1Z6","Uttar Pradesh"),("10AAECR0503Q1ZN","Bihar"),
("12AAECR0503Q1ZJ","Arunachal Pradesh"),("18AAECR0503Q1Z7","Assam"),("19AAECR0503Q1Z5","West Bengal"),
("20AAECR0503Q1ZM","Jharkhand"),("22AAECR0503Q1ZI","Chhattisgarh"),("23AAECR0503Q1ZG","Madhya Pradesh"),
("24AAECR0503Q1ZE","Gujarat"),("27AAECR0503Q1Z8","Maharashtra"),("29AAECR0503Q1Z4","Karnataka"),
("32AAECR0503Q1ZH","Kerala"),("33AAECR0503Q1ZF","Tamil Nadu"),("36AAECR0503Q1Z9","Telangana"),
("37AAECR0503Q1Z7","Andhra Pradesh")]
MYG = REG("MY GSTN")
r = 6
for g_, st in GSTINS:
    r += 1
    ws.cell(r, 1, st); ws.cell(r, 2, g_)
    ws.cell(r, 3).value = '=COUNTIFS(%s,$B%d)' % (MYG, r)
    ws.cell(r, 4).value = '=COUNTIFS(%s,$B%d,%s,"LATE*")' % (MYG, r, STA)
    ws.cell(r, 5).value = '=SUMIFS(%s,%s,$B%d,%s,"LATE*")' % (TAXC, MYG, r, STA)
    ws.cell(r, 6).value = '=SUMIFS(%s,%s,$B%d)' % (INT, MYG, r)
    ws.cell(r, 7).value = '=COUNTIFS(%s,$B%d,%s,"No dates*")' % (MYG, r, STA)
    ws.cell(r, 5).number_format = NF; ws.cell(r, 6).number_format = NF
r += 1
ws.cell(r, 1, "Total").font = TOT
for c in range(3, 8):
    ws.cell(r, c).value = "=SUM(%s7:%s%d)" % (L(c), L(c), r - 1)
    ws.cell(r, c).font = TOT
    if c in (5, 6): ws.cell(r, c).number_format = NF
for c_w, w in zip("ABCDEFGH", [22, 18, 8, 10, 16, 14, 12, 36]):
    ws.column_dimensions[c_w].width = w
ws.freeze_panes = "A7"
# INDEX row
ix = wb["INDEX"]
have = {S(ix.cell(rr2, 2).value) for rr2 in range(5, ix.max_row + 1)}
if "RCM Time of Supply & Interest" not in have:
    thin = Side(style="thin"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
    r2 = ix.max_row + 1
    ix.cell(r2, 1, "RCM"); ix.cell(r2, 2, "RCM Time of Supply & Interest")
    x = ix.cell(r2, 7); x.value = '=HYPERLINK("#\'%s\'!A1","ToS & Interest")' % NM
    x.font = Font(color="0563C1", underline="single")
    for c in range(1, 11): ix.cell(r2, c).border = BD
wb.save(P)
print("step 3 built: 8 register columns (%s..%s) + '%s' summary" % (a, s_, NM))
