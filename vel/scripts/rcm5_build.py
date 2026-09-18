"""RCM step 5: 'Rate-wise summary-RCM' in last year's format (State | MY GSTN | GST Rate |
sums), live SUMIFS over the register. GSTIN-keyed (conso state names have typos); rate
criteria use tolerance bands (float dirt in source rates)."""
import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f = open(P, 'r+b'); f.close()
HF=Font(bold=True,color="FFFFFF"); HB=PatternFill("solid",fgColor="1F4E79"); TOT=Font(bold=True); NF="#,##0.00"
res = pd.read_pickle("rcm_register.pkl")
res["rate"] = pd.to_numeric(res["GST Rate"], errors="coerce").round(2)
GSTIN2ST = {"01AAECR0503Q1ZM":"Jammu & Kashmir","03AAECR0503Q1ZI":"Punjab","06AAECR0503Q1ZC":"Haryana",
"08AAECR0503Q1Z8":"Rajasthan","09AAECR0503Q1Z6":"Uttar Pradesh","10AAECR0503Q1ZN":"Bihar",
"12AAECR0503Q1ZJ":"Arunachal Pradesh","18AAECR0503Q1Z7":"Assam","19AAECR0503Q1Z5":"West Bengal",
"20AAECR0503Q1ZM":"Jharkhand","22AAECR0503Q1ZI":"Chhattisgarh","23AAECR0503Q1ZG":"Madhya Pradesh",
"24AAECR0503Q1ZE":"Gujarat","27AAECR0503Q1Z8":"Maharashtra","29AAECR0503Q1Z4":"Karnataka",
"32AAECR0503Q1ZH":"Kerala","33AAECR0503Q1ZF":"Tamil Nadu","36AAECR0503Q1Z9":"Telangana",
"37AAECR0503Q1Z7":"Andhra Pradesh"}
combos = []
for (g, rate), _grp in res.groupby(["MY GSTN", "rate"], dropna=False):
    combos.append((S(g), None if pd.isna(rate) else float(rate)))
combos = sorted(set(combos), key=lambda t: (GSTIN2ST.get(t[0], "zzz-HOIS"), t[1] if t[1] is not None else 99))
wb = openpyxl.load_workbook(P)
rr = wb["RCM Register"]
RH = {}
for c in range(1, rr.max_column+1):
    h = S(rr.cell(5,c).value)
    if h and h not in RH: RH[h] = c
R0, R1 = 6, 6+3181-1
REG = lambda n: "'RCM Register'!$%s$%d:$%s$%d" % (L(RH[n]), R0, L(RH[n]), R1)
NM = "Rate-wise summary-RCM"
if NM in wb.sheetnames: del wb[NM]
pos = wb.sheetnames.index("RCM GL") + 1
ws = wb.create_sheet(NM, pos)
ws.row_dimensions[1].height = 21
ws.cell(2,1,"Vikran Engineering Limited").font = TOT
ws.cell(3,1,"Rate-wise Summary").font = TOT
ws.cell(4,1,"From RCM Register - FY 2025-26 (live)").font = Font(italic=True)
hdrs = ["State","MY GSTN","GST Rate","Sum of Taxable Value as per SAP","Sum of IGST AS PER SAP","Sum of CGST AS PER SAP","Sum of SGST AS PER SAP"]
for c,h in enumerate(hdrs,1):
    x=ws.cell(5,c,h); x.font=HF; x.fill=HB
r = 5
for g, rate in combos:
    r += 1
    st = GSTIN2ST.get(g, "HOIS / unmapped (ISD)" if not g else g)
    ws.cell(r,1,st); ws.cell(r,2,g if g else "-")
    ws.cell(r,3, rate if rate is not None else "blank")
    gc = '"%s"' % g if g else '""'
    if rate is None:
        rc1, rc2 = '"="', None
    for i, meas in enumerate(("Taxable Value as per SAP","IGST AS PER SAP","CGST AS PER SAP","SGST AS PER SAP")):
        if rate is not None:
            ws.cell(r,4+i).value = '=SUMIFS(%s,%s,%s,%s,">=%s",%s,"<=%s")' % (
                REG(meas), REG("MY GSTN"), gc, REG("GST Rate"), rate-0.01, REG("GST Rate"), rate+0.01)
        else:
            ws.cell(r,4+i).value = '=SUMIFS(%s,%s,%s,%s,"=")' % (REG(meas), REG("MY GSTN"), gc, REG("GST Rate"))
        ws.cell(r,4+i).number_format = NF
r += 1
ws.cell(r,1,"Grand Total").font = TOT
for c in range(4,8):
    ws.cell(r,c).value = "=SUM(%s6:%s%d)" % (L(c),L(c),r-1)
    ws.cell(r,c).number_format = NF; ws.cell(r,c).font = TOT
for c_,w in zip("ABCDEFG",[22,18,9,22,18,18,18]): ws.column_dimensions[c_].width = w
ws.freeze_panes = "A6"
ix = wb["INDEX"]
have = {S(ix.cell(rr2,2).value) for rr2 in range(5, ix.max_row+1)}
if "RCM Rate-wise Summary" not in have:
    thin=Side(style="thin"); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
    r2 = ix.max_row+1
    ix.cell(r2,1,"RCM"); ix.cell(r2,2,"RCM Rate-wise Summary")
    x=ix.cell(r2,7); x.value='=HYPERLINK("#\'%s\'!A1","Rate-wise Summary")' % NM
    x.font=Font(color="0563C1",underline="single")
    for c in range(1,11): ix.cell(r2,c).border=BD
wb.save(P)
print("built %s: %d rows" % (NM, len(combos)))
