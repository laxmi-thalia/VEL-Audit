"""RCM step 6: RCM Register vs GSTR-2B (Reverse Charge items). Match on
(supplier GSTIN, invoice no), fallback (supplier GSTIN, amount). Vice-versa both sides:
'RCM vs 2B' sheet holds the 688 2B items with status; register gains 'Found in 2B'."""
import pandas as pd, re, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None or (isinstance(v,float) and pd.isna(v)) else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f=open(P,'r+b'); f.close()
b2 = pd.read_pickle("b2rcm.pkl")
res = pd.read_pickle("rcm_register.pkl")
def ninv(v):
    s = S(v).upper()
    return re.sub(r"[^A-Z0-9]", "", s)
res["vg"] = res["GSTN"].map(lambda v: S(v).upper())
res["inv"] = res["Inv. No."].map(ninv)
res["tax1"] = pd.to_numeric(res["Taxable Value as per SAP"], errors="coerce").fillna(0)
reg_doc = res[res["vg"]!=""].groupby(["vg","inv"]).agg(rtax=("tax1","sum")).reset_index()
key1 = {(r.vg, r.inv): r.rtax for r in reg_doc.itertuples() if r.inv}
amt_key = {}
for r in res[res["vg"]!=""].itertuples():
    k = (r.vg, round(abs(r.tax1)))
    amt_key.setdefault(k, 0); amt_key[k]+=1
b2["vg"] = b2["Supplier GSTIN"].map(lambda v: S(v).upper())
b2["inv"] = b2["Doc No"].map(ninv)
def match(row):
    k1 = (row["vg"], row["inv"])
    if k1 in key1:
        d = row["Taxable Value (Net)"] - key1[k1]
        return ("Matched with RCM Register", "" if abs(d)<=1 else "Taxable differs by %.2f" % d)
    k2 = (row["vg"], round(abs(row["Taxable Value (Net)"])))
    if k2 in amt_key: return ("Matched by vendor+amount", "invoice number differs between 2B and books")
    return ("NOT IN RCM REGISTER", "")
mm = b2.apply(match, axis=1, result_type="expand")
b2["Match Status"] = mm[0]; b2["Match Remarks"] = mm[1]
print("2B-side statuses:", dict(b2["Match Status"].value_counts()))
# register-side stamp
b2_keys = set(zip(b2["vg"], b2["inv"]))
b2_amt = set(zip(b2["vg"], b2["Taxable Value (Net)"].map(lambda v: round(abs(v)))))
def reg_status(r):
    if r.vg == "": return "URD/self-invoice - not in 2B by design"
    if (r.vg, r.inv) in b2_keys and r.inv: return "Matched in 2B"
    if (r.vg, round(abs(r.tax1))) in b2_amt: return "Matched in 2B (by vendor+amount)"
    return "Not found in 2B (registered vendor)"
res["b2stat"] = [reg_status(r) for r in res.itertuples()]
print("register-side:", dict(res["b2stat"].value_counts()))
res[["b2stat"]].to_pickle("rcm_b2_stamps.pkl")

HF=Font(bold=True,color="FFFFFF"); HB=PatternFill("solid",fgColor="1F4E79"); TOT=Font(bold=True); NF="#,##0.00"
wb = openpyxl.load_workbook(P)
NM = "RCM vs 2B"
if NM in wb.sheetnames: del wb[NM]
pos = wb.sheetnames.index("Rate-wise summary-RCM") + 1
ws = wb.create_sheet(NM, pos)
ws.row_dimensions[1].height = 21
ws.cell(2,1,"VIKRAN ENGINEERING LIMITED").font = TOT
ws.cell(3,1,"RCM Register vs GSTR-2B (Reverse Charge = Yes line items). 2B carries only REGISTERED suppliers' RCM invoices - URD/self-invoiced RCM never appears in 2B.").font = TOT
COLS = ["Company GSTIN","Tax Period","Doc Type","Doc No","Doc Date","Supplier GSTIN","Supplier Name",
        "Supplier State","Place of Supply","Taxable Value (Net)","IGST (Net)","CGST (Net)","SGST (Net)",
        "ITC Eligible","Source","Match Status","Match Remarks"]
for c,h in enumerate(COLS,1):
    x=ws.cell(5,c,h); x.font=HF; x.fill=HB
r=5
for row in b2.itertuples():
    r+=1
    vals=[row._1 if hasattr(row,'_1') else None]
    data=[S(getattr(row,"_"+str(list(b2.columns).index(cn)+1),None)) for cn in []]  # unused
    d = {cn: getattr(row, cn.replace(" ","_").replace("(","").replace(")","").replace("/","_"), None) for cn in []}
    src = b2.loc[row.Index]
    out=[src["Company GSTIN"],src["Tax Period"],src["Doc Type"],src["Doc No"],src["Doc Date"],src["Supplier GSTIN"],
         src["Supplier Name"],src["Supplier State"],src["Place of Supply"],src["Taxable Value (Net)"],src["IGST (Net)"],
         src["CGST (Net)"],src["SGST (Net)"],src["ITC Eligible"],src["Source"],src["Match Status"],src["Match Remarks"]]
    for c,v in enumerate(out,1):
        if isinstance(v,pd.Timestamp): v=v.to_pydatetime()
        if isinstance(v,float) and pd.isna(v): v=None
        cell=ws.cell(r,c,v)
        if c in (10,11,12,13): cell.number_format=NF
        if c in (2,5): cell.number_format="DD.MM.YYYY"
last=r
for c in (10,11,12,13):
    ws.cell(4,c).value="=SUBTOTAL(9,%s6:%s%d)"%(L(c),L(c),last)
    ws.cell(4,c).number_format=NF; ws.cell(4,c).font=TOT
ws.auto_filter.ref="A5:Q%d"%last
ws.freeze_panes="A6"
for c_,w in zip(range(1,18),[17,11,12,18,11,17,26,18,18,14,12,12,12,11,10,26,30]): ws.column_dimensions[L(c_)].width=w
# register stamp column
rr = wb["RCM Register"]
RH={}
for c in range(1, rr.max_column+1):
    h=S(rr.cell(5,c).value)
    if h and h not in RH: RH[h]=c
cB2 = RH.get("Found in 2B", max(RH.values())+1)
x=rr.cell(5,cB2,"Found in 2B"); x.font=HF; x.fill=HB
rr.column_dimensions[L(cB2)].width=30
stamps=pd.read_pickle("rcm_b2_stamps.pkl")
for i,v in enumerate(stamps["b2stat"]):
    rr.cell(6+i,cB2).value=v
# INDEX
ix=wb["INDEX"]
have={S(ix.cell(r2,2).value) for r2 in range(5,ix.max_row+1)}
if "RCM Register vs GSTR-2B" not in have:
    thin=Side(style="thin"); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
    r2=ix.max_row+1
    ix.cell(r2,1,"RCM"); ix.cell(r2,2,"RCM Register vs GSTR-2B")
    x=ix.cell(r2,7); x.value='=HYPERLINK("#\'%s\'!A1","RCM vs 2B")'%NM
    x.font=Font(color="0563C1",underline="single")
    for c in range(1,11): ix.cell(r2,c).border=BD
wb.save(P)
print("RCM vs 2B sheet: %d rows | register stamp col %s" % (last-5, L(cB2)))
