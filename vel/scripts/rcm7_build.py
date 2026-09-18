"""RCM step 7: Trial Balance scrutiny for RCM applicability. Aggregate expense accounts,
flag RCM-indicative heads by word-boundary keyword rules, cross-check against the register's
offsetting accounts. Every flag is an AI proposal for CA confirmation."""
import openpyxl, re, pandas as pd
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f=open(P,'r+b'); f.close()
src = openpyxl.load_workbook("vel_fs.xlsx", read_only=True, data_only=True)["Trial Balance"]
agg = {}
for i, r in enumerate(src.iter_rows(values_only=True)):
    if i == 0 or len(r) < 11: continue
    a = S(r[0])
    if not re.match(r"^4\d{9}$", a): continue
    nm = S(r[9]); amt = pd.to_numeric(r[10], errors="coerce")
    k = (a, nm)
    agg[k] = agg.get(k, 0.0) + (0 if pd.isna(amt) else float(amt))
print("expense accounts:", len(agg))
RULES = [
    ("Rent (residential dwelling 9(3)/9(4)?)", r"\brent\b|guest\s*house"),
    ("Goods Transport Agency", r"\btransport\b|\bfreight\b|\bgta\b|\bcarriage\b|\blr\b"),
    ("Legal services (advocate)", r"\blegal\b|\badvocate\b"),
    ("Security services", r"\bsecurity\b"),
    ("Director services (non-salary)", r"\bdirector\b"),
    ("Govt royalty / statutory", r"\broyalt\w*|statutory\s*deduction|licen[cs]e\s*fee|seigniorage"),
    ("Sponsorship", r"\bsponsor"),
    ("Arbitral tribunal", r"\barbitr"),
    ("Ocean/sea freight (import)", r"ocean|sea\s*freight"),
    ("Possible import of services", r"\bforeign\b|\bimport\b"),
]
res = pd.read_pickle("rcm_register.pkl")
reg_offs = set(str(int(float(v))) for v in res["Offsetting Account"].dropna().astype(str) if S(v))
flagged = []
for (a, nm), amt in sorted(agg.items()):
    low = nm.lower()
    hits = [lbl for lbl, pat in RULES if re.search(pat, low)]
    if not hits: continue
    inreg = "Yes - offsetting acct in register" if a in reg_offs else "NOT in RCM register"
    flagged.append((a, nm, round(amt, 2), "; ".join(hits), inreg))
print("flagged heads:", len(flagged), "| of which not in register:", sum(1 for x in flagged if "NOT" in x[4]))
HF=Font(bold=True,color="FFFFFF"); HB=PatternFill("solid",fgColor="1F4E79"); TOT=Font(bold=True); NF="#,##0.00"
wb = openpyxl.load_workbook(P)
NM = "TB Scrutiny-RCM"
if NM in wb.sheetnames: del wb[NM]
pos = wb.sheetnames.index("RCM vs 2B") + 1
ws = wb.create_sheet(NM, pos)
ws.row_dimensions[1].height = 21
ws.cell(2,1,"VIKRAN ENGINEERING LIMITED").font = TOT
ws.cell(3,1,"Trial Balance scrutiny for RCM applicability - FY 2025-26 expense accounts (4-series), keyword-flagged. ALL flags are AI-proposed: kindly confirm & check (CA)." ).font = TOT
ws.cell(4,1,"Scanned %d expense accounts from 'Trial Balance' (FS with Notes file). Rules: rent/guest house, transport/freight/GTA, legal/advocate, security, director, royalty/statutory/licence, sponsorship, arbitral, ocean freight, foreign/import." % len(agg)).font = Font(italic=True, color="808080")
hdrs = ["Account","Account Name","FY 2025-26 Amount (TB)","RCM indicator (rule hit)","In RCM register?","DPS Remarks"]
for c,h in enumerate(hdrs,1):
    x=ws.cell(6,c,h); x.font=HF; x.fill=HB
r=6
for a,nm,amt,hits,inreg in flagged:
    r+=1
    ws.cell(r,1,a); ws.cell(r,2,nm); ws.cell(r,3,amt); ws.cell(r,4,hits); ws.cell(r,5,inreg)
    ws.cell(r,3).number_format=NF
    if "NOT" in inreg:
        for c in range(1,6): ws.cell(r,c).fill = PatternFill("solid", fgColor="FFF2CC")
for c_,w in zip("ABCDEF",[13,40,18,34,28,36]): ws.column_dimensions[c_].width=w
ws.freeze_panes="A7"
ws.auto_filter.ref="A6:F%d"%r
ix=wb["INDEX"]
have={S(ix.cell(r2,2).value) for r2 in range(5,ix.max_row+1)}
if "TB Scrutiny for RCM applicability" not in have:
    thin=Side(style="thin"); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
    r2=ix.max_row+1
    ix.cell(r2,1,"RCM"); ix.cell(r2,2,"TB Scrutiny for RCM applicability")
    x=ix.cell(r2,7); x.value='=HYPERLINK("#\'%s\'!A1","TB Scrutiny")'%NM
    x.font=Font(color="0563C1",underline="single")
    for c in range(1,11): ix.cell(r2,c).border=BD
wb.save(P)
for x in flagged:
    if "NOT" in x[4]: print("  REVIEW:", x[0], x[1][:36], format(x[2],",.2f"), "|", x[3][:40])
print("sheet built: %d flagged rows" % len(flagged))
