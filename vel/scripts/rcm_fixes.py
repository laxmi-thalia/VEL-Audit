"""RCM batch on MASTER (2):
 (a) RCM Register: 'Final 3B Month' (B) -> live =DATE(YEAR(A),MONTH(A),1) (source text '03 June 2025' failed the
     %b parse -> 304 blanks; tz shift put Apr rows at 31.03.2025). 'POS Check' (BG) -> =$BF=$BE (TRUE/FALSE), all rows.
 (b) 'Month wise RCM vs 3B' rebuilt as State | GSTIN | Month grid (S3 Month-on-Month style, autofilter).
 (c) 'TB Scrutiny-RCM' rebuilt with the FULL trial balance (every account, aggregated across PC/period),
     RCM keyword indicator + live 'In RCM register?' (COUNTIF over register Offsetting Account)."""
import re, warnings, datetime as dt, collections
from copy import copy
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
F = "//192.168.1.69/gst folder/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/Clients Data/Financials/VEL Standalone FS Mar-26 with Notes.xlsx"
f = open(P, "r+b"); f.close()
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79"); TOT = Font(bold=True); NF = "#,##0.00"
AMB = PatternFill("solid", fgColor="FFF2CC"); thin = Side(style="thin", color="B0B0B0"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = openpyxl.load_workbook(P)

# ---------------- (a) RCM Register ----------------
rr = wb["RCM Register"]
RH = {}
for c in range(1, rr.max_column + 1):
    h = S(rr.cell(5, c).value)
    if h and h not in RH: RH[h] = c
R0 = 6
R1 = max(r for r in range(R0, rr.max_row + 1) if rr.cell(r, RH["MY GSTN"]).value or rr.cell(r, RH["Business place"]).value)
print("register rows %d..%d" % (R0, R1))
cA, cB, cBE, cBF, cBG = (L(RH[k]) for k in ("3B Month", "Final 3B Month", "As per State", "As per Amounts", "POS Check"))
assert (cA, cB, cBG) == ("A", "B", "BG"), (cA, cB, cBG)
for r in range(R0, R1 + 1):
    rr.cell(r, RH["Final 3B Month"]).value = '=IF(%s%d="","",DATE(YEAR(%s%d),MONTH(%s%d),1))' % (cA, r, cA, r, cA, r)
    rr.cell(r, RH["Final 3B Month"]).number_format = "DD.MM.YYYY"
    rr.cell(r, RH["POS Check"]).value = "=$%s%d=$%s%d" % (cBF, r, cBE, r)
print("Final 3B Month + POS Check formulas written")

# ---------------- (b) Month wise RCM vs 3B ----------------
NM = "Month wise RCM vs 3B"
pos = wb.sheetnames.index(NM); del wb[NM]
ws = wb.create_sheet(NM, pos)
b3 = wb["3B Data"]
trip = []   # (state, gstin, month) in 3B Data order
for r in range(3, b3.max_row + 1):
    if b3.cell(r, 2).value: trip.append((b3.cell(r, 1).value, b3.cell(r, 2).value, b3.cell(r, 3).value))
N3 = 2 + len(trip)
B3H = {S(b3.cell(2, c).value): c for c in range(1, b3.max_column + 1)}
def b3col(prefix, tail): return L(next(c for h, c in B3H.items() if h.startswith(prefix) and h.endswith(tail)))
b3_tax, b3_i, b3_c, b3_s = (b3col("3.1(d)", t) for t in ("Taxable", "IGST", "CGST", "SGST"))
RG = lambda col: "'RCM Register'!$%s$%d:$%s$%d" % (col, R0, col, R1)
B3 = lambda col: "'3B Data'!$%s$3:$%s$%d" % (col, col, N3)
cAR, cAS, cAT, cAU, cD = (L(RH[k]) for k in ("Taxable Value as per SAP", "IGST AS PER SAP", "CGST AS PER SAP", "SGST AS PER SAP", "MY GSTN"))
ws.row_dimensions[1].height = 21
ws.cell(2, 1, "VIKRAN ENGINEERING LIMITED").font = TOT
ws.cell(3, 1, "RCM Register vs GSTR-3B 3.1(d) - State | Month (live). Filter column A for a state. Register side keyed on 'Final 3B Month'; 3B side on '3B Data'. DPS Remarks: type in the last column.").font = Font(italic=True, color="808080")
hdr = ["State", "GSTIN", "Month", "RCM Reg Taxable", "RCM Reg IGST", "RCM Reg CGST", "RCM Reg SGST", "RCM Reg Total Tax",
       "3B 3.1(d) Taxable", "3B 3.1(d) IGST", "3B 3.1(d) CGST", "3B 3.1(d) SGST", "3B 3.1(d) Total Tax",
       "Diff Taxable", "Diff IGST", "Diff CGST", "Diff SGST", "Diff Total Tax", "DPS Remarks (type here)", "Month start (helper)"]
for c, h in enumerate(hdr, 1):
    x = ws.cell(4, c, h); x.font = HF; x.fill = HB; x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
ws.row_dimensions[4].height = 32
MON = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26"]
r = 4
for st, g, m in trip:
    r += 1
    mi = MON.index(S(m)); mdate = dt.datetime(2025 + (1 if mi >= 9 else 0), (mi + 4 - 1) % 12 + 1, 1)
    ws.cell(r, 1, st); ws.cell(r, 2, g); ws.cell(r, 3, m); ws.cell(r, 20, mdate).number_format = "DD.MM.YYYY"
    for j, col in enumerate((cAR, cAS, cAT, cAU)):
        ws.cell(r, 4 + j, "=SUMIFS(%s,%s,$B%d,%s,$T%d)" % (RG(col), RG(cD), r, RG(cB), r))
    ws.cell(r, 8, "=E%d+F%d+G%d" % (r, r, r))
    for j, col in enumerate((b3_tax, b3_i, b3_c, b3_s)):
        ws.cell(r, 9 + j, "=SUMIFS(%s,%s,$B%d,%s,$C%d)" % (B3(col), B3("B"), r, B3("C"), r))
    ws.cell(r, 13, "=J%d+K%d+L%d" % (r, r, r))
    for j in range(5): ws.cell(r, 14 + j, "=%s%d-%s%d" % (L(4 + j), r, L(9 + j), r))
    for c in range(4, 19): ws.cell(r, c).number_format = NF
    ws.cell(r, 19).fill = AMB
RN = r
r += 1
ws.cell(r, 1, "Grand Total").font = TOT
for c in range(4, 19):
    ws.cell(r, c, "=SUM(%s5:%s%d)" % (L(c), L(c), RN)); ws.cell(r, c).font = TOT; ws.cell(r, c).number_format = NF
for c, w in zip(range(1, 21), [20, 18, 9] + [15] * 15 + [36, 12]):
    ws.column_dimensions[L(c)].width = w
ws.column_dimensions["T"].hidden = True
ws.freeze_panes = "D5"; ws.auto_filter.ref = "A4:S%d" % RN
print("Month wise RCM vs 3B: %d rows" % len(trip))

# ---------------- (c) TB Scrutiny-RCM: full TB ----------------
src = openpyxl.load_workbook(F, read_only=True, data_only=True)["Trial Balance"]
agg = collections.OrderedDict()
for i, row in enumerate(src.iter_rows(values_only=True)):
    if i == 0 or not row or row[0] is None: continue
    a = S(row[0])
    if not re.match(r"^\d{10}$", a): continue
    nm = S(row[9]); amt = row[10]
    amt = float(amt) if isinstance(amt, (int, float)) else 0.0
    k = (a, nm); agg[k] = agg.get(k, 0.0) + amt
print("TB accounts:", len(agg))
RULES = [("Rent (residential dwelling 9(3)/9(4)?)", r"\brent\b|guest\s*house"),
         ("Goods Transport Agency", r"\btransport\b|\bfreight\b|\bgta\b|\bcarriage\b|\blr\b"),
         ("Legal services (advocate)", r"\blegal\b|\badvocate\b"), ("Security services", r"\bsecurity\b"),
         ("Director services (non-salary)", r"\bdirector\b"),
         ("Govt royalty / statutory", r"\broyalt\w*|statutory\s*deduction|licen[cs]e\s*fee|seigniorage"),
         ("Sponsorship", r"\bsponsor"), ("Arbitral tribunal", r"\barbitr"), ("Ocean/sea freight (import)", r"ocean|sea\s*freight"),
         ("Possible import of services", r"\bforeign\b|\bimport\b")]
NM2 = "TB Scrutiny-RCM"
pos = wb.sheetnames.index(NM2); del wb[NM2]
ws = wb.create_sheet(NM2, pos)
ws.row_dimensions[1].height = 21
ws.cell(2, 1, "VIKRAN ENGINEERING LIMITED").font = TOT
ws.cell(3, 1, "Trial Balance scrutiny for RCM applicability - FY 2025-26 FULL trial balance (every account, aggregated across profit centres / periods). RCM indicator is keyword-based (AI-proposed): kindly confirm & check (CA).").font = TOT
ws.cell(4, 1, "Source: 'Trial Balance' sheet of VEL Standalone FS Mar-26 with Notes.xlsx (%d accounts). 'In RCM register?' is live: COUNTIF over the register's Offsetting Account column. Rows with a rule hit but no register entry are shaded." % len(agg)).font = Font(italic=True, color="808080")
hdrs = ["Account", "Account Name", "FY 2025-26 Amount (TB)", "Expense (4-series)?", "RCM indicator (rule hit)", "In RCM register?", "DPS Remarks"]
for c, h in enumerate(hdrs, 1):
    x = ws.cell(6, c, h); x.font = HF; x.fill = HB
OFF = "'RCM Register'!$%s$%d:$%s$%d" % (L(RH["Offsetting Account"]), R0, L(RH["Offsetting Account"]), R1)
r = 6; nflag = 0
for (a, nm), amt in agg.items():
    r += 1
    hits = "; ".join(lbl for lbl, pat in RULES if re.search(pat, nm.lower()))
    ws.cell(r, 1, a); ws.cell(r, 2, nm); ws.cell(r, 3, round(amt, 2)).number_format = NF
    ws.cell(r, 4, '=IF(LEFT(A%d,1)="4","Y","N")' % r)
    ws.cell(r, 5, hits)
    ws.cell(r, 6, '=IF(COUNTIF(%s,A%d)+COUNTIF(%s,VALUE(A%d))>0,"Yes - offsetting acct in register","NOT in RCM register")' % (OFF, r, OFF, r))
    if hits: nflag += 1
    for c in range(1, 8): ws.cell(r, c).border = BD
RN2 = r
r += 1
ws.cell(r, 1, "Total").font = TOT; ws.cell(r, 3, "=SUBTOTAL(9,C7:C%d)" % RN2).number_format = NF; ws.cell(r, 3).font = TOT
# conditional shading: rule hit AND not in register (live)
from openpyxl.formatting.rule import FormulaRule
ws.conditional_formatting.add("A7:G%d" % RN2, FormulaRule(formula=['AND($E7<>"",LEFT($F7,3)="NOT")'], fill=AMB))
for c_, w in zip("ABCDEFG", [13, 40, 18, 12, 34, 30, 36]): ws.column_dimensions[c_].width = w
ws.freeze_panes = "A7"; ws.auto_filter.ref = "A6:G%d" % RN2
print("TB Scrutiny: %d accounts, %d keyword hits" % (len(agg), nflag))
# INDEX description refresh
ix = wb["INDEX"]
for r2 in range(2, ix.max_row + 1):
    v = S(ix.cell(r2, 2).value)
    if v == "TB Scrutiny for RCM applicability": ix.cell(r2, 2).value = "TB Scrutiny for RCM applicability (full trial balance)"
wb.save(P); print("saved")
