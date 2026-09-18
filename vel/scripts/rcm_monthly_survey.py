"""Read-only survey: every monthly 'GSTR-3B <State> <Mon>.xlsx' -> sheet 'RCM' under Main Data\<month>\<state>.
Header drift, row counts, SAP taxable/tax per state-month; compare to the current register and 3B 3.1(d)."""
import os, re, warnings, pickle, collections, datetime as dt
warnings.filterwarnings("ignore")
import openpyxl
def S(v): return "" if v is None else str(v).strip()
M = "//192.168.1.69/gst folder/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/Clients Data/05.08.2026 Main Data/"
MONTHS = ["01 April 2025", "02 May 2025", "03 June 2025", "04 July 2025", "05 Aug 2025", "06 Sep 2025", "07 Oct 2025", "08 Nov 2025", "09 Dec 2025", "10 Jan 2026", "11 Feb 2026", "12 Mar 2026"]
MLBL = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26"]
STATES = {"andhra pradesh": "37AAECR0503Q1Z7", "arunachal pradesh": "12AAECR0503Q1ZJ", "assam": "18AAECR0503Q1Z7", "bihar": "10AAECR0503Q1ZN",
          "chattisgarh": "22AAECR0503Q1ZI", "chhattisgarh": "22AAECR0503Q1ZI", "gujrat": "24AAECR0503Q1ZE", "gujarat": "24AAECR0503Q1ZE",
          "haryana": "06AAECR0503Q1ZC", "jammu & kashmir": "01AAECR0503Q1ZM", "jharkhand": "20AAECR0503Q1ZM", "karnataka": "29AAECR0503Q1Z4",
          "kerala": "32AAECR0503Q1ZH", "madhya pradesh": "23AAECR0503Q1ZG", "maharashtra": "27AAECR0503Q1Z8", "punjab": "03AAECR0503Q1ZI",
          "rajasthan": "08AAECR0503Q1Z8", "tamil nadu": "33AAECR0503Q1ZF", "telangana": "36AAECR0503Q1Z9", "uttar pradesh": "09AAECR0503Q1Z6",
          "west bengal": "19AAECR0503Q1Z5", "isd": "HOIS"}
EXP_H = None
out = {}          # (gstin, month) -> dict(file, header, rows(list))
drift = []; missing = []
for mi, mf in enumerate(MONTHS):
    for d in sorted(os.listdir(M + mf)):
        key = d.strip().lower()
        if key not in STATES: continue
        g = STATES[key]
        folder = M + mf + "/" + d + "/"
        cands = [x for x in os.listdir(folder) if x.lower().startswith("gstr-3b") and x.lower().endswith(".xlsx") and not x.startswith("~$")]
        if not cands: missing.append((MLBL[mi], d, "no GSTR-3B xlsx")); continue
        fn = cands[0]
        try:
            wb = openpyxl.load_workbook(folder + fn, read_only=True, data_only=True)
        except Exception as e:
            missing.append((MLBL[mi], d, "open failed: %s" % str(e)[:60])); continue
        if "RCM" not in wb.sheetnames:
            missing.append((MLBL[mi], d, "no RCM sheet; sheets=%s" % wb.sheetnames[:8])); wb.close(); continue
        ws = wb["RCM"]
        rows = list(ws.iter_rows(values_only=True))
        # header row = first row whose first cell is 'Business place'
        hr = next((i for i, r in enumerate(rows[:10]) if r and S(r[0]).lower() == "business place"), None)
        if hr is None:
            missing.append((MLBL[mi], d, "header 'Business place' not found in first 10 rows")); wb.close(); continue
        H = [S(h) for h in rows[hr]]
        while H and H[-1] == "": H.pop()
        if EXP_H is None: EXP_H = H
        elif H != EXP_H: drift.append((MLBL[mi], d, [(i, a, b) for i, (a, b) in enumerate(zip(H, EXP_H)) if a != b][:5], len(H), len(EXP_H)))
        iDoc = H.index("Document Number") if "Document Number" in H else 10
        data = [r[:len(H)] for r in rows[hr + 1:] if r and r[iDoc] not in (None, "") and S(r[0]) != ""]
        out[(g, MLBL[mi])] = {"file": mf + "/" + d + "/" + fn, "header": H, "rows": data}
        wb.close()
    print("month %s done: %d state files" % (MLBL[mi], sum(1 for k in out if k[1] == MLBL[mi])), flush=True)
print("\nEXPECTED HEADER (%d cols):" % len(EXP_H), EXP_H)
print("\nHEADER DRIFT:", len(drift))
for x in drift: print("  ", x)
print("\nMISSING / UNREADABLE:", len(missing))
for x in missing: print("  ", x)
# totals per state-month
iT = EXP_H.index("Taxable Value"); iI = EXP_H.index("IGST"); iC = EXP_H.index("CGST"); iS = EXP_H.index("SGST")
def num(v):
    try: return float(v or 0)
    except Exception: return 0.0
summary = {}
for (g, m), rec in out.items():
    rs = rec["rows"]
    summary[(g, m)] = (len(rs), sum(num(r[iT]) for r in rs), sum(num(r[iI]) for r in rs), sum(num(r[iC]) for r in rs), sum(num(r[iS]) for r in rs))
pickle.dump({"header": EXP_H, "data": out, "summary": summary, "drift": drift, "missing": missing}, open("rcm_monthly.pkl", "wb"))
# compare with register + 3B
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
wb = openpyxl.load_workbook(P, read_only=True, data_only=True)
reg = collections.defaultdict(lambda: [0, 0.0])
for r in wb["RCM Register"].iter_rows(min_row=6, values_only=True):
    if not (r[3] or r[2]): continue
    a = r[0]; m = a.strftime("%b-%y") if isinstance(a, dt.datetime) else S(a)
    k = (r[3] or "HOIS", m); reg[k][0] += 1; reg[k][1] += num(r[43])
b3 = {}
for r in wb["3B Data"].iter_rows(min_row=3, values_only=True):
    if r[1]: b3[(r[1], r[2])] = num(r[11])
wb.close()
print("\n%-16s %-7s %6s %16s | %6s %16s | %16s | %14s" % ("GSTIN", "Month", "mRows", "monthly taxable", "regRow", "register taxable", "3B 3.1(d)", "monthly-3B"))
gst = sorted({k[0] for k in summary} | {k[0] for k in reg})
big = []
for g in gst:
    for m in MLBL:
        s = summary.get((g, m)); rg = reg.get((g, m), [0, 0.0]); b = b3.get((g, m), 0.0)
        mt = s[1] if s else 0.0; mr = s[0] if s else 0
        d = mt - b
        if abs(d) > 1 or abs(mt - rg[1]) > 1 or s is None:
            print("%-16s %-7s %6d %16.2f | %6d %16.2f | %16.2f | %14.2f %s" % (g, m, mr, mt, rg[0], rg[1], b, d, "" if s else "<-- NO FILE"))
tot_m = sum(v[1] for v in summary.values()); tot_r = sum(v[1] for v in reg.values()); tot_b = sum(b3.values())
print("\nTOTALS: monthly sheets %.2f | current register %.2f | 3B 3.1(d) %.2f | monthly-3B %.2f" % (tot_m, tot_r, tot_b, tot_m - tot_b))
print("rows: monthly %d | register %d" % (sum(v[0] for v in summary.values()), sum(v[0] for v in reg.values())))
