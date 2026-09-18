import openpyxl
from openpyxl.utils import get_column_letter
p = "//server/GST FOLDER/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/DPS Workings/1. Sales/Sales Register_FY 2025-26 Format.xlsx"
wb = openpyxl.load_workbook(p, data_only=False)
ws = wb["SR_2025-26"]
print(f"max_col={ws.max_column}")
for c in range(1, ws.max_column+1):
    v = ws.cell(4, c).value
    if v is not None:
        print(f"  {get_column_letter(c)}: {v!r}")
