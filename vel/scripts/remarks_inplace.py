"""Q2 - remarks typed IN the reco, beside the difference; totality view auto-concatenates that
state's month remarks. Applied to Step 2 (SR vs GSTR-1) and Step 3 (1 vs 3B + SR vs 3B).
The separate 'Remarks' sheets go away; AUTO explanations move into the typed cells as defaults."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79"); AMB = PatternFill("solid", fgColor="FFF2CC")
def S(v): return "" if v is None else str(v).strip()
def locked(p):
    try: f = open(p, "r+b"); f.close(); return False
    except Exception: return True
MONTHS = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26"]

def concat_formula(mm_name, state_cell, key_col, rem_col, first, last):
    """12 IFERROR terms: 'Mon: remark; ' for each month with a non-empty typed remark."""
    terms = []
    for mo in MONTHS:
        idx = "INDEX('%s'!$%s$%d:$%s$%d,MATCH(%s&\"|%s\",'%s'!$%s$%d:$%s$%d,0))" % (
            mm_name, rem_col, first, rem_col, last, state_cell, mo, mm_name, key_col, first, key_col, last)
        terms.append('IFERROR(IF(LEN(%s)>0,"%s: "&%s&"; ",""),"")' % (idx, mo, idx))
    return "=TRIM(" + "&".join(terms) + ")"

# ---------------- Step 2 ----------------
p = r"C:\Users\pawar\Downloads\VEL_Step2_Reco_GSTR1_DRAFT.xlsx"
if locked(p): print("LOCKED:", p)
else:
    wb = openpyxl.load_workbook(p)
    mm = wb["Month-on-Month"]; mm.sheet_state = "visible"
    last = mm.max_row
    AUTO = {("Bihar", "Jan-26"): "34 no-IRN credit notes in GSTR-1 only - not in books, not in 3B",
            ("Telangana", "May-25"): "e-invoiced docs not uploaded to GSTR-1; tax paid via 3B",
            ("Telangana", "Jun-25"): "e-invoiced docs not uploaded to GSTR-1; tax paid via 3B",
            ("Telangana", "Aug-25"): "e-invoiced docs not uploaded to GSTR-1; tax paid via 3B",
            ("Madhya Pradesh", "Dec-25"): "Rs 9 advance rounding"}
    # col Q: typed remark (seed with AUTO); col R: hidden key
    mm.cell(3, 17, "Remark (type here, beside the difference)"); mm.cell(3, 18, "key")
    for c in (17, 18): mm.cell(3, c).font = HF; mm.cell(3, c).fill = HB
    for r in range(4, last + 1):
        st, mo = S(mm.cell(r, 1).value), S(mm.cell(r, 3).value)
        if not st: continue
        cur = S(mm.cell(r, 17).value)
        if cur.startswith("="): cur = ""
        mm.cell(r, 17).value = cur or AUTO.get((st, mo), None)
        mm.cell(r, 17).fill = AMB
        mm.cell(r, 18).value = '=A%d&"|"&C%d' % (r, r)
    mm.column_dimensions["Q"].width = 46; mm.column_dimensions["R"].hidden = True
    mm.auto_filter.ref = "A3:R%d" % last
    ws = wb["SR vs GSTR-1"]
    ws.cell(54, 35).value = "Remarks (auto from Month-on-Month typed remarks)"
    for i2 in range(19):
        r = 55 + i2
        ws.cell(r, 35).value = concat_formula("Month-on-Month", "$B%d" % r, "R", "Q", 4, last)
    if "Remarks" in wb.sheetnames: del wb["Remarks"]
    wb.save(p); print("Step 2: typed remark col Q on Month-on-Month (visible again), totality auto-concat, Remarks sheet removed")

# ---------------- Step 3 ----------------
p = r"C:\Users\pawar\Downloads\VEL_Step3_Reco_GSTR1_vs_3B_DRAFT.xlsx"
if locked(p): print("LOCKED:", p)
else:
    wb = openpyxl.load_workbook(p)
    mm = wb["Month-on-Month"]; last = mm.max_row
    AUTO = {("Bihar", "Aug-25"): "3B taxable mis-keyed (5,73,80,517 as 57,38,017); tax correct - see Exceptions A",
            ("Bihar", "Jan-26"): "34 no-IRN GSTR-1-only credit notes not taken in 3B - see Exceptions B",
            ("Telangana", "May-25"): "e-invoiced docs missing from GSTR-1; in 3B, tax paid - Exceptions C",
            ("Telangana", "Jun-25"): "e-invoiced docs missing from GSTR-1; in 3B, tax paid - Exceptions C",
            ("Telangana", "Aug-25"): "e-invoiced docs missing from GSTR-1; in 3B, tax paid - Exceptions C",
            ("Punjab", "May-25"): "Re 1 CGST rounding"}
    mm.cell(3, 16, "Remark (type here, beside the difference)"); mm.cell(3, 17, "key")
    for c in (16, 17): mm.cell(3, c).font = HF; mm.cell(3, c).fill = HB
    for r in range(4, last + 1):
        st, mo = S(mm.cell(r, 1).value), S(mm.cell(r, 3).value)
        if not st: continue
        cur = S(mm.cell(r, 16).value)
        if cur.startswith("="): cur = ""
        mm.cell(r, 16).value = cur or AUTO.get((st, mo), None)
        mm.cell(r, 16).fill = AMB
        mm.cell(r, 17).value = '=A%d&"|"&C%d' % (r, r)
    mm.column_dimensions["P"].width = 46; mm.column_dimensions["Q"].hidden = True
    mm.auto_filter.ref = "A3:Q%d" % last
    ws = wb["GSTR-1 vs GSTR-3B"]
    ws.cell(32, 15).value = "Remarks (auto from Month-on-Month typed remarks)"
    for r in range(33, 50):
        ws.cell(r, 15).value = concat_formula("Month-on-Month", "$B%d" % r, "Q", "P", 4, last)
    if "Remarks" in wb.sheetnames: del wb["Remarks"]
    wb.save(p); print("Step 3: typed remark col P on Month-on-Month, grid auto-concat, Remarks sheet removed")
