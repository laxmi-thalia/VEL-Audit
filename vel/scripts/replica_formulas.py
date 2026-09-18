"""Write the ITC Summary (129 cols) and Table 13 & 6A1 formulas explicitly - one formula per column, last year's
logic, this year's sheets/ranges. Replaces the regex-rolled formulas from replica_itc.py (which mangled refs)."""
import openpyxl, warnings, time
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
warnings.filterwarnings("ignore")
def S(v): return "" if v is None else str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
open(P, "r+b").close()
t0 = time.time(); wb = openpyxl.load_workbook(P); print("loaded %.0fs" % (time.time() - t0))
reg = wb["ITC Register 2025-26"]; RH = {S(reg.cell(5, c).value): L(c) for c in range(1, reg.max_column + 1)}; NREG = max(r for r in range(6, reg.max_row + 1) if reg.cell(r, 4).value)
r27 = wb["ITC Register 2026-27"]; R27 = {S(r27.cell(4, c).value): L(c) for c in range(1, r27.max_column + 1)}; N27 = max((r for r in range(5, r27.max_row + 1) if r27.cell(r, 4).value), default=5)
b3 = wb["3B Data"]; B3 = {S(b3.cell(2, c).value): L(c) for c in range(1, b3.max_column + 1)}; N3 = max(r for r in range(3, b3.max_row + 1) if b3.cell(r, 2).value)
def b3c(pre, tail): return next(v for k, v in B3.items() if k.startswith(pre) and k.endswith(tail))
b2 = wb["GSTR-2B Apr25-Aug26"]; BH = {S(b2.cell(2, c).value): L(c) for c in range(1, b2.max_column + 1)}; NB = max(r for r in range(3, b2.max_row + 1) if b2.cell(r, 1).value)
ex = wb["T6A1 Extract - 24-25"]; EN = max(r for r in range(5, ex.max_row + 1) if ex.cell(r, 1).value)
tc = wb["Tax comp report"]; TH = {S(tc.cell(7, c).value): L(c) for c in range(1, tc.max_column + 1)}; NTC = max(r for r in range(8, tc.max_row + 1) if tc.cell(r, 2).value)
lyl = wb["LY 24-25 claims"]; LH = {S(lyl.cell(2, c).value): L(c) for c in range(1, lyl.max_column + 1)}; NL = max(r for r in range(3, lyl.max_row + 1) if lyl.cell(r, 1).value)
RG = lambda h: "'ITC Register 2025-26'!$%s$6:$%s$%d" % (RH[h], RH[h], NREG)
R7 = lambda h: "'ITC Register 2026-27'!$%s$5:$%s$%d" % (R27[h], R27[h], N27)
B3R = lambda col: "'3B Data'!$%s$3:$%s$%d" % (col, col, N3)
B2R = lambda h: "'GSTR-2B Apr25-Aug26'!$%s$3:$%s$%d" % (BH[h], BH[h], NB)
EXR = lambda col: "'T6A1 Extract - 24-25'!$%s$5:$%s$%d" % (col, col, EN)
TCR = lambda col: "'Tax comp report'!$%s$8:$%s$%d" % (col, col, NTC)
LYR = lambda h: "'LY 24-25 claims'!$%s$3:$%s$%d" % (LH[h], LH[h], NL)
H3 = ("IGST", "CGST", "SGST"); EXC = {"IGST": "P", "CGST": "Q", "SGST": "R"}; B2C = {"IGST": "IGST (Net)", "CGST": "CGST (Net)", "SGST": "SGST (Net)"}
sm = wb["ITC Summary"]
def a(col, r): return "%s%d" % (col, r)
def trio(base, r): return [a(L(CI(base) + j), r) for j in range(3)]
for i in range(19):
    r = 6 + i; C = "$C%d" % r
    for j, h in enumerate(H3):
        ex_ = EXC[h]; b2h = B2C[h]
        F = {}
        F["G"] = '=SUMIFS(%s,%s,%s,%s,"ITC dated 24-25 in 2B of 24-25 availed in 25-26")' % (EXR(ex_), EXR("I"), C, EXR("G"))
        F["J"] = '=SUMIFS(%s,%s,%s,%s,"ITC dated 24-25 in 2B of 25-26 availed in 25-26")' % (EXR(ex_), EXR("I"), C, EXR("G"))
        F["M"] = '=SUMIFS(%s,%s,%s,%s,"Correction Entries- ITC dated 24-25 reversed in 25-26")' % (EXR(ex_), EXR("I"), C, EXR("G"))
        F["P"] = '=SUMIFS(%s,%s,%s,%s,"Table 6A1 of GSTR-9 - Unclaimed")' % (B2R(b2h), B2R("Company GSTIN"), C, B2R("GSTR-9/9C"))
        F["S"] = '=SUMIFS(%s,%s,%s,%s,"RCM",%s,DATE(2025,4,1))' % (RG(h), RG("VEL GSTIN"), C, RG("Category"), RG("3B Claim  Month"))
        F["V"] = "=%s+%s+%s+%s+%s" % (trio("G", r)[j], trio("J", r)[j], trio("P", r)[j], trio("M", r)[j], trio("S", r)[j])
        F["Y"] = "=SUMIFS(%s,%s,%s)-%s-%s+%s" % (B3R(b3c("4A(5)", h)), B3R("B"), C, trio("AH", r)[j], trio("V", r)[j], trio("S", r)[j])
        F["AB"] = "=SUMIFS(%s,%s,%s)-%s" % (B3R(b3c("4A(3)", h)), B3R("B"), C, trio("S", r)[j])
        F["AE"] = "=SUMIFS(%s,%s,%s)" % (B3R(b3c("4A(4)", h)), B3R("B"), C)
        F["AH"] = "=SUMIFS(%s,%s,%s)-%s-%s+%s" % (B3R(b3c("4D(1)", h)), B3R("B"), C, trio("G", r)[j], trio("M", r)[j], trio("DP", r)[j])
        F["AK"] = "=%s-%s-%s-%s-%s-%s" % (trio("D", r)[j], trio("V", r)[j], trio("Y", r)[j], trio("AB", r)[j], trio("AE", r)[j], trio("AH", r)[j])
        F["AN"] = "=-(SUMIFS(%s,%s,%s)+SUMIFS(%s,%s,%s))" % (B3R(b3c("4B(1)", h)), B3R("B"), C, B3R(b3c("4B(2)", h)), B3R("B"), C)   # Octa reports 4B negative
        F["AQ"] = "=%s+%s+%s+%s-%s" % (trio("Y", r)[j], trio("AB", r)[j], trio("AE", r)[j], trio("AH", r)[j], trio("AN", r)[j])
        F["AT"] = "=SUMIFS(%s,%s,%s)" % (B3R(b3c("4(C)", h)), B3R("B"), C)
        F["AW"] = "=%s-%s+%s" % (trio("AQ", r)[j], trio("AT", r)[j], trio("V", r)[j])
        F["AZ"] = '=SUMIFS(%s,%s,%s,%s,"Yes")' % (B2R(b2h), B2R("Company GSTIN"), C, B2R("Table 8A"))
        F["BC"] = "=%s" % trio("Y", r)[j]
        F["BF"] = '=SUMIFS(%s,%s,%s,%s,"Table 8C of GSTR-9")' % (B2R(b2h), B2R("Company GSTIN"), C, B2R("GSTR-9/9C"))
        F["BI"] = "=%s-%s-%s" % (trio("AZ", r)[j], trio("BC", r)[j], trio("BF", r)[j])
        F["BL"] = "=SUMIFS(%s,%s,%s)" % (TCR(("AS", "AT", "AU")[j]), TCR("AR"), C)
        F["BO"] = "=%s+%s" % (trio("BI", r)[j], trio("BL", r)[j])
        F["BS"] = "=SUMIF(%s,%s,%s)" % (TCR("B"), C, TCR(("V", "W", "X")[j]))
        F["BV"] = "=SUMIF(%s,%s,%s)" % (TCR("B"), C, TCR(("Y", "Z", "AA")[j]))
        F["BY"] = "=SUMIF(%s,%s,%s)" % (TCR("B"), C, TCR(("AB", "AC", "AD")[j]))
        F["CB"] = "=%s+%s+%s+%s" % (trio("BI", r)[j], trio("BS", r)[j], trio("BV", r)[j], trio("BY", r)[j])
        F["CE"] = '=SUMIFS(%s,%s,%s,%s,"Unclaimed")+SUMIFS(%s,%s,%s,%s,"Permanent Reversals")' % (B2R(b2h), B2R("Company GSTIN"), C, B2R("Table 13"), B2R(b2h), B2R("Company GSTIN"), C, B2R("Permanent Reversals"))
        F["CH"] = '=SUMIFS(%s,%s,%s,%s,"13")+%s' % (R7(h), R7("VEL GSTN"), C, R7("GSTR 9_Reporting"), trio("CE", r)[j])
        F["CK"] = "=%s-%s" % (trio("CH", r)[j], trio("CE", r)[j])
        F["CN"] = '=SUMIFS(%s,%s,%s,%s,"NA")+%s' % (RG(h), RG("VEL GSTIN"), C, RG("GSTR 9C_Reporting"), trio("CV", r)[j])      # 12A = booked 25-26 claimed 25-26 (NA) + 12C
        F["CR"] = '=SUMIFS(%s,%s,%s,%s,"12B")' % (RG(h), RG("VEL GSTIN"), C, RG("GSTR 9C_Reporting"))
        F["CV"] = '=SUMIFS(%s,%s,%s,%s,"12C")' % (R7(h), R7("VEL GSTN"), C, R7("GSTR 9C_Reporting"))
        F["CZ"] = "=%s+%s-%s" % (trio("CN", r)[j], trio("CR", r)[j], trio("CV", r)[j])
        F["DD"] = "=%s" % trio("AQ", r)[j]
        F["DH"] = "=%s-%s" % (trio("CZ", r)[j], trio("DD", r)[j])
        F["DP"] = '=SUMIFS(%s,%s,%s,%s,"Add to 4D1")' % (TCR(("S", "T", "U")[j]), TCR("B"), C, TCR("R"))
        F["DS"] = "=SUMIFS(%s,%s,%s,%s,$DS$2,%s,$DT$2)+SUMIFS(%s,%s,%s,%s,$DS$2,%s,$DU$2)" % (B2R(b2h), B2R("Company GSTIN"), C, B2R("Table 8A"), B2R("GSTR-9/9C"), B2R(b2h), B2R("Company GSTIN"), C, B2R("Table 8A"), B2R("GSTR-9/9C"))
        for base, f in F.items(): sm[a(L(CI(base) + j), r)] = f
    sm["B%d" % r] = "=LEFT(C%d,2)" % r
    for base in ("CQ", "CU", "CY", "DC", "DG", "DK"):
        sm["%s%d" % (base, r)] = "=SUM(%s%d:%s%d)" % (L(CI(base) - 3), r, L(CI(base) - 1), r)
    sm["BR%d" % r] = "=SUM(BO%d:BQ%d)" % (r, r)
    sm["DL%d" % r] = "=ROUND(V%d+W%d+X%d,0)" % (r, r, r)
    sm["DM%d" % r] = "=DL%d-DK%d" % (r, r)
    sm["DN%d" % r] = '="The difference of Rs ."&DL%d&"/-arises due to the revised Form GSTR-9 format for FY 2025-26, which excludes ITC of FY 2024-25 availed in FY 2025-26 "&"(reported in Table 6A1) from the auto-population into Table 7J and consequently Table 12E of GSTR-9C"&$DN$4' % r
    sm["DO%d" % r] = None
    for base in ("DV", "DW", "DX"): sm["%s%d" % (base, r)] = None
sm["DS2"] = "Yes - Inv dt 25-26 amended in 26-27"; sm["DT2"] = 'Considered in Table 6B but amended to "0" in FY 26-27'; sm["DU2"] = 'Considered in Table 6B but amended to "Different value" in FY 26-27'
sm["DN4"] = ". The variance is purely format-driven and not attributable to any incorrect availment or reporting of ITC."
print("ITC Summary formulas rewritten")
# Table 13 & 6A1 differences (rows 8..26 -> summary rows 6..24)
t = wb["Table 13 & 6A1 differences"]
for i in range(19):
    r = 8 + i; sr = 6 + i
    for j in range(3):
        t["%s%d" % (L(3 + j), r)] = '=SUMIFS(%s,%s,$A%d,%s,"13")' % (LYR(H3[j]), LYR("VEL GSTIN"), r, LYR("GSTR 9_Reporting"))
        t["%s%d" % (L(6 + j), r)] = "='ITC Summary'!%s%d" % (L(CI("V") + j), sr)
        t["%s%d" % (L(9 + j), r)] = "='ITC Summary'!%s%d" % (L(CI("P") + j), sr)
        t["%s%d" % (L(12 + j), r)] = "=(%s%d-%s%d)-%s%d" % (L(6 + j), r, L(9 + j), r, L(3 + j), r)
        t["%s%d" % (L(19 + j), r)] = "='ITC Summary'!%s%d" % (L(CI("J") + j), sr)
        t["%s%d" % (L(22 + j), r)] = "='ITC Summary'!%s%d" % (L(CI("M") + j), sr)
        t["%s%d" % (L(28 + j), r)] = "='ITC Summary'!%s%d" % (L(CI("S") + j), sr)
        t["%s%d" % (L(31 + j), r)] = "=%s%d-%s%d-%s%d-%s%d-%s%d" % (L(12 + j), r, L(19 + j), r, L(22 + j), r, L(25 + j), r, L(28 + j), r)
    t["O%d" % r] = "=ROUND(SUM(L%d:N%d),0)" % (r, r)
    t["Q%d" % r] = '=IF(ABS(O%d)<1,"Matched",IF(ABS(AE%d)+ABS(AF%d)+ABS(AG%d)<1000,"Reasons identified","Reasons not identifiable"))' % (r, r, r, r)
print("Table 13 & 6A1 formulas rewritten")
wb.save(P); print("saved %.0fs" % (time.time() - t0))
