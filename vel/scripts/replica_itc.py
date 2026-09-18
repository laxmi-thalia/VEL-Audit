"""Faithful replica of last year's ITC sheets (VEL_GSTR 9_9C FY 24-25.xlsb) rolled to FY 2025-26, from ly_itc_dump.pkl:
  ITC Summary (129 cols)          <- every header/fill/nf/width/merge + row-5 formulas, sheet refs rolled
  T6A1 Extract - 24-25 (19 cols)  <- last year's layout, our stamped rows, Source/Remarks vocabulary rolled
  ITC Register 2026-27 (49 cols)  <- last year's 'ITC Register 2025-26' (next-year claims = Table 13 / 12C detail)
  Table 13 & 6A1 differences      <- last year's 33-col layout + formulas
  GSTR-2B Apr25-Aug26             <- Table 8A / Table 13 / 8C / GSTR-9-9C columns in last year's vocabulary
Year strings rolled: 23-24->24-25, 24-25->25-26, 25-26->26-27 (labels only), Mar-24/Apr-24 -> Mar-25/Apr-25."""
import os, re, pickle, warnings, datetime as dt
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
def S(v): return "" if v is None else str(v).strip()
def num(v):
    try: return float(v)
    except Exception: return 0.0
SP = os.path.dirname(os.path.abspath(__file__))
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
open(P, "r+b").close()
D = pickle.load(open(os.path.join(SP, "ly_itc_dump.pkl"), "rb"))
def roll(s):
    if not isinstance(s, str): return s
    s = s.replace("2025-26", "2026-27").replace("2024-25", "2025-26").replace("2023-24", "2024-25")
    s = re.sub(r"\b25-26\b", "26-27", s); s = re.sub(r"\b24-25\b", "25-26", s); s = re.sub(r"\b23-24\b", "24-25", s)
    s = s.replace("Mar-24", "Mar-25").replace("Apr-24", "Apr-25").replace("01 Apr 2024", "01 Apr 2025")
    return s
def rgb(bgr):
    if bgr is None: return None
    b, g, r = (bgr >> 16) & 255, (bgr >> 8) & 255, bgr & 255
    return "%02X%02X%02X" % (r, g, b)
def apply_style(cell, info):
    cell.font = Font(bold=info["bold"], color=rgb(info["fc"]) if info["fc"] else None, size=info["fs"] or 11, name=info["fn"] or "Calibri")
    if info["fill"] is not None: cell.fill = PatternFill("solid", fgColor=rgb(info["fill"]))
    if info["nf"] and info["nf"] != "General": cell.number_format = info["nf"]
    ha = {-4108: "center", -4131: "left", -4152: "right"}.get(info["ha"])
    cell.alignment = Alignment(horizontal=ha, vertical="center", wrap_text=info["wrap"])
wb = openpyxl.load_workbook(P)
SHEETMAP = {"'T6A1 Extract - 23-24'": "'T6A1 Extract - 24-25'", "'GSTR-2B Apr 24-Oct 25'": "'GSTR-2B Apr25-Aug26'", "'Consolidated GSTR-3B Extract'": "'3B Data'",
            "'ITC Register 2024-25'": "'ITC Register 2025-26'", "'ITC Register 2025-26'": "'ITC Register 2026-27'"}
reg = wb["ITC Register 2025-26"]; RH = {S(reg.cell(5, c).value): L(c) for c in range(1, reg.max_column + 1)}
b3 = wb["3B Data"]; B3 = {S(b3.cell(2, c).value): L(c) for c in range(1, b3.max_column + 1)}
def b3c(pre, tail): return next(v for k, v in B3.items() if k.startswith(pre) and k.endswith(tail))
b2 = wb["GSTR-2B Apr25-Aug26"]; BH = {S(b2.cell(2, c).value): L(c) for c in range(1, b2.max_column + 1)}
NB = max(r for r in range(3, b2.max_row + 1) if b2.cell(r, 1).value)
N3 = max(r for r in range(3, b3.max_row + 1) if b3.cell(r, 2).value)
NREG = max(r for r in range(6, reg.max_row + 1) if reg.cell(r, 4).value)
tc = wb["Tax comp report"]; NTC = max(r for r in range(8, tc.max_row + 1) if tc.cell(r, 2).value)
MAPS = {
 "'ITC Register 2025-26'": {"W": RH["IGST"], "X": RH["CGST"], "Y": RH["SGST"], "D": RH["VEL GSTIN"], "F": RH["Category"], "E": RH["3B Claim  Month"], "AA": RH["GSTR 9C_Reporting"]},
 "'3B Data'": {"A": "B", "I": b3c("4A(3)", "IGST"), "J": b3c("4A(3)", "CGST"), "K": b3c("4A(3)", "SGST"), "L": b3c("4A(4)", "IGST"), "M": b3c("4A(4)", "CGST"), "N": b3c("4A(4)", "SGST"),
               "O": b3c("4A(5)", "IGST"), "P": b3c("4A(5)", "CGST"), "Q": b3c("4A(5)", "SGST"), "R": b3c("4B(1)", "IGST"), "S": b3c("4B(1)", "CGST"), "T": b3c("4B(1)", "SGST"),
               "U": b3c("4B(2)", "IGST"), "V": b3c("4B(2)", "CGST"), "W": b3c("4B(2)", "SGST"), "X": b3c("4(C)", "IGST"), "Y": b3c("4(C)", "CGST"), "Z": b3c("4(C)", "SGST"),
               "AA": b3c("4D(1)", "IGST"), "AB": b3c("4D(1)", "CGST"), "AC": b3c("4D(1)", "SGST")},
 "'GSTR-2B Apr25-Aug26'": {"B": BH["Company GSTIN"], "X": BH["IGST (Net)"], "Y": BH["CGST (Net)"], "Z": BH["SGST (Net)"], "I": BH["GSTR-9/9C"], "G": BH["Table 13"], "J": BH["Permanent Reversals"], "E": BH["Table 8A"]},
 "'T6A1 Extract - 24-25'": {}, "'ITC Register 2026-27'": {},
}
# bounded ranges instead of whole columns (41k-row register): sheet -> (first_row, last_row)
BOUNDS = {"'ITC Register 2025-26'": (6, NREG), "'3B Data'": (3, N3), "'GSTR-2B Apr25-Aug26'": (3, NB), "'Tax comp report'": (8, NTC)}
def colmap_formula(f):
    for old, new in SHEETMAP.items():
        if old not in f: continue
        cm = MAPS.get(new, {})
        def sub(m):
            d1, c1, d2, c2 = m.groups(); n1 = cm.get(c1, c1); n2 = cm.get(c2, c2)
            if new in BOUNDS:
                a, b = BOUNDS[new]; return "%s!$%s$%d:$%s$%d" % (new, n1, a, n2, b)
            return "%s!%s%s:%s%s" % (new, d1, n1, d2, n2)
        f = re.sub(re.escape(old) + r"!(\$?)([A-Z]{1,3}):(\$?)([A-Z]{1,3})(?![0-9])", sub, f)
    if "'Tax comp report'" in f:
        a, b = BOUNDS["'Tax comp report'"]
        f = re.sub(r"'Tax comp report'!(\$?)([A-Z]{1,3}):(\$?)([A-Z]{1,3})(?![0-9])", lambda m: "'Tax comp report'!$%s$%d:$%s$%d" % (m.group(2), a, m.group(4), b), f)
    return f
G2ST = {"01AAECR0503Q1ZM": "Jammu & Kashmir", "03AAECR0503Q1ZI": "Punjab", "06AAECR0503Q1ZC": "Haryana", "08AAECR0503Q1Z8": "Rajasthan", "10AAECR0503Q1ZN": "Bihar", "12AAECR0503Q1ZJ": "Arunachal Pradesh", "18AAECR0503Q1Z7": "Assam", "19AAECR0503Q1Z5": "West Bengal", "20AAECR0503Q1ZM": "Jharkhand", "23AAECR0503Q1ZG": "Madhya Pradesh", "24AAECR0503Q1ZE": "Gujarat", "27AAECR0503Q1Z8": "Maharashtra", "32AAECR0503Q1ZH": "Kerala", "33AAECR0503Q1ZF": "Tamil Nadu", "36AAECR0503Q1Z9": "Telangana", "37AAECR0503Q1Z7": "Andhra Pradesh", "09AAECR0503Q1Z6": "Uttar Pradesh", "22AAECR0503Q1ZI": "Chhattisgarh", "29AAECR0503Q1Z4": "Karnataka"}
STATES = [(G2ST[g], g) for g in ["01AAECR0503Q1ZM", "03AAECR0503Q1ZI", "06AAECR0503Q1ZC", "08AAECR0503Q1Z8", "09AAECR0503Q1Z6", "10AAECR0503Q1ZN", "12AAECR0503Q1ZJ", "18AAECR0503Q1Z7", "19AAECR0503Q1Z5", "20AAECR0503Q1ZM", "22AAECR0503Q1ZI", "23AAECR0503Q1ZG", "24AAECR0503Q1ZE", "27AAECR0503Q1Z8", "29AAECR0503Q1Z4", "32AAECR0503Q1ZH", "33AAECR0503Q1ZF", "36AAECR0503Q1Z9", "37AAECR0503Q1Z7"]]
# ---------------------------------------------------------------- 1. 2B classification columns -> last year's vocabulary
c = lambda h: BH[h]
for r in range(3, NB + 1):
    b2.cell(r, CI(c("Table 8A"))).value = ('=IF($%s%d<>"2025-26","No-"&$%s%d,IF($%s%d="Yes","No-25-26 - RCM",IF($%s%d="No","No-25-26 - ITC not available",IF($%s%d="Yes","No-25-26 - Amendment",IF($%s%d<>"2025-26","No-"&$%s%d,"Yes")))))'
                                            % (c("FY (2B period)"), r, c("FY (2B period)"), r, c("Reverse Charge"), r, c("GSTR-9 (8A) ITC Available"), r, c("Is Amendment"), r, c("Doc FY (doc date)"), r, c("Doc FY (doc date)"), r))
    b2.cell(r, CI(c("Table 13"))).value = '=IF(AND($%s%d="2025-26",$%s%d="2026-27"),IF($%s%d="","Unclaimed","Claimed"),"")' % (c("Doc FY (doc date)"), r, c("FY (2B period)"), r, c("3B Claim Month"), r)
    b2.cell(r, CI(c("Table 8C"))).value = '=IF($%s%d="Claimed","Table 8C of GSTR-9","")' % (c("Table 13"), r)
    b2.cell(r, CI(c("GSTR-9/9C"))).value = ('=IF(LEFT($%s%d,3)="6A1","Table 6A1 of GSTR-9 - "&IF(ISNUMBER(SEARCH("Unclaimed",$%s%d)),"Unclaimed","Claimed"),IF($%s%d<>"",$%s%d,IF($%s%d="Yes","Table 6B of GSTR-9","")))'
                                             % (c("6A1 mark"), r, c("6A1 mark"), r, c("Table 8C"), r, c("Table 8C"), r, c("Table 8A"), r))
print("2B classification columns rewritten to last year's vocabulary")
# ---------------------------------------------------------------- 2. ITC Register 2026-27 (last year's 'ITC Register 2025-26' layout)
ly = D["ITC Register 2025-26"]; NM = "ITC Register 2026-27"
if NM in wb.sheetnames: del wb[NM]
ws = wb.create_sheet(NM, wb.sheetnames.index("ITC Register 2025-26") + 1); ws.row_dimensions[1].height = 21
ws.cell(2, 1, "Input Tax Credit (ITC) Register for FY 2026-27 - FY 25-26 dated invoices claimed in FY 26-27 (Table 13 / 12C detail). Built from the FY 26-27 GSTR-2B (Apr-Jul 26); replace with the client's FY 26-27 ITC working when received.").font = Font(bold=True)
LYH = [ly["hdr"][(4, cc)]["v"] for cc in range(1, ly["ncols"] + 1)]
for cc, h in enumerate(LYH, 1):
    x = ws.cell(4, cc, h); apply_style(x, ly["hdr"][(4, cc)]); ws.column_dimensions[L(cc)].width = ly["widths"].get(cc) or 12
b2rows = [r for r in range(3, NB + 1) if S(b2.cell(r, CI(BH["Doc FY (doc date)"])).value) == "2025-26" and S(b2.cell(r, CI(BH["FY (2B period)"])).value) == "2026-27"]
M3B = {4: "01 Apr 2026", 5: "02 May 2026", 6: "03 June 2026", 7: "04 July 2026", 8: "05 Aug 2026", 9: "06 Sep 2026"}
bv = lambda rr, h: b2.cell(rr, CI(BH[h])).value
r = 4
for rr in b2rows:
    r += 1
    g = S(bv(rr, "Company GSTIN")); per = bv(rr, "Tax Period"); dd = bv(rr, "Doc Date")
    vals = {"Company": "VEL", "Business place": "", "STATE NAME": G2ST.get(g, g), "VEL GSTN": g, "3B Claim  Month": M3B.get(per.month, "") if isinstance(per, dt.datetime) else "",
            "Category": "ITC", "Document Type": "NA", "Document Number": "NA", "Posting Date": "NA", "Posting Year": "2026-27", "Invoice No.": bv(rr, "Doc No"), "Invoice Date": dd,
            "Invoice Year": "2025-26", "Vendor Code": "NA", "Vendor GSTIN": bv(rr, "Supplier GSTIN"), "Correct GSTIN": bv(rr, "Supplier GSTIN"), "Vendor Name/RCM Category": bv(rr, "Supplier Name"),
            "Taxable Value": num(bv(rr, "Taxable Value (Net)")), "IGST": num(bv(rr, "IGST (Net)")), "CGST": num(bv(rr, "CGST (Net)")), "SGST": num(bv(rr, "SGST (Net)")),
            "GSTR 9_Reporting": "13", "GSTR 9C_Reporting": "12C", "Available in 2B": "Y", "2B Inv": bv(rr, "Doc No"), "2B Period": per, "Final Remarks": "Matched A",
            "SAP Period": "", "GSTR 2B/6A PERIOD": per, "2B YEAR": "2026-27", "GST CREDIT YES/NO": "Yes", "Company Code": 1000, "Type": "ITC", "Consider in 8A reco": "No"}
    for cc, h in enumerate(LYH, 1):
        cell = ws.cell(r, cc); col = L(cc)
        if h == "Reasons": cell.value = "ITC booked in 2025-26 claimed in 2026-27"
        elif h == "Tax Rate": cell.value = '=IFERROR(W%d/S%d*100,"")' % (r, r)
        elif h == "Total GST": cell.value = "=SUM(T%d:V%d)" % (r, r)
        elif h == "Vendor": cell.value = "=LEFT(D%d,2)" % r
        elif h == "My GSTN": cell.value = "=LEFT(P%d,2)" % r
        elif h == "As per State": cell.value = '=IF(AR%d=AS%d,"Intra State","Inter State")' % (r, r)
        elif h == "As per Amounts": cell.value = '=IF(T%d=0,"Intra State","Inter State")' % r
        elif h == "POS Check": cell.value = "=AT%d=AU%d" % (r, r)
        elif h == "Remarks": cell.value = '=IF(AV%d,"Correct","Check")' % r
        else: cell.value = vals.get(h)
        info = ly["data"].get((5, cc))
        if info and info["nf"] != "General": cell.number_format = info["nf"]
        if isinstance(cell.value, dt.datetime): cell.number_format = "dd-mm-yyyy"
N27 = max(r, 5)
for col in ("S", "T", "U", "V", "W"): ws["%s3" % col] = "=SUBTOTAL(9,%s5:%s%d)" % (col, col, N27)
ws["X3"] = "Table"; ws["Z3"] = "Table"
ws.auto_filter.ref = "A4:AW%d" % N27
print("ITC Register 2026-27:", len(b2rows), "rows")
# ---------------------------------------------------------------- 3. T6A1 Extract - 24-25 (last year's headers/widths; vocabulary rolled)
ex = wb["T6A1 Extract - 24-25"]; lyx = D["T6A1 Extract - 23-24"]
EN = max(r for r in range(5, ex.max_row + 1) if ex.cell(r, 1).value)
for cc in range(1, 20):
    x = ex.cell(4, cc); x.value = roll(lyx["hdr"][(4, cc)]["v"]); apply_style(x, lyx["hdr"][(4, cc)]); ex.column_dimensions[L(cc)].width = lyx["widths"].get(cc) or 12
ex.column_dimensions["H"].width = 24
SRCMAP = {"ITC Register 2025-26 (matched FY 24-25 2B)": "2B vs 3B Reco - 2B merged - Client", "GSTR-2B Apr25-Aug26 (Octa)": "2B vs 3B Reco - 2B merged - Client", "ITC Register 2025-26": "Final ITC Register 2025-26", "RCM Register (fiscal 2024/12 posting, 3B Apr-25)": "Final ITC Register 2025-26"}
REMMAP = {"Correction Entries - ITC dated 24-25 booked/reversed in 25-26 (not in 2B)": "Correction Entries- ITC dated 24-25 reversed in 25-26"}
for r in range(5, EN + 1):
    ex.cell(r, 1).value = SRCMAP.get(S(ex.cell(r, 1).value), ex.cell(r, 1).value)
    v = ex.cell(r, 7).value
    if isinstance(v, str) and v.startswith("='GSTR-2B"):
        ex.cell(r, 7).value = '=IF(ISNUMBER(SEARCH("Unclaimed",%s)),"ITC dated 24-25 in 2B of 25-26 - Unclaimed","ITC dated 24-25 in 2B of 25-26 availed in 25-26")' % v[1:]
    elif isinstance(v, str) and v.startswith("='ITC Register"):
        ex.cell(r, 7).value = '=IF(ISNUMBER(SEARCH("Correction",%s)),"Correction Entries- ITC dated 24-25 reversed in 25-26","ITC dated 24-25 in 2B of 24-25 availed in 25-26")' % v[1:]
    elif isinstance(v, str): ex.cell(r, 7).value = REMMAP.get(v, v)
    for cc in (15, 16, 17, 18): ex.cell(r, cc).number_format = "#,##0"
ex.freeze_panes = "A5"; ex.auto_filter.ref = "A4:S%d" % EN
print("T6A1 Extract re-tagged; rows", EN - 4)
# ---------------------------------------------------------------- 4. ITC Summary (129 cols, last year's exact layout; rows 1..4 -> 2..5)
lys = D["ITC Summary"]
if "ITC Summary" in wb.sheetnames: del wb["ITC Summary"]
sm = wb.create_sheet("ITC Summary", wb.sheetnames.index("T6A1 Extract - 24-25") + 1); sm.row_dimensions[1].height = 21
NC = lys["ncols"]
for (r, cc), info in lys["hdr"].items():
    if 1 <= r <= 4:
        x = sm.cell(r + 1, cc); x.value = roll(info["v"]); apply_style(x, info)
for a in lys["merges"]:
    m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", a)
    if m and int(m.group(2)) <= 4: sm.merge_cells("%s%d:%s%d" % (m.group(1), int(m.group(2)) + 1, m.group(3), int(m.group(4)) + 1))
for cc, w in lys["widths"].items(): sm.column_dimensions[L(cc)].width = w if w else 11.9
for r, h in lys["rowh"].items():
    if r <= 4: sm.row_dimensions[r + 1].height = h
row5 = {cc: lys["data"][(5, cc)] for cc in range(1, NC + 1) if (5, cc) in lys["data"]}
def roll_formula(f, r):
    f = colmap_formula(f)
    f = re.sub(r"(?<![A-Z$'!])(\$?[A-Z]{1,3}\$?)5(?![0-9])", lambda m: m.group(1) + str(r), f)   # same-sheet row 5 -> r
    f = f.replace("'ITC Summary'!$C%d" % r, "$C%d" % r)
    f = f.replace("$DS$1", "$DS$2").replace("$DT$1", "$DT$2").replace("$DU$1", "$DU$2").replace("$DN$3", "$DN$4")
    return roll(f)
BLANK = {"D", "E", "F"}          # 6A gross auto-populate: blank per Priyesh
for i, (st, g) in enumerate(STATES):
    r = 6 + i
    for cc in range(1, NC + 1):
        info = row5.get(cc); cell = sm.cell(r, cc); col = L(cc)
        if cc == 1: cell.value = st
        elif cc == 2: cell.value = "=LEFT(C%d,2)" % r
        elif cc == 3: cell.value = g
        elif col in BLANK: cell.value = None; cell.fill = PatternFill("solid", fgColor="D9D9D9")
        elif col in ("AZ", "BA", "BB"):   # 8A: from 2B (Table 8A = Yes) until the system-generated GSTR-9 arrives
            h = ("IGST (Net)", "CGST (Net)", "SGST (Net)")[cc - CI("AZ")]
            cell.value = "=SUMIFS('GSTR-2B Apr25-Aug26'!$%s$3:$%s$%d,'GSTR-2B Apr25-Aug26'!$%s$3:$%s$%d,$C%d,'GSTR-2B Apr25-Aug26'!$%s$3:$%s$%d,\"Yes\")" % (BH[h], BH[h], NB, BH["Company GSTIN"], BH["Company GSTIN"], NB, r, BH["Table 8A"], BH["Table 8A"], NB)
        elif col == "DN":
            cell.value = ('="The difference of Rs ."&DL%d&"/-arises due to the revised Form GSTR-9 format for FY 2025-26, which excludes ITC of FY 2024-25 availed in FY 2025-26 "&"(reported in Table 6A1) from the auto-population into Table 7J and consequently Table 12E of GSTR-9C"&$DN$4' % r)
        elif col == "DO": cell.value = None
        elif info and info["f"]:
            f = roll_formula(info["f"], r)
            if col == "CN": f = f.replace("-11650", "")      # last year's hard plug dropped (no plugs)
            cell.value = f
        if info: apply_style(cell, info)
        if col in ("AZ", "BA", "BB"): cell.number_format = "#,##0"
for cc in (CI("DS"), CI("DT"), CI("DU"), CI("BS")):
    if (1, cc) in lys["hdr"]: sm.cell(2, cc).value = roll(lys["hdr"][(1, cc)]["v"])
RT = 6 + len(STATES); sm.cell(RT, 1, "Total").font = Font(bold=True)
for cc in range(4, NC + 1):
    info = row5.get(cc); col = L(cc)
    if info and (info["f"] or isinstance(info["v"], (int, float))) and col not in ("DN", "DO", "DL"):
        x = sm.cell(RT, cc, "=SUM(%s6:%s%d)" % (col, col, RT - 1)); x.font = Font(bold=True); x.number_format = info["nf"] if info["nf"] != "General" else "#,##0"
sm.freeze_panes = "D6"
sm.cell(3, CI("AZ")).value = "Total ITC Available Table 8A (from GSTR-2B until the system-generated GSTR-9 is available)"
print("ITC Summary rebuilt: %d cols" % NC)
# ---------------------------------------------------------------- 5. Table 13 & 6A1 differences (last year's 33-col layout)
lyt = D["Table 13 & 6A1 differences"]; NM = "Table 13 & 6A1 differences"
if NM in wb.sheetnames: del wb[NM]
t = wb.create_sheet(NM, wb.sheetnames.index("ITC Summary") + 1); t.row_dimensions[1].height = 21
for (r, cc), info in lyt["hdr"].items():
    if r <= 6:
        x = t.cell(r + 1, cc); x.value = roll(info["v"]); apply_style(x, info)
for a in lyt["merges"]:
    m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", a)
    if m and int(m.group(2)) <= 6: t.merge_cells("%s%d:%s%d" % (m.group(1), int(m.group(2)) + 1, m.group(3), int(m.group(4)) + 1))
for cc, w in lyt["widths"].items(): t.column_dimensions[L(cc)].width = w if w else 12
lyr = {cc: lyt["data"][(7, cc)] for cc in range(1, 34) if (7, cc) in lyt["data"]}
lyl = wb["LY 24-25 claims"]; LH = {S(lyl.cell(2, cc).value): L(cc) for cc in range(1, lyl.max_column + 1)}; NL = max(rr for rr in range(3, lyl.max_row + 1) if lyl.cell(rr, 1).value)
LYR = lambda h: "'LY 24-25 claims'!$%s$3:$%s$%d" % (LH[h], LH[h], NL)
for i, (st, g) in enumerate(STATES):
    r = 8 + i; sr = 6 + i
    for cc in range(1, 34):
        info = lyr.get(cc); cell = t.cell(r, cc)
        if cc == 1: cell.value = g
        elif cc == 2: cell.value = st
        elif cc in (3, 4, 5): cell.value = '=SUMIFS(%s,%s,$A%d,%s,"13")' % (LYR(("IGST", "CGST", "SGST")[cc - 3]), LYR("VEL GSTIN"), r, LYR("GSTR 9_Reporting"))
        elif cc in (6, 7, 8): cell.value = "='ITC Summary'!%s%d" % (L(CI("V") + cc - 6), sr)
        elif cc in (9, 10, 11): cell.value = "='ITC Summary'!%s%d" % (L(CI("P") + cc - 9), sr)
        elif cc in (12, 13, 14): cell.value = "=(%s%d-%s%d)-%s%d" % (L(cc - 6), r, L(cc - 3), r, L(cc - 9), r)
        elif cc == 15: cell.value = "=ROUND(SUM(L%d:N%d),0)" % (r, r)
        elif cc == 17: cell.value = '=IF(ABS(O%d)<1,"Matched",IF(ABS(AE%d)+ABS(AF%d)+ABS(AG%d)<1000,"Reasons identified","Reasons not identifiable"))' % (r, r, r, r)
        elif cc in (19, 20, 21): cell.value = "='ITC Summary'!%s%d" % (L(CI("J") + cc - 19), sr)
        elif cc in (22, 23, 24): cell.value = "='ITC Summary'!%s%d" % (L(CI("M") + cc - 22), sr)
        elif cc in (25, 26, 27): cell.value = 0
        elif cc in (28, 29, 30): cell.value = "='ITC Summary'!%s%d" % (L(CI("S") + cc - 28), sr)
        elif cc in (31, 32, 33): cell.value = "=%s%d-%s%d-%s%d-%s%d-%s%d" % (L(cc - 19), r, L(cc - 12), r, L(cc - 9), r, L(cc - 6), r, L(cc - 3), r)
        if info: apply_style(cell, info)
    for cc in (16, 25, 26, 27): t.cell(r, cc).fill = PatternFill("solid", fgColor="FFF2CC")
RT2 = 8 + len(STATES); t.cell(RT2, 2, "Total").font = Font(bold=True)
for cc in list(range(3, 16)) + list(range(19, 34)):
    x = t.cell(RT2, cc, "=SUM(%s8:%s%d)" % (L(cc), L(cc), RT2 - 1)); x.font = Font(bold=True); x.number_format = (lyr.get(cc) or {}).get("nf") or "#,##0"
t.auto_filter.ref = "A7:AG%d" % RT2
print("Table 13 & 6A1 differences rebuilt (33 cols)")
for nm in ("T12B_T12C differences", "Table 8C vs 13-12 (FY 25-26)"):
    if nm in wb.sheetnames: del wb[nm]
# ---------------------------------------------------------------- INDEX
ix = wb["INDEX"]; thin = Side(style="thin"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
for rr in range(ix.max_row, 4, -1):
    v = S(ix.cell(rr, 2).value)
    if v.startswith("GSTR-9C Table 12C") or v.startswith("Table 8C (FY 25-26) vs"): ix.delete_rows(rr, 1)
have = {S(ix.cell(rr, 2).value) for rr in range(5, ix.max_row + 1)}
desc = "ITC Register FY 2026-27 - FY 25-26 invoices claimed in FY 26-27 (Table 13 / 12C detail)"
if desc not in have:
    rr = ix.max_row + 1; ix.cell(rr, 1, "ITC"); ix.cell(rr, 2, desc)
    x = ix.cell(rr, 7); x.value = '=HYPERLINK("#\'ITC Register 2026-27\'!A1","ITC Register 2026-27")'; x.font = Font(color="0563C1", underline="single")
    for cc in range(1, 11): ix.cell(rr, cc).border = BD
for rr in range(5, ix.max_row + 1):
    v = S(ix.cell(rr, 2).value)
    if v.startswith("ITC Master Summary"): ix.cell(rr, 2).value = "ITC Master Summary - Tables 6A1..8D + 12A-12F/13 reconciliation per GSTIN (last year's 129-col format, live)"
    if v.startswith("Table 13 (FY 24-25)"): ix.cell(rr, 2).value = "Table 13 (FY 24-25) vs Table 6A1 (FY 25-26) differences - last year's format"
wb.save(P); print("saved")
