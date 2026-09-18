"""Replace 'S6 Advances Control' + 'S6 Month-on-Month' with ONE sheet in the CA's 'Advance Format VEL.xlsx'
layout (states across, 3 cols each: Taxable value | CGST | SGST; rows: Opening / Add: Advance received
(12 months) / Less: Advance adjusted (12 months) / Closing / Net GSTR-1 Advances / Diff; Remarks at end).
All figures live SUMIFS over SR_2025-26 (Adv Bucket helper BI, month A, GSTIN C) and GSTR-1 Data.
Opening = closing 31.03.25 per last year's advance working (the format file's cached closing row).
Row 1 reserved; GSTIN key placed in the format's blank row 4 so S9 can look up by GSTIN."""
import datetime as dt, warnings
from copy import copy
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from openpyxl.styles import Font, Alignment
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
F = r"C:\Users\pawar\Downloads\Advance Format VEL.xlsx"
f = open(P, "r+b"); f.close()
fmt = openpyxl.load_workbook(F)["Final"]
fmtv = openpyxl.load_workbook(F, data_only=True)["Final"]
# opening 01.04.25 = last year's closing row 27 (cached values), keyed by state name in row 2
OPEN = {}
for c in range(2, 20, 3):
    st = fmtv.cell(2, c).value
    if st: OPEN[st] = (float(fmtv.cell(27, c).value or 0), float(fmtv.cell(27, c + 1).value or 0), float(fmtv.cell(27, c + 2).value or 0))
STATES = [("Arunachal Pradesh", "12AAECR0503Q1ZJ"), ("Bihar", "10AAECR0503Q1ZN"), ("Gujarat", "24AAECR0503Q1ZE"),
          ("Madhya Pradesh", "23AAECR0503Q1ZG"), ("Karnataka", "29AAECR0503Q1Z4"), ("Uttar Pradesh", "09AAECR0503Q1Z6"),
          ("Punjab", "03AAECR0503Q1ZI"), ("Rajasthan", "08AAECR0503Q1Z8")]
wb = openpyxl.load_workbook(P)
NM = "S6 Advances Control"
pos = wb.sheetnames.index(NM)
del wb[NM]; del wb["S6 Month-on-Month"]
ws = wb.create_sheet(NM, pos)
def sty(src, dst):
    dst.font = copy(src.font); dst.fill = copy(src.fill); dst.border = copy(src.border)
    dst.alignment = copy(src.alignment); dst.number_format = src.number_format
SR = lambda col: "'SR_2025-26'!$%s$6:$%s$27007" % (col, col)
G1 = lambda col: "'GSTR-1 Data'!$%s$3:$%s$2081" % (col, col)
NS = len(STATES); REM = 2 + 3 * NS          # remarks column index
MONTHS = [dt.datetime(2025 if m <= 12 else 2026, m if m <= 12 else m - 12, 1) for m in range(4, 16)]
R_OPEN, R_RECV, R_ADJ, R_CLOSE, R_NET, R_DIFF = 5, 7, 21, 35, 37, 38
# --- header rows (styles from format rows 2/3) ---
for k, (st, g) in enumerate(STATES):
    c0 = 2 + 3 * k
    ws.cell(2, c0, st); ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c0 + 2)
    for j, h in enumerate(("Taxable value", "CGST", "SGST")):
        sty(fmt.cell(2, 2 + j), ws.cell(2, c0 + j)); ws.cell(3, c0 + j, h); sty(fmt.cell(3, 2 + j), ws.cell(3, c0 + j))
        sty(fmt.cell(4, 2 + j), ws.cell(4, c0 + j))
    ws.cell(4, c0, g).font = Font(italic=True, size=8, color="808080")
ws.cell(3, 1, "Particulars"); sty(fmt.cell(3, 1), ws.cell(3, 1)); sty(fmt.cell(2, 1), ws.cell(2, 1))
ws.cell(4, 1, "GSTIN").font = Font(italic=True, size=8, color="808080")
ws.cell(3, REM, "Remarks"); sty(fmt.cell(3, 20), ws.cell(3, REM)); sty(fmt.cell(2, 20), ws.cell(2, REM))
def row_style(r, tmpl_row, label=None, label_nf=None):
    ws.cell(r, 1, label if label is not None else ws.cell(r, 1).value); sty(fmt.cell(tmpl_row, 1), ws.cell(r, 1))
    if label_nf: ws.cell(r, 1).number_format = label_nf
    for k in range(NS):
        for j in range(3): sty(fmt.cell(tmpl_row, 2 + j), ws.cell(r, 2 + 3 * k + j))
    sty(fmt.cell(tmpl_row, 20), ws.cell(r, REM))
# --- Opening ---
row_style(R_OPEN, 5, "Opening as on 01.04.25")
for k, (st, g) in enumerate(STATES):
    o = OPEN.get(st, (0.0, 0.0, 0.0))
    for j in range(3): ws.cell(R_OPEN, 2 + 3 * k + j, o[j])
ws.cell(R_OPEN, 1).comment = Comment("Externally audited input: closing as on 31.03.25 per the FY 24-25 advance working (Advance Format VEL.xlsx, row 'Closing as on 31.03.25'). Not derived in this workbook.", "DPS")
# --- Received block ---
row_style(R_RECV, 7, "Add : Advance received")
for k in range(NS):
    for j in range(3):
        c = 2 + 3 * k + j
        ws.cell(R_RECV, c, "=SUM(%s%d:%s%d)" % (L(c), R_RECV + 1, L(c), R_RECV + 12))
def month_rows(r0, bucket, tmpl):
    for i, m in enumerate(MONTHS):
        r = r0 + i
        row_style(r, tmpl); ws.cell(r, 1, m); ws.cell(r, 1).number_format = "mmm-yy"
        for k, (st, g) in enumerate(STATES):
            c0 = 2 + 3 * k
            for j, col in enumerate(("T", "V", "W")):
                ws.cell(r, c0 + j, '=SUMIFS(%s,%s,%s$4,%s,"%s",%s,TEXT($A%d,"mmm-yy"))' % (SR(col), SR("C"), L(c0), SR("BI"), bucket, SR("A"), r))
                ws.cell(r, c0 + j).number_format = "#,##0"
month_rows(R_RECV + 1, "Received", 8)
# --- Adjusted block ---
row_style(R_ADJ, 16, "Less: Advance adjusted")
for k in range(NS):
    for j in range(3):
        c = 2 + 3 * k + j
        ws.cell(R_ADJ, c, "=SUM(%s%d:%s%d)" % (L(c), R_ADJ + 1, L(c), R_ADJ + 12))
month_rows(R_ADJ + 1, "Adjusted", 17)
# --- Closing / Net GSTR-1 / Diff ---
row_style(R_CLOSE, 27, "Closing as on 31.03.26")
row_style(R_NET, 29, "Net GSTR-1 Advances")
row_style(R_DIFF, 30, "Diff")
for k, (st, g) in enumerate(STATES):
    c0 = 2 + 3 * k
    for j, g1col in enumerate(("V", "X", "Y")):
        c = c0 + j; cl = L(c)
        ws.cell(R_CLOSE, c, "=%s%d+%s%d+%s%d" % (cl, R_OPEN, cl, R_RECV, cl, R_ADJ))
        ws.cell(R_NET, c, '=SUMIFS(%s,%s,%s$4,%s,"SalesSummary-Net (summary)",%s,"Advance*")' % (G1(g1col), G1("A"), L(c0), G1("AL"), G1("AM")))
        ws.cell(R_NET, c).number_format = "#,##0"
        ws.cell(R_DIFF, c, "=%s%d-%s%d-%s%d" % (cl, R_NET, cl, R_RECV, cl, R_ADJ)); ws.cell(R_DIFF, c).number_format = "#,##0"
# widths from the format's repeating pattern
ws.column_dimensions["A"].width = fmt.column_dimensions["A"].width
for k in range(NS):
    for j, w in enumerate((12.6, 11.5, 11.5)): ws.column_dimensions[L(2 + 3 * k + j)].width = w
ws.column_dimensions[L(REM)].width = 65
ws.freeze_panes = "B5"
# --- S9 links: opening (row 35) / closing (row 36) by GSTIN on row 4 ---
s9 = wb["S9 Sales Reco"]
for c in range(2, 21):
    cl = L(c)
    s9.cell(35, c).value = "=-IFERROR(INDEX('%s'!$B$%d:$%s$%d,MATCH(%s$7,'%s'!$B$4:$%s$4,0)),0)" % (NM, R_OPEN, L(REM - 1), R_OPEN, cl, NM, L(REM - 1))
    s9.cell(36, c).value = "=IFERROR(INDEX('%s'!$B$%d:$%s$%d,MATCH(%s$7,'%s'!$B$4:$%s$4,0)),0)" % (NM, R_CLOSE, L(REM - 1), R_CLOSE, cl, NM, L(REM - 1))
# --- INDEX: drop the Month-on-Month row ---
ix = wb["INDEX"]
for r in range(ix.max_row, 1, -1):
    if str(ix.cell(r, 2).value or "").strip() == "Advances Control - Month on Month":
        ix.delete_rows(r, 1); print("INDEX row removed:", r)
for r in range(2, ix.max_row + 1):
    if str(ix.cell(r, 2).value or "").strip() == "Advances Control Account":
        ix.cell(r, 2).value = "Advances Control Account (CA format - state-wise, month rows)"
wb.save(P); print("saved; S6 rebuilt with %d states, remarks col %s" % (NS, L(REM)))
