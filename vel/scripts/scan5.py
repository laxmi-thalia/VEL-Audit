import openpyxl, os, re, hashlib, sys

BASE = "//server/GST FOLDER/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/Clients Data/05.08.2026 Main Data"
months = sorted(d for d in os.listdir(BASE) if re.match(r"^\d\d ", d))
start, end = int(sys.argv[1]), int(sys.argv[2])

def hdr_sig(hdr):
    return hashlib.sha1("|".join(str(h) for h in hdr).encode()).hexdigest()[:12]

for m in months[start:end]:
    d = os.path.join(BASE, m)
    cands = [f for f in os.listdir(d)
             if re.search(r"clear ?tax", f, re.I) and re.search(r"sale", f, re.I)
             and not f.startswith("~$")]
    print(f"\n### {m}   candidates={cands}")
    for f in cands:
        p = os.path.join(d, f)
        if f.lower().endswith(".xls"):
            print(f"   {f}: LEGACY .xls  (needs xlrd/convert) -- skipped this pass")
            continue
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        except Exception as e:
            print(f"   {f}: LOAD FAILED {type(e).__name__}: {e}")
            continue
        print(f"   {f}: sheets={wb.sheetnames}")
        for ws in wb.worksheets:
            if not (ws.title.lower().startswith("ct_") or "detail" in ws.title.lower()):
                continue
            hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            hdr = [h for h in hdr if h is not None]
            print(f"      -> {ws.title!r}: data_rows={ws.max_row-1} cols={ws.max_column} "
                  f"headers={len(hdr)} sig={hdr_sig(hdr)}")
        wb.close()
