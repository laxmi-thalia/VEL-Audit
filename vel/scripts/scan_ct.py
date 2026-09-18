import openpyxl, os, glob, re, json
from openpyxl.utils import get_column_letter

BASE = "//server/GST FOLDER/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/Clients Data/05.08.2026 Main Data"
months = sorted(d for d in os.listdir(BASE) if re.match(r"^\d\d ", d))
report = {}
for m in months:
    d = os.path.join(BASE, m)
    cands = [f for f in os.listdir(d)
             if re.search(r"clear ?tax", f, re.I) and re.search(r"sale", f, re.I)
             and not f.startswith("~$")]
    report[m] = cands
    print(f"{m}: {cands}")
print("\n" + "="*70)
for m in months:
    for f in report[m]:
        p = os.path.join(BASE, m, f)
        if not f.lower().endswith(".xlsx"):
            print(f"\n### {m} / {f}  -> LEGACY .xls, openpyxl cannot read")
            continue
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        except Exception as e:
            print(f"\n### {m} / {f} -> LOAD FAILED: {type(e).__name__}: {e}")
            continue
        det = [s for s in wb.sheetnames if "detail" in s.lower() or s.lower().startswith("ct_")]
        print(f"\n### {m} / {f}")
        print(f"    sheets: {wb.sheetnames}")
        if det:
            ws = wb[det[0]]
            hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            hdr = [h for h in hdr if h is not None]
            print(f"    detail sheet {det[0]!r}: rows={ws.max_row} cols={ws.max_column} headers={len(hdr)}")
            print(f"    HDRHASH={hash(tuple(hdr))}")
        wb.close()
