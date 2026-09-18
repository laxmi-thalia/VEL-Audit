"""Step 1 fixes from CA review session 2 (excluding the flagging/gate layer).

Applied IN PLACE on the existing register workbook:
 1. Supplier Legal Name (D) autofilled from own GSTIN.
 2. Place of Supply (L) zero-padded to a 2-digit text code.
 3. NEW column inserted after P: 'Actual Rate Charged (%)' = Total GST / Taxable.
    P 'GST rate check' now tolerates rounding (|actual - stated| <= 0.1).
 4. POS check rewritten per the CA's rule (code difference 0 -> CGST/SGST, else IGST),
    sign-agnostic so credit notes stop failing; zero-tax rows -> NIL.
 5. 'GSTR-1 Type', 'Matched with GSTR-1', 'GSTR-1 Month' are now populated FROM the
    GSTR-1 match (Step 2 results), not derived from the register's own doc type.
"""
import os, warnings; warnings.filterwarnings("ignore")
import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from copy import copy

SP = os.path.dirname(os.path.abspath(__file__))
P = r"C:\Users\pawar\Downloads\VEL_Sales_Register_FY2025-26_DRAFT.xlsx"
R0, R1 = 5, 27006

def s(x):
    return "" if x is None or (isinstance(x, float) and pd.isna(x)) else str(x).strip()

# ---------- GSTR-1 match results -> per document (type, status, month) ----------
M = pd.read_pickle(os.path.join(SP, "step2_match2.pkl"))
G = pd.read_pickle(os.path.join(SP, "gstr1.pkl"))
g1idx = {}
for _, r in G.iterrows():
    g1idx[(s(r["Company GSTIN"]).upper(), s(r["Doc No"]).upper())] = (s(r["Doc Type"]), r["Tax Period"])
TYPE_LABEL = {"Invoice": "B2B", "Credit Note": "CDNR (Credit Note)", "Debit Note": "CDNR (Debit Note)"}
doc = {}
for _, r in M.iterrows():
    if r["status"] == "IN GSTR-1 ONLY":
        continue
    key = (s(r["gstin"]).upper(), s(r["docno"]).upper(), s(r["reg_type"]))
    if r["status"] == "IN BOOKS ONLY":
        doc[key] = ("Not in GSTR-1", "Not in GSTR-1", "")
        continue
    g = g1idx.get((s(r["gstin"]).upper(), s(r["g1_docno"]).upper()))
    gtype = TYPE_LABEL.get(g[0], g[0]) if g else ""
    month = pd.to_datetime(g[1]).strftime("%b-%y") if g and pd.notna(g[1]) else ""
    doc[key] = (gtype, "Matched (%s)" % r["matched_by"], month)

# ---------- open register ----------
wb = openpyxl.load_workbook(P)
ws = wb["SR_2025-26"]

# supplier names by own GSTIN (from rows that have one)
name_by_gstin = {}
for r in ws.iter_rows(min_row=R0, max_row=R1, min_col=3, max_col=4, values_only=True):
    if s(r[0]) and s(r[1]) and s(r[0]) not in name_by_gstin:
        name_by_gstin[s(r[0])] = s(r[1])

# ---------- insert the new column after P (col 16) ----------
ws.insert_cols(17)
hdr_src = ws.cell(4, 16)
h = ws.cell(4, 17, "Actual Rate Charged (%)")
h.font = copy(hdr_src.font); h.fill = copy(hdr_src.fill); h.alignment = copy(hdr_src.alignment); h.border = copy(hdr_src.border)
ws.column_dimensions["Q"].width = 14

# new letters after the insert:
# A Month B State C GSTIN D Supplier E DocDate F DocNo G InvType H DocType I SupplyType J RecipGSTIN K RecipName
# L POS M ItemPrice N Gross O Rate P RateCheck Q ACTUAL R Taxable S IGST T CGST U SGST V Cess W TotalGST X TotalInv
# Y IRN Z IRNlen AA Status AB CancelDate AC DocStatus AD ShipGSTIN AE ShipName AF ItemDesc AG G/S AH HSN AI DescUpl
# AJ Qty AK UoM AL GLName AM PC AN CC AO VEL2 AP Cust2 AQ POScheck AR GSTR1Type AS MatchedG1 AT G1Month
# AU MatchedGL AV MatchedFS AW GSTR9 AX Queries AY SourceMonth AZ BusinessPlace BA SAPMatchedBy
C = {"D": 4, "H": 8, "I": 9, "J": 10, "L": 12, "P": 16, "Q": 17, "R": 18, "W": 23, "X": 24, "Y": 25, "Z": 26,
     "AG": 33, "AH": 34, "AO": 41, "AP": 42, "AQ": 43, "AR": 44, "AS": 45, "AT": 46, "AW": 49}

filled_names = padded = 0
for i in range(R0, R1 + 1):
    gstin = s(ws.cell(i, 3).value)
    dtc = s(ws.cell(i, C["H"]).value)
    sup = s(ws.cell(i, C["I"]).value)
    docno = s(ws.cell(i, 6).value)
    is_adv = dtc.startswith("MOB")
    # 1. supplier name autofill
    if not s(ws.cell(i, C["D"]).value) and gstin in name_by_gstin:
        ws.cell(i, C["D"]).value = name_by_gstin[gstin]; filled_names += 1
    # 2. place of supply -> 2-digit text
    pos = ws.cell(i, C["L"]).value
    if pos not in (None, ""):
        ps = s(pos)
        if ps.replace(".", "", 1).isdigit():
            ps = str(int(float(ps))).zfill(2)
            if ws.cell(i, C["L"]).value != ps:
                padded += 1
            ws.cell(i, C["L"]).value = ps
            ws.cell(i, C["L"]).number_format = "@"
    # 3. actual rate + tolerant check
    ws.cell(i, C["Q"]).value = "=IF(N(R{0})=0,\"\",ROUND(W{0}/R{0}*100,2))".format(i)
    ws.cell(i, C["Q"]).number_format = "0.00"
    ws.cell(i, C["P"]).value = "=IF(Q{0}=\"\",\"\",IF(ABS(Q{0}-O{0})<=0.1,\"OK\",\"CHECK\"))".format(i)
    # re-point every formula column to the new letters
    ws.cell(i, C["W"]).value = "=SUM(S{0}:U{0})".format(i)
    ws.cell(i, C["X"]).value = "=SUM(R{0}:V{0})".format(i)
    ws.cell(i, C["Z"]).value = "=IF(Y{0}=\"\",\"\",LEN(Y{0}))".format(i)
    ws.cell(i, C["AG"]).value = "=IF(AH{0}=\"\",\"\",IF(LEFT(AH{0}&\"\",2)=\"99\",\"S\",\"G\"))".format(i)
    ws.cell(i, C["AO"]).value = "=LEFT(C{0},2)".format(i)
    ws.cell(i, C["AP"]).value = "=IF(J{0}=\"\",\"\",LEFT(J{0},2))".format(i)
    # 4. POS check per CA rule, sign-agnostic
    ws.cell(i, C["AQ"]).value = (
        "=IF(OR(LEFT(H{0},3)=\"MOB\",J{0}=\"\"),\"NA\","
        "IF(AND(S{0}=0,T{0}=0,U{0}=0),\"NIL\","
        "IF(VALUE(AO{0})-VALUE(AP{0})=0,"
        "IF(AND(S{0}=0,OR(T{0}<>0,U{0}<>0)),\"OK\",\"CHECK\"),"
        "IF(AND(S{0}<>0,T{0}=0,U{0}=0),\"OK\",\"CHECK\"))))").format(i)
    # 5. GSTR-1 type / matched / month FROM GSTR-1
    if is_adv:
        vals = ("Advance (Table 11, summary)", "Summary level - see month-on-month", "")
    elif dtc == "INV" and sup == "B2C":
        vals = ("B2CS (summary)", "Summary level - see month-on-month", "")
    else:
        vals = doc.get((gstin.upper(), docno.upper(), dtc), ("", "Not in GSTR-1", ""))
    ws.cell(i, C["AR"]).value, ws.cell(i, C["AS"]).value, ws.cell(i, C["AT"]).value = vals
    ws.cell(i, C["AW"]).value = ("=IF(H{0}=\"INV\",IF(I{0}=\"B2C\",\"4A\",\"4B\"),"
                                 "IF(H{0}=\"CRN\",\"4I\",IF(H{0}=\"DBN\",\"4J\",\"4F\")))").format(i)

# row-2 SUBTOTALs on the shifted letters
for col in ["R", "S", "T", "U", "V", "W", "AJ"]:
    ws["%s2" % col] = "=SUBTOTAL(9,%s%d:%s%d)" % (col, R0, col, R1)
for col in ["Q"]:
    ws["%s2" % col] = None
wb.save(P)
print("saved | supplier names filled:", filled_names, "| POS codes zero-padded:", padded,
      "| documents with GSTR-1 lookup:", len(doc))
