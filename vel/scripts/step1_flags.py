"""Step 1 data-quality flags + gate (CA session 2).

Three outputs, per the CA: (1) the offending CELL is highlighted in the register,
(2) a remark on the row in a new 'Data Flags' column, (3) a listing sheet 'Step 1 Flags'
whose summary is LIVE (COUNTIF over the register) and which carries the GATE:
  READY FOR STEP 2 when no BLOCKING flag remains, or when the CA types YES in the override cell.

BLOCKING (data needed to reconcile is missing / contradictory):
  - Document Date missing                   [exempt: advances]
  - Recipient GSTIN missing on B2B document [exempt: advances, B2C]
  - IRN missing on B2B document             [exempt: advances, B2C]  "as good as invoice not issued"
  - Advance row tagged B2B / doc-type vs supply-type contradiction
  - Taxable Value missing
REVIEW (CA judgment, does not block):
  - IRN length <> 64 (when present)
  - GST rate charged <> stated rate (rate check = CHECK)
  - HSN missing on a non-advance row
  - Quantity / UoM missing on a GOODS row (services legitimately blank)
  - Recipient name missing (GSTIN present)
"""
import collections, warnings; warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from copy import copy

P = r"C:\Users\pawar\Downloads\VEL_Sales_Register_FY2025-26_DRAFT.xlsx"
R0, R1 = 5, 27006
RED = PatternFill("solid", fgColor="FFC7CE"); AMB = PatternFill("solid", fgColor="FFEB9C")
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79")
def S(v): return "" if v is None else str(v).strip()

wv = openpyxl.load_workbook(P, read_only=True, data_only=True)["SR_2025-26"]
cached = {i: r for i, r in enumerate(wv.iter_rows(min_row=R0, max_row=R1, values_only=True), start=R0)}
wb = openpyxl.load_workbook(P); ws = wb["SR_2025-26"]
H = {S(ws.cell(4, c).value): c for c in range(1, ws.max_column + 1)}

# new column 'Data Flags' right after 'Sample POs'
cFlag = H["Sample POs"] + 1
ws.insert_cols(cFlag)
h = ws.cell(4, cFlag, "Data Flags"); src = ws.cell(4, cFlag - 1)
h.font = copy(src.font); h.fill = copy(src.fill); h.alignment = copy(src.alignment); h.border = copy(src.border)
ws.column_dimensions[get_column_letter(cFlag)].width = 48
H = {S(ws.cell(4, c).value): c for c in range(1, ws.max_column + 1)}   # re-read after insert
c = H
CHECKS = [  # (name, level, column highlighted)
 ("Document Date missing", "BLOCKING", "Document Date"),
 ("Recipient GSTIN missing on B2B document", "BLOCKING", "Recipient GSTIN"),
 ("IRN missing on B2B document", "BLOCKING", "Invoice Reference No"),
 ("Advance / B2B tag contradiction", "BLOCKING", "Supply Type Code"),
 ("Taxable Value missing", "BLOCKING", "Taxable Value"),
 ("IRN length <> 64", "REVIEW", "Invoice Reference No"),
 ("GST rate charged <> stated rate", "REVIEW", "GST rate check"),
 ("HSN missing", "REVIEW", "HSN or SAC Code"),
 ("Quantity / UoM missing on goods row", "REVIEW", "Quantity for uploading"),
 ("Recipient name missing", "REVIEW", "Recipient Legal Name"),
]
LEVEL = {n: l for n, l, _ in CHECKS}; COLOF = {n: col for n, _, col in CHECKS}
detail = []; counts = collections.Counter()
for i in range(R0, R1 + 1):
    v = lambda name: ws.cell(i, c[name]).value
    cv = lambda name: cached[i][c[name] - 1 - (1 if c[name] > cFlag else 0)]   # cached row predates the insert
    dtc, sup = S(v("Document Type Code")), S(v("Supply Type Code"))
    is_adv = dtc.startswith("MOB"); is_b2c = (dtc == "INV" and sup == "B2C")
    rg = S(v("Recipient GSTIN")); irn = S(v("Invoice Reference No"))
    flags = []
    if not is_adv and v("Document Date") in (None, ""): flags.append("Document Date missing")
    if not is_adv and not is_b2c and not rg: flags.append("Recipient GSTIN missing on B2B document")
    if not is_adv and not is_b2c and not irn: flags.append("IRN missing on B2B document")
    if (is_adv and sup in ("B2B", "B2C")) or (dtc in ("INV", "CRN", "DBN") and sup.startswith("MOB")):
        flags.append("Advance / B2B tag contradiction")
    if v("Taxable Value") in (None, ""): flags.append("Taxable Value missing")
    if irn and len(irn) != 64: flags.append("IRN length <> 64")
    if cv("GST rate check") == "CHECK": flags.append("GST rate charged <> stated rate")
    if not is_adv and v("HSN or SAC Code") in (None, ""): flags.append("HSN missing")
    if cv("GOOD (G) or SERVICE(S)") == "G" and (v("Quantity for uploading") in (None, "", 0) or v("Unit of Measurement for uploading") in (None, "")):
        flags.append("Quantity / UoM missing on goods row")
    if rg and not S(v("Recipient Legal Name")): flags.append("Recipient name missing")
    if flags:
        ws.cell(i, cFlag).value = "; ".join(flags)
        worst = "BLOCKING" if any(LEVEL[f] == "BLOCKING" for f in flags) else "REVIEW"
        ws.cell(i, cFlag).fill = RED if worst == "BLOCKING" else AMB
        for f in flags:
            ws.cell(i, c[COLOF[f]]).fill = RED if LEVEL[f] == "BLOCKING" else AMB
            counts[f] += 1
            detail.append((LEVEL[f], f, i, S(v("My State")), S(v("My GSTIN")), S(v("Document Number")), dtc, v("Taxable Value")))

# ---- listing sheet with LIVE summary + gate
if "Step 1 Findings" in wb.sheetnames: del wb["Step 1 Findings"]
if "Step 1 Flags" in wb.sheetnames: del wb["Step 1 Flags"]
sh = wb.create_sheet("Step 1 Flags", 1)
FL = get_column_letter(cFlag)
rng = "'SR_2025-26'!$%s$%d:$%s$%d" % (FL, R0, FL, R1)
sh["A1"] = "STEP 1 — DATA QUALITY GATE"; sh["A1"].font = Font(bold=True, size=13)
sh["A2"] = "Counts below are LIVE (COUNTIF over the register's 'Data Flags' column). Clear a flag in the register and it drops out here."
sh["A4"] = "Blocking flags outstanding"; sh["B4"] = "=SUM(B9:B13)"
sh["A5"] = "Review flags outstanding";   sh["B5"] = "=SUM(B14:B18)"
sh["A6"] = "CA override (type YES to proceed despite blocking flags)"; sh["B6"] = ""
sh["A7"] = "GATE"; sh["B7"] = '=IF(OR(B4=0,UPPER(B6)="YES"),"READY FOR STEP 2","BLOCKED - fix the blocking rows or override")'
for r in (4, 5, 6, 7): sh.cell(r, 1).font = Font(bold=True)
sh["B7"].font = Font(bold=True, size=12)
sh["A8"] = "Check"; sh["B8"] = "Rows"; sh["C8"] = "Level"; sh["D8"] = "Exempt"
for x in sh[8]: x.font = HF; x.fill = HB
EXEMPT = {"Document Date missing": "advances", "Recipient GSTIN missing on B2B document": "advances, B2C",
          "IRN missing on B2B document": "advances, B2C", "Advance / B2B tag contradiction": "-", "Taxable Value missing": "-",
          "IRN length <> 64": "blank IRN", "GST rate charged <> stated rate": "rounding within 0.1", "HSN missing": "advances",
          "Quantity / UoM missing on goods row": "services (SAC 99)", "Recipient name missing": "rows without a GSTIN"}
r = 9
for name, level, _ in CHECKS:
    sh.cell(r, 1, name); sh.cell(r, 2, '=COUNTIF(%s,"*%s*")' % (rng, name)); sh.cell(r, 3, level); sh.cell(r, 4, EXEMPT[name])
    sh.cell(r, 2).fill = RED if level == "BLOCKING" else AMB; r += 1
r += 1
sh.cell(r, 1, "DETAIL — one line per flag (static snapshot at build time)").font = Font(bold=True); r += 1
for j, t in enumerate(["Level", "Check", "Register row", "State", "My GSTIN", "Document No", "Doc Type", "Taxable"], 1):
    x = sh.cell(r, j, t); x.font = HF; x.fill = HB
r += 1
for d in sorted(detail, key=lambda d: (d[0] != "BLOCKING", d[1], d[2])):
    for j, val in enumerate(d, 1): sh.cell(r, j, val)
    sh.cell(r, 1).fill = RED if d[0] == "BLOCKING" else AMB; r += 1
for col_, w in zip("ABCDEFGH", [44, 12, 14, 20, 18, 18, 12, 16]): sh.column_dimensions[col_].width = w
sh.freeze_panes = "A9"
wb.save(P)
print("flags written | rows flagged:", sum(1 for i in range(R0, R1 + 1) if ws.cell(i, cFlag).value))
for name, level, _ in CHECKS: print("  %-8s %-42s %d" % (level, name, counts[name]))
