import os, warnings; warnings.filterwarnings("ignore")
import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill
SP=os.path.dirname(os.path.abspath(__file__))
R=pd.read_pickle(os.path.join(SP,"register.pkl")); G=pd.read_pickle(os.path.join(SP,"gstr1.pkl"))
M=pd.read_pickle(os.path.join(SP,"step2_match2.pkl"))
TY={"Invoice":"INV","Credit Note":"CRN","Debit Note":"DBN"}
def s(x): return "" if pd.isna(x) else str(x).strip().upper()
NOG1={"West Bengal"}            # states with no GSTR-1 file downloaded
# ---- month-on-month, GSTIN key, month falls back to the source file month
R["adv"]=R["dtc"].astype(str).str.startswith("MOB"); rb=R[~R["adv"]].copy()
mdate=pd.to_datetime(rb["ddate"],errors="coerce")
msrc =pd.to_datetime(rb["src"],format="%b-%y",errors="coerce")
rb["month"]=mdate.fillna(msrc).dt.strftime("%Y-%m")
rb["g"]=rb["gstin"].map(s)
fb=rb.groupby(["g","month","dtc"]).agg(taxable=("tax","sum"),igst=("igst","sum"),cgst=("cgst","sum"),sgst=("sgst","sum")).reset_index(); fb["source"]="Books"
G["dtc"]=G["Doc Type"].map(TY); G["g"]=G["Company GSTIN"].map(s)
G["month"]=pd.to_datetime(G["Tax Period"],errors="coerce").dt.strftime("%Y-%m")
for c in ["Taxable Value (Net)","IGST (Net)","CGST (Net)","SGST (Net)"]: G[c]=pd.to_numeric(G[c],errors="coerce").fillna(0)
fg=G.groupby(["g","month","dtc"]).agg(taxable=("Taxable Value (Net)","sum"),igst=("IGST (Net)","sum"),cgst=("CGST (Net)","sum"),sgst=("SGST (Net)","sum")).reset_index(); fg["source"]="GSTR-1"
lab=R.groupby(R["gstin"].map(s))["state"].first().to_dict()
for k,v in G.groupby("g")["__state"].first().items(): lab.setdefault(k,v)
piv=pd.concat([fb,fg]).pivot_table(index=["g","month","dtc"],columns="source",values=["taxable","igst","cgst","sgst"],aggfunc="sum",fill_value=0).reset_index()
piv.columns=[a if not b else f"{a}_{b}" for a,b in piv.columns]
for f in ["taxable","igst","cgst","sgst"]:
    for src in ["Books","GSTR-1"]:
        if f"{f}_{src}" not in piv: piv[f"{f}_{src}"]=0.0
    piv[f"{f}_DIFF"]=(piv[f"{f}_Books"]-piv[f"{f}_GSTR-1"]).round(2)
piv.insert(0,"state",piv["g"].map(lab))
piv["note"]=piv.apply(lambda r: "GSTR-1 report not downloaded - not a real difference"
                      if r["state"] in NOG1 else ("" if abs(r["taxable_DIFF"])<1 else "difference to investigate"),axis=1)
piv=piv.sort_values(["state","month","dtc"]); piv.to_pickle(os.path.join(SP,"mom2.pkl"))
real=piv[(piv["taxable_DIFF"].abs()>=1) & (~piv["state"].isin(NOG1))]
print(f"month-on-month rows {len(piv)} | diffs {int((piv['taxable_DIFF'].abs()>=1).sum())} | excluding WB: {len(real)}")

P=r"C:\Users\pawar\Downloads\VEL_Step2_Reco_GSTR1_DRAFT.xlsx"
wb=openpyxl.load_workbook(P)
for n in ["Month-on-Month","Exceptions","Reconciliation Summary","Matched","Pivot - Month on Month"]:
    if n in wb.sheetnames: del wb[n]
H=Font(bold=True,color="FFFFFF"); F=PatternFill("solid",fgColor="1F4E79")
OK=PatternFill("solid",fgColor="C6EFCE"); BAD=PatternFill("solid",fgColor="FFC7CE"); WARN=PatternFill("solid",fgColor="FFF2CC")
ws=wb.create_sheet("Reconciliation Summary",0)
tb=M["reg_taxable"].sum(); tg=M["g1_taxable"].sum()
for r in [["STEP 2 - Sales Register vs GSTR-1 (FY 2025-26)","",""],["","",""],
          ["Books - taxable (excluding advances)",round(tb,2),""],
          ["GSTR-1 - taxable (net of amendments)",round(tg,2),""],
          ["DIFFERENCE",round(tb-tg,2),""],["","",""],["RECONCILED BY:","documents","taxable"]]:
    ws.append(r)
for st_ in ["MATCHED","IN BOOKS ONLY","IN GSTR-1 ONLY"]:
    X=M[M.status==st_]
    v=X["reg_taxable"].sum() if st_=="IN BOOKS ONLY" else (X["g1_taxable"].sum() if st_=="IN GSTR-1 ONLY" else 0)
    ws.append([st_,len(X),round(v,2)])
ws.append(["","",""])
for k,v in M[M.status.str.startswith("MATCHED")]["matched_by"].value_counts().items(): ws.append([f"  matched by {k}",int(v),""])
for c in ws[1]: c.font=H; c.fill=F
ws.column_dimensions["A"].width=52; ws.column_dimensions["B"].width=18; ws.column_dimensions["C"].width=20
e=wb.create_sheet("Exceptions")
e.append(["Status","State","My GSTIN","Books Doc No","Books type","Books taxable","GSTR-1 Doc No","GSTR-1 taxable","Assessment"])
for _,r in M[M.status!="MATCHED"].sort_values(["status","state"]).iterrows():
    a=("West Bengal - GSTR-1 report not downloaded; already proven filed via Tax Comparison Report" if r["state"]=="West Bengal"
       else "Credit note in GSTR-1, no IRN, no entry in books - REAL" if r["status"]=="IN GSTR-1 ONLY"
       else "e-invoiced (IRN exists) but never declared in GSTR-1 - REAL" )
    e.append([r["status"],r["state"],r["gstin"],r["docno"],r["reg_type"],r["reg_taxable"],r["g1_docno"],r["g1_taxable"],a])
for c in e[1]: c.font=H; c.fill=F
for row in e.iter_rows(min_row=2): 
    for c in row: c.fill = WARN if "West Bengal" in str(row[1].value) else BAD
for col,w in zip("ABCDEFGHI",[18,18,20,18,11,17,18,17,64]): e.column_dimensions[col].width=w
e.freeze_panes="A2"
m=wb.create_sheet("Month-on-Month")
m.append(["State","My GSTIN","GSTR-1 Month","Doc Type","Taxable Books","Taxable GSTR-1","Taxable DIFF","IGST DIFF","CGST DIFF","SGST DIFF","Note"])
for _,r in piv.iterrows():
    m.append([r["state"],r["g"],r["month"],r["dtc"],round(r["taxable_Books"],2),round(r["taxable_GSTR-1"],2),
              r["taxable_DIFF"],r["igst_DIFF"],r["cgst_DIFF"],r["sgst_DIFF"],r["note"]])
for c in m[1]: c.font=H; c.fill=F
for row in m.iter_rows(min_row=2):
    row[6].fill = OK if abs(row[6].value or 0)<1 else (WARN if "not downloaded" in str(row[10].value) else BAD)
for col,w in zip("ABCDEFGHIJK",[20,20,14,10,18,18,16,14,14,14,52]): m.column_dimensions[col].width=w
m.freeze_panes="A2"
mt=wb.create_sheet("Matched")
mt.append(["State","My GSTIN","Books Doc No","GSTR-1 Doc No","Type","Books taxable","GSTR-1 taxable","Matched by"])
for _,r in M[M.status=="MATCHED"].iterrows():
    mt.append([r["state"],r["gstin"],r["docno"],r["g1_docno"],r["reg_type"],r["reg_taxable"],r["g1_taxable"],r["matched_by"]])
for c in mt[1]: c.font=H; c.fill=F
for col,w in zip("ABCDEFGH",[20,20,18,18,8,17,17,18]): mt.column_dimensions[col].width=w
mt.freeze_panes="A2"
wb.save(P); print("workbook rebuilt")
