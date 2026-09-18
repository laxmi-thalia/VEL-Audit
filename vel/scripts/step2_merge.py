"""Step 2 inside the register workbook (CA session 2).
Sheets added: 'GSTR-1 Data' (portal rows), 'SR vs GSTR-1' (CA 3-block layout, every cell a
live SUMIFS over SR_2025-26 / GSTR-1 Data), 'Month-on-Month 1' (bird's-eye: state -> 12
months down the rows, Books & GSTR-1 taxable+I/C/S, live), 'Step 2 Summary' (live, with a
pointer to the exception listing), 'Step 2 Exceptions' (listing). No 'Matched' sheet: the
per-row result lives in the register's 'Matched with GSTR-1' column."""
import os, warnings; warnings.filterwarnings("ignore")
import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
SP = os.path.dirname(os.path.abspath(__file__))
P = r"C:\Users\pawar\Downloads\VEL_Sales_Register_FY2025-26_DRAFT.xlsx"
R0, R1 = 5, 27006
def s(x): return "" if x is None or (isinstance(x, float) and pd.isna(x)) else str(x).strip()
G = pd.read_pickle(os.path.join(SP, "gstr1.pkl")); Su = pd.read_pickle(os.path.join(SP, "g1_summary.pkl"))
M = pd.read_pickle(os.path.join(SP, "step2_match2.pkl"))
for c_ in ["Taxable Value (Net)", "IGST (Net)", "CGST (Net)", "SGST (Net)"]:
    G[c_] = pd.to_numeric(G[c_], errors="coerce").fillna(0); Su[c_] = pd.to_numeric(Su[c_], errors="coerce").fillna(0)
GSTINS = [("01AAECR0503Q1ZM", "Jammu & Kashmir"), ("03AAECR0503Q1ZI", "Punjab"), ("06AAECR0503Q1ZC", "Haryana"),
 ("08AAECR0503Q1Z8", "Rajasthan"), ("10AAECR0503Q1ZN", "Bihar"), ("12AAECR0503Q1ZJ", "Arunachal Pradesh"),
 ("18AAECR0503Q1Z7", "Assam"), ("19AAECR0503Q1Z5", "West Bengal"), ("20AAECR0503Q1ZM", "Jharkhand"),
 ("23AAECR0503Q1ZG", "Madhya Pradesh"), ("24AAECR0503Q1ZE", "Gujarat"), ("27AAECR0503Q1Z8", "Maharashtra"),
 ("32AAECR0503Q1ZH", "Kerala"), ("33AAECR0503Q1ZF", "TamilNadu"), ("36AAECR0503Q1Z9", "Telangana"),
 ("37AAECR0503Q1Z7", "Andhra Pradesh"), ("09AAECR0503Q1Z6", "Uttar Pradesh"), ("22AAECR0503Q1ZI", "Chhattisgarh"),
 ("29AAECR0503Q1Z4", "Karnataka")]
MONTHS = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26"]
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79"); GF = PatternFill("solid", fgColor="DDEBF7")
TOT = Font(bold=True); thin = Side(style="thin", color="B0B0B0"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
OK = PatternFill("solid", fgColor="C6EFCE"); BAD = PatternFill("solid", fgColor="FFC7CE")

wb = openpyxl.load_workbook(P); ws = wb["SR_2025-26"]
H = {s(ws.cell(4, c).value): L(c) for c in range(1, ws.max_column + 1)}
# register column letters used by the SUMIFS
cM, cG_, cH, cI = H["Month"], H["My GSTIN"], H["Document Type Code"], H["Supply Type Code"]
cR, cS, cT, cU = H["Taxable Value"], H["IGST Amount"], H["CGST Amount"], H["SGST Amount"]
SRm = {0: cR, 1: cS, 2: cT, 3: cU}
def SR(col): return "'SR_2025-26'!$%s$%d:$%s$%d" % (col, R0, col, R1)
for n in ["GSTR-1 Data", "SR vs GSTR-1", "Month-on-Month 1", "Step 2 Summary", "Step 2 Exceptions"]:
    if n in wb.sheetnames: del wb[n]

# ---------- GSTR-1 Data ----------
g1 = wb.create_sheet("GSTR-1 Data")
g1.append(["My GSTIN", "Tax Period", "Source", "Raw Type", "Document Number", "Taxable (Net)", "IGST (Net)", "CGST (Net)", "SGST (Net)", "Bucket (formula)", "Month (formula)"])
for c in g1[1]: c.font = HF; c.fill = HB
BF = ('=IF(D{n}="Invoice","B2B",IF(D{n}="Credit Note","Credit note",IF(D{n}="Debit Note","Debit note",'
      'IF(D{n}="B2CS Sales","B2C",IF(D{n}="Advance Received","Advance received","Advance adjusted")))))')
j = 1
for _, r in G.iterrows():
    j += 1; g1.append([s(r["Company GSTIN"]), r["Tax Period"], "Sales-Net (document)", s(r["Doc Type"]), s(r["Doc No"]),
                       r["Taxable Value (Net)"], r["IGST (Net)"], r["CGST (Net)"], r["SGST (Net)"], BF.format(n=j), '=TEXT(B{0},"mmm-yy")'.format(j)])
for _, r in Su.iterrows():
    j += 1; g1.append([s(r["Company GSTIN"]), r["Tax Period"], "SalesSummary-Net (summary)", s(r["Summary Type"]), "",
                       r["Taxable Value (Net)"], r["IGST (Net)"], r["CGST (Net)"], r["SGST (Net)"], BF.format(n=j), '=TEXT(B{0},"mmm-yy")'.format(j)])
NG = j
for c_, w in zip("ABCDEFGHIJK", [18, 12, 24, 16, 18, 15, 12, 12, 12, 18, 12]): g1.column_dimensions[c_].width = w
g1.freeze_panes = "A2"
G1m = {0: "F", 1: "G", 2: "H", 3: "I"}
def G1(col): return "'GSTR-1 Data'!$%s$2:$%s$%d" % (col, col, NG)

# books bucket SUMIFS over the register (REV split by sign = CA rule)
def books(mi, gcell, bucket, extra=""):
    m = SRm[mi]
    base = "SUMIFS(%s,%s,%s%s" % (SR(m), SR(cG_), gcell, extra)
    if bucket == "B2B": return base + ',%s,"INV",%s,"<>B2C")' % (SR(cH), SR(cI))
    if bucket == "B2C": return base + ',%s,"INV",%s,"B2C")' % (SR(cH), SR(cI))
    if bucket == "Credit note": return base + ',%s,"CRN")' % SR(cH)
    if bucket == "Debit note": return base + ',%s,"DBN")' % SR(cH)
    if bucket == "Advance received": return base + ',%s,"MOB ADV REC")+' % SR(cH) + base + ',%s,"MOB ADV REV",%s,">0")' % (SR(cH), SR(cR))
    if bucket == "Advance adjusted": return base + ',%s,"MOB ADV ADJ")+' % SR(cH) + base + ',%s,"MOB ADV REV",%s,"<0")' % (SR(cH), SR(cR))
def g1sum(mi, gcell, bucket, extra=""):
    return 'SUMIFS(%s,%s,%s,%s,"%s"%s)' % (G1(G1m[mi]), G1("A"), gcell, G1("J"), bucket, extra)

# ---------- SR vs GSTR-1 (CA format) ----------
ws2 = wb.create_sheet("SR vs GSTR-1", 2)
ws2["A1"] = "Vikran Engineering Limited"; ws2["A1"].font = Font(bold=True, size=13)
ws2["A2"] = "Sales Register vs GSTR-1 (FY 2025-26) — every cell is a live SUMIFS over SR_2025-26 and GSTR-1 Data"
ws2["A3"] = "Step 1 gate:"; ws2["B3"] = "='Step 1 Flags'!B7"; ws2["B3"].font = Font(bold=True)
GROUPS = ["B2B", "B2C", "Credit note", "Debit note", "Advance received", "Advance adjusted", "Total liability"]
MEAS = ["Taxable value", "IGST", "CGST", "SGST"]
def header(r0, title):
    ws2.cell(r0, 1, title).font = Font(bold=True, size=12)
    for gi, gn in enumerate(GROUPS):
        c = 3 + gi * 4; ws2.cell(r0 + 1, c, gn).font = Font(bold=True)
        ws2.cell(r0 + 1, c).alignment = Alignment(horizontal="center")
        ws2.merge_cells(start_row=r0 + 1, start_column=c, end_row=r0 + 1, end_column=c + 3); ws2.cell(r0 + 1, c).fill = GF
    ws2.cell(r0 + 2, 1, "GSTIN"); ws2.cell(r0 + 2, 2, "State")
    for gi in range(7):
        for mi, mn in enumerate(MEAS): ws2.cell(r0 + 2, 3 + gi * 4 + mi, mn)
    for c in range(1, 31):
        x = ws2.cell(r0 + 2, c); x.font = HF; x.fill = HB; x.alignment = Alignment(horizontal="center", wrap_text=True)
def block(r0, title, kind):
    header(r0, title); first = r0 + 3
    for i2, (g, st) in enumerate(GSTINS):
        r = first + i2; ws2.cell(r, 1, g); ws2.cell(r, 2, st)
        for gi, gn in enumerate(GROUPS):
            for mi in range(4):
                c = 3 + gi * 4 + mi
                if gn == "Total liability": f_ = "=" + "+".join("%s%d" % (L(3 + k * 4 + mi), r) for k in range(6))
                elif kind == "G1": f_ = "=" + g1sum(mi, "$A%d" % r, gn)
                else: f_ = "=" + books(mi, "$A%d" % r, gn)
                ws2.cell(r, c, f_); ws2.cell(r, c).number_format = "#,##0.00"; ws2.cell(r, c).border = BD
        ws2.cell(r, 1).border = BD; ws2.cell(r, 2).border = BD
    tr = first + len(GSTINS); ws2.cell(tr, 1, "Total").font = TOT
    for c in range(3, 31):
        ws2.cell(tr, c, "=SUM(%s%d:%s%d)" % (L(c), first, L(c), tr - 1)); ws2.cell(tr, c).font = TOT
        ws2.cell(tr, c).number_format = "#,##0.00"; ws2.cell(tr, c).border = BD
    return first, tr
f1, t1 = block(5, "As per GSTR-1", "G1"); f2, t2 = block(t1 + 2, "Sales Register", "SR")
r0 = t2 + 2; header(r0, "Difference (Sales Register - GSTR-1)"); f3 = r0 + 3
for i2, (g, st) in enumerate(GSTINS):
    r = f3 + i2; ws2.cell(r, 1, g); ws2.cell(r, 2, st)
    for c in range(3, 31):
        ws2.cell(r, c, "=%s%d-%s%d" % (L(c), f2 + i2, L(c), f1 + i2)); ws2.cell(r, c).number_format = "#,##0.00"; ws2.cell(r, c).border = BD
t3 = f3 + len(GSTINS); ws2.cell(t3, 1, "Total").font = TOT
for c in range(3, 31):
    ws2.cell(t3, c, "=SUM(%s%d:%s%d)" % (L(c), f3, L(c), t3 - 1)); ws2.cell(t3, c).font = TOT; ws2.cell(t3, c).number_format = "#,##0.00"
ws2.column_dimensions["A"].width = 20; ws2.column_dimensions["B"].width = 20
for c in range(3, 31): ws2.column_dimensions[L(c)].width = 15
ws2.freeze_panes = "C8"

# ---------- Month-on-Month 1 (bird's-eye) ----------
mm = wb.create_sheet("Month-on-Month 1", 3)
mm["A1"] = "Sales Register vs GSTR-1 — month on month, state by state (live). Books = SR_2025-26 by Month; GSTR-1 = tax period."
mm["A1"].font = Font(bold=True)
hdrs = ["State", "GSTIN", "Month", "Books Taxable", "Books IGST", "Books CGST", "Books SGST",
        "GSTR-1 Taxable", "GSTR-1 IGST", "GSTR-1 CGST", "GSTR-1 SGST", "Diff Taxable", "Diff IGST", "Diff CGST", "Diff SGST"]
mm.append([]); mm.append(hdrs)
for c in mm[3]: c.font = HF; c.fill = HB; c.alignment = Alignment(horizontal="center", wrap_text=True)
r = 3
for g, st in GSTINS:
    first = r + 1
    for mo in MONTHS:
        r += 1; mm.cell(r, 1, st); mm.cell(r, 2, g); mm.cell(r, 3, mo)
        for mi in range(4):
            mcol = SRm[mi]
            mm.cell(r, 4 + mi, "=SUMIFS(%s,%s,$B%d,%s,$C%d)" % (SR(mcol), SR(cG_), r, SR(cM), r))
            mm.cell(r, 8 + mi, "=SUMIFS(%s,%s,$B%d,%s,$C%d)" % (G1(G1m[mi]), G1("A"), r, G1("K"), r))
            mm.cell(r, 12 + mi, "=%s%d-%s%d" % (L(4 + mi), r, L(8 + mi), r))
        for c in range(4, 16): mm.cell(r, c).number_format = "#,##0.00"
    r += 1; mm.cell(r, 1, st + " — Total").font = TOT
    for c in range(4, 16):
        mm.cell(r, c, "=SUM(%s%d:%s%d)" % (L(c), first, L(c), r - 1)); mm.cell(r, c).font = TOT; mm.cell(r, c).number_format = "#,##0.00"
        mm.cell(r, c).fill = GF
    r += 1
for c_, w in zip("ABCDEFGHIJKLMNO", [20, 18, 9] + [15] * 12): mm.column_dimensions[c_].width = w
mm.freeze_panes = "D4"

# ---------- Exceptions listing ----------
ex = wb.create_sheet("Step 2 Exceptions", 4)
ex.append(["Status", "State", "My GSTIN", "Books Doc No", "Books type", "Books taxable", "GSTR-1 Doc No", "GSTR-1 type", "GSTR-1 taxable", "Assessment"])
for c in ex[1]: c.font = HF; c.fill = HB
for _, rr in M[M.status != "MATCHED"].sort_values(["status", "state"]).iterrows():
    a = ("Credit note declared in GSTR-1 with no IRN; absent from books AND from 3B (see Step 3) - CA to explain"
         if rr["status"] == "IN GSTR-1 ONLY" else "E-invoiced (valid IRN) but never reported in GSTR-1; included in 3B, tax paid (see Step 3)")
    ex.append([rr["status"], rr["state"], rr["gstin"], rr["docno"], rr["reg_type"], rr["reg_taxable"], rr["g1_docno"], rr["g1_type"], rr["g1_taxable"], a])
NE = ex.max_row
for c_, w in zip("ABCDEFGHIJ", [18, 18, 20, 18, 11, 16, 18, 11, 16, 60]): ex.column_dimensions[c_].width = w
ex.freeze_panes = "A2"

# ---------- Summary (live, with pointers) ----------
su = wb.create_sheet("Step 2 Summary", 2)
MS = H["Matched with GSTR-1"]
rows = [["STEP 2 — Sales Register vs GSTR-1 (FY 2025-26) — all figures live", ""],
 ["Step 1 gate", "='Step 1 Flags'!B7"], [],
 ["Register lines matched to a GSTR-1 document (IRN / GSTIN+DocNo / Customer+Amount)", '=COUNTIF(%s,"Matched*")' % SR(MS)],
 ["Register lines reconciled at summary level (advances, B2C)", '=COUNTIF(%s,"Summary level*")' % SR(MS)],
 ["Register lines NOT in GSTR-1", '=COUNTIF(%s,"Not in GSTR-1")' % SR(MS), "-> filter SR_2025-26 column '%s' = Not in GSTR-1, or see 'Step 2 Exceptions'" % "Matched with GSTR-1"],
 [],
 ["Documents in books only", '=COUNTIF(\'Step 2 Exceptions\'!A2:A%d,"IN BOOKS ONLY")' % NE, '=SUMIFS(\'Step 2 Exceptions\'!F2:F%d,\'Step 2 Exceptions\'!A2:A%d,"IN BOOKS ONLY")' % (NE, NE)],
 ["Documents in GSTR-1 only", '=COUNTIF(\'Step 2 Exceptions\'!A2:A%d,"IN GSTR-1 ONLY")' % NE, '=SUMIFS(\'Step 2 Exceptions\'!I2:I%d,\'Step 2 Exceptions\'!A2:A%d,"IN GSTR-1 ONLY")' % (NE, NE)],
 ["   listing of every document above", "'Step 2 Exceptions' sheet (one row per document, with assessment)"], [],
 ["Total liability — Sales Register", "='SR vs GSTR-1'!AA%d" % t2],
 ["Total liability — GSTR-1", "='SR vs GSTR-1'!AA%d" % t1],
 ["Difference", "='SR vs GSTR-1'!AA%d" % t3, "-> 'SR vs GSTR-1' Difference block shows which state / which group"],
]
for rr in rows: su.append(rr)
su["A1"].font = Font(bold=True, size=12)
for r_ in (13, 14, 15): su.cell(r_, 2).number_format = "#,##0.00"
su.cell(8, 3).number_format = "#,##0.00"; su.cell(9, 3).number_format = "#,##0.00"
su.column_dimensions["A"].width = 70; su.column_dimensions["B"].width = 26; su.column_dimensions["C"].width = 70
wb.save(P)
print("Step 2 merged into register | sheets:", wb.sheetnames)
print("CA blocks at rows %d-%d, %d-%d, %d-%d | MoM rows: %d | exceptions: %d" % (f1, t1, f2, t2, f3, t3, mm.max_row, NE - 1))
