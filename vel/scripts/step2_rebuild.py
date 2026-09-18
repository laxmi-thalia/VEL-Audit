import os, warnings; warnings.filterwarnings("ignore")
import pandas as pd, openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
SP=os.path.dirname(os.path.abspath(__file__))
M=pd.read_pickle(os.path.join(SP,"step2_match2.pkl")); piv=pd.read_pickle(os.path.join(SP,"mom2.pkl"))
H=Font(bold=True,color="FFFFFF"); F=PatternFill("solid",fgColor="1F4E79")
OK=PatternFill("solid",fgColor="C6EFCE"); BAD=PatternFill("solid",fgColor="FFC7CE"); AMB=PatternFill("solid",fgColor="FFF2CC")
wb=Workbook(); ws=wb.active; ws.title="Reconciliation Summary"
tb=round(M["reg_taxable"].sum(),2); tg=round(M["g1_taxable"].sum(),2)
rows=[["STEP 2 - Sales Register vs GSTR-1 (FY 2025-26)","",""],["","",""],
 ["Books - taxable value (excluding advances)",tb,""],
 ["GSTR-1 - taxable value (net of amendments)",tg,""],
 ["DIFFERENCE",round(tb-tg,2),""],["","",""],["RECONCILED BY:","documents","taxable"]]
for st in ["MATCHED","IN BOOKS ONLY","IN GSTR-1 ONLY"]:
    X=M[M.status==st]
    v=X["reg_taxable"].sum() if st=="IN BOOKS ONLY" else (X["g1_taxable"].sum() if st=="IN GSTR-1 ONLY" else 0)
    rows.append([st,len(X),round(v,2)])
rows+=[["","",""],["Matched documents agree on taxable value to the rupee","",""]]
for k,v in M[M.status.str.startswith("MATCHED")]["matched_by"].value_counts().items():
    rows.append([f"   matched by {k}",int(v),""])
rows+=[["","",""],
 ["NOTE: advances excluded - GSTR-1 carries them at summary level only","",""],
 ["NOTE: West Bengal has no GSTR-1 file downloaded; its books rows cannot match","",""]]
for r in rows: ws.append(r)
for c in ws[1]: c.font=H; c.fill=F
for col,w in zip("ABC",[62,18,20]): ws.column_dimensions[col].width=w

e=wb.create_sheet("Exceptions")
e.append(["Status","State","My GSTIN","Books Doc No","Books type","Books taxable","GSTR-1 Doc No","GSTR-1 type","GSTR-1 taxable","Assessment"])
for _,r in M[M.status!="MATCHED"].sort_values(["status","state"]).iterrows():
    a=("West Bengal - GSTR-1 report not downloaded; NOT a real difference" if r["state"]=="West Bengal"
       else "Credit note declared in GSTR-1, no IRN, absent from books - CA to explain" if r["status"]=="IN GSTR-1 ONLY"
       else "E-invoiced (valid IRN) but never reported in GSTR-1 - CA to explain")
    e.append([r["status"],r["state"],r["gstin"],r["docno"],r["reg_type"],r["reg_taxable"],
              r["g1_docno"],r["g1_type"],r["g1_taxable"],a])
for c in e[1]: c.font=H; c.fill=F
for row in e.iter_rows(min_row=2):
    row[9].fill = AMB if "NOT a real" in str(row[9].value) else BAD
for col,w in zip("ABCDEFGHIJ",[18,18,20,18,11,16,18,11,16,58]): e.column_dimensions[col].width=w
e.freeze_panes="A2"

m=wb.create_sheet("Month-on-Month")
m.append(["State","My GSTIN","GSTR-1 Month","Doc Type","Taxable Books","Taxable GSTR-1","Taxable DIFF","IGST DIFF","CGST DIFF","SGST DIFF","Note"])
for _,r in piv.iterrows():
    m.append([r["state"],r["g"],r["month"],r["dtc"],round(r["taxable_Books"],2),round(r["taxable_GSTR-1"],2),
              r["taxable_DIFF"],r["igst_DIFF"],r["cgst_DIFF"],r["sgst_DIFF"],r["note"]])
for c in m[1]: c.font=H; c.fill=F
for row in m.iter_rows(min_row=2):
    row[6].fill = OK if abs(row[6].value or 0)<1 else (AMB if row[10].value else BAD)
for col,w in zip("ABCDEFGHIJK",[20,20,14,10,18,18,16,14,14,14,52]): m.column_dimensions[col].width=w
m.freeze_panes="A2"

mt=wb.create_sheet("Matched")
mt.append(["State","My GSTIN","Books Doc No","GSTR-1 Doc No","Type","Books taxable","GSTR-1 taxable","Matched by"])
for _,r in M[M.status.str.startswith("MATCHED")].iterrows():
    mt.append([r["state"],r["gstin"],r["docno"],r["g1_docno"],r["reg_type"],r["reg_taxable"],r["g1_taxable"],r["matched_by"]])
for c in mt[1]: c.font=H; c.fill=F
for col,w in zip("ABCDEFGH",[20,20,20,20,8,18,18,18]): mt.column_dimensions[col].width=w
mt.freeze_panes="A2"
OUT=r"C:\Users\pawar\Downloads\VEL_Step2_Reco_GSTR1_DRAFT.xlsx"
wb.save(OUT); print("WROTE", OUT)
print("  matched:", int((M.status=='MATCHED').sum()), "| books only:", int((M.status=='IN BOOKS ONLY').sum()),
      "| gstr1 only:", int((M.status=='IN GSTR-1 ONLY').sum()))
print("  month-on-month rows:", len(piv), "| with diff:", int((piv['taxable_DIFF'].abs()>=1).sum()))

# --- register: fill the one derivable blank Supply Type (buyer GSTIN present -> B2B)
P=r"C:\Users\pawar\Downloads\VEL_Sales_Register_FY2025-26_DRAFT.xlsx"
w2=openpyxl.load_workbook(P); s2=w2["SR_2025-26"]; n=0
for i in range(5,27007):
    if s2.cell(i,9).value in (None,"") and s2.cell(i,10).value not in (None,""):
        s2.cell(i,9).value="B2B"; n+=1
w2.save(P); print(f"register: filled {n} blank Supply Type cell(s) from buyer GSTIN")
