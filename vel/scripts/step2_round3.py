"""Step 2 file — Kolshet Road call changes:
 1. Match Status columns on SR Data (K) and GSTR-1 Data (L)  [values; the pairing is the IRN join]
 2. 47 zero-value 'IN GSTR-1 ONLY' rows relabelled NIL DUPLICATE (their CN is matched); Summary updated
 3. Auto-explanation on 'SR vs GSTR-1' Difference block: books-not-in-GSTR-1 + GSTR-1-not-in-books
    + Unexplained residual (live SUMIFS)  — the '4,93,072 why?' ask
 4. Auto-explanation + Remark columns on Month-on-Month (long); sheet then HIDDEN (pivot is the view)
 5. 'Remarks' sheet: one place to write; both views display by State+Month key
"""
import os, warnings; warnings.filterwarnings("ignore")
import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
SP = os.path.dirname(os.path.abspath(__file__))
P = r"C:\Users\pawar\Downloads\VEL_Step2_Reco_GSTR1_DRAFT.xlsx"
def S(v): return "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip()
R = pd.read_pickle(os.path.join(SP, "register.pkl")); G = pd.read_pickle(os.path.join(SP, "gstr1.pkl"))
Su = pd.read_pickle(os.path.join(SP, "g1_summary.pkl")); M = pd.read_pickle(os.path.join(SP, "step2_match2.pkl"))
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79")

# ---- doc-level status maps
stat = {}
for _, r in M.iterrows():
    k = (S(r["gstin"]).upper(), S(r["docno"]).upper(), S(r["reg_type"]))
    if r["status"] == "IN BOOKS ONLY": stat[k] = "Not in GSTR-1"
    elif r["status"].startswith("MATCHED"): stat[k] = "Matched"
g1stat = {}
matched_g1 = {(S(r["gstin"]).upper(), S(r["g1_docno"]).upper()) for _, r in M[M.status.str.startswith("MATCHED")].iterrows()}
only = M[M.status == "IN GSTR-1 ONLY"]
val_only = {(S(r["gstin"]).upper(), S(r["g1_docno"]).upper()) for _, r in only[only["g1_taxable"].abs() >= 0.005].iterrows()}
nil_only = {(S(r["gstin"]).upper(), S(r["g1_docno"]).upper()) for _, r in only[only["g1_taxable"].abs() < 0.005].iterrows()}

wb = openpyxl.load_workbook(P)
sr = wb["SR Data"]; g1 = wb["GSTR-1 Data"]; ws = wb["SR vs GSTR-1"]; mm = wb["Month-on-Month"]
NSR = sr.max_row; NG = g1.max_row

# ---- 1. SR Data col K: Match Status
sr.cell(1, 11, "Match Status").font = HF; sr.cell(1, 11).fill = HB
for i in range(2, NSR + 1):
    g_, d_, t_ = S(sr.cell(i, 1).value).upper(), S(sr.cell(i, 3).value).upper(), S(sr.cell(i, 4).value)
    sup = S(sr.cell(i, 5).value)
    if t_.startswith("MOB") or (t_ == "INV" and sup == "B2C"): v = "Summary level"
    else: v = stat.get((g_, d_, t_), "Matched")
    sr.cell(i, 11, v)
sr.column_dimensions["K"].width = 16

# ---- GSTR-1 Data col L: Match Status  (order = G rows then Su rows, as built)
g1.cell(1, 12, "Match Status").font = HF; g1.cell(1, 12).fill = HB
i = 1
zero_nil = 0
for _, r in G.iterrows():
    i += 1
    key = (S(r["Company GSTIN"]).upper(), S(r["Doc No"]).upper())
    tv = float(pd.to_numeric(r["Taxable Value (Net)"], errors="coerce") or 0)
    if key in nil_only and abs(tv) < 0.005 and S(r["IRN"]).lower() in ("", "nan"):
        v = "Nil duplicate of a matched CN"; zero_nil += 1
    elif key in val_only: v = "Not in books"
    else: v = "Matched"
    g1.cell(i, 12, v)
for _ in range(len(Su)):
    i += 1; g1.cell(i, 12, "Summary level")
g1.column_dimensions["L"].width = 24
SRr = lambda c: "'SR Data'!$%s$2:$%s$%d" % (c, c, NSR)
G1r = lambda c: "'GSTR-1 Data'!$%s$2:$%s$%d" % (c, c, NG)

# ---- 2. Exceptions relabel
ex = wb["Exceptions"]
nil_n = 0
for r in range(2, ex.max_row + 1):
    if S(ex.cell(r, 1).value) == "IN GSTR-1 ONLY" and abs(float(ex.cell(r, 9).value or 0)) < 0.005:
        ex.cell(r, 1).value = "NIL DUPLICATE IN GSTR-1"
        ex.cell(r, 10).value = ("Zero-value second entry of a credit note that IS in the books and already matched (same doc no, "
                                "real CN filed in an earlier period with IRN). No liability effect - CA to confirm why re-reported at nil.")
        nil_n += 1

# ---- 5. Remarks sheet (single place to write; both views display)
if "Remarks" in wb.sheetnames: del wb["Remarks"]
rm = wb.create_sheet("Remarks", 4)
rm.append(["State", "Month", "Remark", "Key (auto)"])
for c in rm[1]: c.font = HF; c.fill = HB
AUTO = [("Bihar", "Jan-26", "AUTO: 34 credit notes in GSTR-1 only (no IRN, Rs 1,13,98,101.92) - not in books, not taken in 3B; see Step 2 Exceptions"),
        ("Telangana", "May-25", "AUTO: e-invoiced documents not uploaded to GSTR-1; tax paid via 3B (see Step 3)"),
        ("Telangana", "Jun-25", "AUTO: e-invoiced documents not uploaded to GSTR-1; tax paid via 3B (see Step 3)"),
        ("Telangana", "Aug-25", "AUTO: e-invoiced documents not uploaded to GSTR-1; tax paid via 3B (see Step 3)"),
        ("Madhya Pradesh", "Dec-25", "AUTO: Rs 9 advance rounding"),
        ("Bihar", "All", "AUTO: Jan-26 GSTR-1-only credit notes explain the whole state difference"),
        ("Telangana", "All", "AUTO: 9 documents missing from GSTR-1 (Rs 54,78,582.96 taxable; CGST/SGST Rs 4,93,072.52 each = 9%)")]
for k, (st, mo, txt) in enumerate(AUTO, start=2):
    rm.append([st, mo, txt, "=A%d&\"|\"&B%d" % (k, k)])
for k in range(len(AUTO) + 2, len(AUTO) + 42):   # spare keyed rows for the CA
    rm.append(["", "", "", "=A%d&\"|\"&B%d" % (k, k)])
NR = rm.max_row
for c_, w in zip("ABCD", [20, 10, 110, 24]): rm.column_dimensions[c_].width = w
rm.freeze_panes = "A2"
KEYRNG = "Remarks!$D$2:$D$%d" % NR; REMRNG = "Remarks!$C$2:$C$%d" % NR
def rem_lookup(keyexpr):
    return '=IFERROR(INDEX(%s,MATCH(%s,%s,0)),"")' % (REMRNG, keyexpr, KEYRNG)

# ---- 3. SR vs GSTR-1 difference block: auto-explanation columns AF..AI (32..35)
f3, t3 = 55, 74
hdr = ["Auto: books docs not in GSTR-1 (taxable)", "Auto: GSTR-1 docs not in books (taxable)", "UNEXPLAINED residual (taxable)", "Remark (whole state, from 'Remarks')"]
for j, t in enumerate(hdr, start=32):
    x = ws.cell(f3 - 1, j, t); x.font = HF; x.fill = HB; x.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.column_dimensions[L(j)].width = 22
for r in range(f3, t3):
    ws.cell(r, 32, '=SUMIFS(%s,%s,$A%d,%s,"Not in GSTR-1")' % (SRr("F"), SRr("A"), r, SRr("K")))
    ws.cell(r, 33, '=-SUMIFS(%s,%s,$A%d,%s,"Not in books")' % (G1r("F"), G1r("A"), r, G1r("L")))
    ws.cell(r, 34, "=AA%d-AF%d-AG%d" % (r, r, r))
    ws.cell(r, 35, rem_lookup('$B%d&"|All"' % r))
    for j in (32, 33, 34): ws.cell(r, j).number_format = "#,##0.00"
ws.cell(t3, 32, "=SUM(AF%d:AF%d)" % (f3, t3 - 1)); ws.cell(t3, 33, "=SUM(AG%d:AG%d)" % (f3, t3 - 1))
ws.cell(t3, 34, "=SUM(AH%d:AH%d)" % (f3, t3 - 1))
for j in (32, 33, 34): ws.cell(t3, j).font = Font(bold=True); ws.cell(t3, j).number_format = "#,##0.00"

# ---- 4. Month-on-Month: auto-explanation + remark, then hide (pivot is the view)
last = mm.max_row
for j, t in enumerate(["Auto: books docs not in GSTR-1", "Auto: GSTR-1 docs not in books", "UNEXPLAINED residual", "Remark (from 'Remarks')"], start=17):
    x = mm.cell(3, j, t); x.font = HF; x.fill = HB; x.alignment = Alignment(horizontal="center", wrap_text=True)
    mm.column_dimensions[L(j)].width = 20
for r in range(4, last + 1):
    mm.cell(r, 17, '=SUMIFS(%s,%s,$B%d,%s,$C%d,%s,$D%d,%s,"Not in GSTR-1")' % (SRr("F"), SRr("A"), r, SRr("B"), r, SRr("J"), r, SRr("K")))
    mm.cell(r, 18, '=-SUMIFS(%s,%s,$B%d,%s,$C%d,%s,$D%d,%s,"Not in books")' % (G1r("F"), G1r("A"), r, G1r("K"), r, G1r("J"), r, G1r("L")))
    mm.cell(r, 19, "=M%d-Q%d-R%d" % (r, r, r))
    mm.cell(r, 20, rem_lookup('$A%d&"|"&$C%d' % (r, r)))
    for j in (17, 18, 19): mm.cell(r, j).number_format = "#,##0.00"
mm.auto_filter.ref = "A3:T%d" % last
mm.sheet_state = "hidden"

# ---- Summary: split the GSTR-1-only line
su = wb["Summary"]
for r in range(1, su.max_row + 1):
    if S(su.cell(r, 1).value) == "Documents in GSTR-1 only":
        NE = ex.max_row
        su.cell(r, 1).value = "Documents in GSTR-1 only (with value - the real gap)"
        su.cell(r, 2).value = '=COUNTIF(Exceptions!A2:A%d,"IN GSTR-1 ONLY")' % NE
        su.cell(r, 3).value = '=SUMIFS(Exceptions!I2:I%d,Exceptions!A2:A%d,"IN GSTR-1 ONLY")' % (NE, NE)
        su.insert_rows(r + 1)
        su.cell(r + 1, 1).value = "Nil duplicate entries in GSTR-1 (zero value; their CN is matched)"
        su.cell(r + 1, 2).value = '=COUNTIF(Exceptions!A2:A%d,"NIL DUPLICATE IN GSTR-1")' % NE
        su.cell(r + 1, 3).value = "no liability effect"
        break
wb.save(P)
print("saved | SR Data status col K, GSTR-1 status col L (nil dup rows: %d) | exceptions relabelled: %d" % (zero_nil, nil_n))
print("auto-explanation on SR vs GSTR-1 rows %d-%d and Month-on-Month (now hidden) | Remarks sheet with %d auto rows" % (f3, t3 - 1, len(AUTO)))
