import os, warnings; warnings.filterwarnings("ignore")
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
SP=os.path.dirname(os.path.abspath(__file__))
T=pd.read_pickle(os.path.join(SP,"b3.pkl"))
G=pd.read_pickle(os.path.join(SP,"gstr1.pkl")); S=pd.read_pickle(os.path.join(SP,"g1_summary.pkl"))
def s(x): return "" if pd.isna(x) else str(x).strip().upper()
MM={"taxable":"Taxable Value (Net)","igst":"IGST (Net)","cgst":"CGST (Net)","sgst":"SGST (Net)"}
for c in MM.values():
    G[c]=pd.to_numeric(G[c],errors="coerce").fillna(0); S[c]=pd.to_numeric(S[c],errors="coerce").fillna(0)
GSTINS=[("01AAECR0503Q1ZM","Jammu & Kashmir"),("03AAECR0503Q1ZI","Punjab"),("06AAECR0503Q1ZC","Haryana"),
 ("08AAECR0503Q1Z8","Rajasthan"),("10AAECR0503Q1ZN","Bihar"),("12AAECR0503Q1ZJ","Arunachal Pradesh"),
 ("18AAECR0503Q1Z7","Assam"),("19AAECR0503Q1Z5","West Bengal"),("20AAECR0503Q1ZM","Jharkhand"),
 ("23AAECR0503Q1ZG","Madhya Pradesh"),("24AAECR0503Q1ZE","Gujarat"),("27AAECR0503Q1Z8","Maharashtra"),
 ("32AAECR0503Q1ZH","Kerala"),("33AAECR0503Q1ZF","TamilNadu"),("36AAECR0503Q1Z9","Telangana"),
 ("37AAECR0503Q1Z7","Andhra Pradesh"),("09AAECR0503Q1Z6","Uttar Pradesh"),("22AAECR0503Q1ZI","Chhattisgarh"),
 ("29AAECR0503Q1Z4","Karnataka")]
GROUPS=["B2B","B2C","Credit note","Debit note","Advance received","Advance adjusted"]
MEAS=["Taxable value","IGST","CGST","SGST"]
MONTHS=sorted({c.split("|")[0] for c in T.columns if "|taxable" in c},
              key=lambda m: pd.to_datetime(m, format="%b %Y"))
# ---- GSTR-1 six-group figures per GSTIN
TY={"Invoice":"B2B","Credit Note":"Credit note","Debit Note":"Debit note"}
SM={"B2CS Sales":"B2C","Advance Received":"Advance received","Advance Adjusted":"Advance adjusted"}
G["bk"]=G["Doc Type"].map(TY); S["bk"]=S["Summary Type"].map(SM)
g1={}
for src in (G,S):
    key="Company GSTIN"
    for (g,b),x in src.dropna(subset=["bk"]).groupby([src[key].map(s),"bk"]):
        p=g1.get((g,b),[0,0,0,0])
        g1[(g,b)]=[p[i]+x[MM[k]].sum() for i,k in enumerate(["taxable","igst","cgst","sgst"])]
# ---- 3B per GSTIN
T["g"]=T["gstin"].map(s)
b3={}
for _,r in T.iterrows():
    b3[r["g"]]=[sum(r[f"{m}|{k}"] for m in MONTHS) for k in ["taxable","igst","cgst","sgst"]]
b3m={r["g"]:{m:[r[f"{m}|{k}"] for k in ["taxable","igst","cgst","sgst"]] for m in MONTHS} for _,r in T.iterrows()}

wb=Workbook(); HF=Font(bold=True,color="FFFFFF"); HB=PatternFill("solid",fgColor="1F4E79")
GF=PatternFill("solid",fgColor="DDEBF7"); TOT=Font(bold=True)
BAD=PatternFill("solid",fgColor="FFC7CE"); OKF=PatternFill("solid",fgColor="C6EFCE")
thin=Side(style="thin",color="B0B0B0"); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
ws=wb.active; ws.title="GSTR-1 vs GSTR-3B"
ws["A1"]="Vikran Engineering Limited"; ws["A1"].font=Font(bold=True,size=13)
ws["A2"]="GSTR-1 vs GSTR-3B"; ws["A2"].font=Font(bold=True,size=12)
ws["A3"]="FY 2025-26  |  3B figures are table 3.1(a) Outward taxable supplies, summed Apr-25 to Mar-26"

def grid(r0,title,groups,getter,diff_src=None):
    ws.cell(r0,1,title).font=Font(bold=True,size=12)
    for gi,gn in enumerate(groups):
        c=3+gi*4
        ws.cell(r0+1,c,gn).font=Font(bold=True); ws.cell(r0+1,c).alignment=Alignment(horizontal="center")
        ws.merge_cells(start_row=r0+1,start_column=c,end_row=r0+1,end_column=c+3); ws.cell(r0+1,c).fill=GF
    ws.cell(r0+2,1,"GSTIN"); ws.cell(r0+2,2,"State")
    for gi in range(len(groups)):
        for mi,mn in enumerate(MEAS): ws.cell(r0+2,3+gi*4+mi,mn)
    ncol=2+len(groups)*4
    for c in range(1,ncol+1):
        x=ws.cell(r0+2,c); x.font=HF; x.fill=HB; x.alignment=Alignment(horizontal="center",wrap_text=True)
    first=r0+3
    for i,(g,st) in enumerate(GSTINS):
        r=first+i; ws.cell(r,1,g); ws.cell(r,2,st)
        for gi,gn in enumerate(groups):
            for mi in range(4):
                c=3+gi*4+mi
                if diff_src and gn=="Diff":
                    L=get_column_letter(c); a=get_column_letter(3+mi); b=get_column_letter(7+mi)
                    ws.cell(r,c,f"={a}{r}-{b}{r}")
                else:
                    ws.cell(r,c, round(getter(g,gn,mi),2))
                ws.cell(r,c).number_format="#,##0.00"; ws.cell(r,c).border=BD
        ws.cell(r,1).border=BD; ws.cell(r,2).border=BD
    tr=first+len(GSTINS); ws.cell(tr,1,"Total").font=TOT
    for c in range(3,ncol+1):
        L=get_column_letter(c); ws.cell(tr,c,f"=SUM({L}{first}:{L}{tr-1})")
        ws.cell(tr,c).font=TOT; ws.cell(tr,c).number_format="#,##0.00"; ws.cell(tr,c).border=BD
    return first,tr

f1,t1=grid(4,"As per GSTR-1",GROUPS, lambda g,gn,mi: g1.get((g,gn),[0,0,0,0])[mi])
r2=t1+2
def getter2(g,gn,mi):
    if gn=="As per GSTR-1": return sum(g1.get((g,x),[0,0,0,0])[mi] for x in GROUPS)
    if gn=="As per GSTR-3B": return b3.get(g,[0,0,0,0])[mi]
    return 0
f2,t2=grid(r2,"GSTR-1 vs GSTR-3B",["As per GSTR-1","As per GSTR-3B","Diff"],getter2,diff_src=True)
ws.cell(r2+2,15,"Remarks").font=HF; ws.cell(r2+2,15).fill=HB
for i,(g,st) in enumerate(GSTINS):
    r=f2+i
    for c in range(11,15): ws.cell(r,c).fill=OKF
    ws.cell(r,15).border=BD
ws.column_dimensions["A"].width=20; ws.column_dimensions["B"].width=20
for c in range(3,27): ws.column_dimensions[get_column_letter(c)].width=15
ws.column_dimensions["O"].width=56
ws.freeze_panes="C7"

# ---- bird's eye view
be=wb.create_sheet("3B Extract (bird's eye)")
be.append(["State","GSTIN","Measure"]+MONTHS+["Total"])
for c in be[1]: c.font=HF; c.fill=HB
for _,r in T.sort_values("state").iterrows():
    for k,lbl in [("taxable","Taxable value"),("igst","IGST"),("cgst","CGST"),("sgst","SGST")]:
        vals=[round(r[f"{m}|{k}"],2) for m in MONTHS]
        be.append([r["state"],r["gstin"],lbl]+vals+[round(sum(vals),2)])
for col,w in zip("ABC",[22,20,14]): be.column_dimensions[col].width=w
be.freeze_panes="D2"

# ---- month on month 1 vs 3B
mm=wb.create_sheet("Month-on-Month 1 vs 3B")
mm.append(["State","GSTIN","Month","GSTR-1 taxable","GSTR-3B taxable","Difference"])
for c in mm[1]: c.font=HF; c.fill=HB
G["month"]=pd.to_datetime(G["Tax Period"],errors="coerce").dt.strftime("%b %Y")
S["month"]=pd.to_datetime(S["Tax Period"],errors="coerce").dt.strftime("%b %Y")
g1m={}
for src in (G,S):
    for (g,m),x in src.groupby([src["Company GSTIN"].map(s),"month"]):
        g1m[(g,m)]=g1m.get((g,m),0)+x["Taxable Value (Net)"].sum()
for g,st in GSTINS:
    for m in MONTHS:
        a=g1m.get((g,m),0.0); b=b3m.get(g,{}).get(m,[0,0,0,0])[0]
        if abs(a)<0.005 and abs(b)<0.005: continue
        mm.append([st,g,m,round(a,2),round(b,2),round(a-b,2)])
for row in mm.iter_rows(min_row=2):
    row[5].fill = OKF if abs(row[5].value or 0)<1 else BAD
for col,w in zip("ABCDEF",[22,20,12,20,20,18]): mm.column_dimensions[col].width=w
mm.freeze_panes="A2"
OUT=r"C:\Users\pawar\Downloads\VEL_Step3_Reco_GSTR1_vs_3B_DRAFT.xlsx"
wb.save(OUT); print("WROTE",OUT)
print(f"  block1 rows {f1}-{t1} | block2 rows {f2}-{t2} | months {len(MONTHS)} | MoM rows {mm.max_row-1}")
