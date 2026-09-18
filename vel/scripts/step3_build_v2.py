"""Step 3 — GSTR-1 vs GSTR-3B, formula-driven, with EVIDENCE behind every exception.
Sheets: Summary | GSTR-1 vs GSTR-3B (CA 2-block, live) | Month-on-Month (long, live, T/I/C/S) |
Exceptions (evidence: rows + live arithmetic) | Amendment Check | 3B Data | GSTR-1 Data | SR Docs"""
import os, re, warnings; warnings.filterwarnings("ignore")
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
SP = os.path.dirname(os.path.abspath(__file__))
T = pd.read_pickle(os.path.join(SP, "b3.pkl")); G = pd.read_pickle(os.path.join(SP, "gstr1.pkl"))
Su = pd.read_pickle(os.path.join(SP, "g1_summary.pkl")); R = pd.read_pickle(os.path.join(SP, "register.pkl"))
def s(x): return "" if x is None or (isinstance(x, float) and pd.isna(x)) else str(x).strip()
for c_ in ["Taxable Value (Net)", "IGST (Net)", "CGST (Net)", "SGST (Net)"]:
    G[c_] = pd.to_numeric(G[c_], errors="coerce").fillna(0); Su[c_] = pd.to_numeric(Su[c_], errors="coerce").fillna(0)
GSTINS = [("01AAECR0503Q1ZM", "Jammu & Kashmir"), ("03AAECR0503Q1ZI", "Punjab"), ("06AAECR0503Q1ZC", "Haryana"),
 ("08AAECR0503Q1Z8", "Rajasthan"), ("10AAECR0503Q1ZN", "Bihar"), ("12AAECR0503Q1ZJ", "Arunachal Pradesh"),
 ("18AAECR0503Q1Z7", "Assam"), ("19AAECR0503Q1Z5", "West Bengal"), ("20AAECR0503Q1ZM", "Jharkhand"),
 ("23AAECR0503Q1ZG", "Madhya Pradesh"), ("24AAECR0503Q1ZE", "Gujarat"), ("27AAECR0503Q1Z8", "Maharashtra"),
 ("32AAECR0503Q1ZH", "Kerala"), ("33AAECR0503Q1ZF", "TamilNadu"), ("36AAECR0503Q1Z9", "Telangana"),
 ("37AAECR0503Q1Z7", "Andhra Pradesh"), ("09AAECR0503Q1Z6", "Uttar Pradesh"), ("22AAECR0503Q1ZI", "Chhattisgarh"),
 ("29AAECR0503Q1Z4", "Karnataka")]
STATE = dict(GSTINS)
MONTHS = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26"]
B3M = {m: pd.to_datetime(m, format="%b-%y").strftime("%b %Y") for m in MONTHS}   # b3.pkl keys like 'Apr 2025'
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79"); GF = PatternFill("solid", fgColor="DDEBF7")
TOT = Font(bold=True); thin = Side(style="thin", color="B0B0B0"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
AMB = PatternFill("solid", fgColor="FFF2CC"); RED = PatternFill("solid", fgColor="FFC7CE"); GRN = PatternFill("solid", fgColor="C6EFCE")
wb = Workbook()

# ---------------- 3B Data (long, zeros explicit) ----------------
b3 = wb.active; b3.title = "3B Data"
b3.append(["State", "GSTIN", "Month", "Taxable 3.1(a)", "IGST", "CGST", "SGST"])
for c in b3[1]: c.font = HF; c.fill = HB
T["g"] = T["gstin"].map(s)
for g, st in GSTINS:
    row = T[T["g"] == g]
    for m in MONTHS:
        k = B3M[m]
        vals = [float(row.iloc[0].get(f"{k}|{x}", 0) or 0) if len(row) else 0.0 for x in ("taxable", "igst", "cgst", "sgst")]
        b3.append([st, g, m] + vals)
N3 = b3.max_row
for c_, w in zip("ABCDEFG", [20, 18, 9, 16, 14, 14, 14]): b3.column_dimensions[c_].width = w
b3.freeze_panes = "A2"
def B3(col): return "'3B Data'!$%s$2:$%s$%d" % (col, col, N3)
B3m = {0: "D", 1: "E", 2: "F", 3: "G"}

# ---------------- GSTR-1 Data (with IRN + amendment flag) ----------------
g1 = wb.create_sheet("GSTR-1 Data")
g1.append(["My GSTIN", "Tax Period", "Source", "Raw Type", "Document Number", "Taxable (Net)", "IGST (Net)", "CGST (Net)", "SGST (Net)",
           "Bucket (formula)", "Month (formula)", "IRN", "Is Amendment", "Doc Date", "Customer GSTIN"])
for c in g1[1]: c.font = HF; c.fill = HB
BF = ('=IF(D{n}="Invoice","B2B",IF(D{n}="Credit Note","Credit note",IF(D{n}="Debit Note","Debit note",'
      'IF(D{n}="B2CS Sales","B2C",IF(D{n}="Advance Received","Advance received","Advance adjusted")))))')
j = 1
for _, r in G.iterrows():
    j += 1; g1.append([s(r["Company GSTIN"]), r["Tax Period"], "Sales-Net (document)", s(r["Doc Type"]), s(r["Doc No"]),
                       r["Taxable Value (Net)"], r["IGST (Net)"], r["CGST (Net)"], r["SGST (Net)"], BF.format(n=j), '=TEXT(B{0},"mmm-yy")'.format(j),
                       s(r["IRN"]) if s(r["IRN"]).lower() != "nan" else "", s(r["Is Amendment"]), r["Doc Date"], s(r["Customer GSTIN"])])
for _, r in Su.iterrows():
    j += 1; g1.append([s(r["Company GSTIN"]), r["Tax Period"], "SalesSummary-Net (summary)", s(r["Summary Type"]), "",
                       r["Taxable Value (Net)"], r["IGST (Net)"], r["CGST (Net)"], r["SGST (Net)"], BF.format(n=j), '=TEXT(B{0},"mmm-yy")'.format(j),
                       "", s(r.get("Is Amendment", "")), None, ""])
NG = j
for c_, w in zip("ABCDEFGHIJKLMNO", [18, 12, 24, 16, 18, 15, 12, 12, 12, 18, 12, 20, 12, 12, 18]): g1.column_dimensions[c_].width = w
g1.freeze_panes = "A2"
def G1(col): return "'GSTR-1 Data'!$%s$2:$%s$%d" % (col, col, NG)
G1m = {0: "F", 1: "G", 2: "H", 3: "I"}

# ---------------- SR Docs (document level, from the register) ----------------
sd = wb.create_sheet("SR Docs")
sd.append(["My GSTIN", "State", "Document Number", "Doc Type", "Month", "Taxable", "CGST", "IRN"])
for c in sd[1]: c.font = HF; c.fill = HB
R["adv"] = R["dtc"].astype(str).str.startswith("MOB")
D = R[~R["adv"]].groupby(["gstin", "docno", "dtc"], dropna=False).agg(
    tax=("tax", "sum"), cgst=("cgst", "sum"), irn=("irn", "first"), state=("state", "first"), ddate=("ddate", "first"), src=("src", "first")).reset_index()
for _, r in D.iterrows():
    mon = pd.to_datetime(r["ddate"]).strftime("%b-%y") if pd.notna(r["ddate"]) else s(r["src"])
    sd.append([s(r["gstin"]), s(r["state"]), s(r["docno"]), s(r["dtc"]), mon, round(r["tax"], 2), round(r["cgst"], 2), s(r["irn"]) if s(r["irn"]).lower() != "nan" else ""])
ND = sd.max_row
for c_, w in zip("ABCDEFGH", [18, 18, 18, 9, 9, 16, 14, 20]): sd.column_dimensions[c_].width = w
sd.freeze_panes = "A2"
def SDc(col): return "'SR Docs'!$%s$2:$%s$%d" % (col, col, ND)

# ---------------- GSTR-1 vs GSTR-3B (CA 2-block) ----------------
ws = wb.create_sheet("GSTR-1 vs GSTR-3B", 0)
ws["A1"] = "Vikran Engineering Limited"; ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = "GSTR-1 vs GSTR-3B"; ws["A2"].font = Font(bold=True, size=12)
ws["A3"] = "FY 2025-26 | 3B = table 3.1(a) outward taxable supplies, Apr-25..Mar-26 | every cell is a live SUMIFS over 'GSTR-1 Data' / '3B Data'"
GROUPS = ["B2B", "B2C", "Credit note", "Debit note", "Advance received", "Advance adjusted"]
MEAS = ["Taxable value", "IGST", "CGST", "SGST"]
def grid(r0, title, groups, cellf, ncol):
    ws.cell(r0, 1, title).font = Font(bold=True, size=12)
    for gi, gn in enumerate(groups):
        c = 3 + gi * 4; ws.cell(r0 + 1, c, gn).font = Font(bold=True); ws.cell(r0 + 1, c).alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=r0 + 1, start_column=c, end_row=r0 + 1, end_column=c + 3); ws.cell(r0 + 1, c).fill = GF
    ws.cell(r0 + 2, 1, "GSTIN"); ws.cell(r0 + 2, 2, "State")
    for gi in range(len(groups)):
        for mi, mn in enumerate(MEAS): ws.cell(r0 + 2, 3 + gi * 4 + mi, mn)
    for c in range(1, ncol + 1):
        x = ws.cell(r0 + 2, c); x.font = HF; x.fill = HB; x.alignment = Alignment(horizontal="center", wrap_text=True)
    first = r0 + 3
    for i2, (g, st) in enumerate(GSTINS):
        r = first + i2; ws.cell(r, 1, g); ws.cell(r, 2, st)
        for gi, gn in enumerate(groups):
            for mi in range(4):
                c = 3 + gi * 4 + mi; ws.cell(r, c, cellf(r, gn, mi)); ws.cell(r, c).number_format = "#,##0.00"; ws.cell(r, c).border = BD
        ws.cell(r, 1).border = BD; ws.cell(r, 2).border = BD
    tr = first + len(GSTINS); ws.cell(tr, 1, "Total").font = TOT
    for c in range(3, ncol + 1):
        ws.cell(tr, c, "=SUM(%s%d:%s%d)" % (L(c), first, L(c), tr - 1)); ws.cell(tr, c).font = TOT; ws.cell(tr, c).number_format = "#,##0.00"; ws.cell(tr, c).border = BD
    return first, tr
f1, t1 = grid(4, "As per GSTR-1", GROUPS, lambda r, gn, mi: '=SUMIFS(%s,%s,$A%d,%s,"%s")' % (G1(G1m[mi]), G1("A"), r, G1("J"), gn), 26)
def cell2(r, gn, mi):
    if gn == "As per GSTR-1": return "=" + "+".join("%s%d" % (L(3 + k * 4 + mi), r - (t1 + 2 + 3) + f1) for k in range(6))
    if gn == "As per GSTR-3B": return "=SUMIFS(%s,%s,$A%d)" % (B3(B3m[mi]), B3("B"), r)
    return "=%s%d-%s%d" % (L(3 + mi), r, L(7 + mi), r)
f2, t2 = grid(t1 + 2, "GSTR-1 vs GSTR-3B", ["As per GSTR-1", "As per GSTR-3B", "Diff"], cell2, 14)
ws.cell(t1 + 4, 15, "Remarks").font = HF; ws.cell(t1 + 4, 15).fill = HB
REM = {"Bihar": "Aug-25: 3B taxable mis-keyed (tax correct) — repeat of FY 24-25; Jan-26: 81 no-IRN credit notes in GSTR-1 only. See 'Exceptions' for the rows.",
       "Telangana": "9 e-invoiced documents missing from GSTR-1; included in 3B, tax paid. See 'Exceptions'.",
       "Punjab": "Re 1 CGST rounding (May-25)."}
for i2, (g, st) in enumerate(GSTINS):
    r = f2 + i2; ws.cell(r, 15, REM.get(st, "")); ws.cell(r, 15).border = BD
    for c in range(11, 15): ws.cell(r, c).fill = GRN
ws.column_dimensions["A"].width = 20; ws.column_dimensions["B"].width = 20
for c in range(3, 27): ws.column_dimensions[L(c)].width = 15
ws.column_dimensions["O"].width = 70
ws.freeze_panes = "C7"

# ---------------- Month-on-Month (long) ----------------
mm = wb.create_sheet("Month-on-Month", 1)
mm["A1"] = "GSTR-1 vs GSTR-3B — State | Month, with GSTR-1 / 3B / Difference for Taxable, IGST, CGST, SGST (live). Difference = GSTR-1 minus 3B."
mm["A1"].font = Font(bold=True)
mm.append([]); mm.append(["State", "GSTIN", "Month", "GSTR-1 Taxable", "GSTR-1 IGST", "GSTR-1 CGST", "GSTR-1 SGST",
                          "3B Taxable", "3B IGST", "3B CGST", "3B SGST", "Diff Taxable", "Diff IGST", "Diff CGST", "Diff SGST"])
for c in mm[3]: c.font = HF; c.fill = HB; c.alignment = Alignment(horizontal="center", wrap_text=True)
r = 3
for g, st in GSTINS:
    for mo in MONTHS:
        r += 1; mm.cell(r, 1, st); mm.cell(r, 2, g); mm.cell(r, 3, mo)
        for mi in range(4):
            mm.cell(r, 4 + mi, "=SUMIFS(%s,%s,$B%d,%s,$C%d)" % (G1(G1m[mi]), G1("A"), r, G1("K"), r))
            mm.cell(r, 8 + mi, "=SUMIFS(%s,%s,$B%d,%s,$C%d)" % (B3(B3m[mi]), B3("B"), r, B3("C"), r))
            mm.cell(r, 12 + mi, "=%s%d-%s%d" % (L(4 + mi), r, L(8 + mi), r))
        for c in range(4, 16): mm.cell(r, c).number_format = "#,##0.00"
NM = r
from openpyxl.formatting.rule import CellIsRule
mm.conditional_formatting.add("L4:L%d" % NM, CellIsRule(operator="notBetween", formula=["-1", "1"], fill=RED))
mm.auto_filter.ref = "A3:O%d" % NM
for c_, w in zip("ABCDEFGHIJKLMNO", [20, 18, 9] + [15] * 12): mm.column_dimensions[c_].width = w
mm.freeze_panes = "D4"

# ---------------- Exceptions — EVIDENCE ----------------
ex = wb.create_sheet("Exceptions", 2)
BIH, TEL, PUN = "10AAECR0503Q1ZN", "36AAECR0503Q1Z9", "03AAECR0503Q1ZI"
def g1sum(col, gstin, month, extra=""):
    return "SUMIFS(%s,%s,\"%s\",%s,\"%s\"%s)" % (G1(col), G1("A"), gstin, G1("K"), month, extra)
def b3sum(col, gstin, month):
    return "SUMIFS(%s,%s,\"%s\",%s,\"%s\")" % (B3(col), B3("B"), gstin, B3("C"), month)
ex["A1"] = "STEP 3 EXCEPTIONS — every figure below is a live formula over 'GSTR-1 Data' / '3B Data' / 'SR Docs'; the rows behind each claim are listed."
ex["A1"].font = Font(bold=True, size=12)
r = 3
def title(t):
    global r
    ex.cell(r, 1, t).font = Font(bold=True, size=12); r += 1
def hdr(cols):
    global r
    for j, t in enumerate(cols, 1):
        x = ex.cell(r, j, t); x.font = HF; x.fill = HB
    r += 1
FORMULA_RE = re.compile(r"^=['A-Z(\-]")
def line(vals, fmt_from=2, fill=None):
    global r
    for j, v in enumerate(vals, 1):
        if isinstance(v, str) and v.startswith("=") and not FORMULA_RE.match(v):
            v = "'" + v          # text that merely begins with '=' must not become a formula
        x = ex.cell(r, j, v)
        if j >= fmt_from and (isinstance(v, (int, float)) or (isinstance(v, str) and v.startswith("="))): x.number_format = "#,##0.00"
        if fill: x.fill = fill
    r += 1

# ---- A. Bihar Aug-25
title("A. Bihar — Aug-25: 3B taxable value mis-keyed; tax columns correct")
hdr(["Measure", "As per GSTR-1 (Aug-25)", "As per GSTR-3B (Aug-25)", "Difference", "Reading"])
for lbl, gc, bc, note in [("Taxable value", "F", "D", "3B taxable is ~Rs 5.16 crore LOWER than GSTR-1"),
                          ("CGST", "H", "F", "tax agrees to under Re 1"), ("SGST", "I", "G", "tax agrees to under Re 1")]:
    line([lbl, "=" + g1sum(gc, BIH, "Aug-25"), "=" + b3sum(bc, BIH, "Aug-25"), "=B%d-C%d" % (r, r), note])
line(["Why this is a keying error, not a liability gap",
      "Tax cannot be right while taxable is wrong by Rs 5 crore unless only the taxable cell was mis-entered. 3B shows 57,38,017 against GSTR-1's 5,73,80,517 — the same digits with one dropped.", "", "", ""])
line(["Precedent", "FY 24-25 workbook, sheet 'GSTR-1 vs GSTR-3B', Bihar remark: 'Reporting Error in GSTR-3B - Taxes reported and discharged correctly' (Rs 18.09 crore difference that year).", "", "", ""])
line(["Effect on tax", "NONE — tax discharged correctly; disclosure error to be corrected via GSTR-9.", "", "", ""]); r += 1

# ---- B. Bihar Jan-26
title("B. Bihar — Jan-26: credit notes in GSTR-1 with no IRN, absent from books and from 3B")
hdr(["Measure", "As per GSTR-1 (Jan-26)", "As per GSTR-3B (Jan-26)", "Difference", "Reading"])
line(["Taxable value", "=" + g1sum("F", BIH, "Jan-26"), "=" + b3sum("D", BIH, "Jan-26"), "=B%d-C%d" % (r, r), "3B is HIGHER than GSTR-1 by the credit notes below"])
line(["CGST", "=" + g1sum("H", BIH, "Jan-26"), "=" + b3sum("F", BIH, "Jan-26"), "=B%d-C%d" % (r, r), "equals 9% of the taxable difference (check below)"])
rowcg = r - 1; rowtx = r - 2
line(["9% check: CGST diff / taxable diff", "=D%d/D%d" % (rowcg, rowtx), "", "", "0.09 = the credit notes carried 9%+9% tax"])
line(["Bihar credit notes in GSTR-1 with BLANK IRN and a value (Jan-26) — count", '=COUNTIFS(%s,"%s",%s,"Credit Note",%s,"Jan-26",%s,"")' % (G1("A"), BIH, G1("D"), G1("K"), G1("L")), "", "", "34 of the 81 GSTR-1-only CNs; the other 47 are nil (B2 below)"])
line(["   their taxable value", '=SUMIFS(%s,%s,"%s",%s,"Credit Note",%s,"Jan-26",%s,"")' % (G1("F"), G1("A"), BIH, G1("D"), G1("K"), G1("L")), "", "", "equals the Jan-26 difference above"])
line(["   of these, found in the books (SR Docs) — count", "=SUM(H%d:H%d)" % (r + 3, r + 3 + 80), "", "", "0 = none of them exist in the sales register"])
line(["Effect on tax", "NONE — the reduction was never taken in 3B, so no tax was short-paid; GSTR-1 detail is wrong and needs the client's explanation.", "", "", ""])
hdr(["#", "Document Number", "Doc Date", "Customer GSTIN", "Taxable (Net)", "CGST (Net)", "IRN (blank)", "In books? (COUNTIFS over SR Docs)"])
cn = G[(G["Company GSTIN"].map(s) == BIH) & (G["Doc Type"] == "Credit Note") &
       (pd.to_datetime(G["Tax Period"]).dt.strftime("%b-%y") == "Jan-26") & (G["IRN"].map(lambda v: s(v) in ("", "nan")))]
for k, (_, rr) in enumerate(cn.iterrows(), 1):
    line([k, s(rr["Doc No"]), rr["Doc Date"], s(rr["Customer GSTIN"]), rr["Taxable Value (Net)"], rr["CGST (Net)"], "",
          '=COUNTIFS(%s,"%s",%s,B%d)' % (SDc("A"), BIH, SDc("C"), r)], fmt_from=5)
r += 1
title("B2. The other 47 of the 81 GSTR-1-only Bihar credit notes: ZERO value, each reported in two tax periods (no liability effect)")
line(["Bihar credit-note rows in GSTR-1 with zero taxable and blank IRN — row count", '=COUNTIFS(%s,"%s",%s,"Credit Note",%s,0,%s,"")' % (G1("A"), BIH, G1("D"), G1("F"), G1("L")), "", "", "94 rows = 47 document numbers x 2 periods"])
line(["   their taxable value", '=SUMIFS(%s,%s,"%s",%s,"Credit Note",%s,0,%s,"")' % (G1("F"), G1("A"), BIH, G1("D"), G1("F"), G1("L")), "", "", "nil"])
hdr(["#", "Document Number", "Period 1", "Period 2", "Taxable (Net)", "IRN (blank)", "In books? (COUNTIFS over SR Docs)"])
gz = G[(G["Company GSTIN"].map(s) == BIH) & (G["Doc Type"] == "Credit Note") & (G["Taxable Value (Net)"].abs() < 0.005) & (G["IRN"].map(lambda v: s(v) in ("", "nan")))].copy()
gz["m"] = pd.to_datetime(gz["Tax Period"]).dt.strftime("%b-%y"); gz["d"] = gz["Doc No"].map(s)
for k, (d, grp) in enumerate(sorted(gz.groupby("d"), key=lambda x: x[0]), 1):
    per = sorted(grp["m"].unique(), key=lambda m: pd.to_datetime(m, format="%b-%y"))
    line([k, d, per[0], per[1] if len(per) > 1 else "", 0.0, "", '=COUNTIFS(%s,"%s",%s,B%d)' % (SDc("A"), BIH, SDc("C"), r)], fmt_from=5)
r += 1

# ---- C. Telangana
title("C. Telangana — e-invoiced documents missing from GSTR-1, but included in 3B (tax paid)")
hdr(["Month", "As per GSTR-1 taxable", "As per GSTR-3B taxable", "Difference", "CGST difference", "CGST diff / taxable diff", "Reading"])
for mo in ["May-25", "Jun-25", "Aug-25"]:
    line([mo, "=" + g1sum("F", TEL, mo), "=" + b3sum("D", TEL, mo), "=B%d-C%d" % (r, r),
          "=" + g1sum("H", TEL, mo) + "-" + b3sum("F", TEL, mo), "=IFERROR(E%d/D%d,\"\")" % (r, r), "0.09 = 9%% CGST on the missing documents"])
line(["The documents (from the sales register) — all carry a valid 64-char IRN, none found in GSTR-1:", "", "", "", "", "", ""])
hdr(["#", "Document Number", "Type", "Month", "Taxable", "IRN", "IRN length", "In GSTR-1? (COUNTIFS over GSTR-1 Data)"])
tel = D[(D["gstin"].map(s) == TEL)]
telset = set(G[G["Company GSTIN"].map(s) == TEL]["Doc No"].map(s))
tel = tel[~tel["docno"].map(s).isin(telset)]
for k, (_, rr) in enumerate(tel.sort_values("docno").iterrows(), 1):
    mon = pd.to_datetime(rr["ddate"]).strftime("%b-%y") if pd.notna(rr["ddate"]) else ""
    line([k, s(rr["docno"]), s(rr["dtc"]), mon, round(rr["tax"], 2), s(rr["irn"]), "=LEN(F%d)" % r,
          '=COUNTIFS(%s,"%s",%s,B%d)' % (G1("A"), TEL, G1("E"), r)], fmt_from=5)
line(["Effect on tax", "NONE — tax fully paid through 3B; GSTR-1 reporting omission to be disclosed/corrected via GSTR-9.", "", "", "", "", ""]); r += 1

# ---- D. Punjab + tie-out
title("D. Punjab — May-25: Re 1 CGST rounding (taxable identical)")
line(["CGST difference May-25", "=" + g1sum("H", PUN, "May-25") + "-" + b3sum("F", PUN, "May-25"), "", "", ""]); r += 1
title("ANNUAL TIE-OUT (live)")
line(["Bihar Aug-25 keying error (taxable)", "=" + g1sum("F", BIH, "Aug-25") + "-" + b3sum("D", BIH, "Aug-25"), "", "", ""])
line(["Bihar Jan-26 GSTR-1-only credit notes (taxable)", "=" + g1sum("F", BIH, "Jan-26") + "-" + b3sum("D", BIH, "Jan-26"), "", "", ""])
line(["Telangana (3 months, taxable)", "=" + "+".join("(" + g1sum("F", TEL, mo) + "-" + b3sum("D", TEL, mo) + ")" for mo in ["May-25", "Jun-25", "Aug-25"]), "", "", ""])
line(["Sum of the above", "=SUM(B%d:B%d)" % (r - 3, r - 1), "", "", ""])
line(["Annual GSTR-1 minus 3B difference (sheet 'GSTR-1 vs GSTR-3B')", "='GSTR-1 vs GSTR-3B'!K%d" % t2, "", "", ""])
line(["Unexplained residual", "=B%d-B%d" % (r - 2, r - 1), "", "", "rounding only"])
ex.column_dimensions["A"].width = 58; ex.column_dimensions["B"].width = 26
for c_, w in zip("CDEFGH", [24, 18, 20, 24, 16, 34]): ex.column_dimensions[c_].width = w
ex.column_dimensions["E"].width = 44

# ---------------- Amendment Check ----------------
am = wb.create_sheet("Amendment Check", 3)
am["A1"] = "Amendments — invoices of FY 2025-26 corrected in a later return"; am["A1"].font = Font(bold=True, size=12)
am["A3"] = "FY 2025-26 GSTR-1 rows flagged 'Is Amendment = Yes'"; am["B3"] = '=COUNTIF(%s,"Yes")' % G1("M")
am["A4"] = "   their taxable (Net)"; am["B4"] = '=SUMIFS(%s,%s,"Yes")' % (G1("F"), G1("M")); am["B4"].number_format = "#,##0.00"
am["A6"] = "FY 2026-27 GSTR-1 (Apr-26 onward) — amendments pointing at FY 2025-26 documents"
am["B6"] = "PENDING — FY 26-27 GSTR-1 exports not yet in Portal Reports\\GSTR-1 (only Apr 2025-Mar 2026 files present)."
am["A7"] = "Method once available: load the 26-27 Sales-Net rows, keep Is Amendment = Yes, join Original/Revised Doc No to the FY 25-26 register; any hit = a 25-26 invoice changed after year-end (report the before/after values)."
am["A8"] = "Why it matters (CA): 'I showed sales as 100, next year I went and made it 00' — the change sits outside the audit year and is invisible unless the next year's return is read."
am.column_dimensions["A"].width = 78; am.column_dimensions["B"].width = 60

# ---------------- Summary ----------------
su = wb.create_sheet("Summary", 0)
rows = [["STEP 3 — GSTR-1 vs GSTR-3B (FY 2025-26) — all figures live", ""], [],
 ["Total liability — GSTR-1 (all tables)", "='GSTR-1 vs GSTR-3B'!C%d" % t2],
 ["Outward taxable supplies — GSTR-3B 3.1(a)", "='GSTR-1 vs GSTR-3B'!G%d" % t2],
 ["Difference", "='GSTR-1 vs GSTR-3B'!K%d" % t2], [],
 ["States with a taxable difference >= Re 1", '=SUMPRODUCT(--(ABS(\'GSTR-1 vs GSTR-3B\'!K%d:K%d)>=1))' % (f2, t2 - 1)],
 ["State-months with a taxable difference >= Re 1", '=SUMPRODUCT(--(ABS(\'Month-on-Month\'!L4:L%d)>=1))' % NM],
 ["Unexplained residual after the Exceptions tie-out", "=Exceptions!B%d" % (r - 1)], [],
 ["Where to look", "'Exceptions' — rows and live arithmetic behind each difference; 'Month-on-Month' — every state x month; 'Amendment Check' — prior-year amendments (26-27 files pending)"]]
for rr in rows: su.append(rr)
su["A1"].font = Font(bold=True, size=12)
for r_ in (3, 4, 5, 9): su.cell(r_, 2).number_format = "#,##0.00"
su.column_dimensions["A"].width = 56; su.column_dimensions["B"].width = 110
OUT = r"C:\Users\pawar\Downloads\VEL_Step3_Reco_GSTR1_vs_3B_DRAFT.xlsx"
wb.save(OUT); print("WROTE", OUT, "| sheets:", wb.sheetnames)
print("blocks %d-%d, %d-%d | 3B rows %d | GSTR-1 rows %d | SR docs %d | Bihar valued CN %d | zero-value CN docs %d | Telangana docs %d" % (f1, t1, f2, t2, N3 - 1, NG - 1, ND - 1, len(cn), gz["d"].nunique(), len(tel)))
