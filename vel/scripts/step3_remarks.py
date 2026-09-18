"""Step 3 remarks sync: one 'Remarks' sheet (State | Month or 'All' | Remark);
the state grid shows the state's 'All' remark, Month-on-Month shows the state+month remark.
Existing static remarks on the grid become AUTO rows in the table."""
import openpyxl
from openpyxl.styles import Font, PatternFill
P = r"C:\Users\pawar\Downloads\VEL_Step3_Reco_GSTR1_vs_3B_DRAFT.xlsx"
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79")
wb = openpyxl.load_workbook(P)
if "Remarks" in wb.sheetnames: del wb["Remarks"]
rm = wb.create_sheet("Remarks", 4)
rm.append(["State", "Month", "Remark", "Key (auto)"])
for c in rm[1]: c.font = HF; c.fill = HB
AUTO = [
 ("Bihar", "Aug-25", "AUTO: 3B taxable mis-keyed (5,73,80,517 entered as 57,38,017); tax correct to Rs 0.39 - see Exceptions A"),
 ("Bihar", "Jan-26", "AUTO: 34 no-IRN credit notes in GSTR-1 only; not taken in 3B, not in books - see Exceptions B"),
 ("Bihar", "All", "AUTO: Aug-25 keying error + Jan-26 GSTR-1-only credit notes; tax correctly discharged - see Exceptions"),
 ("Telangana", "May-25", "AUTO: e-invoiced documents missing from GSTR-1; in 3B, tax paid - see Exceptions C"),
 ("Telangana", "Jun-25", "AUTO: e-invoiced documents missing from GSTR-1; in 3B, tax paid - see Exceptions C"),
 ("Telangana", "Aug-25", "AUTO: e-invoiced documents missing from GSTR-1; in 3B, tax paid - see Exceptions C"),
 ("Telangana", "All", "AUTO: 9 e-invoiced documents never uploaded to GSTR-1; included in 3B, tax paid - see Exceptions C"),
 ("Punjab", "May-25", "AUTO: Re 1 CGST rounding"),
 ("Punjab", "All", "AUTO: Re 1 CGST rounding (May-25)"),
]
k = 1
for st, mo, txt in AUTO:
    k += 1; rm.append([st, mo, txt, '=A%d&"|"&B%d' % (k, k)])
for k in range(len(AUTO) + 2, len(AUTO) + 42):
    rm.append(["", "", "", '=A%d&"|"&B%d' % (k, k)])
NR = rm.max_row
for c_, w in zip("ABCD", [20, 10, 110, 24]): rm.column_dimensions[c_].width = w
rm.freeze_panes = "A2"
LOOK = lambda key: '=IFERROR(INDEX(Remarks!$C$2:$C$%d,MATCH(%s,Remarks!$D$2:$D$%d,0)),"")' % (NR, key, NR)
# grid block 2 (rows 33..49, col O=15) -> live lookup on State|All
ws = wb["GSTR-1 vs GSTR-3B"]
ws.cell(32, 15).value = "Remarks (write on the 'Remarks' sheet - State + 'All')"
for r in range(33, 50):
    ws.cell(r, 15).value = LOOK('$B%d&"|All"' % r)
# Month-on-Month: add col P lookup State|Month
mm = wb["Month-on-Month"]
mm.cell(3, 16, "Remark (from 'Remarks')").font = HF; mm.cell(3, 16).fill = HB
last = mm.max_row
for r in range(4, last + 1):
    if mm.cell(r, 1).value:
        mm.cell(r, 16).value = LOOK('$A%d&"|"&$C%d' % (r, r))
mm.column_dimensions["P"].width = 60
mm.auto_filter.ref = "A3:P%d" % last
wb.save(P); print("remarks sync added: table rows %d (auto %d), grid col O live, MoM col P live" % (NR - 1, len(AUTO)))
