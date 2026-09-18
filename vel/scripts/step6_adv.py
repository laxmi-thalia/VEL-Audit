"""Step 6 — Advances control account (per state), formula-driven.
Opening + Received - Adjusted (+/- Reversed) = Closing ; net vs GSTR-1 advance summary lines.
Opening balances PENDING (prior-year 'Advance from Customer.xlsx' unreachable - share down);
GL cross-check vs GST Advance.xlsx PENDING for the same reason. Slots are built and marked.
Sheets: Summary | Control Account | Month-on-Month | Adv Data (register advance rows) |
GSTR-1 Adv (summary lines) | Remarks"""
import os, warnings; warnings.filterwarnings("ignore")
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
SP = os.path.dirname(os.path.abspath(__file__))
def S(v): return "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip()
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79"); GF = PatternFill("solid", fgColor="DDEBF7")
TOT = Font(bold=True); AMB = PatternFill("solid", fgColor="FFF2CC")
R = pd.read_pickle(os.path.join(SP, "register.pkl"))
Su = pd.read_pickle(os.path.join(SP, "g1_summary.pkl"))
A = R[R["dtc"].astype(str).str.startswith("MOB")].copy()
A["mon"] = A.apply(lambda r: pd.to_datetime(r["ddate"]).strftime("%b-%y") if pd.notna(r["ddate"]) else S(r["src"]), axis=1)
# bucket by CA rule: REC -> received; ADJ -> adjusted; REV by sign
def bucket(r):
    if r["dtc"] == "MOB ADV REC": return "Received"
    if r["dtc"] == "MOB ADV ADJ": return "Adjusted"
    return "Received" if (r["tax"] or 0) > 0 else "Adjusted"
A["bk"] = A.apply(bucket, axis=1)
for c in ["Taxable Value (Net)", "CGST (Net)", "SGST (Net)", "IGST (Net)"]:
    Su[c] = pd.to_numeric(Su[c], errors="coerce").fillna(0)
Su = Su[Su["Summary Type"].isin(["Advance Received", "Advance Adjusted"])].copy()
Su["mon"] = pd.to_datetime(Su["Tax Period"]).dt.strftime("%b-%y")
Su["bk"] = Su["Summary Type"].map({"Advance Received": "Received", "Advance Adjusted": "Adjusted"})
GSTINS = [("01AAECR0503Q1ZM", "Jammu & Kashmir"), ("03AAECR0503Q1ZI", "Punjab"), ("06AAECR0503Q1ZC", "Haryana"),
 ("08AAECR0503Q1Z8", "Rajasthan"), ("10AAECR0503Q1ZN", "Bihar"), ("12AAECR0503Q1ZJ", "Arunachal Pradesh"),
 ("18AAECR0503Q1Z7", "Assam"), ("19AAECR0503Q1Z5", "West Bengal"), ("20AAECR0503Q1ZM", "Jharkhand"),
 ("23AAECR0503Q1ZG", "Madhya Pradesh"), ("24AAECR0503Q1ZE", "Gujarat"), ("27AAECR0503Q1Z8", "Maharashtra"),
 ("32AAECR0503Q1ZH", "Kerala"), ("33AAECR0503Q1ZF", "TamilNadu"), ("36AAECR0503Q1Z9", "Telangana"),
 ("37AAECR0503Q1Z7", "Andhra Pradesh"), ("09AAECR0503Q1Z6", "Uttar Pradesh"), ("22AAECR0503Q1ZI", "Chhattisgarh"),
 ("29AAECR0503Q1Z4", "Karnataka")]
MONTHS = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26"]
wb = Workbook()

# ---- Adv Data (register advance rows)
ad = wb.active; ad.title = "Adv Data"
ad.append(["State", "My GSTIN", "Month", "Doc Type (raw)", "Bucket", "Document Number", "Taxable", "CGST", "SGST"])
for c in ad[1]: c.font = HF; c.fill = HB
for _, r in A.iterrows():
    ad.append([S(r["state"]), S(r["gstin"]), r["mon"], S(r["dtc"]), r["bk"], S(r["docno"]),
               round(float(r["tax"] or 0), 2), round(float(r["cgst"] or 0), 2), round(float(r["sgst"] or 0), 2)])
NA = ad.max_row
for c_, w in zip("ABCDEFGHI", [18, 18, 9, 14, 10, 18, 15, 13, 13]): ad.column_dimensions[c_].width = w
ad.freeze_panes = "A2"
AD = lambda c: "'Adv Data'!$%s$2:$%s$%d" % (c, c, NA)

# ---- GSTR-1 Adv
ga = wb.create_sheet("GSTR-1 Adv")
ga.append(["My GSTIN", "State", "Month", "Bucket", "Taxable (Net)", "CGST (Net)", "SGST (Net)"])
for c in ga[1]: c.font = HF; c.fill = HB
ST = dict(GSTINS)
for _, r in Su.iterrows():
    g = S(r["Company GSTIN"])
    ga.append([g, ST.get(g, ""), r["mon"], r["bk"], round(r["Taxable Value (Net)"], 2), round(r["CGST (Net)"], 2), round(r["SGST (Net)"], 2)])
NGA = ga.max_row
for c_, w in zip("ABCDEFG", [18, 18, 9, 10, 15, 13, 13]): ga.column_dimensions[c_].width = w
ga.freeze_panes = "A2"
GA = lambda c: "'GSTR-1 Adv'!$%s$2:$%s$%d" % (c, c, NGA)

# ---- Control Account
ws = wb.create_sheet("Control Account", 0)
ws["A1"] = "Vikran Engineering Limited — Advances from customers, control account FY 2025-26 (taxable values)"
ws["A1"].font = Font(bold=True, size=12)
ws["A2"] = ("Opening + Received - |Adjusted| = Closing. Tax is payable on RECEIPT; only what was received (and taxed) may be adjusted. "
            "Adjusted is stored NEGATIVE in the data, so it is shown as |sum|. OPENING = prior-year closing - PENDING (share down); "
            "a NEGATIVE running closing flags adjustment in excess of taxed receipts.")
hdrs = ["State", "GSTIN", "Opening (PENDING prior-year file)", "Received in FY (books)", "Adjusted in FY (books, |sum|)",
        "Closing (O + R - A)", "GSTR-1 net advances (Rec - |Adj|)", "Books net (R - A)", "Books vs GSTR-1 diff", "Remark"]
ws.append([]); ws.append([])
for j, t in enumerate(hdrs, 1):
    x = ws.cell(4, j, t); x.font = HF; x.fill = HB; x.alignment = Alignment(horizontal="center", wrap_text=True)
first = 5
for i2, (g, st_) in enumerate(GSTINS):
    r = first + i2
    ws.cell(r, 1, st_); ws.cell(r, 2, g)
    ws.cell(r, 3, 0); ws.cell(r, 3).fill = AMB   # opening pending
    ws.cell(r, 4, '=SUMIFS(%s,%s,$B%d,%s,"Received")' % (AD("G"), AD("B"), r, AD("E")))
    ws.cell(r, 5, '=-SUMIFS(%s,%s,$B%d,%s,"Adjusted")' % (AD("G"), AD("B"), r, AD("E")))
    ws.cell(r, 6, "=C%d+D%d-E%d" % (r, r, r))
    ws.cell(r, 7, '=SUMIFS(%s,%s,$B%d)' % (GA("E"), GA("A"), r))
    ws.cell(r, 8, "=D%d-E%d" % (r, r))
    ws.cell(r, 9, "=H%d-G%d" % (r, r))
    ws.cell(r, 10, '=IF(F%d<0,"OVER-ADJUSTED vs taxed receipts (pending opening) - check","")' % r)
    for c in range(3, 10): ws.cell(r, c).number_format = "#,##0.00"
tr = first + len(GSTINS)
ws.cell(tr, 1, "Total").font = TOT
for c in range(3, 10):
    ws.cell(tr, c, "=SUM(%s%d:%s%d)" % (L(c), first, L(c), tr - 1)); ws.cell(tr, c).font = TOT; ws.cell(tr, c).number_format = "#,##0.00"
for c_, w in zip("ABCDEFGHIJ", [18, 18, 18, 18, 18, 16, 18, 16, 16, 44]): ws.column_dimensions[c_].width = w
ws.freeze_panes = "C5"

# ---- Month-on-Month
mm = wb.create_sheet("Month-on-Month", 1)
mm["A1"] = "Advances month on month — books vs GSTR-1 (taxable, live)."; mm["A1"].font = Font(bold=True)
mm.append([]); mm.append(["State", "GSTIN", "Month", "Books Received", "Books Adjusted (|sum|)", "GSTR-1 Received", "GSTR-1 Adjusted (|sum|)",
                          "Diff Received", "Diff Adjusted", "Running closing (books, opening PENDING)"])
for c in mm[3]: c.font = HF; c.fill = HB; c.alignment = Alignment(horizontal="center", wrap_text=True)
r = 3
for g, st_ in GSTINS:
    prev = None
    for mo in MONTHS:
        r += 1
        mm.cell(r, 1, st_); mm.cell(r, 2, g); mm.cell(r, 3, mo)
        mm.cell(r, 4, '=SUMIFS(%s,%s,$B%d,%s,$C%d,%s,"Received")' % (AD("G"), AD("B"), r, AD("C"), r, AD("E")))
        mm.cell(r, 5, '=-SUMIFS(%s,%s,$B%d,%s,$C%d,%s,"Adjusted")' % (AD("G"), AD("B"), r, AD("C"), r, AD("E")))
        mm.cell(r, 6, '=SUMIFS(%s,%s,$B%d,%s,$C%d,%s,"Received")' % (GA("E"), GA("A"), r, GA("C"), r, GA("D")))
        mm.cell(r, 7, '=-SUMIFS(%s,%s,$B%d,%s,$C%d,%s,"Adjusted")' % (GA("E"), GA("A"), r, GA("C"), r, GA("D")))
        mm.cell(r, 8, "=D%d-F%d" % (r, r)); mm.cell(r, 9, "=E%d-G%d" % (r, r))
        mm.cell(r, 10, ("=D%d-E%d" % (r, r)) if prev is None else ("=J%d+D%d-E%d" % (prev, r, r)))
        prev = r
        for c in range(4, 11): mm.cell(r, c).number_format = "#,##0.00"
from openpyxl.formatting.rule import CellIsRule
mm.conditional_formatting.add("H4:I%d" % r, CellIsRule(operator="notBetween", formula=["-1", "1"], fill=PatternFill("solid", fgColor="FFC7CE")))
mm.conditional_formatting.add("J4:J%d" % r, CellIsRule(operator="lessThan", formula=["-1"], fill=AMB))
mm.auto_filter.ref = "A3:J%d" % r
for c_, w in zip("ABCDEFGHIJ", [18, 18, 9, 15, 16, 15, 16, 13, 13, 20]): mm.column_dimensions[c_].width = w
mm.freeze_panes = "D4"

# ---- Summary
su = wb.create_sheet("Summary", 0)
rows = [["STEP 6 — Advances control account (FY 2025-26) — all figures live", ""], [],
 ["Books: advances received (taxable)", "='Control Account'!D%d" % tr],
 ["Books: advances adjusted (|taxable|)", "='Control Account'!E%d" % tr],
 ["Books net", "='Control Account'!H%d" % tr],
 ["GSTR-1 net advance lines", "='Control Account'!G%d" % tr],
 ["Books vs GSTR-1 difference", "='Control Account'!I%d" % tr], [],
 ["PENDING 1", "Opening balances from prior-year 'Advance from Customer.xlsx' (share unreachable) - closing column is understated until filled"],
 ["PENDING 2", "Cross-check vs the GST Advance ledger (Clients Data\\GLs\\Outward Tax\\GST Advance.xlsx) - same reason"],
 ["Rule (CA)", "Tax on receipt; adjustment only against taxed receipts. MOB ADV REV split by sign (positive = receipt side)."]]
for rr in rows: su.append(rr)
su["A1"].font = Font(bold=True, size=12)
for r_ in range(3, 8): su.cell(r_, 2).number_format = "#,##0.00"
su.column_dimensions["A"].width = 40; su.column_dimensions["B"].width = 110

# ---- Remarks
rm = wb.create_sheet("Remarks")
rm.append(["State", "Month", "Remark", "Key (auto)"])
for c in rm[1]: c.font = HF; c.fill = HB
for k in range(2, 42): rm.append(["", "", "", '=A%d&"|"&B%d' % (k, k)])
for c_, w in zip("ABCD", [20, 10, 100, 24]): rm.column_dimensions[c_].width = w
OUT = r"C:\Users\pawar\Downloads\VEL_Step6_Advances_Control_DRAFT.xlsx"
wb.save(OUT)
print("WROTE", OUT)
print("adv rows:", NA - 1, "| GSTR-1 adv lines:", NGA - 1, "| states:", len(GSTINS))
