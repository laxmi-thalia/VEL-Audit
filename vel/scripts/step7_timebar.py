"""Step 7 — Credit-note time-of-supply / Section 34(2) time-bar check.
A CN reducing tax must be declared by 30 Nov following the FY of the ORIGINAL supply
(or the annual return date, if earlier). Source: ClearTax 'Original Invoice Number' /
'Preceding Invoice Date' (populated on only 8 of 288 CNs - itself a finding)."""
import os, warnings; warnings.filterwarnings("ignore")
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
SP = os.path.dirname(os.path.abspath(__file__))
def S(v): return "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip()
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79")
RED = PatternFill("solid", fgColor="FFC7CE"); AMB = PatternFill("solid", fgColor="FFEB9C")
D = pd.read_pickle(os.path.join(SP, "cn_rows.pkl"))
docs = D.groupby(["gstin", "docno"]).agg(state=("state", "first"), month=("month", "first"), cndate=("cndate", "first"),
    orig=("orig", "first"), pdate=("pdate", "first"), tax=("tax", "sum")).reset_index()
docs["pdate"] = pd.to_datetime(docs["pdate"], errors="coerce")
docs = docs.sort_values(["state", "docno"])
wb = Workbook()
ws = wb.active; ws.title = "CN Time-bar Check"
ws["A1"] = "Credit notes — time of supply / Section 34(2) time-bar (FY 2025-26)"
ws["A1"].font = Font(bold=True, size=12)
ws["A2"] = ("Rule: a credit note reducing GST must be declared by 30 Nov following the end of the FY of the ORIGINAL supply "
            "(or the annual-return date, if earlier). CNs issued in FY 25-26 against FY 25-26 invoices are within time by definition. "
            "Status is a live formula; fill column G/H for the CNs without a reference and the status updates.")
hdr = ["State", "My GSTIN", "CN Number", "CN Date", "Reported Month", "Taxable", "Original Invoice No", "Original Invoice Date",
       "Original FY", "Declare-by (30 Nov)", "Status (live)"]
ws.append([]); ws.append([])
for j, t in enumerate(hdr, 1):
    x = ws.cell(4, j, t); x.font = HF; x.fill = HB; x.alignment = Alignment(horizontal="center", wrap_text=True)
r = 4
for _, d in docs.iterrows():
    r += 1
    ws.cell(r, 1, S(d["state"])); ws.cell(r, 2, S(d["gstin"])); ws.cell(r, 3, S(d["docno"]))
    ws.cell(r, 4, d["cndate"]); ws.cell(r, 5, S(d["month"])); ws.cell(r, 6, round(float(d["tax"]), 2))
    ws.cell(r, 7, S(d["orig"]) if S(d["orig"]).lower() != "nan" else "")
    ws.cell(r, 8, d["pdate"] if pd.notna(d["pdate"]) else None)
    ws.cell(r, 9, '=IF(H{0}="","",IF(MONTH(H{0})>=4,YEAR(H{0})&"-"&RIGHT(YEAR(H{0})+1,2),YEAR(H{0})-1&"-"&RIGHT(YEAR(H{0}),2)))'.format(r))
    ws.cell(r, 10, '=IF(H{0}="","",DATE(IF(MONTH(H{0})>=4,YEAR(H{0})+1,YEAR(H{0})),11,30))'.format(r))
    ws.cell(r, 11, ('=IF(H{0}="","Original ref not captured in ClearTax - cannot test",'
                    'IF(H{0}>=DATE(2025,4,1),"OK - original invoice in FY 25-26",'
                    'IF(D{0}>J{0},"BEYOND Sec 34(2) WINDOW - GST reduction not permissible; CA to examine",'
                    '"OK - prior-FY original but declared within the window")))').format(r))
    ws.cell(r, 4).number_format = "dd-mmm-yy"; ws.cell(r, 8).number_format = "dd-mmm-yy"
    ws.cell(r, 10).number_format = "dd-mmm-yy"; ws.cell(r, 6).number_format = "#,##0.00"
last = r
ws.conditional_formatting.add("K5:K%d" % last, FormulaRule(formula=['ISNUMBER(SEARCH("BEYOND",$K5))'], fill=RED, stopIfTrue=True))
ws.conditional_formatting.add("K5:K%d" % last, FormulaRule(formula=['ISNUMBER(SEARCH("not captured",$K5))'], fill=AMB))
ws.auto_filter.ref = "A4:K%d" % last
for c_, w in zip("ABCDEFGHIJK", [18, 18, 16, 11, 10, 14, 18, 13, 10, 13, 52]): ws.column_dimensions[c_].width = w
ws.freeze_panes = "A5"
su = wb.create_sheet("Summary", 0)
rows = [["STEP 7 — Credit-note time-bar check (Sec 34(2)) — status column is live", ""], [],
 ["Credit notes in FY 25-26", "=COUNTA('CN Time-bar Check'!C5:C%d)" % last],
 ["   with an original-invoice reference", "=COUNTIF('CN Time-bar Check'!H5:H%d,\"<>\")" % last],
 ["   WITHOUT a reference (cannot test - data gap)", '=COUNTIF(\'CN Time-bar Check\'!K5:K%d,"*not captured*")' % last],
 ["   BEYOND the Sec 34(2) window", '=COUNTIF(\'CN Time-bar Check\'!K5:K%d,"*BEYOND*")' % last],
 ["   taxable of the beyond-window CNs", '=SUMIFS(\'CN Time-bar Check\'!F5:F%d,\'CN Time-bar Check\'!K5:K%d,"*BEYOND*")' % (last, last)], [],
 ["Findings", "3 MP credit notes of Mar-26 cite an invoice dated 28.11.2023 (FY 23-24; window closed 30.11.2024)."],
 ["Data gap", "280 of 288 CNs carry no original-invoice reference in ClearTax - ask the client to fill, then this sheet retests itself."],
 ["Also", "the CA's session-1 ask; the format still lacks an original-invoice column - this sheet works from source data instead."]]
for rr in rows: su.append(rr)
su["A1"].font = Font(bold=True, size=12)
su.cell(7, 2).number_format = "#,##0.00"
su.column_dimensions["A"].width = 46; su.column_dimensions["B"].width = 100
OUT = r"C:\Users\pawar\Downloads\VEL_Step7_CN_TimeBar_DRAFT.xlsx"
wb.save(OUT); print("WROTE", OUT, "| credit notes:", last - 4)
