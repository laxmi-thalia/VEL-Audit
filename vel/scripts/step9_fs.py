"""Step 9 — Sales as per audited financials vs GST turnover (GSTR-9C Table 5 shape), per state.
FS side: the client's own 'Statewise PnL FY25-26.xlsx' Data sheet, rows tagged
'01. Revenue from operations' (reproduces the P&L statement to Rs 0.15), aggregated
account x state and live-SUMIFS'd. GST side: the Step 1 register (billed ex-advances,
advances received/adjusted). Reconciling items are JUDGMENT rows left for the CA
(unbilled/uncertified revenue movement etc.) with a live residual."""
import os, warnings; warnings.filterwarnings("ignore")
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L
SP = os.path.dirname(os.path.abspath(__file__))
def S(v): return "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip()
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79"); GF = PatternFill("solid", fgColor="DDEBF7")
TOT = Font(bold=True); AMB = PatternFill("solid", fgColor="FFF2CC")
PNL = "//192.168.1.69/GST FOLDER/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/Clients Data/Financials/Statewise PnL FY25-26.xlsx"
FIX = {"Chhatisgarh": "Chhattisgarh", "Tamilnadu": "Tamil Nadu"}
d = pd.read_excel(PNL, sheet_name="Data", header=0)
d.columns = [S(c) for c in d.columns]
d["pnl"] = d["PnL"].map(S); d["st"] = d["State"].map(S).map(lambda x: FIX.get(x, x))
rev = d[d["pnl"] == "01. Revenue from operations"]
agg = (-rev.groupby([rev["Account Number"].map(S), "st"])["Total Balance"].sum()).round(2).reset_index()
agg.columns = ["account", "state", "amount"]
oth = round(-d[d["pnl"] == "02. Other Income"]["Total Balance"].sum(), 2)
R = pd.read_pickle(os.path.join(SP, "register.pkl"))
R["adv"] = R["dtc"].astype(str).str.startswith("MOB")
gst = R.groupby([R["state"].map(S), "adv"])["tax"].sum().unstack(fill_value=0).round(2)
gst.columns = ["billed", "advnet"][:len(gst.columns)] if len(gst.columns) == 2 else gst.columns
STATES = ["Andhra Pradesh", "Arunachal Pradesh", "Assam", "Bihar", "Chhattisgarh", "Gujarat", "Haryana", "Jammu & Kashmir",
          "Jharkhand", "Karnataka", "Kerala", "Madhya Pradesh", "Maharashtra", "Punjab", "Rajasthan", "Tamil Nadu",
          "Telangana", "Uttar Pradesh", "West Bengal"]
wb = Workbook()
# ---- FS Revenue Data
fd = wb.active; fd.title = "FS Revenue Data"
fd.append(["Account (revenue head)", "State", "Amount (FS revenue)"])
for c in fd[1]: c.font = HF; c.fill = HB
for _, r in agg.sort_values(["state", "account"]).iterrows():
    fd.append([r["account"], r["state"], r["amount"]])
NF = fd.max_row
for c_, w in zip("ABC", [44, 20, 18]): fd.column_dimensions[c_].width = w
fd.freeze_panes = "A2"
FD = lambda c: "'FS Revenue Data'!$%s$2:$%s$%d" % (c, c, NF)
# ---- GST Turnover Data (from the Step 1 register, documented)
gt = wb.create_sheet("GST Turnover Data")
gt.append(["State", "Billed taxable (ex-advances)", "Advances net (received - adjusted)", "Source"])
for c in gt[1]: c.font = HF; c.fill = HB
for st in STATES:
    billed = float(gst["billed"].get(st, 0)) if "billed" in gst else 0.0
    advn = float(gst["advnet"].get(st, 0)) if "advnet" in gst else 0.0
    gt.append([st, round(billed, 2), round(advn, 2), "VEL_Step1_Sales_Register_FY2025-26 (Taxable Value; advances = MOB rows)"])
NG = gt.max_row
for c_, w in zip("ABCD", [20, 22, 24, 56]): gt.column_dimensions[c_].width = w
GT = lambda c: "'GST Turnover Data'!$%s$2:$%s$%d" % (c, c, NG)
# ---- FS vs GST grid
ws = wb.create_sheet("FS vs GST", 0)
ws["A1"] = "Vikran Engineering Limited — Revenue per audited financials vs turnover per GST (FY 2025-26) — GSTR-9C Table 5 working"
ws["A1"].font = Font(bold=True, size=12)
ws["A2"] = ("FS = state-wise P&L 'Revenue from operations' (client file, reproduced to Rs 0.15). GST = Step 1 register. "
            "Columns G-J are JUDGMENT reconcilers for the CA (unbilled/uncertified movement, other income, non-GST items, timing); "
            "the residual is live and shrinks as they are filled.")
hdr = ["State", "FS Revenue from operations", "GST billed turnover (ex-advances)", "GST advances net", "GST total",
       "FS - GST difference", "Reconciler 1: unbilled / uncertified movement (CA)", "Reconciler 2: other income / non-GST (CA)",
       "Reconciler 3: credit-note & billing timing (CA)", "Reconciler 4: others (CA)", "UNEXPLAINED residual", "Remark"]
ws.append([]); ws.append([])
for j, t in enumerate(hdr, 1):
    x = ws.cell(4, j, t); x.font = HF; x.fill = HB; x.alignment = Alignment(horizontal="center", wrap_text=True)
first = 5
for i2, st in enumerate(STATES):
    r = first + i2
    ws.cell(r, 1, st)
    ws.cell(r, 2, "=SUMIFS(%s,%s,$A%d)" % (FD("C"), FD("B"), r))
    ws.cell(r, 3, "=SUMIFS(%s,%s,$A%d)" % (GT("B"), GT("A"), r))
    ws.cell(r, 4, "=SUMIFS(%s,%s,$A%d)" % (GT("C"), GT("A"), r))
    ws.cell(r, 5, "=C%d+D%d" % (r, r))
    ws.cell(r, 6, "=B%d-E%d" % (r, r))
    for j in (7, 8, 9, 10): ws.cell(r, j).fill = AMB
    ws.cell(r, 11, "=F%d-SUM(G%d:J%d)" % (r, r, r))
    ws.cell(r, 12, "")
    for c in range(2, 12): ws.cell(r, c).number_format = "#,##0.00"
tr = first + len(STATES)
ws.cell(tr, 1, "Total").font = TOT
for c in range(2, 12):
    ws.cell(tr, c, "=SUM(%s%d:%s%d)" % (L(c), first, L(c), tr - 1)); ws.cell(tr, c).font = TOT; ws.cell(tr, c).number_format = "#,##0.00"
ws.conditional_formatting.add("K%d:K%d" % (first, tr), CellIsRule(operator="notBetween", formula=["-1", "1"], fill=PatternFill("solid", fgColor="FFC7CE")))
for c_, w in zip("ABCDEFGHIJKL", [20, 20, 20, 16, 18, 18, 22, 20, 20, 16, 18, 40]): ws.column_dimensions[c_].width = w
ws.freeze_panes = "B5"
# ---- Summary
su = wb.create_sheet("Summary", 0)
rows = [["STEP 9 — Revenue per financials vs GST turnover (FY 2025-26) — GSTR-9C Table 5 working", ""], [],
 ["FS Revenue from operations", "='FS vs GST'!B%d" % tr],
 ["GST total turnover (billed + advances net)", "='FS vs GST'!E%d" % tr],
 ["Difference to reconcile", "='FS vs GST'!F%d" % tr],
 ["Unexplained residual (falls as CA fills reconcilers)", "='FS vs GST'!K%d" % tr], [],
 ["Other income (FS, not state-allocated)", oth],
 ["Known structural cause", "EPC pattern: revenue booked on work done (incl. uncertified), GST on billing - the SAP registers carry 'Uncertified' accrual postings; unbilled-revenue balances are in the Annual Report notes"], [],
 ["Notable", "Maharashtra FS 705.9cr vs GST 328.7cr (HO booking); Kerala FS revenue NEGATIVE 19.28L with zero GST; West Bengal FS NEGATIVE 24.33L vs GST +5.87L"],
 ["Sources", "'FS Revenue Data' (client P&L extract, tagged rows) and 'GST Turnover Data' (Step 1 register) - both drillable"]]
for rr in rows: su.append(rr)
su["A1"].font = Font(bold=True, size=12)
for r_ in (3, 4, 5, 6, 8): su.cell(r_, 2).number_format = "#,##0.00"
su.column_dimensions["A"].width = 48; su.column_dimensions["B"].width = 110
OUT = r"C:\Users\pawar\Downloads\VEL_Step9_FS_vs_GST_DRAFT.xlsx"
wb.save(OUT); print("WROTE", OUT, "| FS data rows:", NF - 1, "| states:", len(STATES))
