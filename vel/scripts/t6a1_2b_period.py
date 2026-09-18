"""T6A1 Extract - 24-25: fill '2B Return Period' (Pawan 18-09). The column carried formulas pointing at a helper key column
$AF that no longer exists (row offset +68 from the replica roll) -> 3,141 of 3,595 rows evaluated blank.
New self-contained LIVE formula per row: key = normalised(GSTIN of supplier & Invoice number) (same normalisation as the 2B
KEY column); look up the FY 24-25 2B first ('GSTR-2B ITC Data', working-file 2B, new helper KEY col AB), then the merged
Apr-25..Aug-26 Octa 2B ('GSTR-2B Apr25-Aug26', KEY col AW -> Tax Period col B), else 'Not in 2B (Apr-24 to Aug-26)'.
RCM rows (self-invoice) get a fixed text. Helper col T 'Source row' refreshed from the GSTR-9 Remarks pointer."""
import re, time, shutil, collections, warnings, openpyxl, win32com.client as win32, pythoncom
warnings.filterwarnings("ignore")
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"
try: open(P, "r+b").close()
except PermissionError: raise SystemExit("MASTER IS OPEN IN EXCEL - close it (without saving) and rerun")
shutil.copy(P, "master2_snapshot_before_t6a1period.xlsx")
wb = openpyxl.load_workbook(P, read_only=True)
ex = wb["T6A1 Extract - 24-25"]; hdr = [c.value for c in next(ex.iter_rows(min_row=4, max_row=4))]; H = {h: i + 1 for i, h in enumerate(hdr) if h}
rows = [(i + 5, r) for i, r in enumerate(ex.iter_rows(min_row=5, values_only=True)) if r[0]]; EN = rows[-1][0]
old = wb["GSTR-2B ITC Data"]; oh = [c.value for c in next(old.iter_rows(min_row=5, max_row=5))]; OH = {h: i + 1 for i, h in enumerate(oh) if h}
ON = 5 + sum(1 for r in old.iter_rows(min_row=6, values_only=True) if r[0])
b2 = wb["GSTR-2B Apr25-Aug26"]; bh = [c.value for c in next(b2.iter_rows(min_row=2, max_row=2))]; BH = {h: i + 1 for i, h in enumerate(bh) if h}
NB = 2 + sum(1 for r in b2.iter_rows(min_row=3, values_only=True) if r[0]); wb.close()
from openpyxl.utils import get_column_letter as L
assert BH["KEY"] == 49 and BH["Tax Period"] == 2 and OH["2B Return Period"] == 4 and OH["GSTIN of supplier"] == 7 and OH["Invoice number"] == 9, (BH.get("KEY"), OH)
KEYCOL = max(OH.values()) + 1; KL = L(KEYCOL)
print("extract rows 5..%d | old 2B rows 6..%d (helper KEY -> col %s) | Apr25-Aug26 rows 3..%d" % (EN, ON, KL, NB))
norm = lambda a: 'UPPER(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(%s," ",""),"-",""),"/",""),".",""),"\'",""),"_",""))' % a
cS, cJ, cM, cG, cT = L(H["Source"]), L(H["GSTIN of supplier"]), L(H["Invoice number"]), L(H["GSTR-9 Remarks"]), L(H["Source row (helper)"])
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); wbx = xl.Workbooks.Open(P); xl.Calculation = -4135
    o = wbx.Worksheets("GSTR-2B ITC Data"); x = o.Cells(5, KEYCOL); x.Value = "KEY (supplier GSTIN + invoice, normalised)"; x.Font.Bold = True; x.Font.Color = 0xFFFFFF; x.Interior.Color = 0xB09784   # #8497B0 band (DPS-added)
    o.Range("%s6:%s%d" % (KL, KL, ON)).Formula = "=" + norm("$G6&$I6"); o.Columns(KEYCOL).ColumnWidth = 30
    e = wbx.Worksheets("T6A1 Extract - 24-25"); key = norm("$%s5&$%s5" % (cJ, cM))
    f = ('=IF($%s5="","",IFERROR(INDEX(\'GSTR-2B ITC Data\'!$D$6:$D$%d,MATCH(%s,\'GSTR-2B ITC Data\'!$%s$6:$%s$%d,0)),'
         'IFERROR(INDEX(\'GSTR-2B Apr25-Aug26\'!$B$3:$B$%d,MATCH(%s,\'GSTR-2B Apr25-Aug26\'!$AW$3:$AW$%d,0)),"Not in 2B (Apr-24 to Aug-26)")))' % (cJ, ON, key, KL, KL, ON, NB, key, NB))
    src = [v[0] for v in e.Range("%s5:%s%d" % (cS, cS, EN)).Value]
    # simpler & exact: write via FormulaR1C1-free relative fill - set row 5 then FillDown, then overwrite RCM rows with text
    e.Range("C5").Formula = f; e.Range("C5:C%d" % EN).FillDown()
    rcm_rows = [r for r, s in zip(range(5, EN + 1), src) if "RCM" in str(s).upper()]
    for r in rcm_rows: e.Cells(r, 3).Value = "RCM self-invoice - not in 2B"
    e.Range("C5:C%d" % EN).NumberFormat = "mmm-yy"
    # totals row 3 (above the header, like the other sheets): SUBTOTAL so filters drive the figures - Pawan 18-09
    assert [hdr[14], hdr[15], hdr[16], hdr[17]] == ["Taxable Value", "IGST", "CGST", "SGST"], hdr[14:18]
    for col in ("O", "P", "Q", "R"):
        t = e.Range("%s3" % col); t.Formula = "=SUBTOTAL(9,%s5:%s%d)" % (col, col, EN); t.Font.Bold = True; t.NumberFormat = "#,##0.00"
    # helper col T: refresh 'reg!<row>' from the GSTR-9 Remarks pointer (Excel already shifted the pointers on the RCM insert)
    gfx = [v[0] for v in e.Range("%s5:%s%d" % (cG, cG, EN)).Formula]; tv = []
    for g in gfx:
        m = re.search(r"'ITC Register 2025-26'!\$[A-Z]+\$(\d+)", str(g)); tv.append(["reg!%s" % m.group(1) if m else None])
    old_t = [v[0] for v in e.Range("%s5:%s%d" % (cT, cT, EN)).Value]
    e.Range("%s5:%s%d" % (cT, cT, EN)).Value = [[t[0] if t[0] else o_] for t, o_ in zip(tv, old_t)]
    xl.Calculation = -4105; xl.CalculateFullRebuild()
    vals = [v[0] for v in e.Range("C5:C%d" % EN).Value]
    cnt = collections.Counter(("LY/CY period" if hasattr(v, "year") or isinstance(v, float) else str(v)[:32]) for v in vals)
    ly = cy = 0
    ok = [v[0] for v in o.Range("%s6:%s%d" % (KL, KL, ON)).Value]; oper = [v[0] for v in o.Range("D6:D%d" % ON).Value]; lyk = {k: p for k, p in zip(ok, oper) if k}
    ek = [v[0] for v in e.Range("%s5:%s%d" % (cJ, cJ, EN)).Value]; em = [v[0] for v in e.Range("%s5:%s%d" % (cM, cM, EN)).Value]
    nrm = lambda s: re.sub(r"[ \-/.'_]", "", str(s if s is not None else "").upper())
    for g, i, v in zip(ek, em, vals):
        if hasattr(v, "year") and (nrm(g) + nrm(i)) in lyk: ly += 1
        elif hasattr(v, "year"): cy += 1
    err = 0
    for w in wbx.Worksheets:
        try: err += w.UsedRange.SpecialCells(-4123, 16).Count
        except Exception: pass
    print("2B Return Period now:", dict(cnt), "| of the dated ones: FY 24-25 2B %d, Apr-25..Aug-26 2B %d" % (ly, cy))
    st = collections.Counter(); sv = [v[0] for v in e.Range("B5:B%d" % EN).Value]
    for s_, v in zip(sv, vals):
        if isinstance(v, str) and v.startswith("Not in 2B"): st[s_] += 1
    print("Not in 2B by state:", dict(st.most_common(8)), "| error cells:", err, "| %.0fs" % (time.time() - t0))
    wbx.Save(); wbx.Close(False)
finally: xl.Quit()
print("saved")
