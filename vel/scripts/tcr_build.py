"""Tax Comparison Report: merge 'ITC (Other than IMPG)' from the portal TCR files
(VEL PAN only), columns till the Shortfall block, month-on-month long format ->
master sheet 'Tax comp report' in last year's layout."""
import os, re, pandas as pd, openpyxl, datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
def S(v):
    if v is None: return ""
    if isinstance(v, float) and pd.isna(v): return ""
    return str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f = open(P, 'r+b'); f.close()
B = "//192.168.1.69/gst folder/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/DPS Workings/Portal Reports/Tax Comparison Reports/"
G2ST = {"01AAECR0503Q1ZM":"Jammu & Kashmir","03AAECR0503Q1ZI":"Punjab","06AAECR0503Q1ZC":"Haryana",
"08AAECR0503Q1Z8":"Rajasthan","09AAECR0503Q1Z6":"Uttar Pradesh","10AAECR0503Q1ZN":"Bihar",
"12AAECR0503Q1ZJ":"Arunachal Pradesh","18AAECR0503Q1Z7":"Assam","19AAECR0503Q1Z5":"West Bengal",
"20AAECR0503Q1ZM":"Jharkhand","22AAECR0503Q1ZI":"Chhattisgarh","23AAECR0503Q1ZG":"Madhya Pradesh",
"24AAECR0503Q1ZE":"Gujarat","27AAECR0503Q1Z8":"Maharashtra","29AAECR0503Q1Z4":"Karnataka",
"32AAECR0503Q1ZH":"Kerala","33AAECR0503Q1ZF":"Tamil Nadu","36AAECR0503Q1Z9":"Telangana",
"37AAECR0503Q1Z7":"Andhra Pradesh"}
MONPAT = re.compile(r"^(Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|Jan|Feb|Mar)-\d{2}$")
rows_out = []
seen_g = set()
for fn in sorted(os.listdir(B)):
    m = re.match(r"^2025-26_([0-9A-Z]{15})_", fn)
    if not m or not fn.lower().endswith(".xlsx") or fn.startswith("~$"): continue
    g = m.group(1)
    if "AAECR0503Q" not in g: continue          # VEL PAN only (subsidiaries excluded)
    seen_g.add(g)
    wb = openpyxl.load_workbook(B + fn, read_only=True, data_only=True)
    if "ITC (Other than IMPG)" not in wb.sheetnames:
        print("  %s: SHEET MISSING" % g); wb.close(); continue
    ws = wb["ITC (Other than IMPG)"]
    n = 0
    for r in ws.iter_rows(values_only=True):
        p = S(r[0]) if r else ""
        if not MONPAT.match(p): continue
        vals = [pd.to_numeric(r[j], errors="coerce") if len(r) > j else None for j in range(1, 13)]
        vals = [0.0 if (v is None or pd.isna(v)) else float(v) for v in vals]
        rows_out.append([G2ST.get(g, g), g, p] + vals)
        n += 1
    wb.close()
    print("  %-16s %s months %d" % (G2ST.get(g, "?"), g, n))
missing = sorted(set(G2ST) - seen_g)
print("VEL GSTINs with NO TCR file:", [G2ST[g] for g in missing])
df = pd.DataFrame(rows_out)
print("merged rows:", len(df))
# ---- build sheet
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79"); TOT = Font(bold=True); NF = "#,##0.00"
wb = openpyxl.load_workbook(P)
NM = "Tax comp report"
if NM in wb.sheetnames: del wb[NM]
pos = wb.sheetnames.index("RCM Paid vs ITC Claimed") + 1
ws = wb.create_sheet(NM, pos)
ws.row_dimensions[1].height = 21
ws.cell(2, 1, "VIKRAN ENGINEERING LIMITED").font = TOT
ws.cell(3, 1, "Tax Comparison Report & reasons for differences - ITC (Other than IMPG): merged from the portal 'Tax liability and ITC comparison' files (VEL GSTINs), columns till the Shortfall block.").font = TOT
# group header row 5 + sub header row 6 (last year's two-tier layout)
ws.cell(5, 3, "Tax Period").font = TOT
g1 = ws.cell(5, 4, "ITC claimed in GSTR-3B excluding ITC Reversal [Table 4A(4)+4A(5)-4B(2)]"); g1.font = HF; g1.fill = HB
g2 = ws.cell(5, 8, "ITC auto-drafted in GSTR-2B during the month"); g2.font = HF; g2.fill = HB
g3 = ws.cell(5, 12, "Shortfall (-) /Excess (+) in ITC [GSTR-3B - GSTR-2B]"); g3.font = HF; g3.fill = HB
ws.merge_cells(start_row=5, start_column=4, end_row=5, end_column=7)
ws.merge_cells(start_row=5, start_column=8, end_row=5, end_column=11)
ws.merge_cells(start_row=5, start_column=12, end_row=5, end_column=15)
sub = ["State", "GSTIN", "Month", "IGST", "CGST", "SGST/UTGST", "CESS", "IGST", "CGST", "SGST/UTGST", "CESS", "IGST", "CGST", "SGST/UTGST", "CESS"]
for c, h in enumerate(sub, 1):
    x = ws.cell(6, c, h); x.font = HF; x.fill = HB
r = 6
for row in rows_out:
    r += 1
    for c, v in enumerate(row, 1):
        cell = ws.cell(r, c, v)
        if c >= 4: cell.number_format = NF
last = r
r += 1
ws.cell(r, 1, "Grand Total").font = TOT
for c in range(4, 16):
    ws.cell(r, c).value = "=SUM(%s7:%s%d)" % (L(c), L(c), last)
    ws.cell(r, c).number_format = NF; ws.cell(r, c).font = TOT
# subtotal row 4 (respects filters)
for c in range(4, 16):
    ws.cell(4, c).value = "=SUBTOTAL(9,%s7:%s%d)" % (L(c), L(c), last)
    ws.cell(4, c).number_format = NF; ws.cell(4, c).font = TOT
ws.auto_filter.ref = "A6:O%d" % last
ws.freeze_panes = "D7"
for c_, w in zip(range(1, 16), [18, 18, 9] + [14] * 12): ws.column_dimensions[L(c_)].width = w
# INDEX row
ix = wb["INDEX"]
have = {S(ix.cell(r2, 2).value) for r2 in range(5, ix.max_row + 1)}
if "Tax Comparison Report (ITC other than IMPG)" not in have:
    thin = Side(style="thin"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
    r2 = ix.max_row + 1
    ix.cell(r2, 1, "ITC"); ix.cell(r2, 2, "Tax Comparison Report (ITC other than IMPG)")
    x = ix.cell(r2, 7); x.value = '=HYPERLINK("#\'%s\'!A1","Tax comp report")' % NM
    x.font = Font(color="0563C1", underline="single")
    for c in range(1, 11): ix.cell(r2, c).border = BD
wb.save(P)
tots = df[df.columns[3:]].sum()
print("built %s: %d rows | claimed IGST total %.2f | shortfall IGST total %.2f" % (NM, len(rows_out), tots.iloc[0], tots.iloc[8]))
