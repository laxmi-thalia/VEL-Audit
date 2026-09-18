"""Step F-1 (read-only survey): extract the 'AS PER PORTAL' block (client's computed 3B ITC tables) from every
state-month working file -> tcr_comp.pkl. Rows: 4A3, 4A4, 4A5, 4B1, 4B2, 4C, 4D1, 4D2 x IGST/CGST/SGST."""
import os, re, pickle, warnings, collections
warnings.filterwarnings("ignore")
import openpyxl
def S(v): return "" if v is None else str(v).strip()
def num(v):
    try: return float(v)
    except Exception: return 0.0
M = "//192.168.1.69/gst folder/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/Clients Data/05.08.2026 Main Data/"
MONTHS = ["01 April 2025", "02 May 2025", "03 June 2025", "04 July 2025", "05 Aug 2025", "06 Sep 2025", "07 Oct 2025", "08 Nov 2025", "09 Dec 2025", "10 Jan 2026", "11 Feb 2026", "12 Mar 2026"]
MLBL = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26"]
STATES = {"andhra pradesh": "37AAECR0503Q1Z7", "arunachal pradesh": "12AAECR0503Q1ZJ", "assam": "18AAECR0503Q1Z7", "bihar": "10AAECR0503Q1ZN",
          "chattisgarh": "22AAECR0503Q1ZI", "chhattisgarh": "22AAECR0503Q1ZI", "gujrat": "24AAECR0503Q1ZE", "gujarat": "24AAECR0503Q1ZE",
          "haryana": "06AAECR0503Q1ZC", "jammu & kashmir": "01AAECR0503Q1ZM", "jharkhand": "20AAECR0503Q1ZM", "karnataka": "29AAECR0503Q1Z4",
          "kerala": "32AAECR0503Q1ZH", "madhya pradesh": "23AAECR0503Q1ZG", "maharashtra": "27AAECR0503Q1Z8", "punjab": "03AAECR0503Q1ZI",
          "rajasthan": "08AAECR0503Q1Z8", "tamil nadu": "33AAECR0503Q1ZF", "telangana": "36AAECR0503Q1Z9", "uttar pradesh": "09AAECR0503Q1Z6",
          "west bengal": "19AAECR0503Q1Z5"}
KEYS = [("4A3", "(3) inward supplies liable"), ("4A4", "(4) inward supplies from isd"), ("4A5", "(5) all other itc"), ("4B1", "(1) as per rules"),
        ("4B2", "(2) others"), ("4C", "(c) net itc"), ("4D1", "(1) itc reclaimed"), ("4D2", "(2) ineligible")]
out = {}; missing = []
for mi, mf in enumerate(MONTHS):
    for d in sorted(os.listdir(M + mf)):
        key = d.strip().lower()
        if key not in STATES: continue
        g = STATES[key]; folder = M + mf + "/" + d + "/"
        cands = [x for x in os.listdir(folder) if x.lower().startswith("gstr-3b") and x.lower().endswith(".xlsx") and not x.startswith("~$")]
        if not cands: missing.append((MLBL[mi], d, "no GSTR-3B xlsx")); continue
        try: wb = openpyxl.load_workbook(folder + cands[0], read_only=True, data_only=True)
        except Exception as e: missing.append((MLBL[mi], d, "open failed")); continue
        sn = "Computation New" if "Computation New" in wb.sheetnames else ("Computation" if "Computation" in wb.sheetnames else None)
        if not sn: missing.append((MLBL[mi], d, "no Computation sheet: %s" % wb.sheetnames[:6])); wb.close(); continue
        ws = wb[sn]; rows = list(ws.iter_rows(values_only=True, max_col=8))
        # anchor on the '(3) Inward supplies liable' label whichever column it sits in; values = next 3 cols; take the LAST
        # complete block (the AS PER PORTAL block sits below the books block in every layout seen)
        anchors = [(i, c) for i, r in enumerate(rows) for c in range(3) if S(r[c]).lower().startswith("(3) inward supplies liable")]
        rec = {}; anchor_used = None
        for i, c in reversed(anchors):
            block = rows[max(0, i - 3): i + 14]; tmp = {}
            for r in block:
                lab = S(r[c]).lower()
                for k, pat in KEYS:
                    if k not in tmp and lab.startswith(pat): tmp[k] = [num(r[c + 1]), num(r[c + 2]), num(r[c + 3])]
            if len(tmp) >= 6: rec = tmp; anchor_used = (i + 1, c + 1); break
        if not rec: missing.append((MLBL[mi], d, "portal block not found in %s" % sn)); wb.close(); continue
        out[(g, MLBL[mi])] = {"file": mf + "/" + d + "/" + cands[0], "sheet": sn, "vals": rec, "found": sorted(rec), "anchor": anchor_used}
        wb.close()
    print("month %s: %d files" % (MLBL[mi], sum(1 for k in out if k[1] == MLBL[mi])), flush=True)
pickle.dump({"data": out, "missing": missing}, open("tcr_comp.pkl", "wb"))
print("\nextracted:", len(out), "| missing:", len(missing))
for x in missing: print("  ", x)
inc = [(k, v["found"]) for k, v in out.items() if len(v["found"]) < 8]
print("incomplete blocks:", len(inc)); [print("  ", x) for x in inc[:15]]
