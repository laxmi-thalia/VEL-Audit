"""Step F-2 (read-only survey): from every state-month GSTR-3B working file capture, out of the 'Computation New'
(else 'Computation') sheet:
  PART A (CA's 2B->3B working): A_4A5, A_4A4, A_4B2, A_2B (2B current month ITC as per portal), A_matched, A_cf (carry
      forward = 4B2 others), A_reclaim, A_cf_earlier, A_cf_portal, A_net, A_4A3, A_permrev
  PORTAL block (what was filed): P_4A3 P_4A4 P_4A5 P_4B1 P_4B2 P_4C P_4D1 P_4D2
  RECON block of the month: R_2B, R_matched, R_reclaim, R_acc (accounting correction), R_amend (amendment/expense out)
Value columns are anchored on the nearest IGST/CGST/SGST header row above the label (PART A sits under E/F/G, the portal
block under C/D/E); 'Permanent reversal' has no header -> first 3 cells after the label. -> tcr_comp2.pkl"""
import os, re, pickle, warnings, collections
warnings.filterwarnings("ignore")
import openpyxl
def S(v): return "" if v is None else str(v).strip()
def num(v):
    if v is None or v == "": return 0.0
    try: return float(v)
    except Exception: return 0.0
M = "//192.168.1.69/gst folder/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/Clients Data/05.08.2026 Main Data/"
MONTHS = ["01 April 2025", "02 May 2025", "03 June 2025", "04 July 2025", "05 Aug 2025", "06 Sep 2025", "07 Oct 2025", "08 Nov 2025", "09 Dec 2025", "10 Jan 2026", "11 Feb 2026", "12 Mar 2026"]
MLBL = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26"]
MNAME = ["apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec", "jan", "feb", "mar"]
STATES = {"andhra pradesh": "37AAECR0503Q1Z7", "arunachal pradesh": "12AAECR0503Q1ZJ", "assam": "18AAECR0503Q1Z7", "bihar": "10AAECR0503Q1ZN",
          "chattisgarh": "22AAECR0503Q1ZI", "chhattisgarh": "22AAECR0503Q1ZI", "gujrat": "24AAECR0503Q1ZE", "gujarat": "24AAECR0503Q1ZE",
          "haryana": "06AAECR0503Q1ZC", "jammu & kashmir": "01AAECR0503Q1ZM", "jharkhand": "20AAECR0503Q1ZM", "karnataka": "29AAECR0503Q1Z4",
          "kerala": "32AAECR0503Q1ZH", "madhya pradesh": "23AAECR0503Q1ZG", "maharashtra": "27AAECR0503Q1Z8", "punjab": "03AAECR0503Q1ZI",
          "rajasthan": "08AAECR0503Q1Z8", "tamil nadu": "33AAECR0503Q1ZF", "telangana": "36AAECR0503Q1Z9", "uttar pradesh": "09AAECR0503Q1Z6",
          "west bengal": "19AAECR0503Q1Z5"}
A_KEYS = [("A_4A5", r"^4\(a\)\(5\)"), ("A_4A4", r"^4\(a\)\(4\)"), ("A_4B2", r"^4\(b\)\(2\)"), ("A_2B", r"^(gstr ?2b|2b) current month itc as per portal"),
          ("A_matched", r"^matched credit as per books"), ("A_cf", r"^to be carry forward itc as per gstr ?2b of this month"),
          ("A_reclaim", r"^matched credit of earlier month as per gstr ?2b"), ("A_cf_earlier", r"^to be carry forward itc as per gstr ?2b of earlier month"),
          ("A_cf_portal", r"^to be carry forward as per portal"), ("A_net", r"^net itc for the month"), ("A_4A3", r"^3\.1\(d\) 4\(a\)\(3\)"),
          ("A_permrev", r"^permanent reversal")]
P_KEYS = [("P_4A3", r"^\(3\) inward supplies liable"), ("P_4A4", r"^\(4\) inward supplies from isd"), ("P_4A5", r"^\(5\) all other itc"), ("P_4B1", r"^\(1\) as per rules"),
          ("P_4B2", r"^\(2\) others"), ("P_4C", r"^\(c\) net itc"), ("P_4D1", r"^\(1\) itc reclaimed"), ("P_4D2", r"^\(2\) ineligible")]
R_KEYS = [("R_open", r"^opening balance itc as per gstr ?2b"), ("R_2B", r"^add ?: ?gstr ?2b for the month"), ("R_matched", r"^less ?: ?matched itc"), ("R_reclaim", r"^less ?: ?reclaimed"),
          ("R_acc", r"^add ?: ?accounting correction"), ("R_amend", r"^less ?:? ?amendment"), ("R_cf", r"^to be carry forward to next month")]
A_RE = [(k, re.compile(p)) for k, p in A_KEYS]; P_RE = [(k, re.compile(p)) for k, p in P_KEYS]; R_RE = [(k, re.compile(p)) for k, p in R_KEYS]
def is_hdr(r):
    for c in range(0, 6):
        if S(r[c]).upper() == "IGST" and S(r[c + 1]).upper() == "CGST" and S(r[c + 2]).upper().startswith("SGST"): return c
    return None
out = {}; missing = []; t = collections.Counter()
for mi, mf in enumerate(MONTHS):
    for d in sorted(os.listdir(M + mf)):
        key = d.strip().lower()
        if key not in STATES: continue
        g = STATES[key]; folder = M + mf + "/" + d + "/"
        cands = [x for x in os.listdir(folder) if re.match(r"gstr[ -]?3b", x.lower()) and x.lower().endswith(".xlsx") and not x.startswith("~$")]
        if not cands: missing.append((MLBL[mi], d, "no GSTR-3B xlsx")); continue
        try: wb = openpyxl.load_workbook(folder + cands[0], read_only=True, data_only=True)
        except Exception as e: missing.append((MLBL[mi], d, "open failed")); continue
        new = [s for s in wb.sheetnames if re.match(r"^computation new\s*$", s, re.I)]; old = [s for s in wb.sheetnames if re.match(r"^computation\s*$", s, re.I)]
        sn = (new or old or [None])[0]
        if not sn: missing.append((MLBL[mi], d, "no Computation sheet: %s" % wb.sheetnames[:6])); wb.close(); continue
        ws = wb[sn]; rows = list(ws.iter_rows(values_only=True, max_col=10)); rows = [tuple(r) + (None,) * (10 - len(r)) for r in rows]
        rec = {}; meth = {}; hdr_col = None; in_recon = None; recon_title = None
        for i, r in enumerate(rows):
            h = is_hdr(r)
            if h is not None: hdr_col = h
            for c in range(0, 4):
                lab = S(r[c]).lower()
                if not lab: continue
                if lab.startswith("reconciliation of itc carry forward for the month"):
                    in_recon = (MNAME[mi] in lab and (("2025" in lab and mi < 9) or ("2026" in lab and mi >= 9))); recon_title = lab if in_recon else recon_title; continue
                vals_h = [num(r[hdr_col + j]) for j in range(3)] if hdr_col is not None else None
                after = [x for x in r[c + 1:c + 8] if x is not None][:3]; vals_a = [num(x) for x in after] if len(after) == 3 else None
                for k, rx in A_RE:
                    if k not in rec and rx.search(lab):
                        v = vals_a if k == "A_permrev" else vals_h
                        if v is not None: rec[k] = v; meth[k] = i + 1
                for k, rx in P_RE:
                    if rx.search(lab) and vals_h is not None: rec[k] = vals_h; meth[k] = i + 1     # last occurrence wins (portal block is lowest)
                if in_recon:
                    for k, rx in R_RE:
                        if k not in rec and rx.search(lab) and vals_h is not None: rec[k] = vals_h; meth[k] = i + 1
        wb.close()
        if not any(k.startswith("P_") for k in rec): missing.append((MLBL[mi], d, "portal block not found in %r" % sn)); continue
        chk = {}
        if all(k in rec for k in ("A_4A5", "A_4B2", "A_2B")): chk["A: 4A5-4B2=2B"] = max(abs(rec["A_4A5"][j] - rec["A_4B2"][j] - rec["A_2B"][j]) for j in range(3)) < 1.5
        if all(k in rec for k in ("P_4A3", "P_4A4", "P_4A5", "P_4B1", "P_4B2", "P_4C")): chk["P: 4A-4B=4C"] = max(abs(rec["P_4A3"][j] + rec["P_4A4"][j] + rec["P_4A5"][j] - rec["P_4B1"][j] - rec["P_4B2"][j] - rec["P_4C"][j]) for j in range(3)) < 1.5
        for k, ok in chk.items(): t[(k, ok)] += 1
        out[(g, MLBL[mi])] = {"file": mf + "/" + d + "/" + cands[0], "sheet": sn, "vals": rec, "rows": meth, "checks": chk, "recon_title": recon_title}
    print("month %s: %d files" % (MLBL[mi], sum(1 for k in out if k[1] == MLBL[mi])), flush=True)
pickle.dump({"data": out, "missing": missing}, open("tcr_comp2.pkl", "wb"))
print("\nextracted:", len(out), "| missing:", len(missing)); [print("  ", x) for x in missing]
print("identity checks:", dict(t))
print("key coverage:", collections.Counter(k for v in out.values() for k in v["vals"]))
