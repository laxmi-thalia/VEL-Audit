"""Tax comp report - original column structure (A..AO incl Reasons/Impact, analysis blocks,
Difference formulas, 9C-note/Query cols), SAFE formats only. Data = merged ITC (Other than
IMPG) from portal TCR files, VEL PAN, months only, till Shortfall block."""
import os, re, pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
def S(v):
    if v is None: return ""
    if isinstance(v, float) and pd.isna(v): return ""
    return str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f = open(P, 'r+b'); f.close()
B = "//192.168.1.69/gst folder/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/DPS Workings/Portal Reports/Tax Comparison Reports/"
G2ST = {"01AAECR0503Q1ZM":"Jammu & Kashmir","03AAECR0503Q1ZI":"Punjab","08AAECR0503Q1Z8":"Rajasthan",
"09AAECR0503Q1Z6":"Uttar Pradesh","10AAECR0503Q1ZN":"Bihar","12AAECR0503Q1ZJ":"Arunachal Pradesh",
"18AAECR0503Q1Z7":"Assam","19AAECR0503Q1Z5":"West Bengal","20AAECR0503Q1ZM":"Jharkhand",
"22AAECR0503Q1ZI":"Chhattisgarh","23AAECR0503Q1ZG":"Madhya Pradesh","24AAECR0503Q1ZE":"Gujarat",
"27AAECR0503Q1Z8":"Maharashtra","29AAECR0503Q1Z4":"Karnataka","32AAECR0503Q1ZH":"Kerala",
"33AAECR0503Q1ZF":"Tamil Nadu","36AAECR0503Q1Z9":"Telangana"}
MONPAT = re.compile(r"^(Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|Jan|Feb|Mar)-\d{2}$")
MORD = ["Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec","Jan","Feb","Mar"]
rows_out = []
for fn in sorted(os.listdir(B)):
    m = re.match(r"^2025-26_([0-9A-Z]{15})_", fn)
    if not m or not fn.lower().endswith(".xlsx") or fn.startswith("~$"): continue
    g = m.group(1)
    if "AAECR0503Q" not in g: continue
    wb0 = openpyxl.load_workbook(B + fn, read_only=True, data_only=True)
    if "ITC (Other than IMPG)" not in wb0.sheetnames: wb0.close(); continue
    ws0 = wb0["ITC (Other than IMPG)"]
    for r in ws0.iter_rows(values_only=True):
        p = S(r[0]) if r else ""
        if not MONPAT.match(p): continue
        vals = [pd.to_numeric(r[j], errors="coerce") if len(r) > j else None for j in range(1, 13)]
        vals = [0.0 if (v is None or pd.isna(v)) else float(v) for v in vals]
        rows_out.append([G2ST.get(g, g), g, p] + vals)
    wb0.close()
rows_out.sort(key=lambda x: (x[0], MORD.index(x[2][:3])))
N = len(rows_out)
print("data rows:", N)
HF = Font(bold=True, color="FFFFFF"); HB = PatternFill("solid", fgColor="1F4E79")
TOT = Font(bold=True); NF = "#,##0.00"; CTR = Alignment(horizontal="center", wrap_text=True)
wb = openpyxl.load_workbook(P)
NM = "Tax comp report"
if NM in wb.sheetnames: del wb[NM]
pos = wb.sheetnames.index("RCM Paid vs ITC Claimed") + 1
ws = wb.create_sheet(NM, pos)
ws.row_dimensions[1].height = 21
ws.cell(2, 1, "VIKRAN ENGINEERING LIMITED").font = TOT
ws.cell(3, 1, "Tax Comparison Report & reasons for differences").font = TOT
D0 = 8; DN = D0 + N - 1
# row 5: group headers + (cols S..) analysis totals; row 6: totals under data blocks + merged analysis labels; row 7: sub-headers
g1 = ws.cell(5, 4, "ITC claimed in GSTR-3B excluding ITC Reversal [Table 4A(4)+4A(5)-4B(2)]")
g2 = ws.cell(5, 8, "ITC auto-drafted in GSTR-2B during the month")
g3 = ws.cell(5, 12, "Shortfall (-) /Excess (+) in ITC [GSTR-3B - GSTR-2B]")
for x in (g1, g2, g3): x.font = HF; x.fill = HB; x.alignment = CTR
ws.merge_cells(start_row=5, start_column=4, end_row=5, end_column=7)
ws.merge_cells(start_row=5, start_column=8, end_row=5, end_column=11)
ws.merge_cells(start_row=5, start_column=12, end_row=5, end_column=15)
ws.cell(5, 3, "Tax Period").font = TOT
for c in range(19, 34):   # header-area totals over analysis cols
    ws.cell(5, c).value = "=SUM(%s%d:%s%d)" % (L(c), D0, L(c), DN)
    ws.cell(5, c).number_format = NF; ws.cell(5, c).font = TOT
for c in range(4, 16):    # totals over data blocks
    ws.cell(6, c).value = "=SUM(%s%d:%s%d)" % (L(c), D0, L(c), DN)
    ws.cell(6, c).number_format = NF; ws.cell(6, c).font = TOT
LBL = [(19, 21, "Difference to be added to 3B of next year"),
       (22, 24, "Difference in GSTR-2B Amounts"),
       (25, 27, "(Short)/ Excess reported in GSTR-1"),
       (28, 30, "Short/ (Excess) reporting of ITC"),
       (31, 33, "Difference")]
for c1, c2, t in LBL:
    x = ws.cell(6, c1, t); x.font = TOT; x.alignment = CTR
    ws.merge_cells(start_row=6, start_column=c1, end_row=6, end_column=c2)
sub = ["State","GSTIN","Month","IGST","CGST","SGST/UTGST","CESS","IGST","CGST","SGST/UTGST","CESS",
       "IGST","CGST","SGST/UTGST","CESS","Total","Reasons","Impact",
       "IGST","CGST","SGST","IGST","CGST","SGST","IGST","CGST","SGST","IGST","CGST","SGST",
       "IGST","CGST","SGST/UTGST","Total - 4A5 (Excess/ short)","Note to be added in GSTR-9C",
       "Amounts for 9C note","Amounts for 9C note - CN","GSTR-9C - Note 1","GSTR-9C - Note 2",
       "Query","Query Amount"]
for c, h in enumerate(sub, 1):
    x = ws.cell(7, c, h); x.font = HF; x.fill = HB; x.alignment = CTR
for i, row in enumerate(rows_out):
    r = D0 + i
    for c in range(1, 16):
        cell = ws.cell(r, c, row[c - 1])
        if c >= 4: cell.number_format = NF
    ws.cell(r, 16).value = "=SUM(L%d:N%d)" % (r, r)
    ws.cell(r, 31).value = "=L%d-S%d-V%d-Y%d" % (r, r, r, r)
    ws.cell(r, 32).value = "=M%d-T%d-W%d-Z%d" % (r, r, r, r)
    ws.cell(r, 33).value = "=N%d-U%d-X%d-AA%d" % (r, r, r, r)
    ws.cell(r, 34).value = "=SUM(V%d:AD%d)" % (r, r)
    for c in (16, 31, 32, 33, 34): ws.cell(r, c).number_format = NF
for c, w in zip(range(1, 42), [18, 17, 9] + [13] * 12 + [12, 22, 10] + [12] * 15 + [26, 14, 16, 14, 14, 18, 14]):
    ws.column_dimensions[L(c)].width = w
ws.freeze_panes = "D8"
ws.auto_filter.ref = "A7:AO%d" % DN
# INDEX row (restored master lost the earlier one)
ix = wb["INDEX"]
have = {S(ix.cell(r2, 2).value) for r2 in range(5, ix.max_row + 1)}
if "Tax Comparison Report (ITC other than IMPG)" not in have:
    thin = Side(style="thin"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
    r2 = ix.max_row + 1
    ix.cell(r2, 1, "ITC"); ix.cell(r2, 2, "Tax Comparison Report (ITC other than IMPG)")
    x = ix.cell(r2, 7); x.value = '=HYPERLINK("#\'%s\'!A1","Tax comp report")' % NM
    x.font = Font(color="0563C1", underline="single")
    for c in range(1, 11): ix.cell(r2, c).border = BD
wb.save(P)
print("built", NM)
