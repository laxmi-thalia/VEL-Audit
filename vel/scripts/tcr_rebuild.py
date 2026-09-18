"""Rebuild 'Tax comp report' as an EXACT replica of last year's sheet (structure, two-tier
header with live totals, per-row formulas, judgment columns blank) with this year's merged
'ITC (Other than IMPG)' data. Stray helper columns beyond AO dropped (side-scratch in golden)."""
import os, re, pandas as pd, openpyxl, json
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
def S(v):
    if v is None: return ""
    if isinstance(v, float) and pd.isna(v): return ""
    return str(v).strip()
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER.xlsx"
f = open(P, 'r+b'); f.close()
d = json.load(open("tcr_dump.json")); cells = d["cells"]
def cel(r, c): return cells.get("%d,%d" % (r, c), {})
# ---- this year's data (same parse as tcr_build)
B = "//192.168.1.69/gst folder/GST Returns/GST Audit & Annual Return/FY 2025-26/1. Corporate Clients/VEL/DPS Workings/Portal Reports/Tax Comparison Reports/"
G2ST = {"01AAECR0503Q1ZM":"Jammu & Kashmir","03AAECR0503Q1ZI":"Punjab","06AAECR0503Q1ZC":"Haryana",
"08AAECR0503Q1Z8":"Rajasthan","09AAECR0503Q1Z6":"Uttar Pradesh","10AAECR0503Q1ZN":"Bihar",
"12AAECR0503Q1ZJ":"Arunachal Pradesh","18AAECR0503Q1Z7":"Assam","19AAECR0503Q1Z5":"West Bengal",
"20AAECR0503Q1ZM":"Jharkhand","22AAECR0503Q1ZI":"Chhattisgarh","23AAECR0503Q1ZG":"Madhya Pradesh",
"24AAECR0503Q1ZE":"Gujarat","27AAECR0503Q1Z8":"Maharashtra","29AAECR0503Q1Z4":"Karnataka",
"32AAECR0503Q1ZH":"Kerala","33AAECR0503Q1ZF":"Tamil Nadu","36AAECR0503Q1Z9":"Telangana",
"37AAECR0503Q1Z7":"Andhra Pradesh"}
MONPAT = re.compile(r"^(Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|Jan|Feb|Mar)-\d{2}$")
rows_out = []
for fn in sorted(os.listdir(B)):
    m = re.match(r"^2025-26_([0-9A-Z]{15})_", fn)
    if not m or not fn.lower().endswith(".xlsx") or fn.startswith("~$"): continue
    g = m.group(1)
    if "AAECR0503Q" not in g: continue
    wb0 = openpyxl.load_workbook(B + fn, read_only=True, data_only=True)
    if "ITC (Other than IMPG)" not in wb0.sheetnames: wb0.close(); continue
    ws0 = wb0["ITC (Other than IMPG)"]
    for r in ws0.iter_rows(values_only=True):
        p = S(r[0]) if r else ""
        if not MONPAT.match(p): continue
        vals = [pd.to_numeric(r[j], errors="coerce") if len(r) > j else None for j in range(1, 13)]
        vals = [0.0 if (v is None or pd.isna(v)) else float(v) for v in vals]
        rows_out.append([G2ST.get(g, g), g, p] + vals)
    wb0.close()
rows_out.sort(key=lambda x: (x[0], ["Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec","Jan","Feb","Mar"].index(x[2][:3])))
print("data rows:", len(rows_out))
N = len(rows_out)
# ---- rebuild sheet (orig rows 1..6 -> master rows 2..7 due to reserved row 1; data from row 8)
BGR2HEX = lambda x: "%02X%02X%02X" % (x & 0xFF, (x >> 8) & 0xFF, (x >> 16) & 0xFF)
wb = openpyxl.load_workbook(P)
NM = "Tax comp report"
pos = wb.sheetnames.index(NM) if NM in wb.sheetnames else wb.sheetnames.index("RCM Paid vs ITC Claimed") + 1
if NM in wb.sheetnames: del wb[NM]
ws = wb.create_sheet(NM, pos)
ws.row_dimensions[1].height = 21
MAXC = 41
R_OFF = 1                      # original row r -> master row r + 1
D0 = 7 + R_OFF                 # first data row in master = 8
DN = D0 + N - 1
def style_from(x, cell):
    fill = x.get("fill")
    if fill is not None: cell.fill = PatternFill("solid", fgColor=BGR2HEX(fill))
    cell.font = Font(bold=x.get("b", False), color=BGR2HEX(x.get("fc", 0)))
    nf = x.get("nf")
    if nf and nf != "General": cell.number_format = nf
    al = {}
    if x.get("ha") == -4108: al["horizontal"] = "center"
    if x.get("wrap"): al["wrap_text"] = True
    if al: cell.alignment = Alignment(**al)
# header zone: original rows 1..6 -> copy labels/styles; numeric totals -> live SUMs
for r0 in range(1, 7):
    for c in range(1, MAXC + 1):
        x = cel(r0, c)
        if not x: continue
        cell = ws.cell(r0 + R_OFF, c)
        v = x.get("v")
        isnum = False
        if v is not None:
            try: float(v); isnum = True
            except Exception: isnum = False
        if x.get("f") or isnum:
            cell.value = "=SUM(%s%d:%s%d)" % (L(c), D0, L(c), DN)
            cell.number_format = "#,##0.00"
        else:
            cell.value = v
        style_from(x, cell)
# merged label areas (+1 row)
for a in d["merged"]:
    m = re.match(r"\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", a)
    if m:
        ws.merge_cells("%s%d:%s%d" % (m.group(1), int(m.group(2)) + R_OFF, m.group(3), int(m.group(4)) + R_OFF))
# data rows: A-O values, P/AE-AH formulas, judgment cols blank; styles from original row 7
proto = {c: cel(7, c) for c in range(1, MAXC + 1)}
for i, row in enumerate(rows_out):
    r = D0 + i
    for c in range(1, 16):
        cell = ws.cell(r, c, row[c - 1] if c <= len(row) else None)
        if proto.get(c): style_from(proto[c], cell)
        if c >= 4: cell.number_format = "#,##0.00"
    ws.cell(r, 16).value = "=SUM(L%d:N%d)" % (r, r)
    ws.cell(r, 31).value = "=L%d-S%d-V%d-Y%d" % (r, r, r, r)
    ws.cell(r, 32).value = "=M%d-T%d-W%d-Z%d" % (r, r, r, r)
    ws.cell(r, 33).value = "=N%d-U%d-X%d-AA%d" % (r, r, r, r)
    ws.cell(r, 34).value = "=SUM(V%d:AD%d)" % (r, r)
    for c in (16, 31, 32, 33, 34): ws.cell(r, c).number_format = "#,##0.00"
for c_str, w in d["colw"].items():
    c = int(c_str)
    if c <= MAXC: ws.column_dimensions[L(c)].width = w
ws.freeze_panes = "D%d" % D0
ws.auto_filter.ref = "A%d:AO%d" % (D0 - 1, DN)
wb.save(P)
print("rebuilt %s: %d data rows, header rows 2-7, data %d..%d" % (NM, N, D0, DN))
