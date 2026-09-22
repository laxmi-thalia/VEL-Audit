"""'1 –' filter tie between ITC Register 2025-26 and GSTR-2B Apr25-Aug26 (Pawan 22-09, CA Priyesh 21-09 31:13):
 (1) RCM-category register lines whose vendor GSTIN + invoice sit in 2B -> remark 12 'Not applicable – RCM line (in 2B: …)' (values),
     so '1 –' holds ITC lines only on both sheets (the 2B side mirrors the register remark by key);
 (2) helper column '2B pull' (values): 'Yes' on the first Consider line per KEY2, 'No' on a later Consider line with the same KEY2
     (duplicate SAP booking of one invoice) -> 2B_ pulls 0 there, D_ = B_ surfaces the duplicate claim;
 (3) 2B sheet remark lookup: matches that come from ITC Register 2026-27 carry the suffix ' (ITCR 26-27)'.
Single COM session; ITC Summary blocks / golden / Net-ITC must be unchanged; prints both '1 –' subtotals (must tie)."""
import re, time, shutil, collections, warnings, openpyxl, win32com.client as win32, pythoncom
from openpyxl.utils import get_column_letter as L
warnings.filterwarnings("ignore")
P = r"C:\Users\pawar\Downloads\VEL_GST_Audit_FY2025-26_MASTER (2).xlsx"; B = P[:-5] + ".xlsb"
try: open(P, "r+b").close()
except PermissionError: raise SystemExit("LOCKED - close in Excel without saving: " + P)
def xlsb_free():
    try: open(B, "r+b").close(); return True
    except PermissionError: return False
shutil.copy(P, "master2_snapshot_before_tiefix.xlsx")
S = lambda v: "" if v is None else str(v).strip()
BASIS = {"1": "invoice no", "2": "invoice no", "3": "invoice no", "4": "amount & date", "5": "similar invoice + amount", "6": "GSTIN + amount",
         "7": "date + amount", "8": "amount", "9": "FY 24-25 2B"}
RCM12 = "12 – Not applicable – RCM line (in 2B: %s)"
# ---- pre-read (values) ------------------------------------------------------------------------------------------------
wb = openpyxl.load_workbook(P, read_only=True, data_only=True); r = wb["ITC Register 2025-26"]
hd = [c.value for c in next(r.iter_rows(min_row=5, max_row=5))]; H = {h: i for i, h in enumerate(hd) if h}
rows = list(r.iter_rows(min_row=6, values_only=True)); RN = 5 + max(i for i, x in enumerate(rows, 1) if x[3])
rows = rows[:RN - 5]; K, R, C, CAT = H["KEY2 (matched 2B key)"], H["Reco Remarks"], H["Countif"], H["Category"]
new_rem = {}   # excel row -> new remark text
for i, x in enumerate(rows):
    v = S(x[R]); num = v.split(" – ")[0]
    if S(x[CAT]) == "RCM" and S(x[K]) and num in BASIS: new_rem[6 + i] = RCM12 % BASIS[num]
seen = set(); pull = []
for i, x in enumerate(rows):
    if S(x[C]) == "Consider" and S(x[K]):
        if x[K] in seen: pull.append("No")
        else: seen.add(x[K]); pull.append("Yes")
    else: pull.append(None)
print("register rows %d | RCM lines re-labelled 12: %d %s | '2B pull' No: %d" % (len(rows), len(new_rem), dict(collections.Counter(v.split("(in 2B")[1][:14] for v in new_rem.values())), pull.count("No")))
b2 = wb["GSTR-2B Apr25-Aug26"]; BH = {c.value: i for i, c in enumerate(next(b2.iter_rows(min_row=2, max_row=2))) if c.value}
NB = 2 + sum(1 for x in b2.iter_rows(min_row=3, values_only=True) if x[0]); wb.close()
# ---- COM ---------------------------------------------------------------------------------------------------------------
pythoncom.CoInitialize(); xl = win32.DispatchEx("Excel.Application"); xl.Visible = False; xl.DisplayAlerts = False
try:
    t0 = time.time(); wbx = xl.Workbooks.Open(P); xl.Calculation = -4135
    isum = wbx.Worksheets("ITC Summary"); snap = lambda: [round(isum.Cells(25, c).Value or 0, 2) for c in list(range(7, 22)) + [49, 50, 51, 83, 84, 85]]; before = snap()
    rg = wbx.Worksheets("ITC Register 2025-26"); assert S(rg.Cells(5, R + 1).Value) == "Reco Remarks" and rg.Cells(rg.Rows.Count, 4).End(-4162).Row == RN
    for er, txt in new_rem.items(): rg.Cells(er, R + 1).Value = txt
    # helper column after the last header
    last = max(H.values()) + 1; PC = last + 1; hc = rg.Cells(5, PC); src = rg.Cells(5, C + 1)
    hc.Value = "2B pull"; hc.Font.Bold = src.Font.Bold; hc.Font.Color = src.Font.Color; hc.Interior.Color = src.Interior.Color
    rg.Range(rg.Cells(6, PC), rg.Cells(RN, PC)).Value = [[v] for v in pull]; rg.Columns(PC).ColumnWidth = 9
    PL, CF = L(PC), L(C + 1)
    for h in ("2B_IGST", "2B_CGST", "2B_SGST"):
        col = L(H[h] + 1); old = rg.Cells(6, H[h] + 1).Formula
        head = '=IF($%s6="Consider",' % CF; assert old.startswith(head) and old.endswith(',"NA")'), (h, old[:80])
        new = head + 'IF($%s6="No",0,' % PL + old[len(head):-len(',"NA")')] + '),"NA")'
        rg.Range("%s6:%s%d" % (col, col, RN)).Formula = new
    if rg.AutoFilterMode: rg.AutoFilterMode = False
    rg.Range(rg.Cells(5, 1), rg.Cells(RN, PC)).AutoFilter()
    # 2B sheet remark lookup with the 26-27 suffix
    bs = wbx.Worksheets("GSTR-2B Apr25-Aug26"); RC = L(BH["Reco Remarks"] + 1); f = bs.Cells(3, BH["Reco Remarks"] + 1).Formula
    m = re.match(r"^=IFERROR\((INDEX\('ITC Register 2025-26'!.*?\)),IFERROR\((INDEX\('ITC Register 2026-27'!.*?\)),(\"11 – [^\"]*\")\)\)$", f)
    assert m, f[:200]
    bs.Range("%s3:%s%d" % (RC, RC, NB)).Formula = '=IFERROR(%s,IFERROR(%s&" (ITCR 26-27)",%s))' % (m.group(1), m.group(2), m.group(3))
    xl.Calculation = -4105; xl.CalculateFullRebuild()
    # ---- verify ----
    e = 0
    for w in wbx.Worksheets:
        try: e += w.UsedRange.SpecialCells(-4123, 16).Count
        except Exception: pass
    after = snap(); golden = rg.Cells(4, H["Total GST"] + 1).Value
    col = lambda sh, c, r0, r1: [v[0] for v in sh.Range("%s%d:%s%d" % (c, r0, c, r1)).Value]
    rem = col(rg, L(R + 1), 6, RN); cf = col(rg, CF, 6, RN); num = lambda v: v if isinstance(v, (int, float)) else 0.0
    reg2b = [sum(num(v) for v, m_, c_ in zip(col(rg, L(H[h] + 1), 6, RN), rem, cf) if c_ == "Consider" and S(m_).startswith("1 –")) for h in ("2B_IGST", "2B_CGST", "2B_SGST")]
    regb = [sum(num(v) for v, m_, c_ in zip(col(rg, L(H[h] + 1), 6, RN), rem, cf) if c_ == "Consider" and S(m_).startswith("1 –")) for h in ("B_IGST", "B_CGST", "B_SGST")]
    brem = col(bs, RC, 3, NB)
    b2s = [sum(num(v) for v, m_ in zip(col(bs, L(BH[h] + 1), 3, NB), brem) if S(m_).startswith("1 –") and "(ITCR 26-27)" not in S(m_)) for h in ("IGST (Net)", "CGST (Net)", "SGST (Net)")]
    dist = collections.Counter(S(v).split(" – ")[0] + (" (26-27)" if "(ITCR 26-27)" in S(v) else "") for v in brem)
    print("error cells %d | golden %.2f | Net-ITC %s | ITC Summary blocks unchanged %s" % (e, golden, [round(isum.Cells(25, 49 + j).Value or 0, 2) for j in range(3)], before == after))
    print("register '1 –' Consider: B_ %s | 2B_ %s" % ([round(v, 2) for v in regb], [round(v, 2) for v in reg2b]))
    print("2B sheet '1 –' (excl. ITCR 26-27): IGST/CGST/SGST (Net) %s" % [round(v, 2) for v in b2s])
    print("2B sheet remark numbers:", dict(sorted(dist.items())))
    print("register Reco Remarks numbers:", dict(sorted(collections.Counter(S(v).split(" – ")[0] for v in rem).items())))
    assert e == 0 and before == after and abs(golden - 1069969542.15) < 0.01 and all(abs(a - b_) < 1 for a, b_ in zip(reg2b, b2s)), "VERIFICATION FAILED"
    wbx.Save()
    if xlsb_free(): wbx.SaveAs(B, FileFormat=50); print("xlsb exported")
    else: print("xlsb OPEN in Excel - export skipped")
    wbx.Close(False); print("saved (%.0fs)" % (time.time() - t0))
finally: xl.Quit()
