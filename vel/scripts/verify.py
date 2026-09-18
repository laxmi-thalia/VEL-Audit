import openpyxl, collections
p=r"C:\Users\pawar\Downloads\VEL_Sales_Register_FY2025-26_DRAFT.xlsx"
wb=openpyxl.load_workbook(p, read_only=True, data_only=False)
print("sheets:", wb.sheetnames)
ws=wb["SR_2025-26"]
st=collections.Counter(); dt=collections.Counter(); sup=collections.Counter()
blankC=blankQ=rows=0
for r in ws.iter_rows(min_row=5, max_row=27006, min_col=1, max_col=50, values_only=True):
    if r[0] is None and r[2] is None: continue
    rows+=1
    st[r[1]]+=1; dt[r[7]]+=1; sup[r[8]]+=1
    if r[2] in (None,""): blankC+=1
    if r[16] in (None,""): blankQ+=1
print(f"data rows: {rows}")
print(f"states ({len(st)}): {sorted(st)}")
print(f"doc type codes ({len(dt)}): {dict(dt)}")
print(f"supply type codes ({len(sup)}): {dict(sup)}")
print(f"blank My GSTIN: {blankC} | blank Taxable Value: {blankQ}")
wb.close()
